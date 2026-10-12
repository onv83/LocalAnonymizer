import shutil
import threading
from pathlib import Path
from zipfile import ZipFile

import pytest
from docx import Document
from docx.oxml.ns import qn
from openpyxl import Workbook

from local_anonymizer.engine import AnonymizationResult
from local_anonymizer.service import AnonymizerService, _stamp_artifact, _verify_artifact_stamp
from local_anonymizer.storage import MappingError, Storage


def _text_document(service: AnonymizerService, tmp_path: Path, name: str, text: str = "Иванов Сергей Петрович"):
    source = tmp_path / name
    source.write_text(text, encoding="utf-8")
    document_id = service.import_file(source)
    return document_id, service.anonymize(document_id)


@pytest.mark.parametrize("suffix", [".docx", ".xlsx"])
def test_ooxml_stamp_changes_only_core_properties(tmp_path, suffix, monkeypatch):
    path = tmp_path / f"source{suffix}"
    if suffix == ".docx":
        document = Document()
        document.add_paragraph("Иванов Сергей Петрович")
        document.save(path)
    else:
        workbook = Workbook()
        workbook.active["A1"] = "Иванов Сергей Петрович"
        workbook.save(path)

    with ZipFile(path) as archive:
        before = {info.filename: archive.read(info.filename) for info in archive.infolist()}

    def forbidden_save(*args, **kwargs):
        raise AssertionError("artifact stamping must not reserialize the Office package")

    monkeypatch.setattr("docx.document.Document.save", forbidden_save)
    monkeypatch.setattr("openpyxl.workbook.workbook.Workbook.save", forbidden_save)

    _stamp_artifact(path, "document-id", "operation-id")

    with ZipFile(path) as archive:
        after = {info.filename: archive.read(info.filename) for info in archive.infolist()}
    assert before.keys() == after.keys()
    assert before["docProps/core.xml"] != after["docProps/core.xml"]
    assert {name: data for name, data in before.items() if name != "docProps/core.xml"} == {
        name: data for name, data in after.items() if name != "docProps/core.xml"
    }


def test_docx_document_variables_force_manual_review_status(tmp_path):
    service = AnonymizerService(tmp_path / "app")
    source = tmp_path / "docvars.docx"
    document = Document()
    document.add_paragraph("Обычный текст")
    doc_vars = document.settings.element.makeelement(qn("w:docVars"))
    doc_var = document.settings.element.makeelement(qn("w:docVar"))
    doc_var.set(qn("w:name"), "HiddenSetting")
    doc_var.set(qn("w:val"), "SETTINGS_DOCVAR_SECRET_CANARY")
    doc_vars.append(doc_var)
    document.settings.element.append(doc_vars)
    document.save(source)

    document_id = service.import_file(source)
    report = service.anonymize(document_id)

    assert service.storage.get_document(document_id)["status"] == "Требуется ручная проверка"
    assert b"SETTINGS_DOCVAR_SECRET_CANARY" not in Path(report["output_path"]).read_bytes()


def test_text_restore_requires_matching_artifact_stamp_before_creating_output(tmp_path):
    service = AnonymizerService(tmp_path / "app")
    first_id, first = _text_document(service, tmp_path, "first.txt")
    second_id, second = _text_document(service, tmp_path, "second.txt", "Петров Иван Сергеевич")

    edited = tmp_path / "edited.txt"
    edited.write_bytes(Path(first["output_path"]).read_bytes())
    restored = service.restore(first_id, edited)
    assert "Иванов Сергей Петрович" in Path(restored["output_path"]).read_text(encoding="utf-8")

    expected_output = service.storage.documents_dir / second_id / "edited.restored.txt"
    with pytest.raises(MappingError):
        service.restore(second_id, edited)
    assert not expected_output.exists()

    unstamped = tmp_path / "unstamped.txt"
    unstamped.write_text("[ФИО-1]", encoding="utf-8")
    with pytest.raises(MappingError):
        service.restore(first_id, unstamped)


def test_ooxml_restore_requires_matching_core_property_stamp(tmp_path):
    service = AnonymizerService(tmp_path / "app")

    docx_source = tmp_path / "source.docx"
    doc = Document(); doc.add_paragraph("Иванов Сергей Петрович"); doc.save(docx_source)
    docx_id = service.import_file(docx_source)
    docx_report = service.anonymize(docx_id)
    edited_docx = tmp_path / "edited.docx"
    edited = Document(docx_report["output_path"]); edited.add_paragraph("[ФИО-1]"); edited.save(edited_docx)
    assert service.restore(docx_id, edited_docx)["unknown_markers"] == []
    stripped = Document(edited_docx); stripped.core_properties.identifier = None; stripped.save(edited_docx)
    with pytest.raises(MappingError):
        service.restore(docx_id, edited_docx)

    xlsx_source = tmp_path / "source.xlsx"
    workbook = Workbook(); workbook.active["A1"] = "Петров Иван Сергеевич"; workbook.save(xlsx_source)
    xlsx_id = service.import_file(xlsx_source)
    xlsx_report = service.anonymize(xlsx_id)
    with pytest.raises(MappingError):
        service.restore(xlsx_id, docx_report["output_path"])
    assert service.restore(xlsx_id, xlsx_report["output_path"])["unknown_markers"] == []


def test_wrong_docx_stamp_fails_before_user_destination_is_created(tmp_path):
    service = AnonymizerService(tmp_path / "app")
    first = tmp_path / "first.docx"
    first_doc = Document()
    first_doc.add_paragraph("Иванов Сергей Петрович")
    first_doc.save(first)
    first_id = service.import_file(first)
    first_report = service.anonymize(first_id)

    second = tmp_path / "second.docx"
    second_doc = Document()
    second_doc.add_paragraph("Петров Иван Сергеевич")
    second_doc.save(second)
    second_id = service.import_file(second)
    service.anonymize(second_id)

    destination = tmp_path / "chosen" / "must-not-exist.docx"
    with pytest.raises(MappingError, match="не соответствует"):
        service.restore(second_id, first_report["output_path"], destination)
    assert not destination.exists()


def test_restore_processes_verified_private_snapshot_when_external_docx_is_swapped(tmp_path, monkeypatch):
    service = AnonymizerService(tmp_path / "app")
    first = tmp_path / "first.docx"
    first_doc = Document(); first_doc.add_paragraph("FIRST Иванов Сергей Петрович"); first_doc.save(first)
    first_id = service.import_file(first)
    first_report = service.anonymize(first_id)

    second = tmp_path / "second.docx"
    second_doc = Document(); second_doc.add_paragraph("SECOND Петров Иван Сергеевич"); second_doc.save(second)
    second_id = service.import_file(second)
    second_report = service.anonymize(second_id)

    edited = tmp_path / "edited.docx"
    shutil.copy2(first_report["output_path"], edited)
    second_bytes = Path(second_report["output_path"]).read_bytes()

    def verify_then_swap(path, document_id, operation_id):
        result = _verify_artifact_stamp(path, document_id, operation_id)
        edited.write_bytes(second_bytes)
        return result

    monkeypatch.setattr("local_anonymizer.service._verify_artifact_stamp", verify_then_swap)
    destination = tmp_path / "restored.docx"
    before_temp = set(service.storage.temp_dir.iterdir())

    service.restore(first_id, edited, destination)

    assert [p.text for p in Document(destination).paragraphs] == ["FIRST Иванов Сергей Петрович"]
    assert edited.read_bytes() == second_bytes
    assert set(service.storage.temp_dir.iterdir()) == before_temp


def test_export_anonymized_is_status_allowlisted_and_review_requires_confirmation(tmp_path):
    service = AnonymizerService(tmp_path / "app")
    document_id, report = _text_document(service, tmp_path, "source.txt")
    row = service.storage.get_document(document_id)

    for status in ("Импортирован", "Экспорт запрещён", "Неизвестно"):
        service.storage.connection.execute("update documents set status=? where id=?", (status, document_id))
        service.storage.connection.commit()
        with pytest.raises(PermissionError):
            service.export_anonymized(document_id, tmp_path / f"{status}.txt", confirmation=True)

    service.storage.connection.execute("update documents set status=? where id=?", ("Требуется ручная проверка", document_id))
    service.storage.connection.commit()
    with pytest.raises(PermissionError):
        service.export_anonymized(document_id, tmp_path / "review.txt")
    assert service.export_anonymized(document_id, tmp_path / "confirmed.txt", confirmation=True).is_file()

    service.storage.connection.execute("update documents set status=? where id=?", ("Готов", document_id))
    service.storage.connection.commit()
    assert service.export_anonymized(document_id, tmp_path / "ready.txt").is_file()


def test_successful_operation_removes_previous_output_but_failed_commit_preserves_it(tmp_path, monkeypatch):
    service = AnonymizerService(tmp_path / "app")
    document_id, first = _text_document(service, tmp_path, "source.txt")
    first_output = Path(first["output_path"])
    Path(service.storage.get_document(document_id)["owned_source"]).write_text("Петров Иван Сергеевич", encoding="utf-8")
    second = service.anonymize(document_id)
    assert Path(second["output_path"]).is_file()
    assert not first_output.exists()

    second_output = Path(second["output_path"])
    monkeypatch.setattr(service.storage, "commit_operation", lambda *args, **kwargs: (_ for _ in ()).throw(OSError("db fault")))
    with pytest.raises(OSError, match="db fault"):
        service.anonymize(document_id)
    assert second_output.is_file()


def test_same_document_anonymizations_are_serialized_and_only_newest_output_remains(tmp_path):
    service = AnonymizerService(tmp_path / "app")
    source = tmp_path / "source.txt"
    source.write_text("source", encoding="utf-8")
    document_id = service.import_file(source)
    first_entered = threading.Event()
    release_first = threading.Event()
    second_done = threading.Event()
    results = {}
    errors = []

    class Engine:
        def __init__(self, value, *, block=False):
            self.value, self.block = value, block

        def anonymize(self, _text):
            if self.block:
                first_entered.set()
                assert release_first.wait(5)
            return AnonymizationResult(
                text="[ФИО-1]", mapping={"[ФИО-1]": self.value}, total=1
            )

    def run(name, engine, done=None):
        try:
            results[name] = service.anonymize(document_id, engine)
        except Exception as exc:
            errors.append(exc)
        finally:
            if done:
                done.set()

    first = threading.Thread(target=run, args=("first", Engine("FIRST", block=True)))
    second = threading.Thread(target=run, args=("second", Engine("SECOND"), second_done))
    first.start()
    assert first_entered.wait(5)
    second.start()
    second_done.wait(1)
    release_first.set()
    first.join(5); second.join(5)

    assert not errors
    assert not first.is_alive() and not second.is_alive()
    row = service.storage.get_document(document_id)
    assert row["operation_id"] == results["second"]["operation_id"]
    assert service.storage.load_mapping(document_id, row["operation_id"]) == {"[ФИО-1]": "SECOND"}
    outputs = list((service.storage.documents_dir / document_id).glob("anonymized-*"))
    assert outputs == [Path(row["output_path"])]


def test_stamp_failure_removes_new_orphan_and_preserves_previous_operation(tmp_path, monkeypatch):
    service = AnonymizerService(tmp_path / "app")
    document_id, first = _text_document(service, tmp_path, "source.txt")
    before = service.storage.get_document(document_id)
    first_mapping = service.storage.load_mapping(document_id, first["operation_id"])
    Path(before["owned_source"]).write_text("Петров Иван Сергеевич", encoding="utf-8")

    def stamp_then_fail(path, *_args):
        Path(path).write_text("partial stamped artifact", encoding="utf-8")
        raise OSError("stamp fault")

    monkeypatch.setattr("local_anonymizer.service._stamp_artifact", stamp_then_fail)
    with pytest.raises(OSError, match="stamp fault"):
        service.anonymize(document_id)

    after = service.storage.get_document(document_id)
    assert after["operation_id"] == before["operation_id"]
    assert after["output_path"] == before["output_path"]
    assert after["status"] == before["status"]
    assert Path(after["output_path"]).is_file()
    assert service.storage.load_mapping(document_id, after["operation_id"]) == first_mapping
    assert list((service.storage.documents_dir / document_id).glob("anonymized-*")) == [Path(first["output_path"])]


def test_storage_import_db_failure_rolls_back_plaintext_copy(tmp_path):
    source = tmp_path / "source.txt"
    source.write_text("СЕКРЕТНЫЙ ИСХОДНЫЙ ТЕКСТ", encoding="utf-8")
    storage = Storage(tmp_path / "app")
    storage.connection.executescript("""
        create trigger reject_import before insert on documents
        begin select raise(fail, 'db fault'); end;
    """)

    with pytest.raises(Exception, match="db fault"):
        storage.import_document(source)

    assert storage.list_documents() == []
    assert list(storage.documents_dir.iterdir()) == []


def test_doc_conversion_failure_removes_database_row_and_plaintext_folder(tmp_path, monkeypatch):
    source = tmp_path / "legacy.doc"
    source.write_bytes(b"PLAINTEXT LEGACY DOCUMENT")
    service = AnonymizerService(tmp_path / "app")

    def fail_conversion(_source, _target):
        raise RuntimeError("conversion fault")

    monkeypatch.setattr("local_anonymizer.service.convert_doc_to_docx", fail_conversion)
    with pytest.raises(RuntimeError, match="conversion fault"):
        service.import_file(source)

    assert service.storage.list_documents() == []
    assert list(service.storage.documents_dir.iterdir()) == []


def test_doc_conversion_database_failure_removes_converted_and_original_copies(tmp_path, monkeypatch):
    source = tmp_path / "legacy.doc"
    source.write_bytes(b"PLAINTEXT LEGACY DOCUMENT")
    service = AnonymizerService(tmp_path / "app")
    service.storage.connection.executescript("""
        create trigger reject_conversion before update on documents
        begin select raise(fail, 'conversion db fault'); end;
    """)

    monkeypatch.setattr(
        "local_anonymizer.service.convert_doc_to_docx",
        lambda _source, target: Path(target).write_bytes(b"CONVERTED PLAINTEXT"),
    )
    with pytest.raises(Exception, match="conversion db fault"):
        service.import_file(source)

    assert service.storage.list_documents() == []
    assert list(service.storage.documents_dir.iterdir()) == []
