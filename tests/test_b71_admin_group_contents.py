"""Synthetic contents headings versus explicit irreversible secrets."""
import pytest
from local_anonymizer.engine import Anonymizer


@pytest.mark.parametrize("heading", ["Domain Admins12", "Domain Admins....12", "Enterprise Admins....37", "Schema Admins12"])
def test_group_heading_with_page_is_not_password(heading):
    source = f"Оглавление\n{heading}\nСледующий раздел"
    result = Anonymizer().anonymize(source)
    assert result.text == source


@pytest.mark.parametrize("secret", ["Admins12", "Admins....12", "admin123456", "password456789"])
def test_password_tokens_stay_irreversible(secret):
    result = Anonymizer().anonymize(f"Пароль: {secret}")
    assert secret not in result.text
    assert secret not in result.mapping.values()


def test_group_headings_docx_runs_and_table_restore(tmp_path):
    from docx import Document
    from local_anonymizer.documents import DocxProcessor

    source, output, restored = [tmp_path / n for n in ("source.docx", "out.docx", "restored.docx")]
    doc = Document()
    doc.add_paragraph("Оглавление")
    paragraph = doc.add_paragraph()
    for part in ("Domain ", "Admins", "....12"):
        paragraph.add_run(part).bold = True
    doc.add_table(rows=1, cols=1).cell(0, 0).text = "Schema Admins12"
    doc.save(source)
    original = source.read_bytes()
    processor = DocxProcessor()
    result = processor.anonymize(source, output)
    out = Document(output)
    assert out.paragraphs[1].text == "Domain Admins....12"
    assert out.tables[0].cell(0, 0).text == "Schema Admins12"
    processor.restore(output, restored, result.mapping, package_archive=result.package_archive)
    assert restored.read_bytes() == original


@pytest.mark.parametrize("separator", [" ", "\n", "\ufdd0"])
def test_inline_group_like_secret_is_not_exempted(separator):
    result = Anonymizer().anonymize(f"Пароль:{separator}Domain Admins12")
    assert "Admins12" not in result.text
