import ast
from pathlib import Path
from zipfile import ZipFile

import pytest
from docx import Document

from local_anonymizer.documents import DocxProcessor
from scripts import release_verify as rv
from scripts.release_verify import build_acceptance_source


def test_release_acceptance_document_contains_visible_graphic_and_round_trips_exactly(tmp_path):
    source = tmp_path / "acceptance-source.docx"
    anonymized = tmp_path / "acceptance-anonymized.docx"
    restored = tmp_path / "acceptance-restored.docx"

    image_bytes = build_acceptance_source(
        source,
        Path("acceptance_sample.txt"),
        Path("docs/images/acceptance-diagram.png"),
    )

    source_document = Document(source)
    assert len(source_document.inline_shapes) == 1
    source_text = "\n".join(
        paragraph.text for paragraph in source_document.paragraphs
    )
    assert "qzxsafe.test11" in source_text
    assert "go.qzxsafe.test14" in source_text
    assert "порт 8080" in source_text
    with ZipFile(source) as package:
        embedded_images = [
            package.read(name)
            for name in package.namelist()
            if name.startswith("word/media/")
        ]
    assert embedded_images == [image_bytes]

    result = DocxProcessor().anonymize(source, anonymized)
    anonymized_text = "\n".join(paragraph.text for paragraph in Document(anonymized).paragraphs)
    assert "qzxsafe.test" not in anonymized_text
    assert any(
        "[ДОМЕН-" in paragraph.text and paragraph.text.endswith("]11")
        for paragraph in Document(anonymized).paragraphs
    )
    assert any(
        "[ДОМЕН-" in paragraph.text and paragraph.text.endswith("]14")
        for paragraph in Document(anonymized).paragraphs
    )
    assert "порт [ПОРТ-" in anonymized_text
    assert "8080" not in anonymized_text
    with ZipFile(anonymized) as package:
        assert not any(name.startswith("word/media/") for name in package.namelist())
        assert "[СХЕМА-1]".encode("utf-8") in package.read("word/document.xml")

    DocxProcessor().restore(
        anonymized,
        restored,
        result.mapping,
        package_archive=result.package_archive,
    )

    assert restored.read_bytes() == source.read_bytes()
    with ZipFile(restored) as package:
        restored_images = [
            package.read(name)
            for name in package.namelist()
            if name.startswith("word/media/")
        ]
    assert restored_images == [image_bytes]


def test_release_verifier_uses_explicit_checks_that_survive_python_optimization():
    verifier = Path("scripts/release_verify.py")
    tree = ast.parse(verifier.read_text(encoding="utf-8"))

    assert not [node.lineno for node in ast.walk(tree) if isinstance(node, ast.Assert)]


def test_category_control_verifier_uses_explicit_checks_that_survive_python_optimization():
    verifier = Path("scripts/verify_category_controls.py")
    tree = ast.parse(verifier.read_text(encoding="utf-8"))

    assert not [node.lineno for node in ast.walk(tree) if isinstance(node, ast.Assert)]


def test_release_package_verification_rejects_original_image_bytes_hidden_outside_media(tmp_path):
    source = tmp_path / "source.docx"
    anonymized = tmp_path / "anonymized.docx"
    image_bytes = build_acceptance_source(
        source,
        Path("acceptance_sample.txt"),
        Path("docs/images/acceptance-diagram.png"),
    )
    DocxProcessor().anonymize(source, anonymized)
    with ZipFile(anonymized, "a") as package:
        package.writestr("customXml/leaked-original.bin", image_bytes)

    with pytest.raises(rv.ReleaseVerificationError, match="original graphic bytes"):
        rv.verify_anonymized_docx(anonymized, image_bytes)


def test_acceptance_source_rejects_an_unreviewed_synthetic_image(tmp_path):
    changed_image = tmp_path / "changed.png"
    changed_image.write_bytes(Path("docs/images/acceptance-diagram.png").read_bytes() + b"changed")

    with pytest.raises(rv.ReleaseVerificationError, match="reviewed synthetic asset"):
        build_acceptance_source(
            tmp_path / "source.docx",
            Path("acceptance_sample.txt"),
            changed_image,
        )


def test_release_pair_verification_rejects_missing_and_mutated_mapping(tmp_path):
    source = tmp_path / "source.docx"
    build_acceptance_source(
        source,
        Path("acceptance_sample.txt"),
        Path("docs/images/acceptance-diagram.png"),
    )
    service = rv.AnonymizerService(tmp_path / "appdata")
    document_id = service.import_file(source)
    service.anonymize(document_id)

    assert rv.verify_pair_mapping_fail_closed(service, document_id, tmp_path) == {
        "missing_mapping_rejected": True,
        "mutated_mapping_rejected": True,
    }
