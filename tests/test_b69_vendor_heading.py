"""Synthetic regressions for a column heading mistaken for a vendor."""
from local_anonymizer.engine import Anonymizer


def test_software_vendor_heading_does_not_seed_prose_replacements():
    source = "Производитель ПО\nРекомендации по настройке. Под управлением ПО. После проверки порт доступен."
    result = Anonymizer(disabled_categories={"СОКРАЩЕНИЯ"}).anonymize(source)
    assert result.text == source
    assert "ПРОИЗВОДИТЕЛЬ" not in result.counts


def test_vendor_heading_in_docx_does_not_corrupt_body_or_table(tmp_path):
    from docx import Document
    from local_anonymizer.documents import DocxProcessor

    source = tmp_path / "synthetic-vendor.docx"
    output = tmp_path / "synthetic-vendor-output.docx"
    restored = tmp_path / "synthetic-vendor-restored.docx"
    doc = Document()
    doc.add_paragraph("Рекомендации по настройке. После проверки порт доступен.")
    table = doc.add_table(rows=2, cols=2)
    table.cell(0, 0).text = "Производитель ПО"
    table.cell(0, 1).text = "Назначение"
    table.cell(1, 0).text = "ExampleVendor"
    table.cell(1, 1).text = "Система резервного копирования"
    doc.save(source)
    original = source.read_bytes()
    processor = DocxProcessor(Anonymizer(disabled_categories={"СОКРАЩЕНИЯ"}))
    result = processor.anonymize(source, output)
    out = Document(output)
    assert out.paragraphs[0].text == doc.paragraphs[0].text
    assert out.tables[0].cell(0, 0).text == "Производитель ПО"
    processor.restore(output, restored, result.mapping, package_archive=result.package_archive)
    assert source.read_bytes() == restored.read_bytes() == original


def test_explicit_synthetic_vendor_remains_detectable():
    source = "Производитель ExampleVendor\nОборудование ExampleVendor работает."
    result = Anonymizer(disabled_categories={"СОКРАЩЕНИЯ"}).anonymize(source)
    assert result.text == "Производитель [ПРОИЗВОДИТЕЛЬ-1]\nОборудование [ПРОИЗВОДИТЕЛЬ-1] работает."
