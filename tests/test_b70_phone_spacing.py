"""Synthetic contact examples; no user document values."""
import pytest

from local_anonymizer.engine import Anonymizer


@pytest.mark.parametrize("gap", [" ", "\t", "\u00a0", "\u202f"])
def test_phone_with_space_after_plus_is_redacted(gap):
    source = f"Контакт: +{gap}7 (000) 123 45 67."
    result = Anonymizer().anonymize(source)
    assert result.text == "Контакт: [ТЕЛЕФОН-1]."
    assert result.mapping["[ТЕЛЕФОН-1]"] == source[len("Контакт: "):-1]


def test_phone_category_can_be_disabled():
    source = "Контакт: + 7 (000) 123 45 67."
    assert Anonymizer(disabled_categories={"ТЕЛЕФОН"}).anonymize(source).text == source


@pytest.mark.parametrize("source", ["Итого + 7 единиц", "Температура + 12 градусов", "+\n7 (000) 123 45 67"])
def test_plus_spacing_does_not_join_unrelated_values(source):
    result = Anonymizer().anonymize(source)
    assert not any(key.startswith("[ТЕЛЕФОН-") for key in result.mapping)


def test_spaced_phone_in_docx_split_runs_and_table_restores_exactly(tmp_path):
    from docx import Document
    from local_anonymizer.documents import DocxProcessor

    source, output, restored = [tmp_path / name for name in ("source.docx", "output.docx", "restored.docx")]
    doc = Document()
    para = doc.add_paragraph()
    para.add_run("Контакт: ").bold = True
    para.add_run("+ ").italic = True
    para.add_run("7 (000) 123 45 67")
    table = doc.add_table(rows=1, cols=2)
    table.cell(0, 0).text = "Телефон"
    table.cell(0, 1).text = "+ 7 (000) 123 45 67"
    doc.save(source)
    original = source.read_bytes()
    processor = DocxProcessor()
    result = processor.anonymize(source, output)
    out = Document(output)
    assert out.paragraphs[0].text == "Контакт: [ТЕЛЕФОН-1]"
    assert out.paragraphs[0].runs[0].bold is True
    assert out.tables[0].cell(0, 1).text == "[ТЕЛЕФОН-1]"
    processor.restore(output, restored, result.mapping, package_archive=result.package_archive)
    assert source.read_bytes() == original
    assert restored.read_bytes() == original
