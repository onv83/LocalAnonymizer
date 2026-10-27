from zipfile import ZIP_DEFLATED, ZipFile

import pytest
from lxml import etree
from openpyxl import Workbook, load_workbook
from openpyxl.styles import Font
from openpyxl.worksheet.datavalidation import DataValidation

from local_anonymizer.formats import XlsxProcessor
from local_anonymizer.xlsx_package import SHEET_NS

@pytest.mark.parametrize('shared', [False, True])
@pytest.mark.parametrize('operation', ['anonymize', 'restore'])
def test_literal_equals_string_never_becomes_formula(tmp_path, shared, operation):
    source, output = tmp_path / 'source.xlsx', tmp_path / 'output.xlsx'
    text = '=WEBSERVICE("' + SECRET + '")'
    wb = Workbook()
    wb.active['A1'] = text
    wb.active['A1'].data_type = 's'
    wb.save(source)
    if shared:
        with ZipFile(source) as archive:
            root = etree.fromstring(archive.read('xl/worksheets/sheet1.xml'))
            types = etree.fromstring(archive.read('[Content_Types].xml'))
        cell = root.find('.//x:c', NS)
        cell.remove(cell.find('x:is', NS))
        cell.set('t', 's')
        etree.SubElement(cell, f'{{{SHEET_NS}}}v').text = '0'
        sst = etree.Element(f'{{{SHEET_NS}}}sst')
        etree.SubElement(etree.SubElement(sst, f'{{{SHEET_NS}}}si'), f'{{{SHEET_NS}}}t').text = text
        etree.SubElement(types, '{%s}Override' % types.nsmap[None], PartName='/xl/sharedStrings.xml', ContentType='application/vnd.openxmlformats-officedocument.spreadsheetml.sharedStrings+xml')
        patch_package(source, {'xl/worksheets/sheet1.xml': etree.tostring(root), 'xl/sharedStrings.xml': etree.tostring(sst), '[Content_Types].xml': etree.tostring(types)})
    args = ({},) if operation == 'restore' else ()
    getattr(XlsxProcessor(), operation)(source, output, *args)
    cell = load_workbook(output).active['A1']
    assert cell.data_type == 's'
    assert cell.value.startswith('=')
    if operation == 'restore':
        assert cell.value == text
    assert cell.value != '[BLOCKED-FORMULA]'
    if operation == 'anonymize':
        assert SECRET.encode() not in package_bytes(output)


@pytest.mark.parametrize('formula', ['=_xlfn.WEBSERVICE(A2)', '=_XLFN.webservice (A2)', '=IF(1,_xlfn.WEBSERVICE(A2),0)'])
def test_prefixed_dangerous_formula_is_blocked(tmp_path, formula):
    source, output = tmp_path / 'source.xlsx', tmp_path / 'output.xlsx'
    wb = Workbook()
    wb.active['A1'] = formula
    wb.save(source)
    result = XlsxProcessor().anonymize(source, output)
    assert not result.ready
    assert load_workbook(output).active['A1'].value == '[BLOCKED-FORMULA]'
    assert b'WEBSERVICE' not in package_bytes(output).upper()


@pytest.mark.parametrize('part', ['customXml/item1.xml', 'xl/pivotCache/pivotCacheRecords1.xml', 'xl/connections.xml', 'xl/tables/table1.xml', 'xl/unknown.bin'])
def test_unsupported_package_surfaces_rejected_atomically(tmp_path, part):
    source, output = tmp_path / 'source.xlsx', tmp_path / 'output.xlsx'
    wb = Workbook()
    wb.save(source)
    patch_package(source, {part: ('<hidden>' + SECRET + '</hidden>').encode()})
    output.write_bytes(b'previous-output')
    with pytest.raises(ValueError, match='XLSX'):
        XlsxProcessor().anonymize(source, output)
    assert output.read_bytes() == b'previous-output'
    assert not list(tmp_path.glob('.output-*.xlsx'))


def test_validation_formula_payload_is_removed_and_blocks_readiness(tmp_path):
    source, output = tmp_path / 'source.xlsx', tmp_path / 'output.xlsx'
    wb = Workbook()
    validation = DataValidation(type='list', formula1='"' + SECRET + '"')
    validation.add('A1')
    wb.active.add_data_validation(validation)
    wb.save(source)
    result = XlsxProcessor().anonymize(source, output)
    assert not result.ready
    assert SECRET.encode() not in package_bytes(output)
    assert not load_workbook(output).active.data_validations.dataValidation


def test_sensitive_sheet_name_rejected_without_publishing(tmp_path):
    source, output = tmp_path / 'source.xlsx', tmp_path / 'output.xlsx'
    wb = Workbook()
    wb.active.title = SECRET
    wb.save(source)
    with pytest.raises(ValueError, match='XLSX'):
        XlsxProcessor().anonymize(source, output)
    assert not output.exists()


@pytest.mark.parametrize('surface', ['extension', 'foreign', 'filter', 'code_name', 'conditional'])
def test_unsupported_worksheet_surfaces_fail_closed(tmp_path, surface):
    source, output = tmp_path / 'source.xlsx', tmp_path / 'output.xlsx'
    wb = Workbook()
    wb.save(source)
    with ZipFile(source) as archive:
        root = etree.fromstring(archive.read('xl/worksheets/sheet1.xml'))
    if surface == 'code_name':
        root.find('x:sheetPr', NS).set('codeName', SECRET)
    else:
        tag = {'extension': f'{{{SHEET_NS}}}extLst', 'foreign': '{urn:unknown}payload', 'filter': f'{{{SHEET_NS}}}autoFilter', 'conditional': f'{{{SHEET_NS}}}conditionalFormatting'}[surface]
        etree.SubElement(root, tag).text = SECRET
    patch_package(source, {'xl/worksheets/sheet1.xml': etree.tostring(root)})
    with pytest.raises(ValueError, match='XLSX'):
        XlsxProcessor().anonymize(source, output)
    assert not output.exists()


@pytest.mark.parametrize('surface', ['app_metadata', 'workbook_extension', 'workbook_code_name'])
def test_hidden_workbook_metadata_cannot_leak(tmp_path, surface):
    source, output = tmp_path / 'source.xlsx', tmp_path / 'output.xlsx'
    Workbook().save(source)
    part = 'docProps/app.xml' if surface == 'app_metadata' else 'xl/workbook.xml'
    with ZipFile(source) as archive:
        root = etree.fromstring(archive.read(part))
    if surface == 'app_metadata':
        etree.SubElement(root, '{%s}Company' % etree.QName(root).namespace).text = SECRET
    elif surface == 'workbook_code_name':
        root.find('x:workbookPr', NS).set('codeName', SECRET)
    else:
        etree.SubElement(root, f'{{{SHEET_NS}}}extLst').text = SECRET
    patch_package(source, {part: etree.tostring(root)})
    if surface == 'app_metadata':
        XlsxProcessor().anonymize(source, output)
        assert SECRET.encode() not in package_bytes(output)
    else:
        with pytest.raises(ValueError, match='XLSX'):
            XlsxProcessor().anonymize(source, output)
        assert not output.exists()


@pytest.mark.parametrize('operation', ['anonymize', 'restore'])
@pytest.mark.parametrize('element,attribute', [
    ('fileSharing', 'userName'), ('fileVersion', 'appName'),
    ('customWorkbookView', 'name'),
])
def test_workbook_identity_metadata_removed_canonically(tmp_path, operation, element, attribute):
    source, output = tmp_path / 'source.xlsx', tmp_path / 'output.xlsx'
    Workbook().save(source)
    with ZipFile(source) as archive:
        root = etree.fromstring(archive.read('xl/workbook.xml'))
    parent = root
    if element == 'customWorkbookView':
        parent = etree.SubElement(root, f'{{{SHEET_NS}}}customWorkbookViews')
    node = etree.SubElement(parent, f'{{{SHEET_NS}}}{element}')
    node.set(attribute, SECRET)
    if element == 'customWorkbookView':
        node.set('guid', '{00000000-0000-0000-0000-000000000000}')
        node.set('windowWidth', '100')
        node.set('windowHeight', '100')
        node.set('activeSheetId', '0')
    patch_package(source, {'xl/workbook.xml': etree.tostring(root)})
    getattr(XlsxProcessor(), operation)(source, output, *([{}] if operation == 'restore' else []))
    assert SECRET.encode() not in package_bytes(output)
    with ZipFile(output) as archive:
        cleaned = etree.fromstring(archive.read('xl/workbook.xml'))
    assert not cleaned.findall(f'.//x:{element}', NS)
    load_workbook(output)


@pytest.mark.parametrize('operation', ['anonymize', 'restore'])
@pytest.mark.parametrize('secret', ['audit.person@example.com', 'other.identity@example.net'])
@pytest.mark.parametrize('surface', ['cellStyle', 'tableStyle', 'defaultTableStyle', 'defaultPivotStyle', 'font', 'numFmt'])
def test_sensitive_style_strings_fail_closed(tmp_path, operation, secret, surface):
    source, output = tmp_path / 'source.xlsx', tmp_path / 'output.xlsx'
    wb = Workbook()
    wb.active['A1'] = 42
    wb.save(source)
    with ZipFile(source) as archive:
        root = etree.fromstring(archive.read('xl/styles.xml'))
    if surface == 'cellStyle':
        root.find('x:cellStyles/x:cellStyle', NS).set('name', secret)
    elif surface == 'tableStyle':
        etree.SubElement(root.find('x:tableStyles', NS), f'{{{SHEET_NS}}}tableStyle', name=secret, count='0')
    elif surface.startswith('default'):
        root.find('x:tableStyles', NS).set(surface, secret)
    elif surface == 'font':
        root.find('x:fonts/x:font/x:name', NS).set('val', secret)
    else:
        etree.SubElement(root.find('x:numFmts', NS), f'{{{SHEET_NS}}}numFmt', numFmtId='164', formatCode='"' + secret + '"0')
    patch_package(source, {'xl/styles.xml': etree.tostring(root)})
    output.write_bytes(b'previous-output')
    with pytest.raises(ValueError, match='XLSX'):
        getattr(XlsxProcessor(), operation)(source, output, *([{}] if operation == 'restore' else []))
    assert output.read_bytes() == b'previous-output'
    assert not list(tmp_path.glob('.output-*.xlsx'))


@pytest.mark.parametrize('operation', ['anonymize', 'restore'])
def test_normal_named_styles_preserve_exact_xml_and_references(tmp_path, operation):
    from openpyxl.styles import NamedStyle
    source, output = tmp_path / 'source.xlsx', tmp_path / 'output.xlsx'
    wb = Workbook()
    wb.add_named_style(NamedStyle(name='Report heading', font=Font(name='Calibri', bold=True)))
    wb.active['A1'] = 42
    wb.active['A1'].style = 'Report heading'
    wb.active['A1'].number_format = '#,##0.00'
    wb.save(source)
    getattr(XlsxProcessor(), operation)(source, output, *([{}] if operation == 'restore' else []))
    with ZipFile(source) as before, ZipFile(output) as after:
        assert before.read('xl/styles.xml') == after.read('xl/styles.xml')
        assert before.read('xl/worksheets/sheet1.xml') == after.read('xl/worksheets/sheet1.xml')
    cell = load_workbook(output).active['A1']
    assert cell.style == 'Report heading'
    assert cell.font.bold and cell.number_format == '#,##0.00'


NS = {"x": SHEET_NS}
SECRET = "audit.person@example.com"


def patch_package(path, updates):
    temporary = path.with_suffix('.tmp')
    with ZipFile(path) as source, ZipFile(temporary, 'w', ZIP_DEFLATED) as target:
        for info in source.infolist():
            target.writestr(info, updates.get(info.filename, source.read(info)))
        for name, value in updates.items():
            if name not in source.namelist():
                target.writestr(name, value)
    temporary.replace(path)


def package_bytes(path):
    with ZipFile(path) as archive:
        return b''.join(archive.read(name) for name in archive.namelist())


@pytest.mark.parametrize('operation', ['anonymize', 'restore'])
def test_formula_caches_removed_without_changing_formula_or_style(tmp_path, operation):
    source, output = tmp_path / 'source.xlsx', tmp_path / 'output.xlsx'
    wb = Workbook()
    wb.active['A1'] = '=1+1'
    wb.active['A1'].font = Font(bold=True)
    wb.active.column_dimensions['A'].width = 29
    wb.save(source)
    with ZipFile(source) as archive:
        root = etree.fromstring(archive.read('xl/worksheets/sheet1.xml'))
    cell = root.find('.//x:c', NS)
    cell.set('t', 'str')
    cell.find('x:v', NS).text = SECRET
    patch_package(source, {'xl/worksheets/sheet1.xml': etree.tostring(root)})
    args = ({},) if operation == 'restore' else ()
    getattr(XlsxProcessor(), operation)(source, output, *args)
    assert SECRET.encode() not in package_bytes(output)
    result = load_workbook(output)
    assert result.active['A1'].value == '=1+1'
    assert result.active['A1'].font.bold
    assert result.active.column_dimensions['A'].width == 29
    assert load_workbook(output, data_only=True).active['A1'].value is None
