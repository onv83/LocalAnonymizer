from pathlib import Path
from zipfile import ZIP_DEFLATED, ZipFile

from openpyxl import Workbook, load_workbook
from openpyxl.comments import Comment
from lxml import etree

from local_anonymizer.formats import process_text_file, restore_text_file, XlsxProcessor
from local_anonymizer.xlsx_package import SHEET_NS, _set_cell_value


def _rewrite_package(path, updates, remove=()):
    temp = path.with_suffix(".tmp.xlsx")
    with ZipFile(path) as source, ZipFile(temp, "w", ZIP_DEFLATED) as target:
        for info in source.infolist():
            if info.filename not in remove:
                target.writestr(info, updates.get(info.filename, source.read(info)))
        for name, payload in updates.items():
            if name not in source.namelist():
                target.writestr(name, payload)
    temp.replace(path)


def test_xlsx_custom_worksheet_part_is_rewritten(tmp_path):
    source, output = tmp_path / "custom.xlsx", tmp_path / "custom-anon.xlsx"
    workbook = Workbook()
    workbook.active["A1"] = "Иванов Сергей Петрович"
    workbook.save(source)
    with ZipFile(source) as archive:
        sheet = archive.read("xl/worksheets/sheet1.xml")
        rels = etree.fromstring(archive.read("xl/_rels/workbook.xml.rels"))
        types = etree.fromstring(archive.read("[Content_Types].xml"))
    worksheet_rel = next(node for node in rels if node.get("Type", "").endswith("/worksheet"))
    worksheet_rel.set("Target", "/xl/worksheets/custom.xml")
    for override in types:
        if override.get("PartName") == "/xl/worksheets/sheet1.xml":
            override.set("PartName", "/xl/worksheets/custom.xml")
    _rewrite_package(
        source,
        {
            "xl/worksheets/custom.xml": sheet,
            "xl/_rels/workbook.xml.rels": etree.tostring(rels, xml_declaration=True, encoding="UTF-8"),
            "[Content_Types].xml": etree.tostring(types, xml_declaration=True, encoding="UTF-8"),
        },
        remove={"xl/worksheets/sheet1.xml"},
    )

    result = XlsxProcessor().anonymize(source, output)

    assert result.ready and load_workbook(output).active["A1"].value == "[ФИО-1]"


def test_xlsx_shared_strings_do_not_retain_anonymized_source_value(tmp_path):
    source, output = tmp_path / "shared.xlsx", tmp_path / "shared-anon.xlsx"
    secret = "Иванов Сергей Петрович"
    workbook = Workbook()
    workbook.active["A1"] = secret
    workbook.save(source)
    with ZipFile(source) as archive:
        sheet = etree.fromstring(archive.read("xl/worksheets/sheet1.xml"))
        rels = etree.fromstring(archive.read("xl/_rels/workbook.xml.rels"))
        types = etree.fromstring(archive.read("[Content_Types].xml"))
    cell = sheet.xpath('.//x:c[@r="A1"]', namespaces={"x": SHEET_NS})[0]
    for child in list(cell):
        cell.remove(child)
    cell.set("t", "s")
    etree.SubElement(cell, f"{{{SHEET_NS}}}v").text = "0"
    rel_ns = rels.nsmap.get(None)
    etree.SubElement(
        rels,
        f"{{{rel_ns}}}Relationship",
        Id="rId999",
        Type="http://schemas.openxmlformats.org/officeDocument/2006/relationships/sharedStrings",
        Target="sharedStrings.xml",
    )
    type_ns = types.nsmap.get(None)
    etree.SubElement(
        types,
        f"{{{type_ns}}}Override",
        PartName="/xl/sharedStrings.xml",
        ContentType="application/vnd.openxmlformats-officedocument.spreadsheetml.sharedStrings+xml",
    )
    shared = (
        f'<sst xmlns="{SHEET_NS}" count="1" uniqueCount="1">'
        f"<si><r><rPr><b/></rPr><t>{secret}</t></r></si></sst>"
    ).encode("utf-8")
    _rewrite_package(
        source,
        {
            "xl/worksheets/sheet1.xml": etree.tostring(sheet, xml_declaration=True, encoding="UTF-8"),
            "xl/_rels/workbook.xml.rels": etree.tostring(rels, xml_declaration=True, encoding="UTF-8"),
            "[Content_Types].xml": etree.tostring(types, xml_declaration=True, encoding="UTF-8"),
            "xl/sharedStrings.xml": shared,
        },
    )

    result = XlsxProcessor().anonymize(source, output)

    assert result.ready and load_workbook(output).active["A1"].value == "[ФИО-1]"
    with ZipFile(output) as archive:
        assert "xl/sharedStrings.xml" not in archive.namelist()
        assert secret.encode("utf-8") not in b"".join(archive.read(name) for name in archive.namelist())


def test_xlsx_direct_replacement_preserves_rich_text_run_properties():
    cell = etree.fromstring(
        f'<c xmlns="{SHEET_NS}" r="A1" t="inlineStr"><is>'
        '<r><rPr><b/></rPr><t>Иванов </t></r>'
        '<r><rPr><i/></rPr><t>Сергей Петрович</t></r>'
        '</is></c>'.encode()
    )

    _set_cell_value(cell, "[ФИО-1]")

    bold = cell.xpath("count(.//x:r[1]/x:rPr/x:b)", namespaces={"x": SHEET_NS})
    italic = cell.xpath("count(.//x:r[2]/x:rPr/x:i)", namespaces={"x": SHEET_NS})
    text = "".join(cell.xpath(".//x:t/text()", namespaces={"x": SHEET_NS}))
    assert bold == 1 and italic == 1
    assert text == "[ФИО-1]"


def test_xlsx_rejects_duplicate_package_members_before_parsing(tmp_path):
    source = tmp_path / "duplicate.xlsx"
    workbook = Workbook()
    workbook.save(source)
    with ZipFile(source, "a", ZIP_DEFLATED) as archive:
        archive.writestr("xl/workbook.xml", b"<hostile/>")

    import pytest

    with pytest.raises(ValueError, match="повторяющиеся ZIP-части"):
        XlsxProcessor().anonymize(source, tmp_path / "out.xlsx")


def test_text_formats_roundtrip_and_unknown_marker(tmp_path):
    src, anon, restored = tmp_path/"a.md", tmp_path/"a.anonymized.md", tmp_path/"a.restored.md"
    src.write_text("# Иванов Сергей Петрович\nemail a@пример.рф", encoding="utf-8")
    result = process_text_file(src, anon)
    anon.write_text(anon.read_text(encoding="utf-8") + "\n[ФИО-999]", encoding="utf-8")
    unknown = restore_text_file(anon, restored, result.mapping)
    assert "Иванов Сергей Петрович" in restored.read_text(encoding="utf-8")
    assert unknown == ["[ФИО-999]"]


def test_xlsx_preserves_formulas_sanitizes_comments_and_roundtrips(tmp_path):
    src, anon, restored = tmp_path/"a.xlsx", tmp_path/"a.anonymized.xlsx", tmp_path/"a.restored.xlsx"
    wb = Workbook(); ws = wb.active; ws["A1"] = "Иванов Сергей Петрович"; ws["A1"].comment = Comment("CANARY", "AUTHOR"); ws["B1"] = "=SUM(1,2)"; wb.properties.creator = "CANARY_AUTHOR"; wb.save(src)
    result = XlsxProcessor().anonymize(src, anon)
    out = load_workbook(anon, data_only=False)
    assert out.active["A1"].value == "[ФИО-1]"
    assert out.active["A1"].comment is None
    assert out.active["B1"].value == "=SUM(1,2)"
    assert out.properties.creator is None
    unknown = XlsxProcessor().restore(anon, restored, result.mapping)
    assert not unknown
    assert load_workbook(restored).active["A1"].value == "Иванов Сергей Петрович"


def test_xlsx_anonymize_and_restore_do_not_reserialize_with_openpyxl(tmp_path, monkeypatch):
    source, anonymized, restored = tmp_path / "source.xlsx", tmp_path / "anon.xlsx", tmp_path / "restored.xlsx"
    workbook = Workbook()
    workbook.active["A1"] = "Иванов Сергей Петрович"
    workbook.active.column_dimensions["A"].width = 27.5
    workbook.save(source)

    def forbidden_save(*args, **kwargs):
        raise AssertionError("openpyxl must not reserialize the OOXML package")

    monkeypatch.setattr("openpyxl.workbook.workbook.Workbook.save", forbidden_save)
    result = XlsxProcessor().anonymize(source, anonymized)
    assert load_workbook(anonymized).active["A1"].value == "[ФИО-1]"
    assert load_workbook(anonymized).active.column_dimensions["A"].width == 27.5
    XlsxProcessor().restore(anonymized, restored, result.mapping)
    assert load_workbook(restored).active["A1"].value == "Иванов Сергей Петрович"


def test_xlsx_vm_inventory_names_are_removed_from_package_and_restore_exactly(tmp_path):
    source = tmp_path / "vm-source.xlsx"
    anonymized = tmp_path / "vm-anonymized.xlsx"
    restored = tmp_path / "vm-restored.xlsx"
    vm_names = ["qzx-monitor0101", "qzx-repo0101", "qzx-print0101", "qzx-fs0101", "qzx-journal0101"]
    workbook = Workbook()
    for row, name in enumerate(vm_names, start=1):
        workbook.active.cell(row=row, column=1, value=name)
    workbook.save(source)

    result = XlsxProcessor().anonymize(source, anonymized)

    with ZipFile(anonymized) as archive:
        package_payload = b"".join(archive.read(name) for name in archive.namelist())
    assert all(name.encode() not in package_payload for name in vm_names)
    assert set(result.mapping.values()) == set(vm_names)
    XlsxProcessor().restore(anonymized, restored, result.mapping)
    restored_sheet = load_workbook(restored).active
    assert [restored_sheet.cell(row=row, column=1).value for row in range(1, 6)] == vm_names


def test_xlsx_contextual_numeric_requisites_are_replaced_and_restored(tmp_path):
    src, anon, restored = tmp_path/"numbers.xlsx", tmp_path/"numbers.anon.xlsx", tmp_path/"numbers.restored.xlsx"
    wb = Workbook(); ws = wb.active; ws.title = "Реквизиты"
    ws["A1"], ws["B1"] = "ИНН", 7707083893
    ws["A2"], ws["B2"] = "Табельный номер", 4512
    ws["A3"], ws["B3"] = "Обычное число", 7707083893
    ws["B1"].number_format = "0"; ws.merge_cells("D1:E1"); wb.save(src)
    result = XlsxProcessor().anonymize(src, anon)
    out = load_workbook(anon, data_only=False)
    assert out.active["B1"].value == "[ИНН-1]"
    assert out.active["B2"].value == "[ТАБЕЛЬНЫЙ-1]"
    assert out.active["B3"].value == 7707083893
    assert out.active["B1"].number_format == "0" and "D1:E1" in [str(x) for x in out.active.merged_cells.ranges]
    XlsxProcessor().restore(anon, restored, result.mapping)
    restored_ws = load_workbook(restored).active
    assert restored_ws["B1"].value == "7707083893" and restored_ws["B2"].value == "4512"
