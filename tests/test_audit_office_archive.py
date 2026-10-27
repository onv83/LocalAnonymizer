"""Package-level regressions for irreversible Office recovery and ODT metadata."""
import base64
from io import BytesIO
from zipfile import ZipFile, ZIP_STORED

import pytest
from docx import Document

from local_anonymizer.documents import DocxProcessor
from local_anonymizer.odt import OdtProcessor, OFFICE_NS, TEXT_NS, META_NS, MANIFEST_NS

SECRET = "SyntheticAuditSecret73!"
CANARY = "audit.canary@example.org"


def write_odt(path, text="ordinary", hidden=None):
    parts = {
        "content.xml": f'<office:document-content xmlns:office="{OFFICE_NS}" xmlns:text="{TEXT_NS}"><office:body><office:text><text:p text:style-name="Body">{text}' + (f'<text:bookmark text:name="{CANARY}"/>' if hidden == "bookmark" else '') + '</text:p></office:text></office:body></office:document-content>',
        "styles.xml": f'<office:document-styles xmlns:office="{OFFICE_NS}"/>',
        "meta.xml": f'<office:document-meta xmlns:office="{OFFICE_NS}" xmlns:meta="{META_NS}"><office:meta>' + (f'<meta:user-defined meta:name="{CANARY}" meta:value-type="string">ordinary</meta:user-defined>' if hidden == "meta" else '') + '</office:meta></office:document-meta>',
    }
    if hidden == "settings":
        parts["settings.xml"] = f'<office:document-settings xmlns:office="{OFFICE_NS}" xmlns:config="urn:oasis:names:tc:opendocument:xmlns:config:1.0"><office:settings><config:config-item config:name="ViewId" config:type="string">{CANARY}</config:config-item></office:settings></office:document-settings>'
    parts["META-INF/manifest.xml"] = f'<manifest:manifest xmlns:manifest="{MANIFEST_NS}"><manifest:file-entry manifest:full-path="/" manifest:media-type="application/vnd.oasis.opendocument.text"/>' + ''.join(f'<manifest:file-entry manifest:full-path="{name}" manifest:media-type="text/xml"/>' for name in parts) + '</manifest:manifest>'
    with ZipFile(path, "w") as package:
        package.writestr("mimetype", "application/vnd.oasis.opendocument.text", compress_type=ZIP_STORED)
        for name, data in parts.items():
            package.writestr(name, data)


def package_payload(path_or_bytes):
    with ZipFile(BytesIO(path_or_bytes) if isinstance(path_or_bytes, bytes) else path_or_bytes) as package:
        return b"\n".join(package.read(name) for name in package.namelist())


@pytest.mark.parametrize("kind", ["docx", "odt"])
def test_irreversible_secret_is_absent_from_recovery_package(tmp_path, kind):
    source, anonymized, restored = [tmp_path / f"{name}.{kind}" for name in ("source", "anon", "restored")]
    text = f"password: {SECRET}"
    if kind == "docx":
        document = Document()
        document.add_paragraph(text)
        document.save(source)
        processor = DocxProcessor()
    else:
        write_odt(source, text)
        processor = OdtProcessor()
    result = processor.anonymize(source, anonymized)
    assert SECRET.encode() not in package_payload(anonymized)
    archive = result.package_archive or {}
    if "exact_source_b64" in archive:
        assert SECRET.encode() not in package_payload(base64.b64decode(archive["exact_source_b64"]))
    processor.restore(anonymized, restored, result.mapping, package_archive=result.package_archive)
    assert SECRET.encode() not in package_payload(restored)


@pytest.mark.parametrize("surface", ["bookmark", "meta", "settings"])
def test_odt_hidden_metadata_is_not_published(tmp_path, surface):
    source, output = tmp_path / "source.odt", tmp_path / "anon.odt"
    write_odt(source, hidden=surface)
    OdtProcessor().anonymize(source, output)
    assert CANARY.encode() not in package_payload(output)


@pytest.mark.parametrize("kind", ["docx", "odt"])
def test_hidden_irreversible_secret_is_not_archived(tmp_path, kind):
    source, output, restored = [tmp_path / f"{name}.{kind}" for name in ("source", "anon", "restored")]
    if kind == "docx":
        document = Document()
        document.add_paragraph("ordinary")
        document.core_properties.comments = f"password: {SECRET}"
        document.save(source)
        processor = DocxProcessor()
    else:
        write_odt(source, hidden="meta")
        with ZipFile(source) as package:
            parts = [(info, package.read(info)) for info in package.infolist()]
        with ZipFile(source, "w") as package:
            for info, payload in parts:
                if info.filename == "meta.xml":
                    payload = payload.replace(b">ordinary<", f">password: {SECRET}<".encode())
                package.writestr(info, payload)
        processor = OdtProcessor()
    result = processor.anonymize(source, output)
    archive = result.package_archive or {}
    if "exact_source_b64" in archive:
        assert SECRET.encode() not in package_payload(base64.b64decode(archive["exact_source_b64"]))
    processor.restore(output, restored, result.mapping, package_archive=result.package_archive)
    assert SECRET.encode() not in package_payload(restored)
