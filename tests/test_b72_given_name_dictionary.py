"""B72: generic vocabulary only; all full-name combinations are synthetic."""
import pytest

from local_anonymizer.engine import Anonymizer


@pytest.mark.parametrize("given", [
    "Айсылу", "Гульчачак", "Доржи", "Баир", "Патимат", "Залина",
    "Аветик", "Гоар", "Нино", "Фёкла", "Фекла", "Семён", "Семен",
])
def test_broad_given_name_dictionary_is_integrated(given):
    source = f"{given} Леснов"
    result = Anonymizer().anonymize(source)
    assert result.text == "[ФИО-1]"
    assert result.mapping == {"[ФИО-1]": source}


@pytest.mark.parametrize("given,surname", [("Айсылу", "Ли"), ("Баир", "Шо"), ("Нино", "Дюма"), ("Аветик", "Крау")])
def test_atypical_surname_both_orders_and_identity(given, surname):
    source = f"{given} {surname}; {surname} {given}"
    result = Anonymizer().anonymize(source)
    assert result.text == "[ФИО-1]; [ФИО-1]"
    assert len(result.mapping) == 1
    assert Anonymizer(disabled_categories={"ФИО"}).anonymize(source).text == source


@pytest.mark.parametrize("source", [
    "Вера Надежда", "Роман Война", "Любовь Победа", "Система Вера",
    "Продукт Роман", "Компания Любовь", "Технический Отдел", "Новый Рабочий",
    "Айсылу, Крау", "Айсылу\nКрау", "Айсылу\ufdd0Крау", "Айсылу\ufdd3Крау",
    "Айсылу читает", "Роман Сергеевич", "Айсылу\nЛеснов",
    "Леснов\nАйсылу", "АО «Айсылу Крау»", "продукт «Роман Крау»",
    "книга «Вера Крау»",
])
def test_contextual_rule_negatives(source):
    assert "[ФИО-" not in Anonymizer().anonymize(source).text


def test_dictionary_resource_and_manifest():
    import hashlib
    import json
    from importlib.resources import files
    from local_anonymizer.given_names import GIVEN_NAMES
    root = files("local_anonymizer").joinpath("assets/given_names")
    raw = root.joinpath("names.txt").read_bytes()
    names = raw.decode("utf-8").splitlines()
    manifest = json.loads(root.joinpath("manifest.json").read_text(encoding="utf-8"))
    assert names == sorted(set(names))
    assert len(names) == manifest["spellings"] == 10283
    assert len(GIVEN_NAMES) == manifest["normalized_keys"] == 10275
    assert hashlib.sha256(raw).hexdigest() == manifest["resource_sha256"]
    assert all(name.isalpha() and name == name.casefold() for name in names)
    assert {"айсылу", "доржи", "патимат", "аветик", "нино", "гульчачак"} <= GIVEN_NAMES
    assert "BSD 3-Clause" in root.joinpath("LICENSE.txt").read_text(encoding="utf-8")
    genders = json.loads(root.joinpath("genders.json").read_text(encoding="utf-8"))
    assert set(genders) == GIVEN_NAMES
    assert set(genders.values()) <= {"m", "f", "fm"}
    from local_anonymizer.given_names import GIVEN_NAME_FORMS, canonical_given_name
    assert len(GIVEN_NAME_FORMS) == 48191
    assert canonical_given_name("Давиду") == "давид"
    # An exact lexical base must not be silently merged into a different person.
    assert canonical_given_name("Романа") == "романа"
    assert canonical_given_name("Роман") == "роман"


def test_docx_split_runs_tables_boundaries_and_exact_restoration(tmp_path):
    from docx import Document
    from local_anonymizer.documents import DocxProcessor
    source, output, restored = [tmp_path / n for n in ("source.docx", "out.docx", "restored.docx")]
    doc = Document()
    p = doc.add_paragraph()
    p.add_run("Айсы").bold = True
    p.add_run("лу Ли; ")
    p.add_run("Ли Айсылу")
    table = doc.add_table(rows=3, cols=2)
    p = table.cell(0, 0).paragraphs[0]
    for text in ("Аве", "тик ", "Крау"):
        p.add_run(text)
    table.cell(0, 1).text = "Крау Аветик"
    table.cell(1, 0).text = "Нино"
    table.cell(1, 1).text = "Дюма"
    table.cell(2, 0).text = "Баир"
    table.cell(2, 0).add_paragraph("Шо")
    doc.save(source)
    original = source.read_bytes()
    processor = DocxProcessor()
    result = processor.anonymize(source, output)
    out = Document(output)
    assert out.paragraphs[0].text == "[ФИО-1]; [ФИО-1]"
    assert out.paragraphs[0].runs[0].bold
    assert out.tables[0].cell(0, 0).text == out.tables[0].cell(0, 1).text == "[ФИО-2]"
    assert out.tables[0].cell(1, 0).text == "Нино"
    assert out.tables[0].cell(1, 1).text == "Дюма"
    assert out.tables[0].cell(2, 0).text == "Баир\nШо"
    processor.restore(output, restored, result.mapping, package_archive=result.package_archive)
    assert restored.read_bytes() == original == source.read_bytes()


def test_different_given_names_do_not_collapse():
    result = Anonymizer().anonymize("Нино Крау; Нина Крау")
    assert result.text == "[ФИО-1]; [ФИО-2]"


@pytest.mark.parametrize("base,forms", [
    ("Роман", ["Романа", "Роману", "Романом", "Романе"]),
    ("Давид", ["Давида", "Давиду", "Давидом", "Давиде"]),
    ("Игорь", ["Игоря", "Игорю", "Игорем", "Игоре"]),
    ("Николай", ["Николая", "Николаю", "Николаем"]),
    ("Дмитрий", ["Дмитрия", "Дмитрию", "Дмитрием", "Дмитрии"]),
    ("Мария", ["Марии", "Марию", "Марией"]),
    ("Дарья", ["Дарьи", "Дарье", "Дарью", "Дарьей"]),
    ("Любовь", ["Любови", "Любовью"]),
    ("Айсылу", ["Айсылу"]),
    ("Гульчачак", ["Гульчачак"]),
])
def test_inflections_link_unambiguous_person(base, forms):
    source = "; ".join([f"{base} Крау"] + [f"Крау {form}" for form in forms])
    result = Anonymizer().anonymize(source)
    # Upstream includes Романа, Игор and Дмитрии as separate bases. Do not
    # guess away that ambiguity merely to get fewer identity markers.
    expected_ids = {"Роман": [1, 2, 3, 1, 4], "Игорь": [1, 1, 1, 1, 2],
                    "Дмитрий": [1, 1, 1, 1, 2]}.get(base, [1] * (len(forms) + 1))
    assert result.text == "; ".join(f"[ФИО-{n}]" for n in expected_ids)
    assert Anonymizer(disabled_categories={"ФИО"}).anonymize(source).text == source


@pytest.mark.parametrize("text", ["роман завершён", "автор романа", "Давидов метод", "Роману, Крау", "прочитали Романа", "Роман Система"])
def test_inflection_negatives(text):
    assert "[ФИО-" not in Anonymizer().anonymize(text).text


def test_inflected_surnames_both_orders_and_distinct_people():
    result = Anonymizer().anonymize("Давид Леснов; Леснову Давиду; Давид Каменский; Давиду Каменскому")
    assert result.text == "[ФИО-1]; [ФИО-1]; [ФИО-2]; [ФИО-2]"


def test_inflected_docx_exact_restoration(tmp_path):
    from docx import Document
    from local_anonymizer.documents import DocxProcessor
    source, output, restored = [tmp_path / n for n in ("cases.docx", "out.docx", "restored.docx")]
    doc = Document()
    p = doc.add_paragraph()
    p.add_run("Давид Леснов; ")
    p.add_run("Леснову Дави").bold = True
    p.add_run("ду")
    doc.add_table(rows=1, cols=1).cell(0, 0).text = "Марии Крау"
    doc.save(source)
    processor = DocxProcessor()
    result = processor.anonymize(source, output)
    out = Document(output)
    assert out.paragraphs[0].text == "[ФИО-1]; [ФИО-1]"
    assert out.tables[0].cell(0, 0).text == "[ФИО-2]"
    processor.restore(output, restored, result.mapping, package_archive=result.package_archive)
    assert restored.read_bytes() == source.read_bytes()


def test_distribution_configuration_includes_dictionary():
    import ast
    import tomllib
    from pathlib import Path
    root = Path(__file__).resolve().parents[1]
    tree = ast.parse((root / "LocalAnonymizer.spec").read_text(encoding="utf-8"))
    datas = next(ast.literal_eval(n.value) for n in tree.body if isinstance(n, ast.Assign)
                 and any(isinstance(t, ast.Name) and t.id == "datas" for t in n.targets))
    assert ("src/local_anonymizer/assets", "local_anonymizer/assets") in datas
    config = tomllib.loads((root / "pyproject.toml").read_text(encoding="utf-8"))
    assert "src/local_anonymizer" in config["tool"]["hatch"]["build"]["targets"]["wheel"]["packages"]
