"""Portable-pair regressions: exercise the public service, not only processors."""
import base64
import io
import json
from zipfile import ZipFile

import pytest
from docx import Document

from local_anonymizer.service import AnonymizerService


@pytest.mark.parametrize("edited", [False, True])
def test_portable_docx_pair_never_restores_irreversible_secret(tmp_path, edited):
    secret = "SyntheticAuditCanary73!"
    email = "audit.person@example.test"
    source = tmp_path / "source.docx"
    document = Document()
    paragraph = document.add_paragraph()
    paragraph.add_run("password: ").bold = True
    paragraph.add_run(secret)
    document.add_paragraph(email)
    document.save(source)
    service = AnonymizerService(tmp_path / "private-library")
    try:
        doc_id = service.import_file(source)
        report = service.anonymize(doc_id)
        assert report["ready"]
        pair = service.export_pair(doc_id, tmp_path / "safe.docx")
        service.storage.delete_document(doc_id)
        if edited:
            document = Document(pair["anonymized_path"])
            document.add_paragraph("Ordinary added text")
            document.save(pair["anonymized_path"])
        restored = tmp_path / "restored.docx"
        service.restore_from_pair(pair["anonymized_path"], restored)
        output = Document(restored)
        text = "\n".join(p.text for p in output.paragraphs)
        assert secret not in text
        assert email in text
        assert output.paragraphs[0].runs[0].bold
        if edited:
            assert "Ordinary added text" in text
        payload = json.loads(pair["recovery_path"].read_text(encoding="utf-8"))
        assert secret not in json.dumps(payload)
        archive = payload["bundle"].get("package_archive") or {}
        if "exact_source_b64" in archive:
            with ZipFile(io.BytesIO(base64.b64decode(archive["exact_source_b64"]))) as package:
                assert all(secret.encode() not in package.read(name) for name in package.namelist())
    finally:
        service.storage.connection.close()
