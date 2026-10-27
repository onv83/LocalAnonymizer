"""Synthetic Word settings regressions; no real-user XML fixtures."""
from zipfile import ZipFile, ZIP_DEFLATED

import pytest
from docx import Document
from lxml import etree

from local_anonymizer.service import AnonymizerService
from local_anonymizer.documents import _sanitize_package

from local_anonymizer.documents import _scrub_unknown_docx_xml

W = "http://schemas.openxmlformats.org/wordprocessingml/2006/main"


def test_word_saved_zoom_is_discarded_without_blocking_restore():
    # Word persists the user's view zoom, not document formatting.
    root = etree.fromstring(
        f'<w:settings xmlns:w="{W}"><w:zoom w:percent="160"/></w:settings>'
    )
    clean, blocked = _scrub_unknown_docx_xml("word/settings.xml", root)
    assert not blocked
    assert len(clean) == 0


@pytest.mark.parametrize("percent", ["10", "100", "160", "500"])
def test_zoom_safe_range_is_removed(percent):
    root = etree.fromstring(
        f'<w:settings xmlns:w="{W}"><w:zoom w:percent="{percent}"/></w:settings>'
    )
    clean, blocked = _scrub_unknown_docx_xml("word/settings.xml", root)
    assert not blocked
    assert len(clean) == 0


@pytest.mark.parametrize("fragment", [
    '<w:zoom w:percent="9"/>',
    '<w:zoom w:percent="501"/>',
    '<w:zoom w:percent="0160"/>',
    '<w:zoom w:percent="+160"/>',
    '<w:zoom w:percent="160%"/>',
    '<w:zoom w:percent="CANARY"/>',
    '<w:zoom w:percent="987654321"/>',
    '<w:zoom w:percent="160" w:val="CANARY"/>',
    '<w:zoom w:percent="160" x:val="CANARY"/>',
    '<w:zoom w:percent="160" percent="CANARY"/>',
    '<w:zoom x:percent="160"/>',
    '<w:zoom w:percent="160">CANARY</w:zoom>',
    '<w:zoom w:percent="160"/>CANARY',
    '<w:zoom w:percent="100"><w:rsids/></w:zoom>',
    '<w:zoom w:percent="160"><x:hidden>CANARY</x:hidden></w:zoom>',
    '<w:compat><w:zoom w:percent="160"/></w:compat>',
    '<x:zoom w:percent="CANARY"/>',
    '<w:compat><w:compatSetting w:name="compatibilityMode" w:uri="CANARY" w:val="15"/></w:compat>',
])
def test_hostile_settings_are_still_blocked(fragment):
    root = etree.fromstring(
        f'<w:settings xmlns:w="{W}" xmlns:x="urn:hostile">{fragment}</w:settings>'
    )
    clean, blocked = _scrub_unknown_docx_xml("word/settings.xml", root)
    assert blocked
    assert b"CANARY" not in etree.tostring(clean)


@pytest.mark.parametrize("part,root_tag,container", [
    ("word/document.xml", "document", "body"),
    ("word/styles.xml", "styles", "style"),
    ("word/webSettings.xml", "webSettings", "divs"),
])
def test_zoom_permission_does_not_cross_parts(part, root_tag, container):
    root = etree.fromstring(
        f'<w:{root_tag} xmlns:w="{W}"><w:{container}>'
        f'<w:zoom w:percent="160"/></w:{container}></w:{root_tag}>'
    )
    clean, blocked = _scrub_unknown_docx_xml(part, root)
    assert blocked
    assert b"160" not in etree.tostring(clean)


def _replace_settings(path, settings):
    with ZipFile(path) as archive:
        parts = {info.filename: archive.read(info.filename) for info in archive.infolist()}
    parts["word/settings.xml"] = settings
    with ZipFile(path, "w", ZIP_DEFLATED) as archive:
        for name, data in parts.items():
            archive.writestr(name, data)


def test_public_pair_restore_after_synthetic_word_zoom_edit(tmp_path):
    service = AnonymizerService(tmp_path / "app")
    source = tmp_path / "source.docx"
    doc = Document()
    doc.add_paragraph("Иванов Сергей Петрович")
    doc.add_paragraph("password: Secret123!")
    table = doc.add_table(rows=2, cols=2)
    table.cell(0, 0).merge(table.cell(1, 0)).text = "Merged"
    table.cell(0, 1).paragraphs[0].add_run("Bold").bold = True
    doc.save(source)
    doc_id = service.import_file(source)
    service.anonymize(doc_id)
    pair = service.export_pair(doc_id, tmp_path / "edited.docx")
    edited = pair["anonymized_path"]
    doc = Document(edited)
    doc.add_paragraph("Word edit remains")
    doc.save(edited)
    with ZipFile(edited) as archive:
        settings = etree.fromstring(archive.read("word/settings.xml"))
    for zoom in list(settings.findall(f"{{{W}}}zoom")):
        settings.remove(zoom)
    etree.SubElement(settings, f"{{{W}}}zoom").set(f"{{{W}}}percent", "160")
    _replace_settings(edited, etree.tostring(settings))
    service.storage.delete_document(doc_id)
    restored = tmp_path / "restored.docx"
    report = service.restore_from_pair(edited, restored)
    assert report["unknown_markers"] == []
    with ZipFile(restored) as archive:
        body = etree.fromstring(archive.read("word/document.xml"))
        payload = b"".join(archive.read(name) for name in archive.namelist())
    texts = "".join(body.itertext())
    assert "Иванов Сергей Петрович" in texts
    assert "Word edit remains" in texts
    assert b"Secret123!" not in payload
    assert len(body.findall(f".//{{{W}}}tbl")) == 1
    assert len(body.findall(f".//{{{W}}}vMerge")) == 2
    assert body.find(f".//{{{W}}}b") is not None
    assert _sanitize_package(restored) == []
