"""Synthetic-only passport regressions through the production engine."""
import pytest
from local_anonymizer import engine


@pytest.fixture
def passport_engine():
    return engine.Anonymizer


@pytest.mark.parametrize("value", ["00 00 000000", "0000 000000", "0000000000", "00\u00a000\u00a0000000", "00  00   000000", "0000 № 000000"])
def test_combined_synthetic_passport(passport_engine, value):
    result = passport_engine().anonymize("Паспорт: " + value)
    assert result.mapping.get("[ПАСПОРТ-1]") == value
    assert value not in result.text


@pytest.mark.parametrize("source", [
    "Паспорт: серия 0000; номер № 000000",
    "Паспорт\nСерия: 0000\nНомер: 000000",
    "Серия паспорта: 0000; номер паспорта: 000000",
    "Паспорт\ufdd0Серия\ufdd00000\ufdd0\ufdd2\ufdd0Номер\ufdd0000000",
])
def test_separated_labeled_fields(passport_engine, source):
    result = passport_engine().anonymize(source)
    values = {v for k, v in result.mapping.items() if k.startswith("[ПАСПОРТ-")}
    assert values == {"0000", "000000"}
    assert "0000" not in result.text


def test_docx_table_split_runs_exact_restore(passport_engine, tmp_path):
    from docx import Document
    from local_anonymizer.documents import DocxProcessor
    source, masked, restored = (tmp_path / name for name in ("source.docx", "masked.docx", "restored.docx"))
    doc = Document()
    table = doc.add_table(rows=6, cols=2)
    for row, label, chunks in zip(table.rows,
        ("Паспорт:", "Серия паспорта:", "Номер паспорта:", "Кем выдан:", "Дата выдачи:", "Код подразделения:"),
        (("00", "\u00a000", " № ", "000000"), ("00", "00"), ("000", "000"),
         ("Тестовый ", "орган"), ("12.03.", "2020"), ("000-", "000"))):
        row.cells[0].text = label
        for i, chunk in enumerate(chunks):
            run = row.cells[1].paragraphs[0].add_run(chunk)
            run.bold = i % 2 == 0
            run.italic = i % 2 == 1
    doc.save(source)
    processor = DocxProcessor(passport_engine())
    result = processor.anonymize(source, masked)
    visible = "\n".join(c.text for r in Document(masked).tables[0].rows for c in r.cells)
    assert "000" not in visible
    assert "Тестовый орган" not in visible
    assert "12.03.2020" not in visible
    assert len([k for k in result.mapping if k.startswith("[ПАСПОРТ-")]) == 6
    assert processor.restore(masked, restored, result.mapping, package_archive=result.package_archive) == []
    assert restored.read_bytes() == source.read_bytes()
    disabled_path = tmp_path / "disabled.docx"
    disabled = DocxProcessor(passport_engine(disabled_categories={"ПАСПОРТ"})).anonymize(source, disabled_path)
    assert not disabled.mapping
    before = [[cell.text for cell in row.cells] for row in Document(source).tables[0].rows]
    after = [[cell.text for cell in row.cells] for row in Document(disabled_path).tables[0].rows]
    assert after == before


@pytest.mark.parametrize("source", [
    "Паспорт: 00\u00a000 № 000000",
    "Паспорт: серия 0000; номер № 000000",
    "Серия паспорта:\ufdd00000; номер паспорта:\ufdd0000000",
])
def test_disabled_passport_is_not_relabelled(passport_engine, source):
    result = passport_engine(disabled_categories={"ПАСПОРТ"}).anonymize(source)
    assert result.text == source
    assert not result.mapping


@pytest.mark.parametrize("source", [
    "Количество: 0000; номер заказа: 000000; дата: 12.03.2020",
    "Серия оборудования: 0000; номер изделия: 000000",
    "Паспорт изделия: 0000; номер заказа: 000000",
    "Паспорт:\ufdd0\ufdd1\ufdd00000 000000",
    "паспорт\ufdd0 00 00 000000",
    "Паспорт: 0000 0000000",
    "Паспорт: 0000 000000я",
    "Паспорт: серия 00000; номер № 0000000",
    "0000 000000; 12.03.2020; 000-000",
])
def test_unrelated_or_malformed_numbers_are_not_passports(passport_engine, source):
    result = passport_engine().anonymize(source)
    assert not any(k.startswith("[ПАСПОРТ-") for k in result.mapping)


@pytest.mark.parametrize("separator", ["; ", "\n", "\ufdd0\ufdd2\ufdd0"])
def test_issuer_date_code_in_passport_context(passport_engine, separator):
    source = separator.join(("Паспорт: 0000 000000", "Кем выдан: Тестовый орган", "Дата выдачи: 12.03.2020", "Код подразделения: 000-000"))
    result = passport_engine().anonymize(source)
    assert {v for k, v in result.mapping.items() if k.startswith("[ПАСПОРТ-")} == {"0000 000000", "Тестовый орган", "12.03.2020", "000-000"}
    restored = result.text
    for marker, original in result.mapping.items():
        restored = restored.replace(marker, original)
    assert restored == source
    disabled = passport_engine(disabled_categories={"ПАСПОРТ"}).anonymize(source)
    assert disabled.text == source
    assert not disabled.mapping


@pytest.mark.parametrize("prefix", ["", "Паспорт изделия: 0000; ", "Паспорт: 0000 000000\n\n", "Паспорт: 0000 000000\ufdd0\ufdd1\ufdd0", "Паспорт: 0000 000000; Договор: "])
def test_issue_metadata_needs_local_passport_context(passport_engine, prefix):
    source = prefix + "Выдан: Тестовый орган; Дата выдачи: 12.03.2020; Код подразделения: 000-000"
    result = passport_engine().anonymize(source)
    passport_values = {v for k, v in result.mapping.items() if k.startswith("[ПАСПОРТ-")}
    assert not passport_values.intersection({"Тестовый орган", "12.03.2020", "000-000"})


@pytest.mark.parametrize("source", ["Паспорт РФ: 0000 000000", "паспорт серия 0000 000000"])
def test_existing_passport_qualifiers(passport_engine, source):
    result = passport_engine().anonymize(source)
    assert "0000" not in result.text
    assert "0000 000000" in result.mapping.values()


def test_exact_spelling_variants_do_not_share_marker(passport_engine):
    source = "Паспорт: 0000 000000; Паспорт: 00\u00a000 000000"
    result = passport_engine().anonymize(source)
    assert len([k for k in result.mapping if k.startswith("[ПАСПОРТ-")]) == 2
    restored = result.text
    for marker, original in result.mapping.items():
        restored = restored.replace(marker, original)
    assert restored == source
