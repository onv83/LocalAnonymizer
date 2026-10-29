"""Synthetic regressions for B68 financial context and technical quantities."""

import pytest


@pytest.mark.parametrize("separator", ["/", " / "])
def test_currency_amount_after_slash_is_fully_hidden(separator):
    source = f"Затраты 1,234 млн руб.{separator}3,456 млн руб."
    result = Anonymizer().anonymize(source)
    assert "3,456" not in result.text
    assert {row["orig"] for row in result.legend if row["cat"] == "СУММА"} == {
        "1,234 млн руб.", "3,456 млн руб."
    }


from local_anonymizer.engine import Anonymizer


def test_slash_amount_docx_split_runs_table_and_exact_restore(tmp_path):
    from docx import Document
    from local_anonymizer.documents import DocxProcessor

    source, output, restored = [tmp_path / name for name in ("source.docx", "output.docx", "restored.docx")]
    text = "1,234 млн руб./3,456 млн руб."
    doc = Document()
    paragraph = doc.add_paragraph()
    for part in ("1,234 млн руб.", "/", "3,456", " млн руб."):
        paragraph.add_run(part).bold = True
    doc.add_table(rows=1, cols=1).cell(0, 0).text = text
    doc.save(source)
    original = source.read_bytes()
    processor = DocxProcessor()
    result = processor.anonymize(source, output)
    out = Document(output)
    assert out.paragraphs[0].text == "[СУММА-1]/[СУММА-2]"
    assert out.tables[0].cell(0, 0).text == "[СУММА-1]/[СУММА-2]"
    processor.restore(output, restored, result.mapping, package_archive=result.package_archive)
    assert restored.read_bytes() == original
    assert source.read_bytes() == original


def test_slash_currency_disablement():
    text = "1,234 млн руб./3,456 млн руб."
    assert Anonymizer(disabled_categories={"СУММА"}).anonymize(text).text == text



@pytest.mark.parametrize("amount", ["2,7 миллиона рублей", "3 миллиона руб.", "623 тысячи рублей"])
def test_spelled_scale_currency_is_one_complete_money_span(amount):
    source = f"Стоимость: {amount}"
    result = Anonymizer().anonymize(source)
    money = [item["orig"] for item in result.legend if item["cat"] == "СУММА"]
    assert money == [amount]
    assert result.text == "Стоимость: [СУММА-1]"


@pytest.mark.parametrize("amount", ["623 тыс.", "954 тыс."])
def test_expenses_without_currency_have_money_context(amount):
    source = f"Расходы составляют {amount}"
    result = Anonymizer().anonymize(source)
    assert [item["orig"] for item in result.legend if item["cat"] == "СУММА"] == [amount]
    assert result.text == "Расходы составляют [СУММА-1]"


@pytest.mark.parametrize("source", [
    "В сумме 512 ГБ", "Суммарный объем 64 ТБ", "В сумме 512 GB",
    "Суммарный объем 64 TB", "Итого 12 серверов", "Расход: 623 тыс. запросов",
    "Оплата до 12.11.2026", "Оплата до 2026-11-12", "Итого 2.7.14",
])
def test_nonfinancial_quantities_after_commercial_words_are_not_money(source):
    result = Anonymizer(disabled_categories={"СОКРАЩЕНИЯ"}).anonymize(source)
    assert "СУММА" not in result.counts
    assert result.text == source


@pytest.mark.parametrize("source", [
    "Расходы составляют 623 тыс.", "Затраты: 954 тыс.",
    "Стоимость: 2,7 миллиона рублей", "Стоимость: 623 тыс. руб.",
])
def test_money_category_disablement_preserves_complete_amount(source):
    result = Anonymizer(disabled_categories={"СУММА"}).anonymize(source)
    assert result.text == source
    assert "СУММА" not in result.counts
    assert not result.mapping


@pytest.mark.parametrize("source", [
    "623 тыс. запросов", "954 тыс. записей", "512 ГБ", "64 ТБ",
    "12.11.2026", "Версия 2.7.14", "Оплата в течение 7 рабочих дней",
])
def test_unlabelled_counts_capacities_and_deadlines_remain_nonfinancial(source):
    result = Anonymizer(disabled_categories={"СОКРАЩЕНИЯ"}).anonymize(source)
    assert result.text == source
    assert "СУММА" not in result.counts
