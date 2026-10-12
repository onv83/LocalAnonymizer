import base64
import csv
import hashlib
from pathlib import Path
from zipfile import ZipFile

import pytest
from docx import Document
from docx.enum.text import WD_ALIGN_PARAGRAPH
from docx.shared import Inches

from local_anonymizer.engine import Anonymizer
from local_anonymizer.service import AnonymizerService, export_mapping_csv


def test_end_to_end_acceptance_docx_anonymize_edit_restore(tmp_path):
    sample = Path("acceptance_sample.txt").read_text(encoding="utf-8")
    source = tmp_path / "user-source.docx"
    doc = Document(); doc.add_paragraph(sample); doc.save(source)
    before = source.read_bytes()
    service = AnonymizerService(tmp_path / "app")
    doc_id = service.import_file(source)
    report = service.anonymize(doc_id)
    assert report["ready"] is True
    anon = Path(report["output_path"])
    anonymized = "\n".join(p.text for p in Document(anon).paragraphs)
    for expected in ("[ФИО-1]", "[ОРГАНИЗАЦИЯ-1]", "[ПОЧТА-1]", "[IP-1]", "[уязвимость]", "(удалено)"):
        assert expected in anonymized
    for public in ("1.7.5.16", "SLA 99,9%", "НДС 20%", "порт 8080"):
        assert public in anonymized
    edited = tmp_path / "edited.docx"
    edit_doc = Document(anon); edit_doc.add_paragraph("Редактура [ФИО-1] и [ФИО-999]"); edit_doc.save(edited)
    restored = service.restore(doc_id, edited)
    restored_text = "\n".join(p.text for p in Document(restored["output_path"]).paragraphs)
    assert "Иванов Сергей Петрович" in restored_text
    assert "[ФИО-999]" in restored_text
    assert restored["unknown_markers"] == ["[ФИО-999]"]
    assert source.read_bytes() == before


def test_recreate_docx_applies_remembered_rule_to_every_repeat_and_removes_source_xml(tmp_path):
    source = tmp_path / "remembered-rule.docx"
    document = Document()
    document.add_paragraph("Яндекс 360 используется. Яндекс поддерживается.")
    document.add_paragraph("Центральный аппарат (АльфаПроект-ИТ, АльфаПроект-ОЦО).")
    paragraph = document.add_paragraph()
    paragraph.add_run("Ян")
    paragraph.add_run("декс используется повторно.")
    document.save(source)

    service = AnonymizerService(tmp_path / "app")
    document_id = service.import_file(source)
    first = service.anonymize(document_id, Anonymizer())
    assert "Яндекс" in "\n".join(p.text for p in Document(first["output_path"]).paragraphs)

    recreated = service.anonymize(
        document_id,
        Anonymizer(user_dictionary={"Яндекс": "СКРЫТО", "АльфаПроект": "СКРЫТО"}),
    )
    recreated_path = Path(recreated["output_path"])
    recreated_text = "\n".join(p.text for p in Document(recreated_path).paragraphs)

    assert "Яндекс" not in recreated_text
    assert "АльфаПроект" not in recreated_text
    assert recreated_text.count("[СКРЫТО-1]") == 3
    assert recreated_text.count("[СКРЫТО-2]") == 2
    with ZipFile(recreated_path) as package:
        xml_payload = b"".join(
            package.read(name) for name in package.namelist() if name.endswith(".xml")
        )
    assert "Яндекс".encode("utf-8") not in xml_payload
    assert "Яндекс".encode("utf-16-le") not in xml_payload
    assert "АльфаПроект".encode("utf-8") not in xml_payload
    assert "АльфаПроект".encode("utf-16-le") not in xml_payload


def test_service_persists_protected_graphics_archive_and_restores_image(tmp_path):
    image_bytes = base64.b64decode(
        "iVBORw0KGgoAAAANSUhEUgAAAAEAAAABCAQAAAC1HAwCAAAAC0lEQVR42mNk+A8AAQUBAScY42YAAAAASUVORK5CYII="
    )
    image = tmp_path / "diagram.png"
    image.write_bytes(image_bytes)
    source = tmp_path / "with-diagram.docx"
    document = Document()
    document.add_paragraph("Иванов Сергей Петрович")
    document.add_picture(str(image), width=Inches(1))
    document.save(source)
    service = AnonymizerService(tmp_path / "app")
    document_id = service.import_file(source)

    anonymized_report = service.anonymize(document_id)
    anonymized = Path(anonymized_report["output_path"])
    database_bytes = (tmp_path / "app" / "library.db").read_bytes()
    assert image_bytes not in database_bytes
    assert b"word/media/" not in database_bytes
    mapping_csv = service.export_mapping(document_id, tmp_path / "mapping.csv")
    assert "word/media/" not in mapping_csv.read_text(encoding="utf-8-sig")
    with ZipFile(anonymized) as package:
        assert not any(name.startswith("word/media/") for name in package.namelist())

    restored_report = service.restore(document_id, anonymized)

    restored = Path(restored_report["output_path"])
    with ZipFile(restored) as package:
        assert image_bytes in (
            package.read(name)
            for name in package.namelist()
            if name.startswith("word/media/")
        )
    assert Document(restored).inline_shapes


def test_service_unedited_stamped_docx_restores_source_bytes_exactly(tmp_path):
    source = tmp_path / "exact-service-source.docx"
    document = Document()
    document.add_paragraph("Иванов Сергей Петрович")
    document.save(source)
    source_bytes = source.read_bytes()

    service = AnonymizerService(tmp_path / "app")
    document_id = service.import_file(source)
    anonymized_report = service.anonymize(document_id)
    anonymized = Path(anonymized_report["output_path"])

    _, archive = service.storage.load_restore_bundle(
        document_id, anonymized_report["operation_id"]
    )
    assert archive["exact_anonymized_sha256"] == hashlib.sha256(
        anonymized.read_bytes()
    ).hexdigest()

    restored_report = service.restore(document_id, anonymized)
    assert Path(restored_report["output_path"]).read_bytes() == source_bytes


def test_service_rejects_reanonymizing_anonymized_artifact(tmp_path):
    source = tmp_path / "source.docx"
    document = Document()
    document.add_paragraph("Иванов Сергей Петрович")
    document.save(source)

    service = AnonymizerService(tmp_path / "app")
    document_id = service.import_file(source)
    report = service.anonymize(document_id)
    anonymized = Path(report["output_path"])

    with pytest.raises(ValueError, match="уже обезличен.*восстанов"):
        service.import_file(anonymized)

    # Protect artifacts imported by an older version and still present after upgrade.
    legacy_document_id = service.storage.import_document(anonymized)
    with pytest.raises(ValueError, match="уже обезличен.*восстанов"):
        service.anonymize(legacy_document_id)


def test_docx_word_edit_restore_to_destination_preserves_new_formatting_and_unknowns(tmp_path):
    source = tmp_path / "formatted-source.docx"
    document = Document()
    heading = document.add_heading("Отчёт Иванов Сергей Петрович", level=1)
    heading.alignment = WD_ALIGN_PARAGRAPH.CENTER
    paragraph = document.add_paragraph()
    paragraph.add_run("Сервер 10.20.30.40 ").bold = True
    paragraph.add_run("password: Secret123!").italic = True
    original_table = document.add_table(rows=1, cols=1)
    original_table.style = "Table Grid"
    original_table.cell(0, 0).text = "Владелец Иванов Сергей Петрович"
    document.save(source)

    service = AnonymizerService(tmp_path / "app")
    document_id = service.import_file(source)
    anonymized_report = service.anonymize(document_id)
    anonymized = Path(anonymized_report["output_path"])
    edited = tmp_path / "edited-in-word.docx"
    edited.write_bytes(anonymized.read_bytes())
    word_document = Document(edited)
    edit_heading = word_document.add_heading(level=2)
    edit_run = edit_heading.add_run("Добавленный раздел")
    edit_run.bold = True
    edit_run.italic = True
    edit_table = word_document.add_table(rows=1, cols=2)
    edit_table.style = "Table Grid"
    edit_table.cell(0, 0).text = "Новая правка"
    edit_table.cell(0, 1).text = "[IP-999]"
    word_document.save(edited)  # ordinary Word-compatible save keeps the stamp

    destination = tmp_path / "exports" / "restored.docx"
    report = service.restore(document_id, edited, destination)
    restored = Document(destination)
    restored_text = "\n".join(
        [p.text for p in restored.paragraphs]
        + [cell.text for table in restored.tables for row in table.rows for cell in row.cells]
    )

    assert report["output_path"] == str(destination)
    assert report["unknown_markers"] == ["[IP-999]"]
    assert "Иванов Сергей Петрович" in restored_text
    assert "10.20.30.40" in restored_text
    assert "Добавленный раздел" in restored_text and "Новая правка" in restored_text
    assert "[IP-999]" in restored_text
    assert "Secret123!" not in restored_text and "(удалено)" in restored_text
    added = next(p for p in restored.paragraphs if p.text == "Добавленный раздел")
    assert added.style.style_id == "Heading2"
    assert added.runs[0].bold is True and added.runs[0].italic is True
    assert restored.tables[-1].style.name == "Table Grid"


def test_csv_export_neutralizes_formula_injection(tmp_path):
    target = tmp_path / "map.csv"
    export_mapping_csv({"[X-1]": "=CMD()", "[X-2]": "+1", "[X-3]": "@evil", "[X-4]": "-2"}, target)
    rows = list(csv.reader(target.open(encoding="utf-8-sig", newline="")))
    assert all(row[1].startswith("'") for row in rows[1:])


@pytest.mark.parametrize(
    "classification",
    ["Совершенно секретно", "Для служебного пользования", "ДСП", "Конфиденциально"],
)
def test_classification_label_is_informational_and_does_not_block_export(
    tmp_path, classification
):
    source = tmp_path / "classified.txt"
    source.write_text(f"{classification}. login admin_db", encoding="utf-8")
    service = AnonymizerService(tmp_path / "app")
    doc_id = service.import_file(source)

    report = service.anonymize(doc_id)

    assert report["status"] == "Готов"
    assert report["ready"] is True
    assert any(
        isinstance(warning, dict)
        and warning.get("level") == "info"
        and "не блокирует экспорт" in warning.get("message", "")
        for warning in report["warnings"]
    )
    exported = service.export_pair(doc_id, tmp_path / "classified-anonymized.txt")
    assert exported["anonymized_path"].is_file()
    assert exported["recovery_path"].is_file()

    archive = service.export_all_ready(tmp_path / "ready.zip")
    from zipfile import ZipFile
    with ZipFile(archive) as zipped:
        names = zipped.namelist()
    assert any(name.endswith("обезличенный.txt") for name in names)
    assert any(name.endswith("информация для восстановления.larecovery") for name in names)


def test_docx_classification_labels_do_not_block_anonymized_pair_export(tmp_path):
    source = tmp_path / "classified.docx"
    document = Document()
    document.add_paragraph(
        "Совершенно секретно. ДСП. Сервер 10.20.30.40, login admin_db."
    )
    document.save(source)
    service = AnonymizerService(tmp_path / "app")
    document_id = service.import_file(source)

    report = service.anonymize(document_id)

    assert report["status"] == "Готов"
    assert report["ready"] is True
    assert {warning.get("kind") for warning in report["warnings"]} >= {"ГОСТАЙНА", "ДСП"}
    assert all(
        warning.get("level") == "info"
        for warning in report["warnings"]
        if warning.get("kind") in {"ГОСТАЙНА", "ДСП"}
    )
    exported = service.export_pair(document_id, tmp_path / "classified-anonymized.docx")
    assert exported["anonymized_path"].is_file()
    assert exported["recovery_path"].is_file()


def test_unsupported_docx_surface_blocks_export_even_with_confirmation(tmp_path):
    source = tmp_path / "unsupported-chart.docx"
    document = Document()
    document.add_paragraph("Обычный текст")
    document.save(source)
    with ZipFile(source, "a") as package:
        package.writestr(
            "word/charts/chart1.xml",
            '<c:chartSpace xmlns:c="http://schemas.openxmlformats.org/drawingml/2006/chart"/>',
        )

    service = AnonymizerService(tmp_path / "app")
    document_id = service.import_file(source)
    report = service.anonymize(document_id)

    assert report["status"] == "Экспорт запрещён"
    assert report["ready"] is False
    assert any(
        isinstance(warning, dict)
        and warning.get("level") == "stop"
        and "word/charts/chart1.xml" in warning.get("message", "")
        for warning in report["warnings"]
    )
    with pytest.raises(PermissionError, match="Экспорт запрещён"):
        service.export_anonymized(
            document_id, tmp_path / "must-not-exist.docx", confirmation=True
        )
    assert not (tmp_path / "must-not-exist.docx").exists()


def test_failed_operation_commit_keeps_previous_output_mapping_and_status(tmp_path, monkeypatch):
    source = tmp_path / "source.txt"; source.write_text("Иванов Сергей Петрович", encoding="utf-8")
    service = AnonymizerService(tmp_path / "app"); doc_id = service.import_file(source)
    first = service.anonymize(doc_id); before = service.storage.get_document(doc_id)
    first_mapping = service.storage.load_mapping(doc_id, first["operation_id"])
    monkeypatch.setattr(service.storage, "commit_operation", lambda *args, **kwargs: (_ for _ in ()).throw(OSError("db fault")))
    source_owned = Path(before["owned_source"]); source_owned.write_text("Петров Иван Сергеевич", encoding="utf-8")
    import pytest
    with pytest.raises(OSError, match="db fault"):
        service.anonymize(doc_id)
    after = service.storage.get_document(doc_id)
    assert after["operation_id"] == before["operation_id"] and after["output_path"] == before["output_path"]
    assert Path(after["output_path"]).is_file()
    assert service.storage.load_mapping(doc_id, after["operation_id"]) == first_mapping
