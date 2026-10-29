"""Finite geography coverage in the application, not just dictionary lookup."""
import pytest
from local_anonymizer.engine import Anonymizer


@pytest.mark.parametrize("cue", ["г. ", "города ", "из "])
def test_geography_cue_retains_city_with_surname_anchors(cue):
    source = "Иванов Сергей Петрович; Иванов Алексей Николаевич; позвать Иванова; " + cue + "Иванова"
    result = Anonymizer().anonymize(source)
    assert "позвать Иванова" in result.text
    assert result.counts.get("ГОРОД") == 1
    assert result.mapping[next(k for k in result.mapping if k.startswith("[ГОРОД-"))] == "Иванова"



@pytest.mark.parametrize("surface", [
    "Нижний Новгород", "Нижнем Новгороде", "Ростова-на-Дону",
    "Санкт-Петербурге", "Томской области", "Красноярским краем",
    "Республике Татарстан", "Ханты-Мансийском автономном округе — Югре",
    "Берлин", "Берлина", "Берлину", "Берлином", "Берлине",
])
def test_exact_location_surface(surface):
    result = Anonymizer().anonymize(f"Пункт: {surface}; далее.")
    assert result.text == "Пункт: [ГОРОД-1]; далее."
    assert result.mapping == {"[ГОРОД-1]": surface}


def test_inflections_keep_exact_reversible_values():
    source = "Москва; Москве; МОСКВА; Москва"
    result = Anonymizer().anonymize(source)
    assert result.text == "[ГОРОД-1]; [ГОРОД-2]; [ГОРОД-3]; [ГОРОД-1]"
    restored = result.text
    for marker, original in result.mapping.items():
        restored = restored.replace(marker, original)
    assert restored == source


@pytest.mark.parametrize("source", ["Томский модуль", "Москваход", "мир и труд", "орёл летит", "видное место", "советский район"])
def test_no_wildcard_geography_or_geography_hints(source):
    result = Anonymizer().anonymize(source)
    assert "ГОРОД" not in result.counts
    assert not any(p.get("kind") == "город" for p in result.parts)
    assert result.text == source


@pytest.mark.parametrize("paranoid", [False, True])
def test_disabled_geography_blocks_fallbacks(paranoid):
    source = "Пункт: Нижнем Новгороде; МОСКВА; Томской области"
    result = Anonymizer(disabled_categories={"ГОРОД"}, paranoid=paranoid).anonymize(source)
    assert result.text == source


def test_canonical_whitelist_preserves_declined_compound():
    source = "Пункт: Нижнем Новгороде; Нижний Новгород"
    assert Anonymizer(whitelist={"Нижний Новгород"}, paranoid=True).anonymize(source).text == source


def test_address_priority():
    result = Anonymizer().anonymize("Адрес: г. Москва, ул. Лесная, д. 8")
    assert result.counts.get("АДРЕС") == 1
    assert "ГОРОД" not in result.counts


def test_docx_runs_table_exact_restoration(tmp_path):
    from docx import Document
    from local_anonymizer.documents import DocxProcessor
    source, output, restored = [tmp_path / n for n in ("geo.docx", "out.docx", "restored.docx")]
    doc = Document()
    p = doc.add_paragraph()
    p.add_run("Нижнем Нов").bold = True
    p.add_run("городе; Москва; Москве")
    table = doc.add_table(rows=2, cols=2)
    p = table.cell(0, 0).paragraphs[0]
    p.add_run("Томской ").italic = True
    p.add_run("области")
    table.cell(0, 1).text = "Ростова-на-Дону"
    table.cell(1, 0).text = "Нижний"
    table.cell(1, 1).text = "Новгород"
    doc.save(source)
    original = source.read_bytes()
    processor = DocxProcessor()
    result = processor.anonymize(source, output)
    out = Document(output)
    assert out.paragraphs[0].text == "[ГОРОД-1]; [ГОРОД-2]; [ГОРОД-3]"
    assert out.paragraphs[0].runs[0].bold
    assert out.tables[0].cell(0, 0).text == "[ГОРОД-4]"
    assert out.tables[0].cell(0, 0).paragraphs[0].runs[0].italic
    assert out.tables[0].cell(0, 1).text == "[ГОРОД-5]"
    assert out.tables[0].cell(1, 0).text == "Нижний"
    assert out.tables[0].cell(1, 1).text == "Новгород"
    processor.restore(output, restored, result.mapping, package_archive=result.package_archive)
    assert restored.read_bytes() == original == source.read_bytes()


def test_distribution_configuration_includes_geography():
    import ast
    import tomllib
    from pathlib import Path
    root = Path(__file__).resolve().parents[1]
    tree = ast.parse((root / "LocalAnonymizer.spec").read_text(encoding="utf-8"))
    datas = next(ast.literal_eval(n.value) for n in tree.body if isinstance(n, ast.Assign)
                 and any(isinstance(t, ast.Name) and t.id == "datas" for t in n.targets))
    assert ("src/local_anonymizer/assets", "local_anonymizer/assets") in datas
    config = tomllib.loads((root / "pyproject.toml").read_text(encoding="utf-8"))
    assert "src/local_anonymizer" in config["tool"]["hatch"]["build"]["targets"]["wheel"]["packages"]
    assets = root / "src/local_anonymizer/assets/russian_geography"
    assert (assets / "locations.json").is_file()
    assert (assets / "NOTICE.md").is_file()
