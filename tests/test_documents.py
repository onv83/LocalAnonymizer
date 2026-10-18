import base64
import copy
from pathlib import Path
from zipfile import ZIP_DEFLATED, ZIP_STORED, ZipFile, ZipInfo

import pytest
from docx import Document
from docx.enum.table import WD_CELL_VERTICAL_ALIGNMENT
from docx.enum.text import WD_ALIGN_PARAGRAPH
from docx.oxml import OxmlElement
from docx.oxml.ns import qn
from docx.opc.constants import RELATIONSHIP_TYPE
from docx.shared import Inches, Pt, RGBColor
from lxml import etree

from local_anonymizer.documents import (
    DocxProcessor,
    RestorationError,
    _trusted_docx_strings,
    _trusted_wml_schema,
)
from local_anonymizer.engine import Anonymizer


def _all_xml(path: Path) -> bytes:
    with ZipFile(path) as z:
        return b"".join(z.read(n) for n in z.namelist() if n.endswith((".xml", ".rels")))


def _rewrite_docx_parts(path: Path, transforms: dict[str, callable]) -> None:
    replacement = path.with_suffix(".rewrite.docx")
    with ZipFile(path) as incoming, ZipFile(replacement, "w", ZIP_DEFLATED) as outgoing:
        for info in incoming.infolist():
            payload = incoming.read(info.filename)
            if info.filename in transforms:
                payload = transforms[info.filename](payload)
            outgoing.writestr(info, payload)
    replacement.replace(path)


def test_docx_anonymize_and_restore_do_not_reserialize_with_python_docx(tmp_path, monkeypatch):
    source = tmp_path / "source.docx"
    anonymized = tmp_path / "anonymized.docx"
    restored = tmp_path / "restored.docx"
    document = Document()
    document.add_paragraph("Иванов Сергей Петрович")
    document.save(source)
    _trusted_wml_schema()
    _trusted_docx_strings()

    def forbidden_save(*args, **kwargs):
        raise AssertionError("python-docx must not reserialize the OOXML package")

    monkeypatch.setattr("docx.document.Document.save", forbidden_save)
    result = DocxProcessor().anonymize(source, anonymized)
    assert "[ФИО-1]" in Document(anonymized).paragraphs[0].text
    DocxProcessor().restore(anonymized, restored, result.mapping)
    assert Document(restored).paragraphs[0].text == "Иванов Сергей Петрович"


def test_docx_new_ru_en_regressions_preserve_formatting_and_restore_exactly(tmp_path):
    source = tmp_path / "new-regressions.docx"
    anonymized = tmp_path / "new-regressions.anonymized.docx"
    restored = tmp_path / "new-regressions.restored.docx"
    document = Document()
    english = document.add_paragraph()
    english.add_run("Contact: Alex Example\n").bold = True
    english.add_run("Address: 10 Test Street, Testville, 123456\n").italic = True
    english.add_run("Software: Example Office Suite").underline = True
    russian = document.add_paragraph()
    russian.add_run("Сервер 3 штуки — г. Москва, ").bold = True
    russian.add_run("ул. Энтузиастов, д.5, к.6").italic = True
    russian.add_run(". После устранения Исполнителем недостатков Акт подписывается.")
    document.save(source)
    source_bytes = source.read_bytes()

    result = DocxProcessor().anonymize(source, anonymized)

    output = Document(anonymized)
    english_text = output.paragraphs[0].text
    russian_text = output.paragraphs[1].text
    assert "Alex Example" not in english_text
    assert "10 Test Street, Testville" not in english_text
    assert "123456" not in english_text
    assert "Example Office Suite" not in english_text
    assert "ул. Энтузиастов, д.5, к.6" not in russian_text
    assert "недостатков" in russian_text
    assert output.paragraphs[0].runs[0].bold is True
    assert output.paragraphs[0].runs[1].italic is True
    assert output.paragraphs[0].runs[-1].underline is True

    DocxProcessor().restore(
        anonymized,
        restored,
        result.mapping,
        package_archive=result.package_archive,
    )
    assert restored.read_bytes() == source_bytes


def test_docx_inventory_geography_and_split_ordinal_address_do_not_leak(tmp_path):
    source = tmp_path / "synthetic-inventory.docx"
    anonymized = tmp_path / "synthetic-inventory.anonymized.docx"
    restored = tmp_path / "synthetic-inventory.restored.docx"
    document = Document()
    table = document.add_table(rows=3, cols=4)
    cells = table.rows[0].cells
    cells[0].text = "Республика Испытаний (Тестоградск)"
    cells[1].text = "Северо-Западный"
    cells[2].text = "КЦ Тестоградск"
    address = cells[3].paragraphs[0]
    address.add_run("г. Тестоградск, ").bold = True
    address.add_run("ул. 7го примера ").italic = True
    address.add_run("68б").underline = True
    cells = table.rows[1].cells
    cells[0].text = "Макетинский автономный округ — Макетия"
    cells[1].text = "Уральский"
    cells[2].text = "КЦ Макетинск"
    address = cells[3].paragraphs[0]
    address.add_run("г. Макетинск, ").bold = True
    address.add_run("ул. 9 января ").italic = True
    address.add_run("67а").underline = True
    cells = table.rows[2].cells
    cells[0].text = "Проверочный край"
    cells[1].text = "ТФО"
    cells[2].text = "Лабораторный"
    address = cells[3].paragraphs[0]
    address.add_run("г. Лабораторный, ").bold = True
    address.add_run("ул. Контрольная ").italic = True
    address.add_run("9").underline = True
    document.save(source)
    source_bytes = source.read_bytes()

    result = DocxProcessor().anonymize(source, anonymized)

    package_xml = _all_xml(anonymized).decode("utf-8", errors="ignore")
    leaked = (
        "Республика Испытаний", "Тестоградск", "Северо-Западный",
        "7го примера", "68б", "Макетинский автономный округ",
        "Макетия", "Уральский", "Макетинск", "9 января", "67а",
        "Проверочный край", "ТФО", "Лабораторный", "Контрольная",
    )
    assert all(value not in package_xml for value in leaked)
    assert sum(marker.startswith("[АДРЕС-") for marker in result.mapping) == 3
    assert sum(marker.startswith("[ГОРОД-") for marker in result.mapping) == 9

    DocxProcessor().restore(
        anonymized,
        restored,
        result.mapping,
        package_archive=result.package_archive,
    )
    assert restored.read_bytes() == source_bytes


def test_docx_ten_column_domain_inventory_subjects_do_not_leak(tmp_path):
    source = tmp_path / "synthetic-domain-inventory.docx"
    anonymized = tmp_path / "synthetic-domain-inventory.anonymized.docx"
    restored = tmp_path / "synthetic-domain-inventory.restored.docx"
    document = Document()
    table = document.add_table(rows=3, cols=10)
    headers = (
        "Субъект Российской Федерации", "Федеральный округ", "Наименование филиала",
        "Местонахождение филиалов", "Порядковый номер", "Имя сервера",
        "Сокращение сервера", "Платформа сервера", "Перечень сетей", "Количество объектов",
    )
    for cell, value in zip(table.rows[0].cells, headers):
        cell.text = value
    rows = (
        ("Испытательный край", "18", "2000"),
        ("Респ. Примерия", "62", "2000"),
    )
    for row_index, (subject, ordinal, count) in enumerate(rows, start=1):
        cells = table.rows[row_index].cells
        paragraph = cells[0].paragraphs[0]
        paragraph.add_run(subject[: len(subject) // 2]).bold = True
        paragraph.add_run(subject[len(subject) // 2 :]).italic = True
        values = (
            "ТФО", "Макетный", "г. Тестоградск, ул. Проверочная, д. 1", ordinal,
            "srv-test-01", "srv", "Windows Server 2019", "192.0.2.0/24", count,
        )
        for cell, value in zip(cells[1:], values):
            cell.text = value
    document.save(source)
    source_bytes = source.read_bytes()

    result = DocxProcessor().anonymize(source, anonymized)

    package_xml = _all_xml(anonymized).decode("utf-8", errors="ignore")
    assert "Испытательный край" not in package_xml
    assert "Респ. Примерия" not in package_xml
    assert "Испытательный край" in result.mapping.values()
    assert "Респ. Примерия" in result.mapping.values()
    assert source.read_bytes() == source_bytes
    DocxProcessor().restore(
        anonymized,
        restored,
        result.mapping,
        package_archive=result.package_archive,
    )
    assert restored.read_bytes() == source_bytes


def test_docx_domain_inventory_redacts_all_location_columns_with_runs_and_hard_breaks(tmp_path):
    source = tmp_path / "synthetic-domain-inventory-all-columns.docx"
    anonymized = tmp_path / "synthetic-domain-inventory-all-columns.anonymized.docx"
    restored = tmp_path / "synthetic-domain-inventory-all-columns.restored.docx"
    document = Document()
    table = document.add_table(rows=2, cols=10)
    headers = (
        ("Субъект Российской", "Федерации"),
        ("Федеральный округ",),
        ("Наименование филиала",),
        ("Местонахождение", "филиалов"),
        ("Порядковый номер",),
        ("Имя сервера",),
        ("Сокращение сервера",),
        ("Платформа сервера",),
        ("Перечень сетей",),
        ("Количество объектов",),
    )
    for cell, parts in zip(table.rows[0].cells, headers):
        paragraph = cell.paragraphs[0]
        paragraph.add_run(parts[0]).bold = True
        for part in parts[1:]:
            paragraph.add_run().add_break()
            paragraph.add_run(part).italic = True

    sensitive = (
        ("Испытательная", "область"),
        ("Тестовый", "федеральный округ"),
        ("Испытательский", "региональный"),
        ("123456, г. Тестоградск,", "ул. Проверочная, д. 12"),
    )
    for cell, parts in zip(table.rows[1].cells[:4], sensitive):
        paragraph = cell.paragraphs[0]
        paragraph.add_run(parts[0]).bold = True
        paragraph.add_run().add_break()
        paragraph.add_run(parts[1]).italic = True
    for cell, value in zip(
        table.rows[1].cells[4:],
        ("18", "srv-test-01", "srv", "Sample OS", "192.0.2.0/24", "2000"),
    ):
        cell.text = value
    document.save(source)
    source_bytes = source.read_bytes()

    result = DocxProcessor().anonymize(source, anonymized)

    anonymized_text = "\n".join(
        cell.text
        for output_table in Document(anonymized).tables
        for output_row in output_table.rows
        for cell in output_row.cells
    )
    for token in (
        "Испытательная", "область", "Тестовый", "федеральный округ",
        "Испытательский", "региональный", "123456", "Тестоградск", "Проверочная",
    ):
        assert token not in anonymized_text
    assert sum(marker.startswith("[ГОРОД-") for marker in result.mapping) == 3
    assert sum(marker.startswith("[АДРЕС-") for marker in result.mapping) == 1
    assert source.read_bytes() == source_bytes

    DocxProcessor().restore(
        anonymized,
        restored,
        result.mapping,
        package_archive=result.package_archive,
    )
    assert restored.read_bytes() == source_bytes


def test_docx_device_inventory_identifier_across_runs_does_not_leak_and_restores(tmp_path):
    source = tmp_path / "synthetic-device-identifier.docx"
    anonymized = tmp_path / "synthetic-device-identifier.anonymized.docx"
    restored = tmp_path / "synthetic-device-identifier.restored.docx"
    identifier = "Q2048ZX0000042"
    document = Document()
    document.add_paragraph("Идентификатор устройства должен быть изменен.")
    document.add_paragraph("Имя принтера или устройства должно быть не длиннее 15 символов.")
    paragraph = document.add_paragraph("Пример: ")
    paragraph.add_run(identifier[:7]).bold = True
    paragraph.add_run(identifier[7:]).italic = True
    document.save(source)
    source_bytes = source.read_bytes()

    result = DocxProcessor().anonymize(source, anonymized)

    anonymized_xml = _all_xml(anonymized).decode("utf-8", errors="ignore")
    assert identifier not in anonymized_xml
    assert identifier in result.mapping.values()
    assert any(marker.startswith("[УЗЕЛ-") for marker in result.mapping)
    assert source.read_bytes() == source_bytes
    DocxProcessor().restore(
        anonymized,
        restored,
        result.mapping,
        package_archive=result.package_archive,
    )
    assert restored.read_bytes() == source_bytes


def test_docx_underscored_dns_fqdn_does_not_leak_and_restores_exactly(tmp_path):
    source = tmp_path / "synthetic-dns-inventory.docx"
    anonymized = tmp_path / "synthetic-dns-inventory.anonymized.docx"
    restored = tmp_path / "synthetic-dns-inventory.restored.docx"
    value = "_locator.synthetic.example"
    document = Document()
    table = document.add_table(rows=2, cols=3)
    headers = (
        "Fully Qualified Domain Name",
        "Integrated Forward Lookup zones",
        "Integrated Reverse Lookup zones",
    )
    for cell, header in zip(table.rows[0].cells, headers):
        cell.text = header
    paragraph = table.rows[1].cells[1].paragraphs[0]
    paragraph.add_run("_locator.").bold = True
    paragraph.add_run("synthetic.example").italic = True
    document.save(source)
    source_bytes = source.read_bytes()

    result = DocxProcessor().anonymize(source, anonymized)

    package_xml = _all_xml(anonymized).decode("utf-8", errors="ignore")
    assert value not in package_xml
    assert value in result.mapping.values()
    assert any(marker.startswith("[УЗЕЛ-") for marker in result.mapping)
    assert source.read_bytes() == source_bytes
    DocxProcessor().restore(
        anonymized,
        restored,
        result.mapping,
        package_archive=result.package_archive,
    )
    assert restored.read_bytes() == source_bytes


def test_docx_complete_postal_address_across_runs_restores_byte_exactly(tmp_path):
    source = tmp_path / "synthetic-postal-address.docx"
    anonymized = tmp_path / "synthetic-postal-address.anonymized.docx"
    restored = tmp_path / "synthetic-postal-address.restored.docx"
    address = (
        "123456, Макетия, Тестовая Республика, г. Тестоградск, "
        "Макетинский район,бульвар имени А.Б. Испытателя, д.12/3"
    )
    document = Document()
    paragraph = document.add_paragraph()
    paragraph.add_run(address[:32]).bold = True
    paragraph.add_run(address[32:79]).italic = True
    paragraph.add_run(address[79:]).underline = True
    document.save(source)
    source_bytes = source.read_bytes()

    result = DocxProcessor().anonymize(source, anonymized)

    assert Document(anonymized).paragraphs[0].text == "[АДРЕС-1]"
    assert result.mapping == {"[АДРЕС-1]": address}
    DocxProcessor().restore(
        anonymized,
        restored,
        result.mapping,
        package_archive=result.package_archive,
    )
    assert restored.read_bytes() == source_bytes


def test_docx_address_record_with_reordered_geography_and_tail_restores_byte_exactly(tmp_path):
    source = tmp_path / "synthetic-reordered-address.docx"
    anonymized = tmp_path / "synthetic-reordered-address.anonymized.docx"
    restored = tmp_path / "synthetic-reordered-address.restored.docx"
    address = (
        "ул. Испытателей, д. 12, корп. 1, лит. А г. Тестоградск, "
        "Макетинская область, 123456"
    )
    document = Document()
    table = document.add_table(rows=1, cols=1)
    paragraph = table.cell(0, 0).paragraphs[0]
    paragraph.add_run(address[:29]).bold = True
    paragraph.add_run(address[29:57]).italic = True
    paragraph.add_run(address[57:]).underline = True
    document.save(source)
    source_bytes = source.read_bytes()

    result = DocxProcessor().anonymize(source, anonymized)

    assert Document(anonymized).tables[0].cell(0, 0).text == "[АДРЕС-1]"
    assert result.mapping == {"[АДРЕС-1]": address}
    DocxProcessor().restore(
        anonymized,
        restored,
        result.mapping,
        package_archive=result.package_archive,
    )
    assert restored.read_bytes() == source_bytes


def test_docx_embedded_complete_address_preserves_surrounding_runs_and_restores_exactly(tmp_path):
    source = tmp_path / "synthetic-embedded-address.docx"
    anonymized = tmp_path / "synthetic-embedded-address.anonymized.docx"
    restored = tmp_path / "synthetic-embedded-address.restored.docx"
    prefix = "Оплата производится через 7 дней; доставка: "
    address = (
        "ул. Испытателей, д. 12, корп. 1, стр. 2, г. Тестоградск, "
        "Макетинская область, 123456"
    )
    suffix = "; после подписания."
    document = Document()
    paragraph = document.add_paragraph()
    paragraph.add_run(prefix).bold = True
    paragraph.add_run(address[:31]).italic = True
    paragraph.add_run(address[31:67]).underline = True
    paragraph.add_run(address[67:])
    paragraph.add_run(suffix).bold = True
    document.save(source)
    source_bytes = source.read_bytes()

    result = DocxProcessor().anonymize(source, anonymized)

    output = Document(anonymized).paragraphs[0].text
    assert output == f"{prefix}[АДРЕС-1]{suffix}"
    assert result.mapping == {"[АДРЕС-1]": address}
    DocxProcessor().restore(
        anonymized,
        restored,
        result.mapping,
        package_archive=result.package_archive,
    )
    assert restored.read_bytes() == source_bytes


def test_docx_contextual_address_list_is_sanitized_and_restores_byte_exactly(tmp_path):
    source = tmp_path / "synthetic-contextual-addresses.docx"
    anonymized = tmp_path / "synthetic-contextual-addresses.anonymized.docx"
    restored = tmp_path / "synthetic-contextual-addresses.restored.docx"
    addresses = (
        "Макетный пер. д.3",
        "Проверочный бульвар д.13 стр.1",
        "1-ый Лабораторный пр. д.1 стр.3",
        "Пробная аллея д.2/д",
        "Малый Контрольный пер. д.7",
        "Большая Испытательная, д.46",
    )
    text = (
        "Площадки расположены по адресам "
        + ", ".join(addresses[:4])
        + ", и "
        + addresses[4]
        + ". Площадка расположена по адресу "
        + addresses[5]
        + "."
    )
    document = Document()
    paragraph = document.add_paragraph()
    paragraph.add_run(text[:47]).bold = True
    paragraph.add_run(text[47:119]).italic = True
    paragraph.add_run(text[119:]).underline = True
    document.save(source)
    source_bytes = source.read_bytes()

    result = DocxProcessor().anonymize(source, anonymized)

    assert source.read_bytes() == source_bytes
    package_xml = _all_xml(anonymized)
    assert all(address.encode() not in package_xml for address in addresses)
    for token in (
        "Макетный", "Проверочный", "Лабораторный",
        "Пробная", "Контрольный", "Большая Испытательная",
    ):
        assert token.encode() not in package_xml
    assert set(addresses) == {
        original
        for marker, original in result.mapping.items()
        if marker.startswith("[АДРЕС-")
    }
    assert Document(anonymized).paragraphs[0].text.count("[АДРЕС-") == len(addresses)
    DocxProcessor().restore(
        anonymized,
        restored,
        result.mapping,
        package_archive=result.package_archive,
    )
    assert restored.read_bytes() == source_bytes


def test_docx_split_ordinal_address_is_one_marker_and_restores_exactly(tmp_path):
    source = tmp_path / "split-ordinal-address.docx"
    anonymized = tmp_path / "split-ordinal-address.anonymized.docx"
    restored = tmp_path / "split-ordinal-address.restored.docx"
    address = "3-ый Проверочный пр. д.7 стр.2"
    document = Document()
    paragraph = document.add_paragraph("Площадка расположена по адресу ")
    paragraph.add_run("3-ый Проверочный ").bold = True
    paragraph.add_run("пр. д.7 ").italic = True
    paragraph.add_run("стр.2").underline = True
    paragraph.add_run(", используется макет.")
    document.save(source)
    source_bytes = source.read_bytes()

    result = DocxProcessor().anonymize(source, anonymized)

    assert source.read_bytes() == source_bytes
    assert result.mapping == {"[АДРЕС-1]": address}
    assert address.encode() not in _all_xml(anonymized)
    assert Document(anonymized).paragraphs[0].text == (
        "Площадка расположена по адресу [АДРЕС-1], используется макет."
    )
    DocxProcessor().restore(
        anonymized,
        restored,
        result.mapping,
        package_archive=result.package_archive,
    )
    assert restored.read_bytes() == source_bytes


def test_docx_address_like_project_prose_in_table_cells_stays_visible(tmp_path):
    source = tmp_path / "synthetic-address-controls.docx"
    anonymized = tmp_path / "synthetic-address-controls.anonymized.docx"
    restored = tmp_path / "synthetic-address-controls.restored.docx"
    controls = (
        "Пр. Завершён, д. 12",
        "Улица Проекта обсуждалась 5 дней.",
    )
    document = Document()
    table = document.add_table(rows=len(controls), cols=1)
    for index, value in enumerate(controls):
        table.cell(index, 0).text = value
    document.save(source)
    source_bytes = source.read_bytes()

    result = DocxProcessor().anonymize(source, anonymized)

    output = Document(anonymized)
    assert tuple(row.cells[0].text for row in output.tables[0].rows) == controls
    assert not any(marker.startswith("[АДРЕС-") for marker in result.mapping)
    DocxProcessor().restore(
        anonymized,
        restored,
        result.mapping,
        package_archive=result.package_archive,
    )
    assert restored.read_bytes() == source_bytes


def test_docx_tab_separated_english_labels_are_redacted_and_restore_exactly(tmp_path):
    source = tmp_path / "tab-labels.docx"
    anonymized = tmp_path / "tab-labels.anonymized.docx"
    restored = tmp_path / "tab-labels.restored.docx"
    document = Document()
    paragraph = document.add_paragraph()
    paragraph.add_run("Contact:").bold = True
    paragraph.add_run().add_tab()
    paragraph.add_run("Alex J. Example")
    paragraph.add_run().add_break()
    paragraph.add_run("Software:").italic = True
    paragraph.add_run().add_tab()
    paragraph.add_run("Example Office Suite")
    document.save(source)
    source_bytes = source.read_bytes()

    result = DocxProcessor().anonymize(source, anonymized)

    output = Document(anonymized).paragraphs[0]
    assert "Alex J. Example" not in output.text
    assert "Example Office Suite" not in output.text
    assert "\t" in output.text
    assert "\n" in output.text
    assert output.runs[0].bold is True
    assert any(run.italic for run in output.runs)

    DocxProcessor().restore(
        anonymized,
        restored,
        result.mapping,
        package_archive=result.package_archive,
    )
    assert restored.read_bytes() == source_bytes


def test_docx_vm_inventory_names_are_removed_from_package_and_restore_exactly(tmp_path):
    source = tmp_path / "vm-source.docx"
    anonymized = tmp_path / "vm-anonymized.docx"
    restored = tmp_path / "vm-restored.docx"
    vm_names = ["qzx-monitor0101", "qzx-repo0101", "qzx-print0101", "qzx-fs0101", "qzx-journal0101"]
    document = Document()
    table = document.add_table(rows=0, cols=1)
    for name in vm_names:
        table.add_row().cells[0].text = name
    document.save(source)

    result = DocxProcessor().anonymize(source, anonymized)

    package_xml = _all_xml(anonymized)
    assert all(name.encode() not in package_xml for name in vm_names)
    assert set(result.mapping.values()) == set(vm_names)
    DocxProcessor().restore(anonymized, restored, result.mapping)
    restored_values = [row.cells[0].text for row in Document(restored).tables[0].rows]
    assert restored_values == vm_names


def test_docx_server_inventory_name_column_is_anonymized_and_restores_byte_exactly(tmp_path):
    source = tmp_path / "server-inventory-source.docx"
    anonymized = tmp_path / "server-inventory-anonymized.docx"
    restored = tmp_path / "server-inventory-restored.docx"
    server_names = ["q1234-qz01", "1234-lb02-2", "zx23"]
    document = Document()
    table = document.add_table(rows=1, cols=4)
    for index, value in enumerate((
        "Наименование", "Программное обеспечение", "IP адрес", "ИТ-сервис",
    )):
        table.rows[0].cells[index].text = value
    for index, name in enumerate(server_names, start=11):
        cells = table.add_row().cells
        for position, value in enumerate((
            name, "Sample Linux", f"192.0.2.{index}", "Тестовый сервис",
        )):
            cells[position].text = value
    document.save(source)
    source_bytes = source.read_bytes()

    result = DocxProcessor().anonymize(source, anonymized)

    package_xml = _all_xml(anonymized)
    assert all(name.encode() not in package_xml for name in server_names)
    assert set(server_names) <= set(result.mapping.values())
    output = Document(anonymized).tables[0]
    assert all(row.cells[0].text.startswith("[УЗЕЛ-") for row in output.rows[1:])
    DocxProcessor().restore(
        anonymized,
        restored,
        result.mapping,
        package_archive=result.package_archive,
    )
    assert restored.read_bytes() == source_bytes


def test_docx_semantic_server_capacity_table_redacts_every_node_name(tmp_path):
    source = tmp_path / "semantic-server-capacity-source.docx"
    anonymized = tmp_path / "semantic-server-capacity-anonymized.docx"
    disabled = tmp_path / "semantic-server-capacity-disabled.docx"
    restored = tmp_path / "semantic-server-capacity-restored.docx"
    node_names = ("qrf73-xnt04", "qrf73-vmk02", "qrf72-bkp01")
    document = Document()
    table = document.add_table(rows=1, cols=6)
    for index, value in enumerate((
        "Наименование", "IP адрес", "Программное обеспечение",
        "CPU", "MEM цел", "HDD SATA",
    )):
        table.rows[0].cells[index].text = value
    for row_index, name in enumerate(node_names, start=1):
        cells = table.add_row().cells
        if row_index == 2:
            cells[0].paragraphs[0].add_run(name[:6]).bold = True
            cells[0].paragraphs[0].add_run(name[6:]).italic = True
        elif row_index == 3:
            cells[0].paragraphs[0].add_run(name[:7])
            cells[0].paragraphs[0].add_run().add_break()
            cells[0].paragraphs[0].add_run(name[7:])
        else:
            cells[0].text = name
        cells[1].text = f"192.0.2.{70 + row_index}"
        cells[2].text = "Synthetic Linux (64-bit)"
        cells[3].text = "32"
        cells[4].text = "64"
        cells[5].text = "144"
    document.save(source)
    source_bytes = source.read_bytes()

    result = DocxProcessor().anonymize(source, anonymized)

    package_xml = _all_xml(anonymized)
    assert all(name.encode() not in package_xml for name in node_names)
    assert all(
        fragment.encode() not in package_xml
        for fragment in ("qrf73-", "xnt04", "vmk02", "qrf72-b", "kp01")
    )
    assert set(node_names) <= {value.replace("\n", "") for value in result.mapping.values()}
    output = Document(anonymized).tables[0]
    assert all(row.cells[0].text.strip().startswith("[УЗЕЛ-") for row in output.rows[1:])

    disabled_result = DocxProcessor(Anonymizer(disabled_categories={"УЗЕЛ"})).anonymize(
        source, disabled,
    )
    disabled_table = Document(disabled).tables[0]
    assert [row.cells[0].text.replace("\n", "") for row in disabled_table.rows[1:]] == list(node_names)
    assert not any(marker.startswith("[УЗЕЛ-") for marker in disabled_result.mapping)

    DocxProcessor().restore(
        anonymized,
        restored,
        result.mapping,
        package_archive=result.package_archive,
    )
    assert restored.read_bytes() == source_bytes


def test_docx_server_inventory_names_split_inside_one_cell_are_one_node_and_restore_exactly(tmp_path):
    source = tmp_path / "split-server-inventory-source.docx"
    anonymized = tmp_path / "split-server-inventory-anonymized.docx"
    restored = tmp_path / "split-server-inventory-restored.docx"
    server_names = ("1234-lb02-2", "q2345-qz03", "zx23", "qzlab-dc01")
    document = Document()
    table = document.add_table(rows=1, cols=4)
    for index, value in enumerate((
        "Наименование", "Программное обеспечение", "IP адрес", "ИТ-сервис",
    )):
        table.rows[0].cells[index].text = value

    cells = table.add_row().cells
    paragraph = cells[0].paragraphs[0]
    paragraph.add_run("1234-")
    paragraph.add_run().add_break()
    paragraph.add_run("lb02-2")
    cells[1].text = "Sample Linux"
    cells[2].text = "192.0.2.41"
    cells[3].text = "Тестовый сервис"

    cells = table.add_row().cells
    cells[0].paragraphs[0].text = "q2345-q"
    cells[0].add_paragraph("z03")
    cells[1].text = "Sample Linux"
    cells[2].text = "192.0.2.42"
    cells[3].text = "Тестовый сервис"

    cells = table.add_row().cells
    paragraph = cells[0].paragraphs[0]
    paragraph.add_run("zx")
    paragraph.add_run().add_break()
    paragraph.add_run("23")
    cells[1].text = "Sample Linux"
    cells[2].text = "192.0.2.43"
    cells[3].text = "Тестовый сервис"

    cells = table.add_row().cells
    paragraph = cells[0].paragraphs[0]
    paragraph.add_run("qzlab-")
    paragraph.add_run().add_break()
    paragraph.add_run("dc01")
    cells[1].text = "Sample Linux"
    cells[2].text = "192.0.2.44"
    cells[3].text = "Тестовый сервис"
    document.save(source)
    source_bytes = source.read_bytes()

    result = DocxProcessor().anonymize(source, anonymized)

    output = Document(anonymized).tables[0]
    assert all(row.cells[0].text.strip().startswith("[УЗЕЛ-") for row in output.rows[1:])
    assert all(name.encode() not in _all_xml(anonymized) for name in server_names)
    assert set(server_names) <= {value.replace("\n", "") for value in result.mapping.values()}
    DocxProcessor().restore(
        anonymized,
        restored,
        result.mapping,
        package_archive=result.package_archive,
    )
    assert restored.read_bytes() == source_bytes


def test_docx_underscored_service_accounts_across_runs_and_hard_break_restore_exactly(tmp_path):
    source = tmp_path / "service-account-source.docx"
    anonymized = tmp_path / "service-account-anonymized.docx"
    restored = tmp_path / "service-account-restored.docx"
    accounts = ("SvcQXZ-DB_Backup", "SvcRTY-File_Sync")
    document = Document()
    paragraph = document.add_paragraph("Примеры:\n")
    paragraph.add_run("SvcQXZ-")
    paragraph.add_run("DB_").bold = True
    paragraph.add_run("Backup")
    paragraph.add_run().add_break()
    paragraph.add_run(accounts[1])
    document.save(source)
    source_bytes = source.read_bytes()

    result = DocxProcessor().anonymize(source, anonymized)

    output = Document(anonymized).paragraphs[0].text
    assert all(account not in output for account in accounts)
    assert output.count("[ЛОГИН-") == 2
    assert set(accounts) <= {value.replace("\n", "") for value in result.mapping.values()}
    assert all(account.encode() not in _all_xml(anonymized) for account in accounts)
    DocxProcessor().restore(
        anonymized,
        restored,
        result.mapping,
        package_archive=result.package_archive,
    )
    assert restored.read_bytes() == source_bytes


def test_docx_service_account_split_inside_token_by_hard_break_has_no_fragment_leak(tmp_path):
    source = tmp_path / "split-service-account-source.docx"
    anonymized = tmp_path / "split-service-account-anonymized.docx"
    restored = tmp_path / "split-service-account-restored.docx"
    prefix = "SvcQXZ-DB_"
    suffix = "Backup"
    account = prefix + suffix
    document = Document()
    paragraph = document.add_paragraph("Пример: ")
    paragraph.add_run(prefix)
    paragraph.add_run().add_break()
    paragraph.add_run(suffix).bold = True
    document.save(source)
    source_bytes = source.read_bytes()

    result = DocxProcessor().anonymize(source, anonymized)

    output = Document(anonymized).paragraphs[0].text
    payload = _all_xml(anonymized)
    assert output.count("[ЛОГИН-") == 1
    assert "[СОКРАЩЕНИЯ-" not in output
    assert prefix not in output
    assert suffix not in output
    assert prefix.encode() not in payload
    assert suffix.encode() not in payload
    assert account not in output.replace("\n", "")
    assert account not in payload.decode("utf-8", errors="ignore").replace("\n", "")
    assert list(result.mapping.values()) == [prefix + "\n" + suffix]
    assert b"<w:br" in payload
    assert source.read_bytes() == source_bytes

    DocxProcessor().restore(
        anonymized,
        restored,
        result.mapping,
        package_archive=result.package_archive,
    )
    assert restored.read_bytes() == source_bytes


def test_docx_same_service_account_with_distinct_break_layouts_restores_exactly(tmp_path):
    source = tmp_path / "service-account-layouts-source.docx"
    anonymized = tmp_path / "service-account-layouts-anonymized.docx"
    restored = tmp_path / "service-account-layouts-restored.docx"
    account = "SvcQXZ-DB_Backup"
    document = Document()
    document.add_paragraph(account)
    split = document.add_paragraph()
    split.add_run("SvcQXZ-DB_")
    split.add_run().add_break()
    split.add_run("Backup")
    document.save(source)
    source_bytes = source.read_bytes()

    result = DocxProcessor().anonymize(source, anonymized)

    output = Document(anonymized)
    assert all(account not in paragraph.text.replace("\n", "") for paragraph in output.paragraphs)
    assert len(result.mapping) == 2
    assert set(result.mapping.values()) == {account, "SvcQXZ-DB_\nBackup"}
    DocxProcessor().restore(
        anonymized,
        restored,
        result.mapping,
        package_archive=result.package_archive,
    )
    assert restored.read_bytes() == source_bytes


def test_docx_hard_line_breaks_do_not_join_server_names_to_preceding_ips(tmp_path):
    source = tmp_path / "line-break-source.docx"
    anonymized = tmp_path / "line-break-anonymized.docx"
    restored = tmp_path / "line-break-restored.docx"
    document = Document()
    cell = document.add_table(rows=1, cols=1).cell(0, 0)
    paragraph = cell.paragraphs[0]
    lines = [
        "ooo-srkv/172.17.60.113",
        "CyberBackUp/172.17.60.120",
        "Cyber2Data/172.17.60.125",
        "Сервера Вне домена",
    ]
    for index, line in enumerate(lines):
        if index:
            paragraph.add_run().add_break()
        paragraph.add_run(line)
    document.save(source)

    # Reproduces the screenshot even when IP replacement was disabled: every
    # slash-delimited hostname must still be independently anonymized.
    processor = DocxProcessor(Anonymizer(disabled_categories={"IP"}))
    result = processor.anonymize(source, anonymized)

    output = Document(anonymized).tables[0].cell(0, 0).paragraphs[0].text
    assert "ooo-srkv" not in output
    assert "CyberBackUp" not in output
    assert "Cyber2Data" not in output
    assert output.count("[УЗЕЛ-") == 3
    assert output.count("\n") == 3
    assert all(address in output for address in ("172.17.60.113", "172.17.60.120", "172.17.60.125"))
    assert {"ooo-srkv", "CyberBackUp", "Cyber2Data"}.issubset(set(result.mapping.values()))

    processor.restore(anonymized, restored, result.mapping)
    restored_text = Document(restored).tables[0].cell(0, 0).paragraphs[0].text
    assert restored_text == "\n".join(lines)


def test_docx_hard_line_breaks_preserve_ip_boundaries(tmp_path):
    source = tmp_path / "line-break-ip-source.docx"
    anonymized = tmp_path / "line-break-ip-anonymized.docx"
    document = Document()
    paragraph = document.add_paragraph()
    for index, line in enumerate((
        "ooo-srkv/172.17.60.113",
        "CyberBackUp/172.17.60.120",
        "Cyber2Data/172.17.60.125",
        "NodeTab/172.17.60.126",
    )):
        if index == 1:
            paragraph.add_run()._r.append(OxmlElement("w:cr"))
        elif index == 2:
            paragraph.add_run().add_break()
        elif index == 3:
            paragraph.add_run().add_tab()
        paragraph.add_run(line)
    document.save(source)

    result = DocxProcessor().anonymize(source, anonymized)

    output = Document(anonymized).paragraphs[0].text
    assert all(value not in output for value in (
        "ooo-srkv", "CyberBackUp", "Cyber2Data", "NodeTab",
        "172.17.60.113", "172.17.60.120", "172.17.60.125", "172.17.60.126",
    ))
    assert output.count("[УЗЕЛ-") == 4
    assert output.count("[IP-") == 4
    assert output.count("\n") == 2
    assert output.count("\t") == 1
    assert result.total == 8
    with ZipFile(anonymized) as package:
        root = etree.fromstring(package.read("word/document.xml"))
    namespace = {"w": "http://schemas.openxmlformats.org/wordprocessingml/2006/main"}
    assert len(root.xpath(".//w:cr", namespaces=namespace)) == 1
    assert len(root.xpath(".//w:br", namespaces=namespace)) == 1
    assert len(root.xpath(".//w:tab", namespaces=namespace)) == 1


def test_docx_contextual_inventory_values_are_removed_and_exactly_restored(tmp_path):
    source = tmp_path / "context-source.docx"
    anonymized = tmp_path / "context-anonymized.docx"
    restored = tmp_path / "context-restored.docx"
    sensitive = [
        "lab-nas", "vm-pc", "ALPHAHQ", "BETA-LAB",
        "Computers_ACME/Production/Site-A", "Servers/Legacy-RDP",
        "Domain Controllers", "NAS_VOG", "Computers_VOG",
        "Example Secure Client", "Northwind Banking Plug-in",
        "1C 8.3.27.1859; New_setup_1C; Setup_1C_*; TessaDesk; Установка LanDocs; deploy_office2016; Kaspersky_AK*",
    ]
    document = Document()

    hosts = document.add_table(rows=1, cols=4)
    for index, value in enumerate(("Имя хоста", "IP-адрес", "vCPU", "RAM")):
        hosts.rows[0].cells[index].text = value
    for values in (("lab-nas", "[IP-74]", "32", "128"), ("vm-pc", "[IP-75]", "16", "32")):
        cells = hosts.add_row().cells
        for index, value in enumerate(values):
            cells[index].text = value

    netbios = document.add_table(rows=2, cols=3)
    for index, value in enumerate(("Параметр", "Площадка 1", "Площадка 2")):
        netbios.rows[0].cells[index].text = value
    for index, value in enumerate(("NetBIOS", "ALPHAHQ", "BETA-LAB")):
        netbios.rows[1].cells[index].text = value

    ous = document.add_table(rows=1, cols=4)
    for index, value in enumerate(("OU или контейнер AD", "Компьютерных УЗ", "Включено", "Отключено")):
        ous.rows[0].cells[index].text = value
    for values in (
        ("Computers_ACME/Production/Site-A", "21", "21", "0"),
        ("Domain Controllers", "2", "2", "0"),
        ("NAS_VOG", "2", "2", "0"),
        ("Computers_VOG", "1", "1", "0"),
        ("Servers/Legacy-RDP", "12", "11", "1"),
    ):
        cells = ous.add_row().cells
        for index, value in enumerate(values):
            cells[index].text = value

    software = document.add_table(rows=2, cols=3)
    software.cell(0, 0).text = "Инвентаризация программного обеспечения"
    software.cell(0, 1).text = "Установок"
    software.cell(0, 2).text = "Доля"
    software.cell(1, 0).text = "Example Secure Client"
    software.cell(1, 1).text = "10"
    software.cell(1, 2).text = "1,6%"
    row = software.add_row().cells
    row[0].text = "Northwind Banking Plug-in"
    row[1].text = "4"
    row[2].text = "0,6%"

    deployment = document.add_table(rows=1, cols=3)
    deployment.cell(0, 0).text = "Развёртывание ПО"
    deployment.cell(0, 1).text = sensitive[-1]
    deployment.cell(0, 2).text = "Установка возможна при наличии поддерживаемых Linux-пакетов"
    document.save(source)
    source_bytes = source.read_bytes()

    result = DocxProcessor().anonymize(source, anonymized)

    package_xml = _all_xml(anonymized)
    assert all(value.encode("utf-8") not in package_xml for value in sensitive)
    assert set(sensitive).issubset(set(result.mapping.values()))

    DocxProcessor().restore(
        anonymized,
        restored,
        result.mapping,
        package_archive=result.package_archive,
    )
    assert restored.read_bytes() == source_bytes


def test_docx_abbreviations_across_runs_tables_and_tabs_restore_byte_exactly(tmp_path):
    source = tmp_path / "abbreviations-source.docx"
    anonymized = tmp_path / "abbreviations-anonymized.docx"
    restored = tmp_path / "abbreviations-restored.docx"
    document = Document()

    paragraph = document.add_paragraph()
    paragraph.add_run("Система ")
    paragraph.add_run("QZ")
    paragraph.add_run("XB").bold = True
    paragraph.add_run().add_tab()
    paragraph.add_run("UCMDB")
    adjacent = document.add_paragraph("Отчет об обследовании СК АО ")
    adjacent.add_run("QZX")
    adjacent.add_run("B1").bold = True
    table = document.add_table(rows=1, cols=2)
    table.cell(0, 0).text = "Интеграция"
    table.cell(0, 1).text = "SSO и NTLM"
    document.save(source)
    source_bytes = source.read_bytes()

    result = DocxProcessor().anonymize(source, anonymized)
    output_document = Document(anonymized)
    output = "\n".join(
        [p.text for p in output_document.paragraphs]
        + [cell.text for row in output_document.tables[0].rows for cell in row.cells]
    )

    assert all(value not in output for value in ("QZXB", "QZXB1", "UCMDB", "SSO", "NTLM"))
    assert set(result.mapping.values()) >= {"QZXB", "QZXB1", "UCMDB", "SSO", "NTLM"}

    DocxProcessor().restore(
        anonymized,
        restored,
        result.mapping,
        package_archive=result.package_archive,
    )
    assert restored.read_bytes() == source_bytes


def test_docx_single_two_letter_abbreviation_spans_runs_and_restores_byte_exactly(tmp_path):
    source = tmp_path / "short-abbreviations-source.docx"
    anonymized = tmp_path / "short-abbreviations-anonymized.docx"
    restored = tmp_path / "short-abbreviations-restored.docx"
    document = Document()

    split = document.add_paragraph()
    split.add_run("Q")
    split.add_run("Z").bold = True
    document.save(source)
    source_bytes = source.read_bytes()

    result = DocxProcessor().anonymize(source, anonymized)

    assert set(result.mapping.values()) == {"QZ"}
    assert b"QZ" not in _all_xml(anonymized)

    DocxProcessor().restore(
        anonymized,
        restored,
        result.mapping,
        package_archive=result.package_archive,
    )
    assert restored.read_bytes() == source_bytes


def test_docx_toc_domains_with_attached_page_numbers_restore_byte_exactly(tmp_path):
    source = tmp_path / "toc-domains-source.docx"
    anonymized = tmp_path / "toc-domains-anonymized.docx"
    restored = tmp_path / "toc-domains-restored.docx"
    document = Document()
    values = (
        ("qzxsafe.", "test", "11"),
        ("go.qzxsafe.", "test", "14"),
        ("rf.qzxsafe.", "test", "19"),
    )
    for prefix, tld, page in values:
        paragraph = document.add_paragraph("Общие данные по домену ")
        paragraph.add_run(prefix)
        paragraph.add_run(tld).bold = True
        paragraph.add_run(page)
    document.save(source)
    source_bytes = source.read_bytes()

    result = DocxProcessor().anonymize(source, anonymized)

    output = [paragraph.text for paragraph in Document(anonymized).paragraphs]
    assert output == [
        "Общие данные по домену [ДОМЕН-1]11",
        "Общие данные по домену [ДОМЕН-2]14",
        "Общие данные по домену [ДОМЕН-3]19",
    ]
    assert result.mapping == {
        "[ДОМЕН-1]": "qzxsafe.test",
        "[ДОМЕН-2]": "go.qzxsafe.test",
        "[ДОМЕН-3]": "rf.qzxsafe.test",
    }

    DocxProcessor().restore(
        anonymized,
        restored,
        result.mapping,
        package_archive=result.package_archive,
    )
    assert restored.read_bytes() == source_bytes


def test_docx_ports_across_runs_tabs_and_endpoints_restore_byte_exactly(tmp_path):
    source = tmp_path / "ports-source.docx"
    anonymized = tmp_path / "ports-anonymized.docx"
    restored = tmp_path / "ports-restored.docx"
    document = Document()

    labelled = document.add_paragraph("Сетевой порт ")
    labelled.add_run("80")
    labelled.add_run("80").bold = True
    tabbed = document.add_paragraph("Port")
    tabbed.add_run().add_tab()
    tabbed.add_run("8443")
    endpoint = document.add_paragraph()
    endpoint.add_run("node.qzxsafe.test:")
    endpoint.add_run("54")
    endpoint.add_run("32").italic = True
    document.save(source)
    source_bytes = source.read_bytes()

    result = DocxProcessor().anonymize(source, anonymized)

    output = [paragraph.text for paragraph in Document(anonymized).paragraphs]
    assert output == [
        "Сетевой порт [ПОРТ-1]",
        "Port\t[ПОРТ-2]",
        "[ДОМЕН-1]:[ПОРТ-3]",
    ]
    assert result.mapping["[ПОРТ-1]"] == "8080"
    assert result.mapping["[ПОРТ-2]"] == "8443"
    assert result.mapping["[ПОРТ-3]"] == "5432"

    DocxProcessor().restore(
        anonymized,
        restored,
        result.mapping,
        package_archive=result.package_archive,
    )
    assert restored.read_bytes() == source_bytes


def test_docx_windows_to_alse_table_redacts_both_software_columns_and_restores(tmp_path):
    source = tmp_path / "windows-alse-source.docx"
    anonymized = tmp_path / "windows-alse-anonymized.docx"
    restored = tmp_path / "windows-alse-restored.docx"
    edited = tmp_path / "windows-alse-edited.docx"
    semantic_restored = tmp_path / "windows-alse-semantic-restored.docx"
    rows = [
        ("Microsoft Office", "LibreOffice, пакет libreoffice", "Да — Astra"),
        (
            "Яндекс Браузер",
            "Mozilla Firefox ESR, пакет firefox-esr",
            "Совместимость подтверждена внутренним тестированием",
        ),
        (
            "Агент администрирования Kaspersky Security Center",
            "Kaspersky Security Center Network Agent for Linux",
            "Нет в перечне Astra; дистрибутив производителя",
        ),
        ("7-Zip", "7-Zip, пакет 7zip", "Доступен в штатном репозитории"),
        (
            "DameWare Mini Remote Control",
            "Remmina и TigerVNC",
            "Ручная установка по инструкции поставщика",
        ),
    ]
    document = Document()
    table = document.add_table(rows=1, cols=3)
    for index, value in enumerate((
        "Исходное Windows-ПО",
        "Вариант для ALSE / Linux",
        "Статус аналогов и наличие в репозитории Astra",
    )):
        table.rows[0].cells[index].text = value
    for values in rows:
        cells = table.add_row().cells
        for index, value in enumerate(values):
            cells[index].text = value

    long_software = "LegacySuite-" + "X" * 260
    cells = table.add_row().cells
    cells[0].text = f"  {long_software}  "
    cells[1].text = ""
    cells[1].add_paragraph("Remmina")
    cells[1].add_paragraph("и TigerVNC")
    cells[1].add_paragraph("")
    cells[2].text = "Проверено отдельным стендом"

    cells = table.add_row().cells
    cells[0].text = ""
    cells[1].text = "OnlyOffice Desktop Editors"
    cells[2].text = "Проходит сертификацию"

    cells = table.add_row().cells
    cells[0].text = "Legacy Browser"
    paragraph = cells[1].paragraphs[0]
    paragraph.clear()
    run = paragraph.add_run("Firefox")
    run.add_break()
    run.add_text("ESR package")
    cells[2].text = "Поддерживается поставщиком"

    cells = table.add_row().cells
    cells[0].merge(cells[1]).text = "Merged layout note"
    cells[2].text = "Не строка из трёх ячеек"

    structured_products = [
        long_software,
        "Remmina\nи TigerVNC",
        "OnlyOffice Desktop Editors",
        "Legacy Browser",
        "Firefox\nESR package",
    ]
    leaked_fragments = [
        long_software,
        "Remmina",
        "TigerVNC",
        "OnlyOffice Desktop Editors",
        "Legacy Browser",
        "Firefox",
        "ESR package",
    ]
    unrelated_values = ("Раздел согласования", "Ответственный отдел", "Да, согласовано")
    unrelated = document.add_table(rows=1, cols=3)
    for index, value in enumerate(unrelated_values):
        unrelated.rows[0].cells[index].text = value
    document.save(source)
    source_bytes = source.read_bytes()

    result = DocxProcessor().anonymize(source, anonymized)

    package_xml = _all_xml(anonymized)
    products = [value for row in rows for value in row[:2]]
    assert all(value.encode("utf-8") not in package_xml for value in products)
    assert all(value.encode("utf-8") not in package_xml for value in leaked_fragments)
    assert set(products + structured_products).issubset(set(result.mapping.values()))
    assert all(not any(marker in value for marker in "\ufdd0\ufdd1\ufdd2\ufdd3") for value in result.mapping.values())
    assert all(status.encode("utf-8") in package_xml for _, _, status in rows)
    assert "Merged layout note".encode("utf-8") in package_xml
    assert all(value.encode("utf-8") in package_xml for value in unrelated_values)
    assert not set(unrelated_values).intersection(result.mapping.values())

    DocxProcessor().restore(
        anonymized,
        restored,
        result.mapping,
        package_archive=result.package_archive,
    )
    assert restored.read_bytes() == source_bytes

    changed = Document(anonymized)
    changed.add_paragraph("Редакторская правка")
    changed.save(edited)
    DocxProcessor().restore(
        edited,
        semantic_restored,
        result.mapping,
        package_archive=result.package_archive,
    )
    semantic_document = Document(semantic_restored)
    semantic_table = semantic_document.tables[0]
    assert semantic_table.rows[6].cells[0].text == f"  {long_software}  "
    assert semantic_table.rows[6].cells[1].text == "\nRemmina\nи TigerVNC\n"
    assert semantic_table.rows[7].cells[0].text == ""
    assert semantic_table.rows[7].cells[1].text == "OnlyOffice Desktop Editors"
    assert semantic_table.rows[8].cells[0].text == "Legacy Browser"
    assert semantic_table.rows[8].cells[1].text == "Firefox\nESR package"
    assert semantic_table.rows[9].cells[0].text == "Merged layout note"
    assert semantic_document.paragraphs[-1].text == "Редакторская правка"


def test_docx_graphic_is_removed_from_anonymized_package_and_restored_exactly(tmp_path):
    source = tmp_path / "graphic-source.docx"
    image = tmp_path / "synthetic.png"
    anonymized = tmp_path / "graphic-anonymized.docx"
    restored = tmp_path / "graphic-restored.docx"
    image_bytes = base64.b64decode(
        "iVBORw0KGgoAAAANSUhEUgAAAAEAAAABCAQAAAC1HAwCAAAAC0lEQVR42mNk+A8AAQUBAScY42YAAAAASUVORK5CYII="
    )
    image.write_bytes(image_bytes)
    document = Document()
    document.add_paragraph("Иванов Сергей Петрович")
    document.add_picture(str(image), width=Inches(1))
    document.save(source)

    result = DocxProcessor().anonymize(source, anonymized)

    with ZipFile(anonymized) as package:
        assert image_bytes not in (package.read(name) for name in package.namelist())
        assert not any(name.startswith("word/media/") for name in package.namelist())
    assert result.package_archive

    DocxProcessor().restore(
        anonymized,
        restored,
        result.mapping,
        package_archive=result.package_archive,
    )

    with ZipFile(restored) as package:
        restored_images = [
            package.read(name)
            for name in package.namelist()
            if name.startswith("word/media/")
        ]
        relationships = etree.fromstring(package.read("word/_rels/document.xml.rels"))
        assert restored_images == [image_bytes]
        assert relationships.xpath(
            "//*[local-name()='Relationship' and contains(@Type, '/image')]"
        )
    restored_document = Document(restored)
    assert restored_document.inline_shapes
    assert restored_document.paragraphs[0].text == "Иванов Сергей Петрович"


def test_docx_unedited_anonymized_package_restores_source_bytes_exactly(tmp_path):
    source = tmp_path / "exact-source.docx"
    image = tmp_path / "synthetic.png"
    anonymized = tmp_path / "exact-anonymized.docx"
    restored = tmp_path / "exact-restored.docx"
    image.write_bytes(base64.b64decode(
        "iVBORw0KGgoAAAANSUhEUgAAAAEAAAABCAQAAAC1HAwCAAAAC0lEQVR42mNk+A8AAQUBAScY42YAAAAASUVORK5CYII="
    ))
    document = Document()
    document.add_paragraph("Сервер QZX-SRV-0101")
    document.add_picture(str(image), width=Inches(1))
    document.save(source)

    result = DocxProcessor().anonymize(source, anonymized)
    DocxProcessor().restore(
        anonymized,
        restored,
        result.mapping,
        package_archive=result.package_archive,
    )

    assert restored.read_bytes() == source.read_bytes()


def test_docx_graphic_restores_original_zip_member_metadata(tmp_path):
    source = tmp_path / "graphic-metadata-source.docx"
    image = tmp_path / "synthetic.png"
    anonymized = tmp_path / "graphic-metadata-anonymized.docx"
    restored = tmp_path / "graphic-metadata-restored.docx"
    image.write_bytes(base64.b64decode(
        "iVBORw0KGgoAAAANSUhEUgAAAAEAAAABCAQAAAC1HAwCAAAAC0lEQVR42mNk+A8AAQUBAScY42YAAAAASUVORK5CYII="
    ))
    document = Document()
    document.add_paragraph("Сервер QZX-SRV-0101")
    document.add_picture(str(image), width=Inches(1))
    document.save(source)

    with ZipFile(source) as package:
        media_name = next(name for name in package.namelist() if name.startswith("word/media/"))

    replacement = source.with_suffix(".metadata.docx")
    with ZipFile(source) as incoming, ZipFile(replacement, "w") as outgoing:
        for original_info in incoming.infolist():
            payload = incoming.read(original_info.filename)
            if original_info.filename != media_name:
                outgoing.writestr(original_info, payload)
                continue
            info = ZipInfo(media_name, date_time=(2003, 4, 5, 6, 7, 8))
            info.compress_type = ZIP_STORED
            info.comment = b"GRAPHIC-METADATA"
            info.extra = b"\x99\x99\x04\x00TEST"
            info.create_system = 3
            info.create_version = 45
            info.extract_version = 20
            info.internal_attr = 1
            info.external_attr = 0o100640 << 16
            outgoing.writestr(info, payload)
    replacement.replace(source)

    with ZipFile(source) as package:
        expected = package.getinfo(media_name)
    result = DocxProcessor().anonymize(source, anonymized)
    DocxProcessor().restore(
        anonymized, restored, result.mapping, package_archive=result.package_archive
    )

    with ZipFile(restored) as package:
        actual = package.getinfo(media_name)
    assert actual.compress_type == expected.compress_type
    assert actual.date_time == expected.date_time
    assert actual.comment == expected.comment
    assert actual.extra == expected.extra
    assert actual.create_system == expected.create_system
    assert actual.create_version == expected.create_version
    assert actual.extract_version == expected.extract_version
    assert actual.internal_attr == expected.internal_attr
    assert actual.external_attr == expected.external_attr


def test_docx_vml_scheme_without_binary_part_round_trips_through_protected_archive(tmp_path):
    source = tmp_path / "vml-source.docx"
    anonymized = tmp_path / "vml-anonymized.docx"
    restored = tmp_path / "vml-restored.docx"
    document = Document()
    run = document.add_paragraph().add_run()
    pict = etree.fromstring(
        b'<w:pict xmlns:w="http://schemas.openxmlformats.org/wordprocessingml/2006/main" '
        b'xmlns:v="urn:schemas-microsoft-com:vml"><v:shape id="synthetic-scheme" '
        b'style="width:100pt;height:40pt"><v:textbox><w:txbxContent><w:p><w:r>'
        b'<w:t>\xd0\x98\xd0\xb2\xd0\xb0\xd0\xbd\xd0\xbe\xd0\xb2 \xd0\xa1\xd0\xb5\xd1\x80\xd0\xb3\xd0\xb5\xd0\xb9 \xd0\x9f\xd0\xb5\xd1\x82\xd1\x80\xd0\xbe\xd0\xb2\xd0\xb8\xd1\x87</w:t></w:r></w:p></w:txbxContent></v:textbox>'
        b'</v:shape></w:pict>'
    )
    run._r.append(pict)
    document.save(source)

    result = DocxProcessor().anonymize(source, anonymized)

    with ZipFile(anonymized) as package:
        document_xml = package.read("word/document.xml")
        assert b"synthetic-scheme" not in document_xml
        assert "Иванов Сергей Петрович".encode() not in document_xml
        assert "[СХЕМА-1]".encode() in document_xml

    DocxProcessor().restore(
        anonymized,
        restored,
        result.mapping,
        package_archive=result.package_archive,
    )

    with ZipFile(restored) as package:
        document_xml = package.read("word/document.xml")
        assert b"synthetic-scheme" in document_xml
        assert "Иванов Сергей Петрович".encode() in document_xml
    Document(restored)


def test_docx_graphics_are_preserved_when_redaction_checkbox_is_off(tmp_path):
    source = tmp_path / "source.docx"
    image = tmp_path / "synthetic.png"
    anonymized = tmp_path / "anonymized.docx"
    image_bytes = base64.b64decode(
        "iVBORw0KGgoAAAANSUhEUgAAAAEAAAABCAQAAAC1HAwCAAAAC0lEQVR42mNk+A8AAQUBAScY42YAAAAASUVORK5CYII="
    )
    image.write_bytes(image_bytes)
    document = Document()
    document.add_paragraph("Обычный текст")
    document.add_picture(str(image), width=Inches(1))
    document.save(source)

    result = DocxProcessor().anonymize(source, anonymized, redact_graphics=False)

    with ZipFile(anonymized) as package:
        media = [package.read(name) for name in package.namelist() if name.startswith("word/media/")]
        assert image_bytes in media
        assert "[СХЕМА-".encode() not in package.read("word/document.xml")
    assert result.package_archive["fragments"] == []


def test_docx_restore_rejects_tampered_graphics_archive_binary(tmp_path):
    source = tmp_path / "source.docx"
    image = tmp_path / "synthetic.png"
    anonymized = tmp_path / "anonymized.docx"
    restored = tmp_path / "restored.docx"
    image.write_bytes(base64.b64decode(
        "iVBORw0KGgoAAAANSUhEUgAAAAEAAAABCAQAAAC1HAwCAAAAC0lEQVR42mNk+A8AAQUBAScY42YAAAAASUVORK5CYII="
    ))
    document = Document()
    document.add_paragraph("Обычный текст")
    document.add_picture(str(image), width=Inches(1))
    document.save(source)
    result = DocxProcessor().anonymize(source, anonymized)
    archive = copy.deepcopy(result.package_archive)
    part_name = next(iter(archive["parts"]))
    archive["parts"][part_name] = base64.b64encode(b"MZ-not-an-image").decode("ascii")

    with pytest.raises(RestorationError, match="графическ"):
        DocxProcessor().restore(
            anonymized, restored, result.mapping, package_archive=archive
        )

    assert not restored.exists()


def test_docx_restore_rejects_unapproved_graphics_zip_metadata(tmp_path):
    source = tmp_path / "source.docx"
    image = tmp_path / "synthetic.png"
    anonymized = tmp_path / "anonymized.docx"
    restored = tmp_path / "restored.docx"
    image.write_bytes(base64.b64decode(
        "iVBORw0KGgoAAAANSUhEUgAAAAEAAAABCAQAAAC1HAwCAAAAC0lEQVR42mNk+A8AAQUBAScY42YAAAAASUVORK5CYII="
    ))
    document = Document()
    document.add_paragraph("Обычный текст")
    document.add_picture(str(image), width=Inches(1))
    document.save(source)
    result = DocxProcessor().anonymize(source, anonymized)
    archive = copy.deepcopy(result.package_archive)
    part_name = next(iter(archive["part_metadata"]))
    archive["part_metadata"][part_name]["compress_type"] = 99

    with pytest.raises(RestorationError, match="сжатие"):
        DocxProcessor().restore(
            anonymized, restored, result.mapping, package_archive=archive
        )

    assert not restored.exists()

    archive["part_metadata"][part_name]["compress_type"] = ZIP_STORED
    archive["part_metadata"][part_name]["extra"] = base64.b64encode(
        b"\x99\x99\x05\x00x"
    ).decode("ascii")
    with pytest.raises(RestorationError, match="метаданные ZIP"):
        DocxProcessor().restore(
            anonymized, restored, result.mapping, package_archive=archive
        )

    assert not restored.exists()


def test_docx_restore_rejects_unapproved_graphics_content_type(tmp_path):
    source = tmp_path / "source.docx"
    image = tmp_path / "synthetic.png"
    anonymized = tmp_path / "anonymized.docx"
    restored = tmp_path / "restored.docx"
    image.write_bytes(base64.b64decode(
        "iVBORw0KGgoAAAANSUhEUgAAAAEAAAABCAQAAAC1HAwCAAAAC0lEQVR42mNk+A8AAQUBAScY42YAAAAASUVORK5CYII="
    ))
    document = Document()
    document.add_paragraph("Обычный текст")
    document.add_picture(str(image), width=Inches(1))
    document.save(source)
    result = DocxProcessor().anonymize(source, anonymized)
    archive = copy.deepcopy(result.package_archive)
    archive["content_types"].append(base64.b64encode(
        b'<Default xmlns="http://schemas.openxmlformats.org/package/2006/content-types" '
        b'Extension="png" ContentType="application/x-msdownload"/>'
    ).decode("ascii"))

    with pytest.raises(RestorationError, match="тип"):
        DocxProcessor().restore(
            anonymized, restored, result.mapping, package_archive=archive
        )

    assert not restored.exists()


def test_docx_restore_rejects_dtd_in_protected_graphics_fragment(tmp_path):
    source = tmp_path / "source.docx"
    image = tmp_path / "synthetic.png"
    anonymized = tmp_path / "anonymized.docx"
    restored = tmp_path / "restored.docx"
    image.write_bytes(base64.b64decode(
        "iVBORw0KGgoAAAANSUhEUgAAAAEAAAABCAQAAAC1HAwCAAAAC0lEQVR42mNk+A8AAQUBAScY42YAAAAASUVORK5CYII="
    ))
    document = Document()
    document.add_paragraph("Обычный текст")
    document.add_picture(str(image), width=Inches(1))
    document.save(source)
    result = DocxProcessor().anonymize(source, anonymized)
    archive = copy.deepcopy(result.package_archive)
    archive["fragments"][0]["xml"] = base64.b64encode(
        b'<!DOCTYPE w:pict [<!ENTITY canary "ARCHIVE_DTD_CANARY">]>'
        b'<w:pict xmlns:w="http://schemas.openxmlformats.org/wordprocessingml/2006/main">'
        b'&canary;</w:pict>'
    ).decode("ascii")

    with pytest.raises(RestorationError, match="XML"):
        DocxProcessor().restore(
            anonymized, restored, result.mapping, package_archive=archive
        )

    assert not restored.exists()


@pytest.mark.parametrize(
    "payload",
    [
        b'<w:object xmlns:w="http://schemas.openxmlformats.org/wordprocessingml/2006/main"/>',
        b'<o:OLEObject xmlns:o="urn:schemas-microsoft-com:office:office" '
        b'xmlns:r="http://schemas.openxmlformats.org/officeDocument/2006/relationships" '
        b'r:id="rId1"/>',
        b'<a:hlinkClick xmlns:a="http://schemas.openxmlformats.org/drawingml/2006/main" '
        b'xmlns:r="http://schemas.openxmlformats.org/officeDocument/2006/relationships" '
        b'r:id="rId1" tooltip="https://example.invalid/payload"/>',
    ],
    ids=["active-object", "embedded-ole", "external-url"],
)
def test_docx_restore_rejects_active_or_external_xml_in_graphics_archive(
    tmp_path, payload
):
    source = tmp_path / "source.docx"
    image = tmp_path / "synthetic.png"
    anonymized = tmp_path / "anonymized.docx"
    restored = tmp_path / "restored.docx"
    image.write_bytes(base64.b64decode(
        "iVBORw0KGgoAAAANSUhEUgAAAAEAAAABCAQAAAC1HAwCAAAAC0lEQVR42mNk+A8AAQUBAScY42YAAAAASUVORK5CYII="
    ))
    document = Document()
    document.add_paragraph("Обычный текст")
    document.add_picture(str(image), width=Inches(1))
    document.save(source)
    result = DocxProcessor().anonymize(source, anonymized)
    archive = copy.deepcopy(result.package_archive)
    fragment = etree.fromstring(base64.b64decode(archive["fragments"][0]["xml"]))
    fragment.append(etree.fromstring(payload))
    archive["fragments"][0]["xml"] = base64.b64encode(
        etree.tostring(fragment)
    ).decode("ascii")

    with pytest.raises(RestorationError, match="XML"):
        DocxProcessor().restore(
            anonymized, restored, result.mapping, package_archive=archive
        )

    assert not restored.exists()


def test_docx_deleted_graphics_placeholder_does_not_restore_dangling_relationship(tmp_path):
    source = tmp_path / "source.docx"
    image = tmp_path / "synthetic.png"
    anonymized = tmp_path / "anonymized.docx"
    edited = tmp_path / "edited.docx"
    restored = tmp_path / "restored.docx"
    image.write_bytes(base64.b64decode(
        "iVBORw0KGgoAAAANSUhEUgAAAAEAAAABCAQAAAC1HAwCAAAAC0lEQVR42mNk+A8AAQUBAScY42YAAAAASUVORK5CYII="
    ))
    document = Document()
    document.add_paragraph("Обычный текст")
    document.add_picture(str(image), width=Inches(1))
    document.save(source)
    result = DocxProcessor().anonymize(source, anonymized)
    changed = Document(anonymized)
    for paragraph in changed.paragraphs:
        if "[СХЕМА-" in paragraph.text:
            paragraph.clear()
            paragraph.add_run("Схема удалена редактором")
    changed.save(edited)

    DocxProcessor().restore(
        edited, restored, result.mapping, package_archive=result.package_archive
    )

    with ZipFile(restored) as package:
        assert not any(name.startswith("word/media/") for name in package.namelist())
        relationships = etree.fromstring(package.read("word/_rels/document.xml.rels"))
        assert not relationships.xpath(
            "//*[local-name()='Relationship' and contains(@Type, '/image')]"
        )
    Document(restored)


def test_docx_deleting_one_of_two_graphics_restores_only_retained_placeholder(tmp_path):
    source = tmp_path / "source.docx"
    first_image = tmp_path / "first.png"
    second_image = tmp_path / "second.png"
    anonymized = tmp_path / "anonymized.docx"
    edited = tmp_path / "edited.docx"
    restored = tmp_path / "restored.docx"
    first_bytes = base64.b64decode(
        "iVBORw0KGgoAAAANSUhEUgAAAAEAAAABCAQAAAC1HAwCAAAAC0lEQVR42mNk+A8AAQUBAScY42YAAAAASUVORK5CYII="
    )
    second_bytes = base64.b64decode(
        "iVBORw0KGgoAAAANSUhEUgAAAAEAAAABCAIAAACQd1PeAAAADUlEQVR42mP8z8BQDwAFgAH/q640GAAAAABJRU5ErkJggg=="
    )
    first_image.write_bytes(first_bytes)
    second_image.write_bytes(second_bytes)
    document = Document()
    document.add_paragraph("Обычный текст")
    document.add_picture(str(first_image), width=Inches(1))
    document.add_picture(str(second_image), width=Inches(1))
    document.save(source)
    result = DocxProcessor().anonymize(source, anonymized)
    assert len(result.package_archive["fragments"]) == 2
    changed = Document(anonymized)
    for paragraph in changed.paragraphs:
        if "[СХЕМА-2]" in paragraph.text:
            paragraph.clear()
            paragraph.add_run("Вторая схема удалена")
    changed.save(edited)

    DocxProcessor().restore(
        edited, restored, result.mapping, package_archive=result.package_archive
    )

    with ZipFile(restored) as package:
        restored_media = [
            package.read(name)
            for name in package.namelist()
            if name.startswith("word/media/")
        ]
        relationships = etree.fromstring(package.read("word/_rels/document.xml.rels"))
        document_xml = etree.fromstring(package.read("word/document.xml"))
    assert restored_media == [first_bytes]
    assert len(relationships.xpath(
        "//*[local-name()='Relationship' and contains(@Type, '/image')]"
    )) == 1
    assert len(document_xml.xpath(".//w:drawing", namespaces={
        "w": "http://schemas.openxmlformats.org/wordprocessingml/2006/main"
    })) == 1
    Document(restored)


def test_docx_graphics_placeholder_does_not_collide_with_existing_text(tmp_path):
    source = tmp_path / "source.docx"
    image = tmp_path / "synthetic.png"
    anonymized = tmp_path / "anonymized.docx"
    restored = tmp_path / "restored.docx"
    image.write_bytes(base64.b64decode(
        "iVBORw0KGgoAAAANSUhEUgAAAAEAAAABCAQAAAC1HAwCAAAAC0lEQVR42mNk+A8AAQUBAScY42YAAAAASUVORK5CYII="
    ))
    document = Document()
    document.add_paragraph("Пользовательский текст [СХЕМА-1]")
    document.add_picture(str(image), width=Inches(1))
    document.save(source)

    result = DocxProcessor().anonymize(source, anonymized)
    assert result.package_archive["fragments"][0]["marker"] == "[СХЕМА-2]"
    DocxProcessor().restore(
        anonymized, restored, result.mapping, package_archive=result.package_archive
    )

    restored_document = Document(restored)
    assert restored_document.paragraphs[0].text == "Пользовательский текст [СХЕМА-1]"
    assert len(restored_document.inline_shapes) == 1


def test_docx_anonymize_preserves_byte_identical_unchanged_xml_parts(tmp_path):
    source = tmp_path / "source-byte-preserving.docx"
    output = tmp_path / "anonymized-byte-preserving.docx"
    document = Document()
    document.add_paragraph("Иванов Сергей Петрович")
    document.save(source)

    expected = {}

    def alternate_xml_serialization(payload):
        declaration_end = payload.find(b"?>")
        assert declaration_end >= 0
        body = payload[declaration_end + 2 :].lstrip(b"\r\n")
        changed = b'<?xml version="1.0" encoding="UTF-8" standalone="yes"?>\r\n' + body
        expected["payload"] = changed
        return changed

    _rewrite_docx_parts(
        source,
        {"customXml/item1.xml": alternate_xml_serialization},
    )

    DocxProcessor().anonymize(source, output)

    with ZipFile(output) as package:
        actual = package.read("customXml/item1.xml")
    assert actual == expected["payload"]


def test_docx_split_runs_tables_headers_and_restore_unknown(tmp_path):
    src = tmp_path / "source.docx"
    anon = tmp_path / "anon.docx"
    restored = tmp_path / "restored.docx"
    doc = Document()
    p = doc.add_paragraph(); p.add_run("Иванов ").bold = True; p.add_run("Сергей ").italic = True; p.add_run("Петрович tail").underline = True
    outer = doc.add_table(1, 1); cell = outer.cell(0, 0); nested = cell.add_table(1, 1); q = nested.cell(0, 0).paragraphs[0]; q.add_run("s.iva"); q.add_run("nov@почта.ru")
    h = doc.sections[0].header.paragraphs[0]; h.add_run("+7 999 "); h.add_run("123-45-67")
    f = doc.sections[0].footer.paragraphs[0]; f.add_run("ООО «Аль"); f.add_run("фа»")
    doc.save(src)

    result = DocxProcessor().anonymize(src, anon)
    assert src.read_bytes() != anon.read_bytes()
    out = Document(anon)
    assert out.paragraphs[0].text == "[ФИО-1] tail"
    assert out.paragraphs[0].runs[0].bold is True
    assert out.paragraphs[0].runs[-1].underline is True
    assert "[ПОЧТА-1]" in outer_text(Document(anon))
    assert "[ТЕЛЕФОН-1]" in Document(anon).sections[0].header.paragraphs[0].text
    assert "[ОРГАНИЗАЦИЯ-1]" in Document(anon).sections[0].footer.paragraphs[0].text

    edited = Document(anon); edited.add_paragraph("неизвестно [ФИО-999]"); edited.save(anon)
    unknown = DocxProcessor().restore(anon, restored, result.mapping)
    restored_doc = Document(restored)
    assert restored_doc.paragraphs[0].text == "Иванов Сергей Петрович tail"
    assert "[ФИО-999]" in restored_doc.paragraphs[-1].text
    assert unknown == ["[ФИО-999]"]


def test_docx_restore_preserves_word_generated_table_row_height(tmp_path):
    source = tmp_path / "source.docx"
    anonymized = tmp_path / "anonymized.docx"
    edited = tmp_path / "edited.docx"
    restored = tmp_path / "restored.docx"
    document = Document()
    document.add_paragraph("Иванов Сергей Петрович")
    document.add_table(rows=1, cols=1).cell(0, 0).text = "Ячейка"
    document.save(source)
    result = DocxProcessor().anonymize(source, anonymized)

    word_saved = Document(anonymized)
    row_properties = OxmlElement("w:trPr")
    row_height = OxmlElement("w:trHeight")
    row_height.set(qn("w:val"), "497")
    row_properties.append(row_height)
    word_saved.tables[0].rows[0]._tr.insert(0, row_properties)
    word_saved.save(edited)

    DocxProcessor().restore(edited, restored, result.mapping)

    restored_row = Document(restored).tables[0].rows[0]._tr
    assert restored_row.xpath("./w:trPr/w:trHeight/@w:val") == ["497"]


def outer_text(doc):
    values = []
    def walk(table):
        for row in table.rows:
            for cell in row.cells:
                values.extend(p.text for p in cell.paragraphs)
                for nested in cell.tables: walk(nested)
    for table in doc.tables: walk(table)
    return "\n".join(values)


def test_docx_sanitizes_metadata_deletions_external_relationships_and_flags_unsupported(tmp_path):
    src, out = tmp_path / "dirty.docx", tmp_path / "clean.docx"
    doc = Document(); doc.add_paragraph("Иванов Сергей Петрович")
    doc.core_properties.author = "CANARY_AUTHOR"
    p = doc.paragraphs[0]._p
    deletion = OxmlElement("w:del"); run = OxmlElement("w:r"); txt = OxmlElement("w:delText"); txt.text = "CANARY_DELETED"; run.append(txt); deletion.append(run); p.append(deletion)
    doc.part.rels.add_relationship("http://schemas.openxmlformats.org/officeDocument/2006/relationships/hyperlink", "https://CANARY_EXTERNAL.invalid", "rExternal", is_external=True)
    doc.save(src)
    report = DocxProcessor().anonymize(src, out)
    xml = _all_xml(out)
    for canary in (b"CANARY_AUTHOR", b"CANARY_DELETED", b"CANARY_EXTERNAL"):
        assert canary not in xml
    Document(out)
    assert report.ready


def test_docx_physically_removes_custom_properties_part_and_package_references(tmp_path):
    source = tmp_path / "custom-properties.docx"
    output = tmp_path / "custom-properties.anonymized.docx"
    document = Document()
    document.add_paragraph("Иванов Сергей Петрович")
    document.save(source)

    content_types_ns = "http://schemas.openxmlformats.org/package/2006/content-types"
    relationships_ns = "http://schemas.openxmlformats.org/package/2006/relationships"
    custom_relationship = (
        "http://schemas.openxmlformats.org/officeDocument/2006/relationships/custom-properties"
    )

    def add_content_type(payload):
        root = etree.fromstring(payload)
        override = etree.SubElement(root, f"{{{content_types_ns}}}Override")
        override.set("PartName", "/docProps/custom.xml")
        override.set(
            "ContentType",
            "application/vnd.openxmlformats-officedocument.custom-properties+xml",
        )
        return etree.tostring(root, xml_declaration=True, encoding="UTF-8", standalone=True)

    def add_relationship(payload):
        root = etree.fromstring(payload)
        relationship = etree.SubElement(root, f"{{{relationships_ns}}}Relationship")
        relationship.set("Id", "rId99")
        relationship.set("Type", custom_relationship)
        relationship.set("Target", "docProps/custom.xml")
        return etree.tostring(root, xml_declaration=True, encoding="UTF-8", standalone=True)

    _rewrite_docx_parts(
        source,
        {
            "[Content_Types].xml": add_content_type,
            "_rels/.rels": add_relationship,
        },
    )
    replacement = source.with_suffix(".with-custom.docx")
    with ZipFile(source) as incoming, ZipFile(replacement, "w", ZIP_DEFLATED) as outgoing:
        for info in incoming.infolist():
            outgoing.writestr(info, incoming.read(info.filename))
        outgoing.writestr(
            "docProps/custom.xml",
            b'<?xml version="1.0" encoding="UTF-8" standalone="yes"?>'
            b'<Properties xmlns="http://schemas.openxmlformats.org/officeDocument/2006/custom-properties" '
            b'xmlns:vt="http://schemas.openxmlformats.org/officeDocument/2006/docPropsVTypes">'
            b'<property fmtid="{D5CDD505-2E9C-101B-9397-08002B2CF9AE}" pid="2" name="Secret">'
            b'<vt:lpwstr>CUSTOM_PROPERTY_CANARY</vt:lpwstr></property></Properties>',
        )
    replacement.replace(source)

    report = DocxProcessor().anonymize(source, output)

    with ZipFile(output) as package:
        names = set(package.namelist())
        relationships = etree.fromstring(package.read("_rels/.rels"))
        content_types = etree.fromstring(package.read("[Content_Types].xml"))
    assert "docProps/custom.xml" not in names
    assert not relationships.xpath(
        "//*[local-name()='Relationship'][@Type=$type or @Target='docProps/custom.xml']",
        type=custom_relationship,
    )
    assert not content_types.xpath(
        "//*[local-name()='Override'][@PartName='/docProps/custom.xml']"
    )
    assert report.ready
    Document(output)


def test_docx_rejects_doctype_without_publishing_output(tmp_path):
    source = tmp_path / "doctype.docx"
    output = tmp_path / "doctype.anonymized.docx"
    document = Document()
    document.add_paragraph("Иванов Сергей Петрович")
    document.save(source)

    def add_doctype(payload):
        declaration_end = payload.index(b"?>") + 2
        return (
            payload[:declaration_end]
            + b'<!DOCTYPE Types [<!ENTITY canary "DOCTYPE_SECRET_CANARY">]>'
            + payload[declaration_end:]
        )

    _rewrite_docx_parts(source, {"[Content_Types].xml": add_doctype})

    with pytest.raises(ValueError, match="DTD"):
        DocxProcessor().anonymize(source, output)

    assert not output.exists()


def test_docx_removes_metadata_comments_without_failing(tmp_path):
    source = tmp_path / "metadata-comment.docx"
    output = tmp_path / "metadata-comment.anonymized.docx"
    document = Document()
    document.add_paragraph("Иванов Сергей Петрович")
    document.save(source)

    def add_comment(payload):
        root = etree.fromstring(payload)
        root.insert(0, etree.Comment("METADATA_COMMENT_SECRET_CANARY"))
        return etree.tostring(root, xml_declaration=True, encoding="UTF-8", standalone=True)

    _rewrite_docx_parts(source, {"docProps/core.xml": add_comment})

    report = DocxProcessor().anonymize(source, output)

    assert report.ready
    assert b"METADATA_COMMENT_SECRET_CANARY" not in _all_xml(output)
    Document(output)


def test_docx_removes_xml_comments_from_non_metadata_parts(tmp_path):
    source = tmp_path / "styles-comment.docx"
    output = tmp_path / "styles-comment.anonymized.docx"
    document = Document()
    document.add_paragraph("Иванов Сергей Петрович")
    document.save(source)

    def add_comment(payload):
        root = etree.fromstring(payload)
        root.insert(0, etree.Comment("STYLES_COMMENT_SECRET_CANARY"))
        return etree.tostring(root, xml_declaration=True, encoding="UTF-8", standalone=True)

    _rewrite_docx_parts(source, {"word/styles.xml": add_comment})

    report = DocxProcessor().anonymize(source, output)

    assert report.ready
    assert b"STYLES_COMMENT_SECRET_CANARY" not in _all_xml(output)
    Document(output)


def test_docx_external_hyperlink_is_unwrapped_without_dangling_relationship_reference(tmp_path):
    src, out = tmp_path / "linked.docx", tmp_path / "clean.docx"
    doc = Document()
    paragraph = doc.add_paragraph("Ссылка: ")
    rel_id = doc.part.relate_to(
        "https://EXTERNAL_RELATIONSHIP_CANARY.invalid/path",
        RELATIONSHIP_TYPE.HYPERLINK,
        is_external=True,
    )
    hyperlink = OxmlElement("w:hyperlink")
    hyperlink.set(qn("r:id"), rel_id)
    run = OxmlElement("w:r")
    text = OxmlElement("w:t")
    text.text = "видимый текст"
    run.append(text)
    hyperlink.append(run)
    paragraph._p.append(hyperlink)
    doc.save(src)

    report = DocxProcessor().anonymize(src, out)

    with ZipFile(out) as package:
        relationships = etree.fromstring(package.read("word/_rels/document.xml.rels"))
        relationship_ids = {relationship.get("Id") for relationship in relationships}
        document_xml = etree.fromstring(package.read("word/document.xml"))
        referenced_ids = set(document_xml.xpath(
            "//@r:id",
            namespaces={"r": "http://schemas.openxmlformats.org/officeDocument/2006/relationships"},
        ))
    assert referenced_ids <= relationship_ids
    assert "EXTERNAL_RELATIONSHIP_CANARY" not in _all_xml(out).decode("utf-8")
    assert "видимый текст" in "".join(document_xml.itertext())
    assert report.ready


def test_docx_preserves_context_valid_dynamic_formatting_values(tmp_path):
    src, out = tmp_path / "formatted.docx", tmp_path / "formatted.anonymized.docx"
    doc = Document()
    section = doc.sections[0]
    section.left_margin = Inches(0.7)
    section.right_margin = Inches(0.8)
    paragraph = doc.add_paragraph("Иванов Сергей Петрович")
    paragraph.alignment = WD_ALIGN_PARAGRAPH.RIGHT
    run = paragraph.runs[0]
    run.font.name = "Arial"
    run.font.size = Pt(13)
    run.font.color.rgb = RGBColor(0x60, 0x5E, 0x5C)
    table = doc.add_table(rows=1, cols=1)
    table.autofit = False
    table.cell(0, 0).width = Inches(2.37)
    table.cell(0, 0).vertical_alignment = WD_CELL_VERTICAL_ALIGNMENT.BOTTOM
    table.cell(0, 0).text = "ячейка"
    doc.save(src)

    DocxProcessor().anonymize(src, out)
    reopened = Document(out)

    assert reopened.sections[0].left_margin == Inches(0.7)
    assert reopened.sections[0].right_margin == Inches(0.8)
    assert reopened.paragraphs[0].alignment == WD_ALIGN_PARAGRAPH.RIGHT
    assert reopened.paragraphs[0].runs[0].font.size == Pt(13)
    assert reopened.paragraphs[0].runs[0].font.color.rgb == RGBColor(0x60, 0x5E, 0x5C)
    assert reopened.tables[0].cell(0, 0).vertical_alignment == WD_CELL_VERTICAL_ALIGNMENT.BOTTOM


def test_docx_preserves_required_dynamic_attributes_across_word_parts(tmp_path):
    src, out = tmp_path / "required.docx", tmp_path / "required.anonymized.docx"
    doc = Document()
    run = doc.add_paragraph("Иванов Сергей Петрович").add_run(" без жирного")
    run.bold = False
    doc.add_table(rows=1, cols=3).cell(0, 0).merge(doc.tables[0].cell(0, 2)).text = "ячейка"
    doc.save(src)

    def mutate(name, changes):
        def apply(payload):
            root = etree.fromstring(payload)
            namespaces = {"w": "http://schemas.openxmlformats.org/wordprocessingml/2006/main",
                          "a": "http://schemas.openxmlformats.org/drawingml/2006/main"}
            for xpath, attribute, value in changes:
                nodes = root.xpath(xpath, namespaces=namespaces)
                assert nodes, (name, xpath)
                nodes[0].set(attribute, value)
            return etree.tostring(root, xml_declaration=True, encoding="UTF-8", standalone=True)
        return apply

    def add_web_div(payload):
        root = etree.fromstring(payload)
        divs = etree.SubElement(root, qn("w:divs"))
        div = etree.SubElement(divs, qn("w:div"))
        div.set(qn("w:id"), "413207870")
        for local, value in (("bodyDiv", "1"), ("marLeft", "0")):
            node = etree.SubElement(div, qn(f"w:{local}"))
            node.set(qn("w:val"), value)
        return etree.tostring(root, xml_declaration=True, encoding="UTF-8", standalone=True)

    def add_style_dynamics(payload):
        root = etree.fromstring(payload)
        ns = {"w": "http://schemas.openxmlformats.org/wordprocessingml/2006/main"}
        style = root.xpath("//w:style[1]", namespaces=ns)[0]
        ppr = style.find(qn("w:pPr"))
        if ppr is None:
            ppr = etree.SubElement(style, qn("w:pPr"))
        num_pr = etree.SubElement(ppr, qn("w:numPr"))
        for local, value in (("ilvl", "7"), ("numId", "41")):
            node = etree.SubElement(num_pr, qn(f"w:{local}"))
            node.set(qn("w:val"), value)
        widow = etree.SubElement(ppr, qn("w:widowControl"))
        widow.set(qn("w:val"), "0")
        rpr = style.find(qn("w:rPr"))
        if rpr is None:
            rpr = etree.SubElement(style, qn("w:rPr"))
        kern = etree.SubElement(rpr, qn("w:kern"))
        kern.set(qn("w:val"), "26")
        layout = etree.SubElement(rpr, qn("w:eastAsianLayout"))
        layout.set(qn("w:id"), "128")
        return etree.tostring(root, xml_declaration=True, encoding="UTF-8", standalone=True)

    _rewrite_docx_parts(src, {
        "word/document.xml": mutate("document", [
            ("//w:pgSz[1]", qn("w:orient"), "landscape"),
        ]),
        "word/numbering.xml": mutate("numbering", [
            ("//w:abstractNum[1]", qn("w:abstractNumId"), "36"),
            ("//w:abstractNumId[1]", qn("w:val"), "36"),
        ]),
        "word/theme/theme1.xml": mutate("theme", [
            ("//a:theme", "name", "Office-Design"),
            ("//a:latin[1]", "typeface", "Arial"),
            ("//a:ln[1]", "w", "6350"),
            ("//a:srgbClr[1]", "val", "FFFFFF"),
        ]),
        "word/styles.xml": add_style_dynamics,
        "word/webSettings.xml": add_web_div,
    })

    DocxProcessor().anonymize(src, out)

    with ZipFile(out) as package:
        document = etree.fromstring(package.read("word/document.xml"))
        numbering = etree.fromstring(package.read("word/numbering.xml"))
        styles = etree.fromstring(package.read("word/styles.xml"))
        theme = etree.fromstring(package.read("word/theme/theme1.xml"))
        web_settings = etree.fromstring(package.read("word/webSettings.xml"))
    ns = {"w": "http://schemas.openxmlformats.org/wordprocessingml/2006/main",
          "a": "http://schemas.openxmlformats.org/drawingml/2006/main"}
    assert document.xpath("//w:gridSpan/@w:val", namespaces=ns) == ["3"]
    assert "0" in document.xpath("//w:b/@w:val", namespaces=ns)
    assert document.xpath("//w:pgSz[1]/@w:orient", namespaces=ns) == ["landscape"]
    assert numbering.xpath("//w:abstractNum[1]/@w:abstractNumId", namespaces=ns) == ["36"]
    assert numbering.xpath("(//w:abstractNumId)[1]/@w:val", namespaces=ns) == ["36"]
    assert theme.xpath("//a:theme/@name", namespaces=ns) == ["Office Theme"]
    assert theme.xpath("(//a:latin)[1]/@typeface", namespaces=ns) == ["Arial"]
    assert theme.xpath("(//a:ln)[1]/@w", namespaces=ns) == ["6350"]
    assert theme.xpath("(//a:srgbClr)[1]/@val", namespaces=ns) == ["FFFFFF"]
    assert "7" in styles.xpath("//w:ilvl/@w:val", namespaces=ns)
    assert "41" in styles.xpath("//w:numId/@w:val", namespaces=ns)
    assert "26" in styles.xpath("//w:kern/@w:val", namespaces=ns)
    assert "128" in styles.xpath("//w:eastAsianLayout/@w:id", namespaces=ns)
    assert web_settings.xpath("//w:bodyDiv/@w:val", namespaces=ns) == ["1"]
    assert web_settings.xpath("//w:div[1]/@w:id", namespaces=ns) == ["413207870"]
    assert web_settings.xpath("//w:marLeft[1]/@w:val", namespaces=ns) == ["0"]


def test_atomic_failure_keeps_previous_output(tmp_path, monkeypatch):
    src, out = tmp_path / "src.docx", tmp_path / "out.docx"
    d = Document(); d.add_paragraph("Иванов Сергей Петрович"); d.save(src)
    out.write_bytes(b"previous")
    monkeypatch.setattr("local_anonymizer.documents.os.replace", lambda *a: (_ for _ in ()).throw(OSError("fail")))
    with pytest.raises(OSError):
        DocxProcessor().anonymize(src, out)
    assert out.read_bytes() == b"previous"


def test_docx_xml_comments_are_removed_without_breaking_anonymization(tmp_path):
    clean = tmp_path / "clean.docx"
    source = tmp_path / "commented.docx"
    output = tmp_path / "anonymized.docx"
    document = Document()
    document.add_paragraph("Иванов Сергей Петрович")
    document.save(clean)

    with ZipFile(clean) as incoming, ZipFile(source, "w") as outgoing:
        for info in incoming.infolist():
            payload = incoming.read(info.filename)
            if info.filename == "word/document.xml":
                payload = payload.replace(b"<w:body>", b"<!--COMMENT_CANARY--><w:body>", 1)
            outgoing.writestr(info, payload)

    result = DocxProcessor().anonymize(source, output)

    assert result.ready is True
    assert not result.warnings
    assert b"COMMENT_CANARY" not in _all_xml(output)
    assert "[ФИО-1]" in "".join(node.text or "" for node in Document(output).element.xpath(".//w:t"))


def test_docx_hyperlink_visible_text_is_anonymized_before_ready(tmp_path):
    src, out = tmp_path / "hyperlink.docx", tmp_path / "hyperlink.anonymized.docx"
    doc = Document(); paragraph = doc.add_paragraph("Контакт: ")
    hyperlink = OxmlElement("w:hyperlink")
    run = OxmlElement("w:r"); text = OxmlElement("w:t"); text.text = "Иванов Сергей Петрович"; run.append(text); hyperlink.append(run)
    paragraph._p.append(hyperlink); doc.save(src)
    result = DocxProcessor().anonymize(src, out)
    assert result.ready
    assert "Иванов Сергей Петрович".encode("utf-8") not in _all_xml(out)
    assert "[ФИО-1]" in "".join(node.text or "" for node in Document(out).element.xpath(".//w:t"))


def test_docx_publication_strips_simple_and_complex_field_instructions_package_wide(tmp_path):
    src = tmp_path / "fields.docx"
    anonymized = tmp_path / "fields.anonymized.docx"
    restored = tmp_path / "fields.restored.docx"
    doc = Document()

    simple = OxmlElement("w:fldSimple")
    simple.set(qn("w:instr"), ' HYPERLINK "https://SIMPLE_FIELD_CANARY.invalid" ')
    simple_run = OxmlElement("w:r")
    simple_text = OxmlElement("w:t")
    simple_text.text = "safe simple result"
    simple_run.append(simple_text)
    simple.append(simple_run)
    doc.add_paragraph()._p.append(simple)

    table_simple = OxmlElement("w:fldSimple")
    table_simple.set(qn("w:instr"), ' HYPERLINK "https://TABLE_FIELD_CANARY.invalid" ')
    table_run = OxmlElement("w:r")
    table_text = OxmlElement("w:t")
    table_text.text = "safe table result"
    table_run.append(table_text)
    table_simple.append(table_run)
    doc.add_table(1, 1).cell(0, 0).paragraphs[0]._p.append(table_simple)

    def add_complex_field(paragraph, first_canary, second_canary, result):
        for field_type in ("begin",):
            run = OxmlElement("w:r")
            field = OxmlElement("w:fldChar")
            field.set(qn("w:fldCharType"), field_type)
            run.append(field)
            paragraph._p.append(run)
        for instruction in (first_canary, second_canary):
            run = OxmlElement("w:r")
            text = OxmlElement("w:instrText")
            text.text = instruction
            run.append(text)
            paragraph._p.append(run)
        for field_type in ("separate",):
            run = OxmlElement("w:r")
            field = OxmlElement("w:fldChar")
            field.set(qn("w:fldCharType"), field_type)
            run.append(field)
            paragraph._p.append(run)
        paragraph.add_run(result)
        run = OxmlElement("w:r")
        field = OxmlElement("w:fldChar")
        field.set(qn("w:fldCharType"), "end")
        run.append(field)
        paragraph._p.append(run)

    add_complex_field(
        doc.sections[0].header.paragraphs[0],
        ' HYPERLINK "https://HEADER_SPLIT_',
        'FIELD_CANARY.invalid" ',
        "safe header result",
    )
    add_complex_field(
        doc.sections[0].footer.paragraphs[0],
        ' HYPERLINK "https://FOOTER_SPLIT_',
        'FIELD_CANARY.invalid" ',
        "safe footer result",
    )
    doc.save(src)

    report = DocxProcessor().anonymize(src, anonymized)
    assert report.ready
    assert report.warnings == []
    anonymized_xml = _all_xml(anonymized)
    for canary in (
        b"SIMPLE_FIELD_CANARY",
        b"TABLE_FIELD_CANARY",
        b"HEADER_SPLIT_",
        b"FOOTER_SPLIT_",
    ):
        assert canary not in anonymized_xml
    for active_field_token in (b"fldSimple", b"instrText", b"fldChar", b"HYPERLINK"):
        assert active_field_token not in anonymized_xml
    for visible_result in (
        b"safe simple result",
        b"safe table result",
        b"safe header result",
        b"safe footer result",
    ):
        assert visible_result in anonymized_xml

    assert DocxProcessor().restore(src, restored, {}) == []
    restored_xml = _all_xml(restored)
    assert b"FIELD_CANARY" not in restored_xml
    assert b"fldSimple" not in restored_xml
    assert b"instrText" not in restored_xml
    assert b"fldChar" not in restored_xml
