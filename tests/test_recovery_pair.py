import json
from pathlib import Path
from zipfile import ZipFile

import pytest
from docx import Document

from local_anonymizer.app import AppApi
from local_anonymizer.service import AnonymizerService
from local_anonymizer.storage import MappingError


def _ready_text(service: AnonymizerService, folder: Path, name: str, text: str):
    source = folder / name
    source.write_text(text, encoding="utf-8")
    document_id = service.import_file(source)
    service.anonymize(document_id)
    return document_id


def test_export_pair_creates_clear_names_and_portable_unencrypted_recovery_file(tmp_path):
    service = AnonymizerService(tmp_path / "app")
    document_id = _ready_text(service, tmp_path, "Файл1.txt", "Иванов Сергей Петрович")
    anonymized = tmp_path / "Файл1 — обезличенный.txt"

    pair = service.export_pair(document_id, anonymized)

    recovery = tmp_path / "Файл1 — информация для восстановления.larecovery"
    assert pair == {"anonymized_path": anonymized, "recovery_path": recovery}
    assert anonymized.is_file() and recovery.is_file()
    assert "Иванов Сергей Петрович" not in anonymized.read_text(encoding="utf-8")
    payload = json.loads(recovery.read_text(encoding="utf-8"))
    assert payload["version"] == 2
    assert payload["bundle"]["mapping"]
    assert "Иванов Сергей Петрович" in recovery.read_text(encoding="utf-8")
    assert "envelope" not in payload
    assert len(payload["checksum_sha256"]) == 64


def test_portable_recovery_does_not_need_windows_dpapi_to_restore(tmp_path, monkeypatch):
    service = AnonymizerService(tmp_path / "app")
    document_id = _ready_text(service, tmp_path, "Файл1.txt", "Иванов Сергей Петрович")
    pair = service.export_pair(document_id, tmp_path / "Файл1 — обезличенный.txt")
    service.storage.delete_document(document_id)
    monkeypatch.setattr(
        service.storage,
        "decode_restore_envelope",
        lambda *args, **kwargs: (_ for _ in ()).throw(AssertionError("DPAPI must not be used")),
    )

    restored = tmp_path / "Файл1 — восстановленный.txt"
    service.restore_from_pair(pair["anonymized_path"], restored)

    assert restored.read_text(encoding="utf-8") == "Иванов Сергей Петрович"


def test_portable_recovery_rejects_modified_bundle_by_checksum(tmp_path):
    service = AnonymizerService(tmp_path / "app")
    document_id = _ready_text(service, tmp_path, "Файл1.txt", "Иванов Сергей Петрович")
    pair = service.export_pair(document_id, tmp_path / "Файл1 — обезличенный.txt")
    payload = json.loads(pair["recovery_path"].read_text(encoding="utf-8"))
    marker = next(iter(payload["bundle"]["mapping"]))
    payload["bundle"]["mapping"][marker] = "Подменённое значение"
    pair["recovery_path"].write_text(json.dumps(payload, ensure_ascii=False), encoding="utf-8")

    with pytest.raises(MappingError, match="контрольная сумма"):
        service.restore_from_pair(
            pair["anonymized_path"],
            tmp_path / "tampered.txt",
            recovery_path=pair["recovery_path"],
        )


def test_pair_restores_without_local_library_record_and_is_found_automatically(tmp_path):
    service = AnonymizerService(tmp_path / "app")
    document_id = _ready_text(service, tmp_path, "Файл1.txt", "Иванов Сергей Петрович")
    anonymized = tmp_path / "Файл1 — обезличенный.txt"
    pair = service.export_pair(document_id, anonymized)
    service.storage.delete_document(document_id)

    restored = tmp_path / "Файл1 — восстановленный.txt"
    report = service.restore_from_pair(anonymized, restored)

    assert report["recovery_path"] == str(pair["recovery_path"])
    assert report["unknown_markers"] == []
    assert "Иванов Сергей Петрович" in restored.read_text(encoding="utf-8")


def test_restore_rejects_wrong_or_tampered_recovery_companion(tmp_path):
    service = AnonymizerService(tmp_path / "app")
    first_id = _ready_text(service, tmp_path, "Первый.txt", "Иванов Сергей Петрович")
    second_id = _ready_text(service, tmp_path, "Второй.txt", "Петров Иван Сергеевич")
    first = service.export_pair(first_id, tmp_path / "Первый — обезличенный.txt")
    second = service.export_pair(second_id, tmp_path / "Второй — обезличенный.txt")

    with pytest.raises(MappingError, match="не соответствует"):
        service.restore_from_pair(
            first["anonymized_path"],
            tmp_path / "wrong.txt",
            recovery_path=second["recovery_path"],
        )

    second["recovery_path"].write_text('{"format":"broken"}', encoding="utf-8")
    with pytest.raises(MappingError):
        service.restore_from_pair(
            first["anonymized_path"],
            tmp_path / "tampered.txt",
            recovery_path=second["recovery_path"],
        )


def test_automatic_pair_lookup_uses_embedded_pair_id_not_only_filename(tmp_path):
    service = AnonymizerService(tmp_path / "app")
    document_id = _ready_text(service, tmp_path, "Файл1.txt", "Иванов Сергей Петрович")
    pair = service.export_pair(document_id, tmp_path / "Файл1 — обезличенный.txt")
    renamed = tmp_path / "любое-имя.larecovery"
    pair["recovery_path"].rename(renamed)

    restored = tmp_path / "restored.txt"
    report = service.restore_from_pair(pair["anonymized_path"], restored)

    assert Path(report["recovery_path"]) == renamed
    assert "Иванов Сергей Петрович" in restored.read_text(encoding="utf-8")


class _Dialogs:
    def __init__(self, *responses):
        self.responses = list(responses)
        self.calls = []

    def create_file_dialog(self, *args, **kwargs):
        self.calls.append((args, kwargs))
        return self.responses.pop(0)


def test_app_exports_two_files_then_confirms_pair_before_restore(tmp_path):
    app_root = tmp_path / "app"
    api = AppApi(root=app_root)
    source = tmp_path / "Файл1.txt"
    source.write_text("Иванов Сергей Петрович", encoding="utf-8")
    document_id = api.import_paths([str(source)])["ids"][0]
    api.anonymize_document(document_id, paranoid=True)

    anonymized = tmp_path / "out" / "Файл1 — обезличенный.txt"
    export_dialog = _Dialogs([str(anonymized)])
    api.set_window(export_dialog)
    exported = api.export_anonymized(document_id, confirmation=True)

    recovery = tmp_path / "out" / "Файл1 — информация для восстановления.larecovery"
    assert exported == {
        "ok": True,
        "path": str(anonymized),
        "recovery_path": str(recovery),
    }
    assert export_dialog.calls[0][1]["save_filename"] == "Файл1 — обезличенный.txt"
    assert anonymized.is_file() and recovery.is_file()

    api.delete_document(document_id)
    restored = tmp_path / "Файл1 — восстановленный.txt"
    restore_dialog = _Dialogs([str(anonymized)], [str(restored)])
    api.set_window(restore_dialog)
    matched = api.choose_restore_pair()
    assert matched == {
        "ok": True,
        "anonymized_name": anonymized.name,
        "recovery_name": recovery.name,
    }
    report = api.restore_matched_file()
    assert report == {"ok": True, "path": str(restored), "unknown_markers": []}
    assert restored.read_text(encoding="utf-8") == "Иванов Сергей Петрович"


def test_export_all_contains_one_anonymized_and_one_recovery_file_per_document(tmp_path):
    service = AnonymizerService(tmp_path / "app")
    _ready_text(service, tmp_path, "Первый.txt", "Мария Орлова")
    _ready_text(service, tmp_path, "Второй.txt", "Иван Петров")

    bundle = service.export_all_ready(tmp_path / "Все пары.zip")

    with ZipFile(bundle) as archive:
        names = archive.namelist()
        assert len([name for name in names if "обезличенный" in name]) == 2
        assert len([name for name in names if name.endswith(".larecovery")]) == 2
        recovery_data = b"".join(
            archive.read(name) for name in names if name.endswith(".larecovery")
        )
    assert "Мария Орлова".encode("utf-8") in recovery_data
    assert "Иван Петров".encode("utf-8") in recovery_data


def test_docx_pair_restores_after_edit_and_without_library_record(tmp_path):
    service = AnonymizerService(tmp_path / "app")
    source = tmp_path / "Файл1.docx"
    document = Document()
    document.add_paragraph("Иванов Сергей Петрович")
    document.save(source)
    document_id = service.import_file(source)
    service.anonymize(document_id)
    pair = service.export_pair(
        document_id, tmp_path / "Файл1 — обезличенный.docx", confirmation=True
    )

    edited = Document(pair["anonymized_path"])
    edited.add_paragraph("Новая редакторская строка [ФИО-1]")
    edited.save(pair["anonymized_path"])
    service.storage.delete_document(document_id)

    restored = tmp_path / "Файл1 — восстановленный.docx"
    report = service.restore_from_pair(pair["anonymized_path"], restored)
    restored_text = "\n".join(p.text for p in Document(restored).paragraphs)
    assert "Новая редакторская строка Иванов Сергей Петрович" in restored_text
    assert report["unknown_markers"] == []
