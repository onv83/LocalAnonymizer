"""Synthetic-only regressions for B68 personal-name detection."""
import pytest

from local_anonymizer.engine import Anonymizer


@pytest.mark.parametrize("name", ["Рамиль Сафин", "Марат Каменский", "Тимур Леснов"])
def test_cyrillic_given_name_and_surname(name):
    result = Anonymizer().anonymize(name)
    assert result.text == "[ФИО-1]"
    assert result.mapping == {"[ФИО-1]": name}


@pytest.mark.parametrize("name", ["Рамиль Сафин", "Марат Каменский", "Тимур Леснов"])
def test_reverse_order_reuses_identity(name):
    reverse = " ".join(reversed(name.split()))
    result = Anonymizer().anonymize(f"{name}; {reverse}")
    assert result.text == "[ФИО-1]; [ФИО-1]"
    assert result.mapping == {"[ФИО-1]": name}


@pytest.mark.parametrize("name", ["Р. А. Сафин", "Сафин Р. А.", "Р.  А.  Сафин", "Сафин Р  . А  ."])
def test_unanchored_initials_are_masked(name):
    result = Anonymizer().anonymize(name)
    assert result.text == "[ФИО-1]"
    assert result.mapping == {"[ФИО-1]": name}


@pytest.mark.parametrize("given", "Азат Айдар Айрат Раиль Рафаэль Дамир Эльдар Давид Рустам Ренат Ринат Арсен Арман Эмиль Амина Динара Лейла Алия Гузель Эльвира Ильнур Ильяс Рустем Искандер Наиль Рашид Равиль Булат Альберт Анвар Азамат Мансур Мадина Гульнара".split())
def test_additional_cyrillic_given_names(given):
    source = f"{given} Сафин; Сафин {given}"
    result = Anonymizer().anonymize(source)
    assert result.text == "[ФИО-1]; [ФИО-1]"


@pytest.mark.parametrize("source", [
    "Рамиль Сафин", "Каменский Марат", "Р.  А.  Сафин",
    "Рамиль Сафин; Сафин Р. А.",
])
def test_disabled_people_remain_unchanged(source):
    result = Anonymizer(disabled_categories={"ФИО"}).anonymize(source)
    assert result.text == source
    assert result.mapping == {}


@pytest.mark.parametrize("source", [
    "Новый Рабочий", "Рабочий Кабинет", "Технический Отдел",
    "Марат согласовал документ", "Версия А. Б.",
    "Рамиль\ufdd0Сафин", "Сафин\ufdd0Марат",
])
def test_non_names_and_unrelated_cells_are_not_joined(source):
    result = Anonymizer().anonymize(source)
    assert result.text == source
    assert result.mapping == {}


def test_same_surname_different_names_and_initials_stay_separate():
    source = "Рамиль Сафин; Марат Сафин; Р. А. Сафин; М. Б. Сафин"
    result = Anonymizer().anonymize(source)
    assert result.text == "[ФИО-1]; [ФИО-2]; [ФИО-3]; [ФИО-4]"
    assert set(result.mapping.values()) == set(source.split("; "))


def test_docx_split_runs_and_table_names_restore_exact_package(tmp_path):
    from docx import Document
    from local_anonymizer.documents import DocxProcessor

    source, output, restored = [tmp_path / n for n in ("source.docx", "out.docx", "restored.docx")]
    doc = Document()
    para = doc.add_paragraph()
    para.add_run("Рам").bold = True
    para.add_run("иль ").italic = True
    para.add_run("Сафин")
    table = doc.add_table(rows=2, cols=2)
    table.cell(0, 0).text = "Контакт"
    table.cell(0, 1).text = "Каменский Марат"
    table.cell(1, 0).text = "Подпись"
    p = table.cell(1, 1).paragraphs[0]
    for text in ("Т. ", " Л.  ", "Леснов"):
        p.add_run(text)
    doc.save(source)
    original = source.read_bytes()
    processor = DocxProcessor()
    result = processor.anonymize(source, output)
    out = Document(output)
    assert out.paragraphs[0].text == "[ФИО-1]"
    assert out.paragraphs[0].runs[0].bold is True
    assert out.tables[0].cell(0, 1).text.startswith("[ФИО-")
    assert out.tables[0].cell(1, 1).text.startswith("[ФИО-")
    assert "Каменский" not in out.tables[0].cell(0, 1).text
    assert "Леснов" not in out.tables[0].cell(1, 1).text
    processor.restore(output, restored, result.mapping, package_archive=result.package_archive)
    assert source.read_bytes() == original
    assert restored.read_bytes() == original
