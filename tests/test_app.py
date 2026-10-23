import base64
import html as html_module
import json
from pathlib import Path
from html.parser import HTMLParser
import re
import subprocess
import sys
from types import SimpleNamespace

import pytest

import local_anonymizer.app as app_module
from local_anonymizer.app import AppApi, asset_path, platform_smoke_check, smoke_check
from local_anonymizer.engine import Anonymizer
from docx import Document
from docx.enum.section import WD_SECTION
from docx.enum.text import WD_ALIGN_PARAGRAPH
from docx.oxml import OxmlElement
from docx.oxml.ns import qn
from docx.shared import Inches, Pt, RGBColor


def test_address_inventory_self_test_requires_exact_complete_mapping():
    assert app_module._address_inventory_self_test() is True

    class RecordingDetector:
        def __init__(self):
            self.source = ""

        def anonymize(self, source):
            self.source = source
            return Anonymizer(disabled_categories={"СОКРАЩЕНИЯ"}).anonymize(source)

    recording = RecordingDetector()
    assert app_module._address_inventory_self_test(recording) is True
    assert "ТФО\ufdd0Лабораторный\ufdd0г. Лабораторный, ул. Проверочная 9" in recording.source

    class PartialDetector:
        def anonymize(self, _source):
            return SimpleNamespace(
                text="[ГОРОД-1]\ufdd0[ГОРОД-2]\ufdd0КЦ [ГОРОД-3]\ufdd0[АДРЕС-1]",
                mapping={"[АДРЕС-1]": "partial"},
            )

    assert app_module._address_inventory_self_test(PartialDetector()) is False


def test_complete_postal_address_self_test_requires_every_variant():
    assert app_module._complete_postal_address_self_test() is True

    class RecordingDetector:
        def __init__(self):
            self.source = ""

        def anonymize(self, source):
            self.source = source
            return Anonymizer().anonymize(source)

    recording = RecordingDetector()
    assert app_module._complete_postal_address_self_test(recording) is True
    assert "ул. Испытателей, д. 12, корп. 1, лит. А г. Тестоградск, 123456" in recording.source
    assert "д. 12,\u00a0лит. А" in recording.source
    assert "ул.Испытателей" in recording.source
    assert "ул. испытателей" in recording.source

    class PartialDetector:
        def anonymize(self, _source):
            return SimpleNamespace(text="[АДРЕС-1]", mapping={"[АДРЕС-1]": "partial"})

    assert app_module._complete_postal_address_self_test(PartialDetector()) is False


def test_address_review_parity_self_test_requires_plain_and_structural_consistency():
    assert app_module._address_review_parity_self_test() is True

    class StalePlainDetector:
        def anonymize(self, source):
            if source.startswith("\ufdd0"):
                return Anonymizer().anonymize(source)
            result = Anonymizer().anonymize(source)
            return SimpleNamespace(
                text="Площадка 2 «Демоинфо» - [АДРЕС-1]",
                mapping=result.mapping,
                hint_count=1,
                parts=[{"t": "hint"}],
            )

    assert app_module._address_review_parity_self_test(StalePlainDetector()) is False


def test_contextual_address_self_test_requires_every_complete_variant():
    assert app_module._contextual_address_self_test() is True

    class RecordingDetector:
        def __init__(self):
            self.source = ""

        def anonymize(self, source):
            self.source = source
            return Anonymizer().anonymize(source)

    recording = RecordingDetector()
    assert app_module._contextual_address_self_test(recording) is True
    assert (
        "ул. Контрольная, д.9, 1-ый Лабораторный пр. д.1 стр.3"
        in recording.source
    )

    class PartialDetector:
        def anonymize(self, _source):
            return SimpleNamespace(
                text="[АДРЕС-1] [АДРЕС-2]",
                mapping={
                    "[АДРЕС-1]": "Макетный пер. д.3",
                    "[АДРЕС-2]": "Большая Испытательная, д.46",
                },
            )

    assert app_module._contextual_address_self_test(PartialDetector()) is False


def test_server_inventory_name_self_test_requires_standalone_and_table_context_patterns():
    assert app_module._server_inventory_name_self_test() is True

    class PartialDetector:
        def anonymize(self, _source):
            return SimpleNamespace(
                text="[УЗЕЛ-1]",
                mapping={"[УЗЕЛ-1]": "s1234-qzx01"},
            )

    assert app_module._server_inventory_name_self_test(PartialDetector()) is False


def test_split_server_inventory_name_self_test_requires_one_marker_across_intra_cell_boundary():
    assert app_module._split_server_inventory_name_self_test() is True

    class PartialDetector:
        def anonymize(self, source):
            return SimpleNamespace(text=source, mapping={})

    assert app_module._split_server_inventory_name_self_test(PartialDetector()) is False


def test_underscored_dns_fqdn_self_test_requires_bounded_node_detection():
    assert app_module._underscored_dns_fqdn_self_test() is True

    class PartialDetector:
        def anonymize(self, source):
            return SimpleNamespace(text=source, mapping={})

    assert app_module._underscored_dns_fqdn_self_test(PartialDetector()) is False


def test_domain_inventory_self_test_requires_all_location_columns_and_boundaries():
    assert app_module._domain_inventory_subject_self_test() is True

    class PartialDetector:
        def anonymize(self, source):
            return SimpleNamespace(text=source, mapping={})

    assert app_module._domain_inventory_subject_self_test(PartialDetector()) is False


def test_device_inventory_identifier_self_test_requires_contextual_node_detection():
    assert app_module._device_inventory_identifier_self_test() is True

    class PartialDetector:
        def anonymize(self, source):
            return SimpleNamespace(text=source, mapping={})

    assert app_module._device_inventory_identifier_self_test(PartialDetector()) is False


def test_service_account_self_test_requires_one_bounded_login_marker():
    assert app_module._service_account_self_test() is True

    class PartialDetector:
        def anonymize(self, source):
            return SimpleNamespace(text=source, mapping={}, counts={})

    assert app_module._service_account_self_test(PartialDetector()) is False


def test_single_two_letter_abbreviation_self_test_requires_unconditional_unicode_behavior():
    assert app_module._single_two_letter_abbreviation_self_test() is True

    class PartialDetector:
        def anonymize(self, _source):
            return SimpleNamespace(text=_source, mapping={}, counts={})

    assert app_module._single_two_letter_abbreviation_self_test(PartialDetector()) is False


def test_bundled_ui_and_api_smoke(tmp_path, monkeypatch):
    monkeypatch.setenv("LOCALAPPDATA", str(tmp_path))
    index = asset_path("index-v1.8.18.html")
    assert index.exists()
    html = index.read_text(encoding="utf-8")
    for label in ("Исходный текст", "Обезличенный текст", "Вернуть исходные значения", "Настройки", "Проверка движка"):
        assert label in html
    assert "http://" not in html and "https://" not in html
    assert smoke_check()["ok"] is True
    summary = AppApi().engine_self_test()
    assert summary["ok"] is True
    assert summary["vm_node_ok"] is True
    assert summary["abbreviation_ok"] is True
    assert summary["abbreviation_unicode_ok"] is True
    assert summary["abbreviation_single_two_letter_ok"] is True
    assert summary["domain_toc_ok"] is True
    assert summary["port_ok"] is True
    assert summary["address_inventory_ok"] is True
    assert summary["address_postal_variants_ok"] is True
    assert summary["address_contextual_variants_ok"] is True
    assert summary["compound_policy_abbreviation_ok"] is True
    assert summary["mixed_separator_compound_abbreviation_ok"] is True
    assert summary["inventory_style_abbreviation_ok"] is True
    assert summary["intra_boundary_semantic_precedence_ok"] is True
    assert summary["hyphenated_abbreviation_hint_precedence_ok"] is True
    assert summary["prefixed_lowercase_compound_abbreviation_ok"] is True
    assert summary["dotted_version_compound_abbreviation_ok"] is True
    assert summary["contextual_location_table_ok"] is True
    assert summary["server_inventory_name_ok"] is True
    assert summary["split_server_inventory_name_ok"] is True
    assert summary["semantic_server_capacity_node_ok"] is True
    assert summary["underscored_dns_fqdn_ok"] is True
    assert summary["domain_inventory_row_ok"] is True
    assert summary["device_inventory_identifier_ok"] is True
    assert summary["service_account_ok"] is True
    assert summary["split_service_account_ok"] is True
    assert summary["graphics_roundtrip_ok"] is True
    assert summary["domain_inventory_subject_ok"] is True
    assert "Иванов" not in str(summary)


def test_platform_smoke_checks_secure_storage_and_native_clipboard_without_reading_it():
    class Protector:
        def __bool__(self):
            return False

        def protect(self, data, purpose):
            return b"protected:" + purpose.encode() + b":" + data

        def unprotect(self, envelope, purpose):
            prefix = b"protected:" + purpose.encode() + b":"
            assert envelope.startswith(prefix)
            return envelope[len(prefix):]

    class Clipboard:
        available = True

        def __bool__(self):
            return False

        def read_text(self):
            raise AssertionError("platform smoke must not read user clipboard")

        def write_text(self, text):
            raise AssertionError("platform smoke must not overwrite user clipboard")

    result = platform_smoke_check(
        platform_name="linux",
        protector=Protector(),
        clipboard=Clipboard(),
    )

    assert result == {
        "ok": True,
        "platform": "linux",
        "secure_storage": True,
        "clipboard": True,
    }


def test_platform_smoke_fails_closed_without_secure_storage_or_clipboard():
    class BrokenProtector:
        def protect(self, data, purpose):
            raise RuntimeError("backend unavailable")

        def unprotect(self, envelope, purpose):
            raise AssertionError("protect must fail first")

    result = platform_smoke_check(
        platform_name="darwin",
        protector=BrokenProtector(),
        clipboard=SimpleNamespace(available=False),
    )

    assert result["ok"] is False
    assert result["secure_storage"] is False
    assert result["clipboard"] is False
    assert result["error"] == "RuntimeError: backend unavailable"


def test_profile_json_export_import_roundtrip_restores_exact_settings(tmp_path):
    profile_file = tmp_path / "Общий.json"

    class Window:
        def create_file_dialog(self, *args, **kwargs):
            return [str(profile_file)]

    api = AppApi(tmp_path / "data")
    api.set_window(Window())
    rules = "Переключение => СКРЫТО\nПоведение => ТЕРМИН"
    assert api.save_profile("Общий", rules, "Оставить\nРазрешено", True)["ok"] is True
    assert api.export_profile("Общий")["ok"] is True

    exported = profile_file.read_text(encoding="utf-8")
    assert '"version": 1' in exported
    assert '"Переключение": "СКРЫТО"' in exported
    api.save_profile("Общий", "", "", False)

    result = api.import_profile()
    assert result["ok"] is True
    assert result["name"] == "Общий"
    assert api.load_profile("Общий") == {
        "version": 1,
        "name": "Общий",
        "dictionary": {"Переключение": "СКРЫТО", "Поведение": "ТЕРМИН"},
        "whitelist": ["Оставить", "Разрешено"],
        "paranoid": True,
    }


def test_cancelled_profile_export_does_not_persist_visible_settings(tmp_path):
    class Window:
        def create_file_dialog(self, *args, **kwargs):
            return None

    api = AppApi(tmp_path / "data")
    api.save_profile("Общий", "Старое => ПРАВИЛО", "Старый список", False)
    before = api.load_profile("Общий")
    api.set_window(Window())

    result = api.export_profile("Общий", "Новое => ПРАВИЛО", "Новый список", True)

    assert result == {"ok": False}
    assert api.load_profile("Общий") == before


def test_profile_json_import_reports_invalid_file_without_overwriting_profile(tmp_path):
    profile_file = tmp_path / "broken.json"
    profile_file.write_text('{"version": 1, "name": "Общий", "dictionary": ', encoding="utf-8")

    class Window:
        def create_file_dialog(self, *args, **kwargs):
            return [str(profile_file)]

    api = AppApi(tmp_path / "data")
    api.save_profile("Общий", "Переключение => СКРЫТО", "Оставить", False)
    before = api.load_profile("Общий")
    api.set_window(Window())

    result = api.import_profile()
    assert result == {"ok": False, "error": "Файл профиля содержит некорректный JSON."}
    assert api.load_profile("Общий") == before

    profile_file.write_text(
        json.dumps(
            {
                "version": 1,
                "name": "Общий",
                "dictionary": [],
                "whitelist": [],
                "paranoid": False,
            },
            ensure_ascii=False,
        ),
        encoding="utf-8",
    )
    result = api.import_profile()
    assert result["ok"] is False
    assert result["error"].startswith("Профиль не импортирован:")
    assert api.load_profile("Общий") == before


def test_uncertain_fragments_use_two_visible_checkboxes_without_duplicate_heading():
    labels = {
        "index-v1.8.18.html": (
            "Искать спорные места для ручной проверки",
            "Автоматически скрывать спорное без подтверждения",
        ),
        "index-en-v1.8.18.html": (
            "Find uncertain items for manual review",
            "Automatically hide uncertain items without confirmation",
        ),
    }
    forbidden_headings = {
        "index-v1.8.18.html": ("<summary>Дополнительные режимы</summary>", "<legend>Дополнительные режимы</legend>"),
        "index-en-v1.8.18.html": ("<summary>Additional modes</summary>", "<legend>Additional modes</legend>"),
    }
    for html_name, expected_labels in labels.items():
        html = asset_path(html_name).read_text(encoding="utf-8")
        assert html.count('id="uncertainMode"') == 1
        assert '<input id="uncertainMode" type="hidden" value="review">' in html
        assert '<select id="uncertainMode"' not in html
        assert '<input class="category" data-categories="HINTS" id="reviewHints" type="checkbox" checked hidden>' in html
        assert '<input id="reviewHintsMode" type="checkbox" checked>' in html
        assert '<input id="paranoid" type="checkbox">' in html
        for label in expected_labels:
            assert label in html
        for heading in forbidden_headings[html_name]:
            assert heading not in html

    for js_name in ("app-v1.8.18.js", "app-en-v1.8.18.js"):
        javascript = asset_path(js_name).read_text(encoding="utf-8")
        assert "const UNCERTAIN_MODES=new Set(['off','review','auto'])" in javascript
        assert "function applyUncertainMode(mode)" in javascript
        assert "$('#reviewHints').checked=normalized!=='off'" in javascript
        assert "$('#reviewHintsMode').checked=normalized==='review'" in javascript
        assert "$('#paranoid').checked=normalized==='auto'" in javascript
        assert "p.uncertain_mode" in javascript
        assert "$('#uncertainMode').value" in javascript


def test_uncertain_mode_runtime_maps_all_three_states_and_legacy_profiles():
    for js_name in ("app-v1.8.18.js", "app-en-v1.8.18.js"):
        javascript = asset_path(js_name).read_text(encoding="utf-8")
        runtime = javascript[:javascript.index("async function applyCurrentDesign")]
        harness = f"""
const elements={{
  '#uncertainMode':{{value:'review'}},
  '#reviewHints':{{checked:true}},
  '#reviewHintsMode':{{checked:true}},
  '#paranoid':{{checked:false}}
}};
const document={{querySelector:(selector)=>elements[selector]}};
{runtime}
function snapshot(mode){{applyUncertainMode(mode);return [elements['#uncertainMode'].value,elements['#reviewHints'].checked,elements['#reviewHintsMode'].checked,elements['#paranoid'].checked];}}
const actual={{
  off:snapshot('off'),
  review:snapshot('review'),
  auto:snapshot('auto'),
  legacyFalse:normalizeUncertainMode(undefined,false,true),
  legacyTrue:normalizeUncertainMode(undefined,true,true),
  invalid:normalizeUncertainMode('unexpected',false,true)
}};
console.log(JSON.stringify(actual));
"""
        completed = subprocess.run(
            ["node", "-e", harness], capture_output=True, text=True, check=False
        )
        assert completed.returncode == 0, completed.stderr
        assert json.loads(completed.stdout) == {
            "off": ["off", False, False, False],
            "review": ["review", True, True, False],
            "auto": ["auto", True, False, True],
            "legacyFalse": "review",
            "legacyTrue": "auto",
            "invalid": "review",
        }


def test_uncertain_mode_checkbox_changes_keep_three_state_model_consistent():
    for js_name in ("app-v1.8.18.js", "app-en-v1.8.18.js"):
        javascript = asset_path(js_name).read_text(encoding="utf-8")
        assert "function syncUncertainModeFromCheckboxes()" in javascript
        assert "$('#reviewHintsMode').onchange=async()=>" in javascript
        assert "$('#paranoid').onchange=async()=>" in javascript
        assert "$('#uncertainMode').onchange" not in javascript
        runtime = javascript[:javascript.index("async function applyCurrentDesign")]
        handlers = "\n".join(
            line for line in javascript.splitlines()
            if line.startswith("$('#reviewHintsMode').onchange=") or line.startswith("$('#paranoid').onchange=")
        )
        harness = f"""
const elements={{
  '#uncertainMode':{{value:'review'}},
  '#reviewHints':{{checked:true}},
  '#reviewHintsMode':{{checked:true}},
  '#paranoid':{{checked:false}}
}};
const document={{querySelector:(selector)=>elements[selector]}};
let dirty=0;function markSettingsDirty(){{dirty++;}}
{runtime}
{handlers}
(async()=>{{
  elements['#paranoid'].checked=true;await elements['#paranoid'].onchange();
  const automatic=[elements['#uncertainMode'].value,elements['#reviewHints'].checked,elements['#reviewHintsMode'].checked,elements['#paranoid'].checked];
  elements['#paranoid'].checked=false;await elements['#paranoid'].onchange();
  const off=[elements['#uncertainMode'].value,elements['#reviewHints'].checked,elements['#reviewHintsMode'].checked,elements['#paranoid'].checked];
  elements['#reviewHintsMode'].checked=true;await elements['#reviewHintsMode'].onchange();
  const review=[elements['#uncertainMode'].value,elements['#reviewHints'].checked,elements['#reviewHintsMode'].checked,elements['#paranoid'].checked];
  console.log(JSON.stringify({{automatic,off,review,dirty}}));
}})();
"""
        completed = subprocess.run(["node", "-e", harness], capture_output=True, text=True, check=False)
        assert completed.returncode == 0, completed.stderr
        assert json.loads(completed.stdout) == {
            "automatic": ["auto", True, False, True],
            "off": ["off", False, False, False],
            "review": ["review", True, True, False],
            "dirty": 3,
        }


def test_profile_api_roundtrip_preserves_three_state_uncertain_mode(tmp_path):
    api = AppApi(tmp_path / "data")
    assert api.save_profile("Точный", "", "", False, "off")["ok"] is True
    assert api.load_profile("Точный")["uncertain_mode"] == "off"


def test_settings_and_manual_review_belong_to_their_preview_columns_in_both_languages():
    class LayoutParser(HTMLParser):
        VOID_TAGS = {"area", "base", "br", "col", "embed", "hr", "img", "input", "link", "meta", "param", "source", "track", "wbr"}

        def __init__(self):
            super().__init__()
            self.stack = []
            self.parents = {}
            self.children = {}

        def handle_starttag(self, tag, attrs):
            attrs = dict(attrs)
            node = attrs.get("id") or (
                "columns" if tag == "main" and "columns" in attrs.get("class", "").split() else None
            )
            parent = next((item for item in reversed(self.stack) if item), None)
            if node:
                self.parents[node] = parent
                self.children.setdefault(parent, []).append(node)
            if tag not in self.VOID_TAGS:
                self.stack.append(node)

        def handle_endtag(self, tag):
            if tag not in self.VOID_TAGS:
                self.stack.pop()

        def is_descendant(self, node, ancestor):
            while node in self.parents:
                node = self.parents[node]
                if node == ancestor:
                    return True
            return False

    for filename in ("index-v1.8.18.html", "index-en-v1.8.18.html"):
        parser = LayoutParser()
        parser.feed(asset_path(filename).read_text(encoding="utf-8"))

        assert parser.parents["sourceColumn"] == "columns"
        assert parser.parents["anonymizedColumn"] == "columns"
        assert parser.is_descendant("sourceHeading", "sourceColumn")
        assert parser.is_descendant("source", "sourceColumn")
        assert parser.parents["anonymizationSettings"] == "sourceColumn"
        assert parser.is_descendant("outputHeading", "anonymizedColumn")
        assert parser.is_descendant("anonymizedPreviewSurface", "anonymizedColumn")
        assert parser.parents["manualReview"] == "anonymizedColumn"
        for review_element in (
            "warning", "reviewCurrent", "reviewPrev", "reviewNext", "reviewDecisionActions",
            "hideCandidate", "keepCandidate", "reviewDecision", "reviewFinish", "resultStatus",
        ):
            assert parser.is_descendant(review_element, "manualReview")


def test_settings_use_a_punto_style_navigation_and_editor_panels_in_both_languages():
    for filename in ("index-v1.8.18.html", "index-en-v1.8.18.html"):
        html = asset_path(filename).read_text(encoding="utf-8")
        assert 'class="settings-workbench"' in html
        assert 'class="settings-nav" role="tablist"' in html
        for panel in ("profile", "rules", "whitelist"):
            assert f'id="settingsTab-{panel}"' in html
            assert f'aria-controls="settingsPanel-{panel}"' in html
            assert f'id="settingsPanel-{panel}"' in html
            assert f'aria-labelledby="settingsTab-{panel}"' in html
        assert 'id="rules"' in html
        assert 'id="whitelist"' in html
        assert html.count('class="settings-nav-chevron" aria-hidden="true"') == 3
        assert html.count('<span class="settings-nav-label">') == 3

    for script_name in ("app-v1.8.18.js", "app-en-v1.8.18.js"):
        javascript = asset_path(script_name).read_text(encoding="utf-8")
        assert "function initSettingsTabs()" in javascript
        assert "tab.addEventListener('click'" in javascript
        assert "tab.addEventListener('keydown'" in javascript
        assert "initSettingsTabs();" in javascript


def test_settings_navigation_chevrons_are_right_facing_and_do_not_steal_accessibility_names():
    css = asset_path("new-design-v1.11.css").read_text(encoding="utf-8")
    assert ".settings-nav-chevron" in css
    assert "margin-inline-start:auto" in css
    assert "flex:0 0 auto" in css
    for filename in ("index-v1.8.18.html", "index-en-v1.8.18.html"):
        html = asset_path(filename).read_text(encoding="utf-8")
        assert html.count('aria-hidden="true">›</span>') == 3


def test_mobile_library_controls_collapse_to_one_bounded_column():
    css = asset_path("new-design-v1.11.css").read_text(encoding="utf-8")
    mobile = css[css.index("@media(max-width:800px)") :]
    assert ".source-library>.compact-library{grid-template-columns:minmax(0,1fr)}" in mobile
    assert ".source-library>.compact-library>*{min-width:0;max-width:100%;width:100%}" in mobile


def test_manual_review_uses_one_toolbar_and_hides_redundant_prompts():
    for filename in ("index-v1.8.18.html", "index-en-v1.8.18.html"):
        html = asset_path(filename).read_text(encoding="utf-8")
        toolbar = re.search(
            r'<div class="review-toolbar">(.*?)</div><div class="review-controls-disclosure">',
            html,
        )
        assert toolbar is not None
        controls = toolbar.group(1)
        ordered_ids = ("reviewPrev", "reviewNext", "hideCandidate", "keepCandidate")
        assert [controls.index(f'id="{control_id}"') for control_id in ordered_ids] == sorted(
            controls.index(f'id="{control_id}"') for control_id in ordered_ids
        )
        assert re.search(r'id="hideCandidate"[^>]+title="[^"]+"', controls)
        assert re.search(r'id="keepCandidate"[^>]+title="[^"]+"', controls)
        assert not re.search(r'<button class="primary"[^>]+id="hideCandidate"', controls)
        assert 'class="warning sr-only" id="warning"' in html
        assert 'id="reviewDecision" class="review-decision sr-only"' in html
        assert re.search(r'id="reviewFinish" class="sr-only"', html)

    for script_name in ("app-v1.8.18.js", "app-en-v1.8.18.js"):
        javascript = asset_path(script_name).read_text(encoding="utf-8")
        assert "const warning=$('#warning'),hasWarnings=items.length>0" in javascript
        assert "warning.classList.toggle('sr-only',!hasWarnings)" in javascript
        assert "warning.textContent=hasWarnings?items.join('\\n'):''" in javascript


def test_manual_review_intro_is_an_accessible_on_demand_tooltip_in_both_languages():
    for filename in ("index-v1.8.18.html", "index-en-v1.8.18.html"):
        html = asset_path(filename).read_text(encoding="utf-8")
        heading = re.search(r'<div class="review-heading">(.*?)<strong id="reviewPos"', html, re.S)
        assert heading is not None
        content = heading.group(1)
        assert "<p>" not in content
        assert re.search(
            r'<h2 id="reviewTitle"[^>]+tabindex="0"[^>]+aria-describedby="reviewIntroHelp"'
            r'[^>]+data-tooltip="[^"]+"[^>]+title="[^"]+"',
            content,
        )
        assert '<span class="review-heading-help" aria-hidden="true">?</span>' in content
        assert re.search(r'<span id="reviewIntroHelp" class="sr-only">[^<]+(?:<b>.*?</b>[^<]*)?</span>', content)

    css = asset_path("new-design-v1.11.css").read_text(encoding="utf-8")
    assert ".review-heading h2:focus-visible::after" in css
    assert ".review-heading h2:hover::after" not in css


def test_manual_review_warning_runtime_hides_only_empty_guidance(tmp_path):
    for script_name in ("app-v1.8.18.js", "app-en-v1.8.18.js"):
        javascript = asset_path(script_name).read_text(encoding="utf-8")
        render_warnings = re.search(r"function renderWarnings\(\)\{.*?\n\}", javascript, re.S)
        assert render_warnings is not None
        harness = f"""
const assert = require('assert');
const classes = new Set();
const warning = {{
  textContent: 'seed',
  classList: {{
    toggle(name, enabled) {{ if (enabled) classes.add(name); else classes.delete(name); }},
    contains(name) {{ return classes.has(name); }}
  }}
}};
const $ = selector => {{ assert.strictEqual(selector, '#warning'); return warning; }};
const warningText = item => typeof item === 'object' ? item.message : String(item);
let last = {{warnings: [], hintCount: 7}};
{render_warnings.group(0)}
renderWarnings();
assert.strictEqual(warning.textContent, '');
assert.strictEqual(warning.classList.contains('sr-only'), true);
assert.strictEqual(warning.classList.contains('stop'), false);
last = {{warnings: [{{message: 'Synthetic warning', level: 'warning'}}], hintCount: 7}};
renderWarnings();
assert.strictEqual(warning.textContent, 'Synthetic warning');
assert.strictEqual(warning.classList.contains('sr-only'), false);
assert.strictEqual(warning.classList.contains('stop'), false);
last = {{warnings: [{{message: 'Synthetic blocker', level: 'stop'}}], hintCount: 7}};
renderWarnings();
assert.strictEqual(warning.textContent, 'Synthetic blocker');
assert.strictEqual(warning.classList.contains('sr-only'), false);
assert.strictEqual(warning.classList.contains('stop'), true);
"""
        script = tmp_path / f"{script_name}.warning-runtime.js"
        script.write_text(harness, encoding="utf-8")
        completed = subprocess.run(["node", str(script)], capture_output=True, text=True)
        assert completed.returncode == 0, completed.stderr


def test_settings_have_one_checkbox_per_actual_replacement_class():
    html = asset_path("index-v1.8.18.html").read_text(encoding="utf-8")
    controls = re.findall(
        r'<input class="category" data-categories="([^"]+)"[^>]*>\s*([^<]+)',
        html,
    )
    categories = [category for category, _label in controls if category != "HINTS"]
    assert len(categories) == len(set(categories))
    assert all("," not in category for category in categories)

    expected = {
        "уязвимость", "ПОЧТА", "ТЕЛЕФОН", "АККАУНТ", "IP", "ПОРТ", "MAC", "IMEI", "ССЫЛКА",
        "ДОМЕН", "УЗЕЛ", "DN", "СТРУКТУРА-AD", "ДИСК", "ПУТЬ", "СЗИ",
        "ПРОГРАММА", "СОКРАЩЕНИЯ", "СКРИПТЫ", "ПРОИЗВОДИТЕЛЬ", "ИНН", "ОГРН", "СНИЛС",
        "КАРТА", "КПП", "СЧЁТ", "БИК", "ПАСПОРТ", "VIN", "КАДАСТР",
        "ГОСНОМЕР", "ТАБЕЛЬНЫЙ", "ДОГОВОР", "ДОКУМЕНТ", "ДАТА-РОЖДЕНИЯ",
        "СУММА", "СТАВКА", "ФИО", "ЛОГИН", "ОРГАНИЗАЦИЯ", "СИСТЕМА",
        "ЗАДАЧА", "ГОРОД", "АДРЕС", "ИНДЕКС", "ДОЛЖНОСТЬ", "СПЕЦКАТЕГОРИЯ", "СКРЫТО",
        "CUSTOM",
    }
    assert set(categories) == expected

    assert "IP-адрес → <code>[IP-N]</code>" in html
    assert "Порт → <code>[ПОРТ-N]</code>" in html
    assert "Домен → <code>[ДОМЕН-N]</code>" in html
    assert "Электронная почта → <code>[ПОЧТА-N]</code>" in html
    assert "Ссылка → <code>[ССЫЛКА-N]</code>" in html
    assert "Уязвимость → <code>[уязвимость]</code>" in html
    assert "Дополнительные правила → <code>[МЕТКА-N]</code>" in html
    javascript = asset_path("app-v1.8.18.js").read_text(encoding="utf-8")
    english_javascript = asset_path("app-en-v1.8.18.js").read_text(encoding="utf-8")
    assert "function updateCategoryMarkerExamples()" in javascript
    assert "updateCategoryMarkerExamples();markSettingsDirty()" in javascript
    assert "СОКРАЩЕНИЯ:'СОКР'" in javascript
    assert "ПОРТ:'ПРТ'" in javascript
    assert "\"Сокращения\":['СОКРАЩЕНИЯ']" in javascript
    assert "СОКРАЩЕНИЯ:'СОКР'" in english_javascript
    assert "ПОРТ:'ПРТ'" in english_javascript
    assert "\"Abbreviations\":['СОКРАЩЕНИЯ']" in english_javascript


def test_disabling_one_replacement_class_does_not_disable_or_enable_another():
    source = "Узел 198.51.100.23 доступен как node.example.test."

    no_ip = Anonymizer(disabled_categories={"IP"}).anonymize(source)
    assert "198.51.100.23" in no_ip.text
    assert "[IP-" not in no_ip.text
    assert "[ДОМЕН-" in no_ip.text

    no_domain = Anonymizer(disabled_categories={"ДОМЕН"}).anonymize(source)
    assert "node.example.test" in no_domain.text
    assert "[ДОМЕН-" not in no_domain.text
    assert "[IP-" in no_domain.text


def test_disabled_outer_class_is_not_partially_replaced_by_fallback_detectors():
    source = "Пишите на user@example.test или откройте https://portal.example.test/path."

    result = Anonymizer(disabled_categories={"ПОЧТА", "ССЫЛКА"}).anonymize(source)

    assert result.text == source
    assert result.mapping == {}


def test_interface_contains_current_collapsed_usage_and_restore_guide_in_both_languages():
    localized = {
        "index-v1.8.18.html": (
            "Как пользоваться программой",
            "«Открыть»",
            "перетащите их в окно программы",
            "откроется автоматически",
            "Исходный текст",
            "Обезличенный текст",
            "Ручная проверка спорных мест",
            "Скрыть и запомнить правило",
            "Оставить и добавить в исключения",
            "Применить настройки к файлу",
            "исходный файл не изменяется",
            "Сохранить документ и файл восстановления",
            "Как восстановить исходный документ",
            "1. Выбрать обезличенный файл…",
            "2. Восстановить файл…",
            "Windows, Linux и macOS",
        ),
        "index-en-v1.8.18.html": (
            "How to use the app",
            "“Open”",
            "drag them into the app window",
            "opens automatically",
            "Source text",
            "Anonymized text",
            "Manual review of uncertain items",
            "Hide and remember rule",
            "Keep and add to exceptions",
            "Apply settings to file",
            "source file is not modified",
            "Save document and recovery file",
            "How to restore the original document",
            "1. Select anonymized file…",
            "2. Restore file…",
            "Windows, Linux, and macOS",
        ),
    }

    for asset_name, required in localized.items():
        html = asset_path(asset_name).read_text(encoding="utf-8")
        assert all(text in html for text in required)
        assert '<details class="card usage-guide">' in html
        assert '<details class="card usage-guide" open>' not in html

    ru_html = asset_path("index-v1.8.18.html").read_text(encoding="utf-8")
    en_html = asset_path("index-en-v1.8.18.html").read_text(encoding="utf-8")
    javascript = asset_path("app-v1.8.18.js").read_text(encoding="utf-8")
    assert "Нажмите <b>«Выбрать исходные файлы…»</b>" not in ru_html
    assert "Выберите документ в <b>«Локальной библиотеке»</b>" not in ru_html
    assert "Click <b>“Select source files…”</b>" not in en_html
    assert "Select the document in the <b>“Local library”</b>" not in en_html
    assert "зашифрован средствами Windows" not in ru_html
    assert "профиле Windows" not in ru_html
    assert "guide.open" not in javascript
    assert 'id="stats">Версия 1.8' in ru_html
    assert 'id="stats">Версия 1.7' not in ru_html


def test_gui_loads_default_profile_rules_on_startup():
    javascript = asset_path("app-v1.8.18.js").read_text(encoding="utf-8")

    assert (
        "showLibrary(await pywebview.api.list_documents());await refreshProfiles('Общий');"
        "await loadProfileIntoEditor('Общий');"
    ) in javascript



def test_native_window_handle_is_private_to_avoid_pywebview_bridge_introspection(tmp_path):
    api = AppApi(root=tmp_path)
    sentinel = object()
    api.set_window(sentinel)
    assert not hasattr(api, "window")
    assert api._window is sentinel


def test_new_design_omits_redundant_workflow_steps_but_keeps_file_state():
    html = asset_path("index-v1.8.18.html").read_text(encoding="utf-8")

    assert 'id="workflowSteps"' not in html
    assert 'class="workflow-steps"' not in html
    assert 'id="fileLifecycleStatus"' in html
    assert 'id="applyNow"' in html
    assert 'id="exportNow"' in html


def test_header_omits_redundant_sensitive_data_subtitle_in_both_languages():
    ru_html = asset_path("index-v1.8.18.html").read_text(encoding="utf-8")
    en_html = asset_path("index-en-v1.8.18.html").read_text(encoding="utf-8")

    assert "Персональные и другие чувствительные данные заменяются маркерами или удаляются." not in ru_html
    assert "Всё выполняется локально." not in ru_html
    assert "Personal and other sensitive data is replaced with markers or removed." not in en_html
    assert "Everything runs locally." not in en_html


def test_open_file_button_is_first_source_action_with_localized_tooltip():
    localized = {
        "index-v1.8.18.html": (
            ">Открыть</button>",
            'title="Открыть документ или несколько файлов и импортировать их в локальную библиотеку"',
        ),
        "index-en-v1.8.18.html": (
            ">Open</button>",
            'title="Open one or more documents and import them into the local library"',
        ),
    }

    for name, (label, tooltip) in localized.items():
        html = asset_path(name).read_text(encoding="utf-8")
        source_bar = html.split('aria-labelledby="sourceHeading">', 1)[1].split("</div>", 1)[0]
        assert source_bar.index('id="dropChoose"') < source_bar.index('id="paste"')
        assert label in source_bar
        assert tooltip in source_bar
        assert html.count('id="dropChoose"') == 1


def test_new_design_styles_define_contrast_actions_targets_and_focus():
    html = asset_path("index-v1.8.18.html").read_text(encoding="utf-8")
    css = asset_path("new-design-v1.11.css").read_text(encoding="utf-8") if asset_path("new-design-v1.11.css").exists() else ""

    assert 'new-design-v1.11.css' in html
    for token in ("--text-primary", "--text-secondary", "--surface", "--border", "--action-primary", "--danger", "--warning", "--success"):
        assert token in css
    assert ":focus-visible" in css
    assert "min-height:40px" in css
    assert 'body[data-design="new"]>details{order:6}' in css
    assert 'body[data-design="legacy"]' in css
    assert "--border:#647c92" in css
    assert "--focus:#0069c0" in css
    assert "summary{min-height:24px" in css


def test_all_collapsible_sections_show_right_or_down_disclosure_arrow():
    for name in ("index-v1.8.18.html", "index-en-v1.8.18.html"):
        html = asset_path(name).read_text(encoding="utf-8")
        assert 'new-design-v1.11.css' in html

    css = asset_path("new-design-v1.11.css").read_text(encoding="utf-8")
    assert 'body[data-design="new"] details>summary{list-style:none' in css
    assert 'body[data-design="new"] details>summary::-webkit-details-marker{display:none}' in css
    assert 'body[data-design="new"] details:not([open])>summary::before{content:"▶"}' in css
    assert 'body[data-design="new"] details[open]>summary::before{content:"▼"}' in css


def test_source_pane_omits_redundant_visible_label_and_storage_notice():
    localized_copy = {
        "index-v1.8.18.html": (
            "Исходный текст или текст выбранного документа",
            "Исходник остаётся на компьютере; перед передачей нужна ручная проверка.",
        ),
        "index-en-v1.8.18.html": (
            "Source text or text from the selected document",
            "The source stays on this computer. Review it manually before sharing.",
        ),
    }

    for name, (accessible_label, storage_notice) in localized_copy.items():
        html = asset_path(name).read_text(encoding="utf-8")
        assert f'<label for="source" class="sr-only">{accessible_label}</label>' in html
        assert 'class="sr-label"' not in html
        assert 'id="events"' not in html
        assert storage_notice not in html
        assert '<div class="drop">' not in html

    css = asset_path("new-design-v1.11.css").read_text(encoding="utf-8")
    assert '.sr-label{' not in css


def test_empty_undo_row_collapses_only_when_empty(tmp_path):
    edge_candidates = (
        Path(r"C:\Program Files (x86)\Microsoft\Edge\Application\msedge.exe"),
        Path(r"C:\Program Files\Microsoft\Edge\Application\msedge.exe"),
    )
    edge = next((candidate for candidate in edge_candidates if candidate.exists()), None)
    if edge is None:
        pytest.skip("Microsoft Edge is required for the Windows layout regression")

    for html_name in ("index-v1.8.18.html", "index-en-v1.8.18.html"):
        html = asset_path(html_name).read_text(encoding="utf-8")
        assert 'new-design-v1.11.css' in html
        assert 'id="undoChips" class="undo-chips"' in html

    for script_name in ("app-v1.8.18.js", "app-en-v1.8.18.js"):
        javascript = asset_path(script_name).read_text(encoding="utf-8")
        assert "$('#undoChips').innerHTML=" in javascript
        assert "$('#undoChips').onclick=" in javascript

    links = "".join(
        f'<link rel="stylesheet" href="{asset_path(name).as_uri()}">'
        for name in ("style-v1.1.css", "new-design-v1.11.css")
    )
    for width in (320, 1430):
        probe = tmp_path / f"empty-undo-row-{width}.html"
        probe.write_text(
            "<!doctype html><html><head><meta charset=\"utf-8\">"
            + links
            + "</head><body data-design=\"new\"><section class=\"card pane\">"
            + "<div id=\"richOutput\" class=\"rich-output\"></div>"
            + "<div id=\"undoChips\" class=\"undo-chips\"></div>"
            + "<div id=\"chips\" class=\"chips\"><span>People: 0</span></div>"
            + "</section><pre id=\"probe\"></pre><script>"
            + "const rich=document.querySelector('#richOutput'),undo=document.querySelector('#undoChips'),chips=document.querySelector('#chips');"
            + "const metrics=()=>({undoDisplay:getComputedStyle(undo).display,undoHeight:undo.getBoundingClientRect().height,gap:chips.getBoundingClientRect().top-rich.getBoundingClientRect().bottom});"
            + "const empty=metrics();undo.innerHTML='<button class=\"undo-chip\">Undo</button>';const populated=metrics();"
            + "document.querySelector('#probe').textContent=JSON.stringify({empty,populated,overflow:document.documentElement.scrollWidth>document.documentElement.clientWidth});"
            + "</script></body></html>",
            encoding="utf-8",
        )
        profile = tmp_path / f"edge-profile-empty-undo-{width}"
        result = subprocess.run(
            [
                str(edge),
                "--headless=new",
                "--disable-gpu",
                "--allow-file-access-from-files",
                f"--user-data-dir={profile}",
                f"--window-size={width},923",
                "--dump-dom",
                probe.as_uri(),
            ],
            check=True,
            capture_output=True,
            text=True,
            encoding="utf-8",
            errors="replace",
            timeout=60,
        )
        match = re.search(r'<pre id="probe">(.*?)</pre>', result.stdout, re.DOTALL)
        assert match is not None
        metrics = json.loads(html_module.unescape(match.group(1)))
        assert metrics["empty"]["undoDisplay"] == "none", (width, metrics)
        assert metrics["empty"]["undoHeight"] == 0, (width, metrics)
        assert metrics["empty"]["gap"] <= 8.5, (width, metrics)
        assert metrics["populated"]["undoDisplay"] != "none", (width, metrics)
        assert metrics["populated"]["undoHeight"] >= 28, (width, metrics)
        assert metrics["overflow"] is False, (width, metrics)


def test_entity_summary_uses_the_free_space_above_output_preview_in_both_languages():
    for html_name in ("index-v1.8.18.html", "index-en-v1.8.18.html"):
        html = asset_path(html_name).read_text(encoding="utf-8")
        output = html[html.index('aria-labelledby="outputHeading"'):html.index('id="manualReview"')]

        chips = output.index('id="chips"')
        alignment = output.index('id="outputPreviewAlignment"')
        preview = output.index('id="anonymizedPreviewSurface"')
        undo = output.index('id="undoChips"')

        assert chips < alignment < preview < undo

    css = asset_path("new-design-v1.11.css").read_text(encoding="utf-8")
    assert 'body[data-design="new"] .pane>.chips{' in css
    assert "flex-wrap:nowrap" in css
    assert 'body[data-design="new"] .pane>.chips span{' in css


def test_new_design_uses_compact_interface_text_without_scaling_document_preview():
    css = asset_path("new-design-v1.11.css").read_text(encoding="utf-8")

    assert 'body[data-design="new"]{display:flex;flex-direction:column;background:var(--surface-raised);color:var(--text-primary);font-size:13px;' in css
    assert 'body[data-design="new"] textarea{font-size:13px;' in css
    assert 'body[data-design="new"] h2{font-size:16px;' in css
    assert 'body[data-design="new"] .docx-preview{' not in css


def test_new_design_aligns_remaining_header_controls_on_one_level():
    css = asset_path("new-design-v1.11.css").read_text(encoding="utf-8")
    ru_html = asset_path("index-v1.8.18.html").read_text(encoding="utf-8")
    en_html = asset_path("index-en-v1.8.18.html").read_text(encoding="utf-8")

    assert 'body[data-design="new"] .language-switch{display:flex;align-items:center;gap:6px;white-space:nowrap}' in css
    assert 'body[data-design="new"] .header-actions select,body[data-design="new"] #themeToggle{height:34px;min-height:34px;padding:4px 9px;box-sizing:border-box}' in css
    assert 'body[data-design="new"] #themeToggle{white-space:nowrap}' in css
    for html in (ru_html, en_html):
        assert 'class="language-switch"' in html
        assert 'class="design-switch"' not in html


def test_compact_header_reclaims_height_for_document_previews():
    for name in ("index-v1.8.18.html", "index-en-v1.8.18.html"):
        html = asset_path(name).read_text(encoding="utf-8")
        assert 'new-design-v1.11.css' in html

    css = asset_path("new-design-v1.11.css").read_text(encoding="utf-8")
    assert 'body[data-design="new"]{display:flex;flex-direction:column;background:var(--surface-raised);color:var(--text-primary);font-size:13px;line-height:1.45;max-width:1920px;margin:auto;padding:8px 20px 90px}' in css
    assert 'body[data-design="new"] header{order:0;position:sticky;top:0;z-index:80;background:var(--surface-raised);padding:3px 0 5px;margin-bottom:6px;' in css
    assert 'grid-template-columns:minmax(0,1fr) auto;grid-template-areas:"title actions" "explanation explanation";align-items:center;gap:2px 12px' in css
    assert 'body[data-design="new"] header>div:first-child{display:contents}' in css
    assert 'body[data-design="new"] header p{grid-area:explanation;margin:0;min-width:0;font-size:11px;line-height:1.2;white-space:normal;overflow:visible}' in css
    assert 'body[data-design="new"] .header-actions #stats{font-size:12px;line-height:1.2;white-space:nowrap}' in css
    assert 'body[data-design="new"] .header-actions select,body[data-design="new"] #themeToggle{height:34px;min-height:34px;padding:4px 9px;box-sizing:border-box}' in css
    assert 'body[data-design="new"] header h1{white-space:normal;overflow-wrap:anywhere;line-height:1.2}' in css
    assert 'body[data-design="new"] .header-actions #stats{white-space:normal;overflow-wrap:anywhere;text-align:left}' in css
    assert '@media(min-width:801px) and (min-height:761px) and (max-height:950px){body[data-design="new"] .pane .editor-stack,body[data-design="new"] .rich-output,body[data-design="new"] .docx-preview-host{height:min(calc(36vh + 46px),390px);min-height:346px}}' in css
    assert '@media(min-width:801px) and (max-height:760px){body[data-design="new"] .pane .editor-stack,body[data-design="new"] .rich-output,body[data-design="new"] .docx-preview-host{height:min(22vh,220px);min-height:170px}}' in css


def test_drop_prompt_is_removed_but_window_drop_import_remains_available():
    for name, prompt in (
        ("index-v1.8.18.html", "Перетащите файлы сюда"),
        ("index-en-v1.8.18.html", "Drag files here"),
    ):
        html = asset_path(name).read_text(encoding="utf-8")
        assert '<div class="drop">' not in html
        assert prompt not in html
        assert 'id="dropChoose"' in html

    for name in ("app-v1.8.18.js", "app-en-v1.8.18.js"):
        script = asset_path(name).read_text(encoding="utf-8")
        assert "$('#dropChoose').onclick=$('#choose').onclick;" in script
        assert "document.addEventListener('drop',async e=>" in script


def test_compact_header_renders_complete_without_horizontal_overflow(tmp_path):
    edge_candidates = (
        Path(r"C:\Program Files (x86)\Microsoft\Edge\Application\msedge.exe"),
        Path(r"C:\Program Files\Microsoft\Edge\Application\msedge.exe"),
    )
    edge = next((candidate for candidate in edge_candidates if candidate.exists()), None)
    if edge is None:
        pytest.skip("Microsoft Edge is required for the Windows layout regression")

    statuses = {
        "index-v1.8.18.html": "Найдено замен: 374 · Требуют решения: 83 · Проверено: 0 из 83",
        "index-en-v1.8.18.html": "Replacements found: 374 · Need review: 83 · Reviewed: 0 of 83",
    }
    for html_name, status in statuses.items():
        page = asset_path(html_name).read_text(encoding="utf-8")
        header = page[page.index("<header>"):page.index("</header>") + len("</header>")]
        stylesheets = re.findall(r'<link rel="stylesheet" href="([^"]+)">', page)
        links = "".join(
            f'<link rel="stylesheet" href="{asset_path(name).as_uri()}">'
            for name in stylesheets
        )
        for width in (320, 1430):
            probe = tmp_path / f"header-{html_name}-{width}.html"
            body_style = ' style="width:305px;max-width:305px;margin:0"' if width == 320 else ""
            probe.write_text(
                "<!doctype html><html><head><meta charset=\"utf-8\">"
                + links
                + f"</head><body data-design=\"new\"{body_style}>"
                + header
                + "<pre id=\"probe\"></pre><script>"
                + f"document.querySelector('#stats').textContent={json.dumps(status, ensure_ascii=False)};"
                + "const root=document.documentElement,title=document.querySelector('header h1'),"
                + "explanation=document.querySelector('header p'),stats=document.querySelector('#stats'),"
                + "headerBox=document.querySelector('header');"
                + "document.querySelector('#probe').textContent=JSON.stringify({"
                + "docOverflow:root.scrollWidth>root.clientWidth,"
                + "bodyOverflow:document.body.scrollWidth>document.body.clientWidth,"
                + "titleClipped:title.scrollWidth>title.clientWidth+1||title.scrollHeight>title.clientHeight+1,"
                + "explanationClipped:explanation?(explanation.scrollWidth>explanation.clientWidth+1||explanation.scrollHeight>explanation.clientHeight+1):false,"
                + "statsClipped:stats.scrollWidth>stats.clientWidth+1||stats.scrollHeight>stats.clientHeight+1,"
                + "headerHeight:headerBox.getBoundingClientRect().height});"
                + "</script></body></html>",
                encoding="utf-8",
            )
            profile = tmp_path / f"edge-profile-{html_name}-{width}"
            result = subprocess.run(
                [
                    str(edge),
                    "--headless=new",
                    "--disable-gpu",
                    "--allow-file-access-from-files",
                    f"--user-data-dir={profile}",
                    f"--window-size={width},923",
                    "--dump-dom",
                    probe.as_uri(),
                ],
                check=True,
                capture_output=True,
                text=True,
                encoding="utf-8",
                errors="replace",
                timeout=60,
            )
            match = re.search(r'<pre id="probe">(.*?)</pre>', result.stdout, re.DOTALL)
            assert match is not None
            metrics = json.loads(html_module.unescape(match.group(1)))
            assert metrics["docOverflow"] is False, (html_name, width, metrics)
            assert metrics["bodyOverflow"] is False, (html_name, width, metrics)
            assert metrics["titleClipped"] is False, (html_name, width, metrics)
            assert metrics["explanationClipped"] is False, (html_name, width, metrics)
            assert metrics["statsClipped"] is False, (html_name, width, metrics)
            if width == 1430:
                assert metrics["headerHeight"] <= 64, (html_name, width, metrics)


def test_dark_new_design_stat_chips_have_explicit_readable_colors():
    css = asset_path("new-design-v1.11.css").read_text(encoding="utf-8")

    assert 'html[data-theme="dark"] body[data-design="new"] .chips span{' in css
    assert "background:var(--surface-raised)" in css
    assert "color:var(--text-primary)" in css
    assert 'html[data-theme="dark"] body[data-design="new"] .chips .yellow{' in css
    assert "background:#fff5c5" in css
    assert "color:#725a00" in css


def test_new_design_current_review_text_has_explicit_readable_color():
    css = asset_path("new-design-v1.11.css").read_text(encoding="utf-8")

    assert 'body[data-design="new"] #reviewCurrent{color:var(--text-primary)}' in css


def test_new_design_manual_review_uses_dark_theme_surfaces_without_inner_scroll():
    css = asset_path("new-design-v1.11.css").read_text(encoding="utf-8")

    assert 'body[data-design="new"] .review-candidate{background:var(--surface-raised);color:var(--text-primary);border-color:var(--border)}' in css
    assert 'body[data-design="new"] .review-candidate{height:auto;min-height:46px;max-height:none;overflow:visible;padding:6px 8px;box-sizing:border-box;grid-template-columns:auto minmax(0,2fr) minmax(0,1fr);gap:4px 10px;align-items:center}' in css
    assert 'body[data-design="new"] .review-decision,body[data-design="new"] .review-finish{background:var(--surface-raised);color:var(--text-secondary)}' in css


def test_settings_and_manual_review_are_below_their_respective_preview_panes():
    css = asset_path("new-design-v1.11.css").read_text(encoding="utf-8")

    for name in ("index-v1.8.18.html", "index-en-v1.8.18.html"):
        html = asset_path(name).read_text(encoding="utf-8")
        source_start = html.index('id="sourceColumn"')
        output_start = html.index('id="anonymizedColumn"')
        settings_start = html.index('id="anonymizationSettings"')
        review_start = html.index('id="manualReview"')
        assert source_start < settings_start < output_start
        assert output_start < review_start
        assert settings_start > html.index('id="sourceHeading"')
        assert review_start > html.index('id="anonymizedPreviewSurface"')
        assert '<div class="review-controls-disclosure">' in html
        assert '<fieldset class="review-controls" aria-label=' in html
        for control_id in ("reviewHints", "reviewHintsMode", "paranoid", "compactMarkers", "mono"):
            assert f'id="{control_id}"' in html

    assert 'grid-template-columns:minmax(260px,.85fr) minmax(420px,1.15fr)' in css
    assert 'body[data-design="new"] .workspace-column{container-type:inline-size;display:flex;min-width:0;flex-direction:column;gap:8px}' in css
    assert 'body[data-design="new"] .workspace-column>.pane,body[data-design="new"] .workspace-column>.review,body[data-design="new"] .workspace-column>.settings-panel{width:100%;min-width:0;margin:0}' in css
    assert 'body[data-design="new"] .review-context,body[data-design="new"] .review-toolbar{grid-column:1/-1;min-width:0}' in css
    assert 'body[data-design="new"] .review-context{display:grid;grid-template-columns:minmax(0,1fr);gap:6px}' in css
    assert 'body[data-design="new"] .pane .editor-stack,body[data-design="new"] .rich-output,body[data-design="new"] .docx-preview-host{height:min(34vh,380px);min-height:250px}' in css
    assert '@media(min-width:801px) and (min-height:761px) and (max-height:950px){body[data-design="new"] .pane .editor-stack,body[data-design="new"] .rich-output,body[data-design="new"] .docx-preview-host{height:min(calc(36vh + 46px),390px);min-height:346px}}' in css
    assert '@media(min-width:801px) and (max-height:760px){body[data-design="new"] .pane .editor-stack,body[data-design="new"] .rich-output,body[data-design="new"] .docx-preview-host{height:min(22vh,220px);min-height:170px}}' in css
    assert 'body[data-design="new"] .review-candidate{height:auto;min-height:46px;max-height:none;overflow:visible;padding:6px 8px;box-sizing:border-box;grid-template-columns:auto minmax(0,2fr) minmax(0,1fr);gap:4px 10px;align-items:center}' in css
    assert 'body[data-design="new"] .review-toolbar{display:grid;grid-template-columns:repeat(4,minmax(0,1fr));gap:6px;align-items:stretch}' in css
    assert 'body[data-design="new"] .review-navigation,body[data-design="new"] .review-actions{display:contents}' in css
    assert 'body[data-design="new"] .review-toolbar button{width:100%;min-width:0;min-height:34px;margin:0;padding:5px 8px;white-space:normal;overflow-wrap:anywhere;line-height:1.2}' in css


def test_review_candidate_and_toolbar_span_the_output_review_width():
    css = asset_path("new-design-v1.11.css").read_text(encoding="utf-8")

    assert 'body[data-design="new"] .columns{column-gap:14px;row-gap:6px}' in css
    assert 'body[data-design="new"] .workspace-column{container-type:inline-size;display:flex;min-width:0;flex-direction:column;gap:8px}' in css
    assert 'body[data-design="new"] .review-context,body[data-design="new"] .review-toolbar{grid-column:1/-1;min-width:0}' in css


def test_review_checkbox_modes_follow_the_four_button_toolbar_in_both_languages():
    for name in ("index-v1.8.18.html", "index-en-v1.8.18.html"):
        html = asset_path(name).read_text(encoding="utf-8")
        review = html.split('<section id="manualReview"', 1)[1].split("</section>", 1)[0]

        disclosure_start = review.index('<div class="review-controls-disclosure">')
        toolbar_start = review.index('<div class="review-toolbar">')
        decision_start = review.index('<div id="reviewDecision" class=')

        assert toolbar_start < disclosure_start < decision_start
        toolbar = review[toolbar_start:disclosure_start]
        assert toolbar.index('id="reviewPrev"') < toolbar.index('id="reviewNext"')
        assert toolbar.index('id="reviewNext"') < toolbar.index('id="hideCandidate"')
        assert toolbar.index('id="hideCandidate"') < toolbar.index('id="keepCandidate"')


def test_review_action_buttons_use_tooltips_and_are_grouped_beside_each_other():
    expectations = (
        (
            "index-v1.8.18.html",
            "Заменить все одинаковые вхождения маркерами и добавить это значение в дополнительные правила",
            "Оставить текущий фрагмент и добавить это значение в белый список",
            "Все одинаковые вхождения будут заменены маркерами; правило будет добавлено в текущие настройки.",
            "Текущий фрагмент останется без изменения; значение будет добавлено в белый список.",
        ),
        (
            "index-en-v1.8.18.html",
            "Replace all matching occurrences with markers and add this value to additional rules",
            "Keep the current item and add its value to the allowlist",
            "All matching occurrences will be replaced with markers; the rule will be added to the current settings.",
            "The current item will remain unchanged; its value will be added to the allowlist.",
        ),
    )

    for asset_name, hide_title, keep_title, hide_help, keep_help in expectations:
        html = asset_path(asset_name).read_text(encoding="utf-8")
        toolbar = html.split('<div class="review-toolbar">', 1)[1].split(
            '<div class="review-controls-disclosure">', 1
        )[0]
        navigation = toolbar.split('<div class="review-navigation">', 1)[1].split("</div>", 1)[0]
        actions = toolbar.split('<div class="review-actions" id="reviewDecisionActions">', 1)[1].split(
            '<div class="review-controls-disclosure">', 1
        )[0]

        assert 'id="hideCandidate"' not in navigation
        assert (
            f'id="hideCandidate" title="{hide_title}" '
            'aria-describedby="hideCandidateHelp"'
        ) in actions
        assert (
            f'<span id="hideCandidateHelp" class="sr-only">{hide_help}</span>'
        ) in actions
        assert toolbar.index('id="reviewPrev"') < toolbar.index('id="reviewNext"')
        assert toolbar.index('id="reviewNext"') < toolbar.index('id="hideCandidate"')
        assert toolbar.index('id="hideCandidate"') < toolbar.index('id="keepCandidate"')
        assert (
            f'id="keepCandidate" title="{keep_title}" '
            'aria-describedby="keepCandidateHelp"'
        ) in actions
        assert (
            f'<span id="keepCandidateHelp" class="sr-only">{keep_help}</span>'
        ) in actions
        assert "<small>" not in actions
        assert 'id="reviewFinish" class="sr-only"' in html


def test_populated_review_keeps_heading_and_remaining_action_compact():
    css = asset_path("new-design-v1.11.css").read_text(encoding="utf-8")

    base_css = css.split("@media", 1)[0]
    assert 'body[data-design="new"] .workflow-status{position:fixed;left:20px;right:20px;bottom:0;z-index:70;' in base_css
    assert '.workflow-status{position:static' not in css
    assert 'body[data-design="new"]{padding:10px 10px 210px}' in css
    assert 'body[data-design="new"] .workflow-status{left:10px;right:10px;flex-wrap:wrap}' in css
    assert 'body[data-design="new"] .review-heading,body[data-design="new"] .review-controls-disclosure{grid-column:1/-1}' in css
    assert 'body[data-design="new"] .review-heading{margin:0;gap:8px;align-items:center;display:grid;grid-template-columns:minmax(0,1fr) auto}' in css
    assert 'body[data-design="new"] .review-heading>div{min-width:0}' in css
    assert 'body[data-design="new"] .review-heading h2{position:relative;margin:0;white-space:normal;overflow-wrap:anywhere;outline-offset:3px}' in css
    assert 'body[data-design="new"] .review-toolbar{display:grid;grid-template-columns:repeat(4,minmax(0,1fr));gap:6px;align-items:stretch}' in css
    assert 'body[data-design="new"] .review-navigation,body[data-design="new"] .review-actions{display:contents}' in css
    assert 'body[data-design="new"] .review-toolbar button{width:100%;min-width:0;min-height:34px;' in css
    assert '@media(max-width:800px){body[data-design="new"] .review{grid-template-columns:1fr}' in css
    assert '@container(max-width:519px){body[data-design="new"] .review-toolbar{grid-template-columns:repeat(2,minmax(0,1fr))}' in css


def test_restore_steps_stack_on_narrow_screens_in_both_languages():
    css = asset_path("new-design-v1.11.css").read_text(encoding="utf-8")
    ru = asset_path("index-v1.8.18.html").read_text(encoding="utf-8")
    en = asset_path("index-en-v1.8.18.html").read_text(encoding="utf-8")

    assert '@media(max-width:620px){body[data-design="new"] .restore-steps{grid-template-columns:minmax(0,1fr)}}' in css
    assert 'class="restore-steps"' in ru
    assert 'class="restore-steps"' in en


def test_new_design_hides_empty_review_result_status():
    css = asset_path("new-design-v1.11.css").read_text(encoding="utf-8")

    assert 'body[data-design="new"] #resultStatus:empty{display:none}' in css


def test_new_design_has_labels_headings_live_status_and_sensitive_artifact_warnings():
    html = asset_path("index-v1.8.18.html").read_text(encoding="utf-8")

    assert '<label for="source"' in html
    assert re.search(r'<h2 id="reviewTitle"[^>]+>Ручная проверка спорных мест ', html)
    assert 'id="fileLifecycleStatus"' in html and 'role="status"' in html
    assert "локальный архив документов и ключей восстановления" in html.lower()
    assert "профиль может содержать исходные чувствительные значения" in html.lower()
    for field in ("library", "profile"):
        assert f'<label for="{field}"' in html
    for status in ("fileStatus", "restoreStatus", "testResult", "warning"):
        assert f'id="{status}"' in html
    assert 'id="richOutput"' in html and 'role="region"' in html
    assert '<h2 id="sourceHeading">Исходный текст</h2>' in html
    assert '<h2 id="outputHeading">Обезличенный текст</h2>' in html


def test_new_design_uses_plain_language_stats_review_context_and_operation_states():
    html = asset_path("index-v1.8.18.html").read_text(encoding="utf-8")
    javascript = asset_path("app-v1.8.18.js").read_text(encoding="utf-8")

    assert 'id="reviewMeta"' in html
    for text in ("Найдено замен:", "Требуют решения:", "Проверено:", "Создаётся предпросмотр", "Обрабатывается файл"):
        assert text in javascript
    assert "вхожд." in javascript
    assert "Правила: «" in javascript


def test_review_copy_stays_compact_when_a_candidate_appears_in_both_languages():
    ru_html = asset_path("index-v1.8.18.html").read_text(encoding="utf-8")
    en_html = asset_path("index-en-v1.8.18.html").read_text(encoding="utf-8")
    ru_js = asset_path("app-v1.8.18.js").read_text(encoding="utf-8")
    en_js = asset_path("app-en-v1.8.18.js").read_text(encoding="utf-8")

    assert "Автопоиск не уверен — выберите действие." not in ru_html
    assert "Detection is uncertain — choose an action." not in en_html
    assert "`«${all[reviewIndex].textContent.trim()}» · ${all[reviewIndex].dataset.kind" in ru_js
    assert "`${occurrences} вхожд. · Правила: «${$('#profile').value||'Общий'}»`" in ru_js
    assert "`“${all[reviewIndex].textContent.trim()}” · ${hintKindText(all[reviewIndex].dataset.kind)" in en_js
    assert "`${occurrences} match(es) · Rules: “${profileDisplayName($('#profile').value||'Общий')}”`" in en_js

    css = asset_path("new-design-v1.11.css").read_text(encoding="utf-8")
    assert 'body[data-design="new"] .review-candidate .review-help{grid-column:3;margin:0}' in css
    assert 'body[data-design="new"] .review-candidate small' not in css
    assert 'body[data-design="new"] .review-context,body[data-design="new"] .review-toolbar{grid-column:1/-1;min-width:0}' in css


@pytest.mark.parametrize("asset_name", ["app-v1.8.18.js", "app-en-v1.8.18.js"])
def test_review_navigation_preserves_formatted_mode_and_targets_both_previews(asset_name):
    javascript = asset_path(asset_name).read_text(encoding="utf-8")
    focus_mark = re.search(
        r"function focusMark\(i\)\{(?P<body>.*?)\n\}", javascript, re.S
    )

    assert focus_mark is not None
    assert "setAnonymizedPreviewMode(false)" not in focus_mark.group("body")
    assert "focusFormattedReview(part,sourceOffset)" in focus_mark.group("body")
    assert "function focusFormattedReview(part,sourceOffset)" in javascript
    assert "textRangeForSourceOffset(sourceRoot,needle,$('#source').value,sourceOffset,globalOccurrence)" in javascript
    assert "data-preview-anchor" in javascript
    assert "CSS.highlights.set('current-review'" in javascript


def test_formatted_review_current_match_has_visible_highlight_style():
    css = asset_path("style-v1.1.css").read_text(encoding="utf-8")
    assert "::highlight(current-review)" in css
    assert ".preview-review-current" in css


def test_formatted_preview_status_and_buttons_have_deterministic_rows():
    for name in ("style-v1.1.css", "style.css"):
        css = asset_path(name).read_text(encoding="utf-8")
        assert ".preview-switch{display:grid;grid-template-columns:repeat(5,minmax(0,1fr));gap:4px;margin:-3px 0 7px}" in css
        assert ".preview-completeness{grid-column:1/-1;min-width:0;color:#58708a;padding:0 4px}" in css
        assert ".preview-switch>button{min-width:0;overflow-wrap:anywhere}" in css


def test_formatted_preview_controls_keep_one_button_row_while_loading_and_ready(tmp_path: Path):
    edge = next(
        (
            candidate
            for candidate in (
                Path(r"C:\Program Files (x86)\Microsoft\Edge\Application\msedge.exe"),
                Path(r"C:\Program Files\Microsoft\Edge\Application\msedge.exe"),
            )
            if candidate.exists()
        ),
        None,
    )
    if edge is None:
        pytest.skip("Microsoft Edge is required for the Windows preview-controls regression")

    base_css_uri = asset_path("style-v1.1.css").as_uri()
    design_css_uri = asset_path("new-design-v1.11.css").as_uri()
    for index, html_name in enumerate(("index-v1.8.18.html", "index-en-v1.8.18.html")):
        source = asset_path(html_name).read_text(encoding="utf-8")
        controls = [
            re.search(rf'<div id="{control_id}" class="preview-switch" hidden>.*?</div>', source, re.S)
            for control_id in ("docxPreviewControls", "anonymizedPreviewControls")
        ]
        assert all(controls)
        fixture = tmp_path / f"preview-controls-{index}.html"
        fixture.write_text(
            f'''<!doctype html><html><head><meta charset="utf-8"><link rel="stylesheet" href="{base_css_uri}"><link rel="stylesheet" href="{design_css_uri}"></head>
<body data-design="new"><div id="case">{controls[0].group(0)}{controls[1].group(0)}</div><pre id="probe"></pre><script>
const rows=[];const root=document.getElementById('case');
for(const width of [390,700,900]){{for(const state of ['loading','ready']){{root.style.width=width+'px';for(const controls of root.querySelectorAll('.preview-switch')){{controls.hidden=false;controls.querySelector('.preview-completeness').textContent=state==='loading'?'Formatted preview is being prepared…':'Formatted preview is ready · tables: 44. Text preview matches the selected settings.';const status=controls.querySelector('.preview-completeness').getBoundingClientRect(),buttons=[...controls.querySelectorAll('button')].map(button=>button.getBoundingClientRect());rows.push({{width,state,id:controls.id,count:buttons.length,buttonTops:[...new Set(buttons.map(box=>Math.round(box.top)))],statusBottom:Math.round(status.bottom),buttonTop:Math.round(Math.min(...buttons.map(box=>box.top))),overflow:controls.scrollWidth-controls.clientWidth,buttonOverflow:Math.max(...buttons.map(box=>box.right-controls.getBoundingClientRect().right))}})}}}}}}
document.getElementById('probe').textContent=JSON.stringify(rows);</script></body></html>''',
            encoding="utf-8",
        )
        profile = tmp_path / f"edge-profile-preview-controls-{index}"
        result = subprocess.run(
            [str(edge), "--headless=new", "--disable-gpu", "--allow-file-access-from-files", f"--user-data-dir={profile}", "--window-size=1200,900", "--dump-dom", fixture.as_uri()],
            capture_output=True,
            text=True,
            timeout=45,
            check=True,
        )
        match = re.search(r'<pre id="probe">(.*?)</pre>', result.stdout, re.S)
        assert match, result.stdout[-2000:]
        metrics = json.loads(html_module.unescape(match.group(1)))
        assert len(metrics) == 12
        for metric in metrics:
            assert metric["count"] == 5, metric
            assert len(metric["buttonTops"]) == 1, metric
            assert metric["statusBottom"] <= metric["buttonTop"], metric
            assert metric["overflow"] <= 1, metric
            assert metric["buttonOverflow"] <= 1, metric


def test_review_navigation_runs_inside_formatted_previews_in_both_languages(tmp_path: Path):
    edge = next(
        (
            candidate
            for candidate in (
                Path(r"C:\Program Files (x86)\Microsoft\Edge\Application\msedge.exe"),
                Path(r"C:\Program Files\Microsoft\Edge\Application\msedge.exe"),
            )
            if candidate.exists()
        ),
        None,
    )
    if edge is None:
        pytest.skip("Microsoft Edge is required for the Windows formatted-navigation regression")

    harness_template = r'''<!doctype html><meta charset="utf-8"><title>pending</title><iframe id="app" style="width:__WIDTH__px;height:800px"></iframe><pre id="probe">pending</pre><script>
const frame=document.getElementById('app');frame.src=__INDEX__;
function finish(value){document.getElementById('probe').textContent=JSON.stringify(value);document.title=value.ok?'PASS':'FAIL';}
frame.addEventListener('load',()=>setTimeout(()=>{try{
 const w=frame.contentWindow,d=w.document,needle='SyntheticCandidate',lineA=`Context A ${needle} tail A`,lineB=`Context B ${needle} tail B`;
 const blocks=Array.from({length:18},(_,i)=>`<p data-preview-anchor="body/p${i}">${i===4?lineA:i===15?lineB:'Synthetic paragraph '+i}</p>`).join('');
 d.getElementById('docxPreview').innerHTML=blocks;d.getElementById('anonymizedDocxPreview').innerHTML=blocks;
 d.getElementById('source').value=`${lineB}\n${'filler\n'.repeat(10)}${lineA}`;const first=d.getElementById('source').value.indexOf(needle),second=d.getElementById('source').value.lastIndexOf(needle);
 d.getElementById('richOutput').innerHTML=`<span class="mark-hint" data-s0="${first}">${needle}</span>${' filler '.repeat(40)}<span class="mark-hint" data-s0="${second}">${needle}</span>`;
 d.getElementById('docxPreview').hidden=false;d.getElementById('textPreview').hidden=true;d.getElementById('showFormatted').setAttribute('aria-pressed','true');d.getElementById('anonymizedDocxPreview').hidden=false;d.getElementById('richOutput').hidden=true;d.getElementById('showAnonymizedFormatted').setAttribute('aria-pressed','true');
 const anchors=()=>Array.from(w.CSS.highlights.get('current-review')||[]).map(range=>range.startContainer.parentElement.closest('[data-preview-anchor]').dataset.previewAnchor);
 d.getElementById('reviewNext').click();d.getElementById('reviewNext').click();const nextOffset=d.querySelector('#richOutput .mark-current')?.dataset.s0,nextAnchors=anchors();
 d.getElementById('reviewPrev').click();const previousOffset=d.querySelector('#richOutput .mark-current')?.dataset.s0,previousAnchors=anchors();
 const formatted=!d.getElementById('docxPreview').hidden&&!d.getElementById('anonymizedDocxPreview').hidden&&d.getElementById('richOutput').hidden;
 finish({ok:formatted&&nextOffset===String(second)&&nextAnchors.length===2&&nextAnchors.every(anchor=>anchor==='body/p4')&&previousOffset===String(first)&&previousAnchors.length===2&&previousAnchors.every(anchor=>anchor==='body/p15'),formatted,nextOffset,previousOffset,nextAnchors,previousAnchors});
}catch(error){finish({ok:false,error:String(error)});}},1000));setTimeout(()=>{if(document.title==='pending')finish({ok:false,error:'timeout'});},4500);</script>'''

    for index, html_name in enumerate(("index-v1.8.18.html", "index-en-v1.8.18.html")):
        for workspace_width in (390, 700, 900):
            fixture = tmp_path / f"formatted-review-navigation-{index}-{workspace_width}.html"
            fixture.write_text(
                harness_template.replace("__INDEX__", json.dumps(asset_path(html_name).as_uri())).replace("__WIDTH__", str(workspace_width)),
                encoding="utf-8",
            )
            profile = tmp_path / f"edge-profile-formatted-review-{index}-{workspace_width}"
            result = subprocess.run(
                [str(edge), "--headless=new", "--disable-gpu", "--allow-file-access-from-files", f"--user-data-dir={profile}", "--virtual-time-budget=6000", "--dump-dom", fixture.as_uri()],
                capture_output=True,
                text=True,
                timeout=45,
                check=True,
            )
            match = re.search(r'<pre id="probe">(.*?)</pre>', result.stdout, re.S)
            assert match, result.stdout[-2000:]
            outcome = json.loads(html_module.unescape(match.group(1)))
            assert outcome["ok"], {"language": html_name, "width": workspace_width, **outcome}


def test_selected_review_is_visible_after_initial_navigation_and_formatted_ready(tmp_path: Path):
    edge = next(
        (
            candidate
            for candidate in (
                Path(r"C:\Program Files (x86)\Microsoft\Edge\Application\msedge.exe"),
                Path(r"C:\Program Files\Microsoft\Edge\Application\msedge.exe"),
            )
            if candidate.exists()
        ),
        None,
    )
    if edge is None:
        pytest.skip("Microsoft Edge is required for the Windows formatted-review visibility regression")

    harness_template = r'''<!doctype html><meta charset="utf-8"><title>pending</title><iframe id="app" style="width:__WIDTH__px;height:800px"></iframe><pre id="probe">pending</pre><script>
const frame=document.getElementById('app');frame.src=__INDEX__;
function finish(value){document.getElementById('probe').textContent=JSON.stringify(value);document.title=value.ok?'PASS':'FAIL';}
frame.addEventListener('load',()=>setTimeout(async()=>{try{
 const w=frame.contentWindow,d=w.document,needles=['SyntheticCandidateTop','SyntheticCandidateMiddle','SyntheticCandidateBottom'];
 const words=Array.from({length:360},(_,i)=>`filler${i}`);words.splice(8,0,needles[0]);words.splice(180,0,needles[1]);words.splice(350,0,needles[2]);
 const paragraph=words.join(' '),sourceText=paragraph+'\n'+paragraph,offsets=needles.map(needle=>paragraph.length+1+paragraph.indexOf(needle));
 const blocks=()=>`<p data-preview-anchor="body/p0" style="line-height:24px;margin:0;overflow-wrap:anywhere">${paragraph}</p><p data-preview-anchor="body/p1" style="line-height:24px;margin:0;overflow-wrap:anywhere">${paragraph}</p>`;
 const rich=needles.map((needle,index)=>`<span class="mark-hint" data-kind="name" data-s0="${offsets[index]}">${needle}</span>`).join('');
 const left=d.getElementById('docxPreview'),right=d.getElementById('anonymizedDocxPreview');
 d.getElementById('source').value=sourceText;left.innerHTML=blocks();right.innerHTML=blocks();
 left.hidden=false;d.getElementById('textPreview').hidden=true;d.getElementById('showFormatted').setAttribute('aria-pressed','true');
 right.hidden=false;d.getElementById('richOutput').hidden=true;d.getElementById('showAnonymizedFormatted').setAttribute('aria-pressed','true');
 const exactRange=(pane,needle)=>{const walker=d.createTreeWalker(pane,w.NodeFilter.SHOW_TEXT);let node,seen=0;while(node=walker.nextNode()){const index=node.nodeValue.indexOf(needle);if(index>=0){if(seen++===1){const range=d.createRange();range.setStart(node,index);range.setEnd(node,index+needle.length);return range;}}}return null;};
 const visible=(pane,needle)=>{const range=exactRange(pane,needle);if(!range)return false;const p=pane.getBoundingClientRect(),t=range.getBoundingClientRect();return t.top>=p.top+1&&t.bottom<=p.bottom-1;};
 const settle=()=>new Promise(resolve=>w.setTimeout(resolve,80));
 const currentNeedle=()=>needles.find(needle=>d.getElementById('reviewCurrent').textContent.includes(needle));
 const snapshot=label=>{const needle=currentNeedle();return {label,needle,left:Boolean(needle)&&visible(left,needle),right:Boolean(needle)&&visible(right,needle),leftTop:left.scrollTop,rightTop:right.scrollTop};};
 const checks=[];
 d.getElementById('richOutput').innerHTML=rich;w.refreshFormattedReviewFocus();await settle();checks.push(snapshot('initial'));
 d.getElementById('reviewNext').click();await settle();checks.push(snapshot('next-1'));
 d.getElementById('reviewNext').click();await settle();checks.push(snapshot('next-2'));
 d.getElementById('reviewPrev').click();await settle();checks.push(snapshot('previous'));
 w.refreshFormattedReviewFocus();d.getElementById('reviewNext').click();await settle();checks.push(snapshot('stale-frame-next'));
 [left,right].forEach(pane=>pane.insertAdjacentHTML('afterbegin','<div style="height:240px" aria-hidden="true"></div>'));await settle();checks.push(snapshot('delayed-layout'));
 right.innerHTML='';w.refreshFormattedReviewFocus();await settle();checks.push({label:'pending-source',needle:currentNeedle(),left:visible(left,currentNeedle())});
 right.innerHTML=blocks();w.refreshFormattedReviewFocus();await settle();checks.push(snapshot('ready'));
 [left,right].forEach(pane=>[...pane.querySelectorAll('[data-preview-anchor]')].at(-1).removeAttribute('data-preview-anchor'));w.refreshFormattedReviewFocus();await settle();checks.push(snapshot('missing-target-anchor-fallback'));
 [left,right].forEach(pane=>{const paragraphs=[...pane.querySelectorAll('p')];paragraphs[1].dataset.previewAnchor='body/p1';paragraphs[0].removeAttribute('data-preview-anchor');});w.refreshFormattedReviewFocus();await settle();checks.push(snapshot('missing-earlier-anchor-fallback'));
 const visited=new Set(checks.slice(0,3).map(check=>check.needle));
 finish({ok:visited.size===3&&checks.every(check=>check.left&&check.right!==false),checks,visited:[...visited]});
}catch(error){finish({ok:false,error:String(error)});}},1000));setTimeout(()=>{if(document.title==='pending')finish({ok:false,error:'timeout'});},6000);</script>'''

    for index, html_name in enumerate(("index-v1.8.18.html", "index-en-v1.8.18.html")):
        for workspace_width in (390, 700, 900):
            fixture = tmp_path / f"formatted-review-visible-{index}-{workspace_width}.html"
            fixture.write_text(
                harness_template.replace("__INDEX__", json.dumps(asset_path(html_name).as_uri())).replace("__WIDTH__", str(workspace_width)),
                encoding="utf-8",
            )
            profile = tmp_path / f"edge-profile-formatted-visible-{index}-{workspace_width}"
            result = subprocess.run(
                [str(edge), "--headless=new", "--disable-gpu", "--allow-file-access-from-files", f"--user-data-dir={profile}", "--virtual-time-budget=7500", "--dump-dom", fixture.as_uri()],
                capture_output=True,
                text=True,
                timeout=45,
                check=True,
            )
            match = re.search(r'<pre id="probe">(.*?)</pre>', result.stdout, re.S)
            assert match, result.stdout[-2000:]
            outcome = json.loads(html_module.unescape(match.group(1)))
            assert outcome["ok"], {"language": html_name, "width": workspace_width, **outcome}


def test_formatted_review_focus_is_bound_to_render_lifecycle_without_timeout():
    for name in ("app-v1.8.18.js", "app-en-v1.8.18.js"):
        script = asset_path(name).read_text(encoding="utf-8")
        render_output = script[script.index("function renderOutput()") : script.index("function marks()")]
        refresh_focus = script[
            script.index("function refreshFormattedReviewFocus()") : script.index("function reviewDisabledReason()")
        ]
        formatted_start = script.index("if(hasFormattedDocumentPreview(formattedPreview?.format))")
        formatted_ready = script[
            formatted_start : script.index("}else{", formatted_start)
        ]

        assert "refreshFormattedReviewFocus();" in render_output
        assert "refreshFormattedReviewFocus();" in formatted_ready
        assert "cancelDocxPreviewScrollReset();cancelAnonymizedPreviewScrollReset();" in script
        assert "requestAnimationFrame(()=>{reviewFocusFrame=requestAnimationFrame(()=>{reviewFocusFrame=null;focusMark(reviewIndex);});})" in refresh_focus
        assert "setTimeout" not in refresh_focus
        assert "function holdReviewScrollSync()" in script
        assert "if(reviewScrollSyncLocked){programmaticPreviewScrollTargets.delete(source);return;}" in script
        assert "if(reviewScrollSyncLocked)return;" in script
        assert "refreshFormattedReviewFocus();return;" in script


def test_manual_review_candidate_removes_reason_copy_and_reclaims_its_column_in_both_languages():
    for name in ("index-v1.8.18.html", "index-en-v1.8.18.html"):
        html = asset_path(name).read_text(encoding="utf-8")
        candidate = re.search(r'<div class="review-candidate">(.*?)</div>', html, re.S)
        assert candidate is not None
        content = candidate.group(1)
        assert "<small>" not in content
        assert re.search(
            r'<span class="review-label">.*?</span>\s*'
            r'<span id="reviewCurrent"[^>]*>.*?</span>\s*'
            r'<span id="reviewMeta" class="review-help"></span>\s*$',
            content,
            re.S,
        )

    css = asset_path("new-design-v1.11.css").read_text(encoding="utf-8")
    assert "grid-template-columns:auto minmax(0,2fr) minmax(0,1fr);gap:4px 10px" in css
    assert ".review-candidate small" not in css


def test_production_manual_review_dom_reclaims_reason_column_at_responsive_widths(tmp_path: Path):
    base_css_uri = asset_path("style-v1.1.css").as_uri()
    css_uri = asset_path("new-design-v1.11.css").as_uri()
    edge = next(
        (
            candidate
            for candidate in (
                Path(r"C:\Program Files (x86)\Microsoft\Edge\Application\msedge.exe"),
                Path(r"C:\Program Files\Microsoft\Edge\Application\msedge.exe"),
            )
            if candidate.exists()
        ),
        None,
    )
    if edge is None:
        pytest.skip("Microsoft Edge is required for the Windows layout regression")
    cases = (
        ("index-v1.8.18.html", "«Синтетический фрагмент» · возможное имя", "2 вхожд. · Правила: «Общий»"),
        ("index-en-v1.8.18.html", '"Synthetic fragment" · possible name', '2 occurrences · Rules: "General"'),
    )

    for index, (html_name, current_text, meta_text) in enumerate(cases):
        source = asset_path(html_name).read_text(encoding="utf-8")
        manual_review = re.search(
            r'<section id="manualReview".*?</section>', source, re.S
        )
        assert manual_review is not None
        fixture = tmp_path / f"production-review-{index}.html"
        fixture.write_text(
            f'''<!doctype html><html><head><meta charset="utf-8"><link rel="stylesheet" href="{base_css_uri}"><link rel="stylesheet" href="{css_uri}"></head>
<body data-design="new"><main id="workspace" class="workspace-column">{manual_review.group(0)}</main><pre id="probe"></pre><script>
const rows=[];const workspace=document.getElementById('workspace');const candidate=document.querySelector('.review-candidate');
document.getElementById('reviewCurrent').textContent={json.dumps(current_text)};document.getElementById('reviewMeta').textContent={json.dumps(meta_text)};
for(const width of [390,700,900]){{workspace.style.width=width+'px';void candidate.offsetHeight;const help=document.getElementById('reviewMeta'),style=getComputedStyle(help),box=candidate.getBoundingClientRect();rows.push({{width,hasReason:Boolean(candidate.querySelector('small')),childColumns:[...candidate.children].map(child=>getComputedStyle(child).gridColumnStart),helpLines:Math.round(help.getBoundingClientRect().height/parseFloat(style.lineHeight)),candidateHeight:box.height,overflow:candidate.scrollWidth-candidate.clientWidth}})}}
document.getElementById('probe').textContent=JSON.stringify(rows);</script></body></html>''',
            encoding="utf-8",
        )
        profile = tmp_path / f"edge-profile-production-review-{index}"
        result = subprocess.run(
            [str(edge), "--headless=new", "--disable-gpu", "--allow-file-access-from-files", f"--user-data-dir={profile}", "--window-size=1200,900", "--dump-dom", fixture.as_uri()],
            capture_output=True,
            text=True,
            timeout=45,
            check=True,
        )
        match = re.search(r'<pre id="probe">(.*?)</pre>', result.stdout, re.S)
        assert match, result.stdout[-2000:]
        metrics = json.loads(html_module.unescape(match.group(1)))
        assert [metric["width"] for metric in metrics] == [390, 700, 900]
        for metric in metrics:
            assert not metric["hasReason"], metric
            assert metric["overflow"] <= 1, metric
            assert metric["helpLines"] <= 2, metric
        assert metrics[-1]["childColumns"] == ["1", "2", "3"], metrics[-1]
        assert metrics[-1]["candidateHeight"] < 50, metrics[-1]


def test_review_candidate_keeps_active_height_for_off_review_and_auto_modes(tmp_path: Path):
    edge = next(
        (
            candidate
            for candidate in (
                Path(r"C:\Program Files (x86)\Microsoft\Edge\Application\msedge.exe"),
                Path(r"C:\Program Files\Microsoft\Edge\Application\msedge.exe"),
            )
            if candidate.exists()
        ),
        None,
    )
    if edge is None:
        pytest.skip("Microsoft Edge is required for the review-height regression")

    state_text = {
        "index-v1.8.18.html": {
            "review": ("«SyntheticCandidate» · возможное имя собственное", "12 вхожд. · Правила: «Общий»"),
            "auto": ("Автоматическое скрытие включено", ""),
            "off": ("Поиск спорных мест выключен", ""),
        },
        "index-en-v1.8.18.html": {
            "review": ("“SyntheticCandidate” · possible proper name", "12 matches · Rules: “General”"),
            "auto": ("Automatic hiding is enabled", ""),
            "off": ("Uncertain-item detection is off", ""),
        },
    }
    harness_template = r'''<!doctype html><meta charset="utf-8"><title>pending</title><iframe id="app" style="width:__WIDTH__px;height:700px;border:0"></iframe><pre id="probe">pending</pre><script>
const frame=document.getElementById('app');frame.src=__INDEX__;
function finish(value){document.getElementById('probe').textContent=JSON.stringify(value);document.title=value.ok?'PASS':'FAIL';}
frame.addEventListener('load',()=>setTimeout(()=>{try{
 const d=frame.contentDocument,candidate=d.querySelector('.review-candidate'),toolbar=d.querySelector('.review-toolbar'),current=d.getElementById('reviewCurrent'),meta=d.getElementById('reviewMeta'),states=__STATES__,metrics={};
 for(const [mode,text] of Object.entries(states)){
   current.textContent=text[0];meta.textContent=text[1];
   const box=candidate.getBoundingClientRect(),toolbarBox=toolbar.getBoundingClientRect();
   metrics[mode]={height:box.height,toolbarTop:toolbarBox.top,verticalOverflow:candidate.scrollHeight-candidate.clientHeight,horizontalOverflow:candidate.scrollWidth-candidate.clientWidth};
 }
 const heights=Object.values(metrics).map(value=>value.height),tops=Object.values(metrics).map(value=>value.toolbarTop);
 const ok=Math.max(...heights)-Math.min(...heights)<=0.5&&Math.max(...tops)-Math.min(...tops)<=0.5&&Object.values(metrics).every(value=>value.verticalOverflow<=1&&value.horizontalOverflow<=1);
 finish({ok,metrics});
}catch(error){finish({ok:false,error:String(error)});}},800));setTimeout(()=>{if(document.title==='pending')finish({ok:false,error:'timeout'});},4000);</script>'''

    for language, texts in state_text.items():
        for workspace_width in (390, 700, 900):
            fixture = tmp_path / f"review-height-{language}-{workspace_width}.html"
            fixture.write_text(
                harness_template.replace("__INDEX__", json.dumps(asset_path(language).as_uri()))
                .replace("__WIDTH__", str(workspace_width))
                .replace("__STATES__", json.dumps(texts, ensure_ascii=False)),
                encoding="utf-8",
            )
            profile = tmp_path / f"edge-profile-review-height-{language}-{workspace_width}"
            result = subprocess.run(
                [str(edge), "--headless=new", "--disable-gpu", "--allow-file-access-from-files", f"--user-data-dir={profile}", "--virtual-time-budget=5000", "--dump-dom", fixture.as_uri()],
                capture_output=True,
                text=True,
                timeout=45,
                check=True,
            )
            match = re.search(r'<pre id="probe">(.*?)</pre>', result.stdout, re.S)
            assert match, result.stdout[-2000:]
            outcome = json.loads(html_module.unescape(match.group(1)))
            assert outcome["ok"], {"language": language, "width": workspace_width, **outcome}

    for script_name in ("app-v1.8.18.js", "app-en-v1.8.18.js"):
        script = asset_path(script_name).read_text(encoding="utf-8")
        update_review = script[script.index("function updateReviewPosition()") : script.index("function restoreReplacement(")]
        assert "reviewDisabledSummary" in update_review
        assert ".title=hasCurrent?'':reason" in update_review
        assert ".ariaLabel=hasCurrent?reviewCurrent.textContent:reason" in update_review


def test_production_workspace_lower_panels_share_one_horizontal_line(tmp_path: Path):
    base_css_uri = asset_path("style-v1.1.css").as_uri()
    css_uri = asset_path("new-design-v1.11.css").as_uri()
    edge = next(
        (
            candidate
            for candidate in (
                Path(r"C:\Program Files (x86)\Microsoft\Edge\Application\msedge.exe"),
                Path(r"C:\Program Files\Microsoft\Edge\Application\msedge.exe"),
            )
            if candidate.exists()
        ),
        None,
    )
    if edge is None:
        pytest.skip("Microsoft Edge is required for the Windows layout regression")

    for index, (html_name, javascript_name) in enumerate(
        (("index-v1.8.18.html", "app-v1.8.18.js"), ("index-en-v1.8.18.html", "app-en-v1.8.18.js"))
    ):
        source = asset_path(html_name).read_text(encoding="utf-8")
        workspace = re.search(r'<main class="columns">.*?</main>', source, re.S)
        assert workspace is not None
        javascript = asset_path(javascript_name).read_text(encoding="utf-8")
        start = javascript.index("const sourcePreviewAlignment=")
        end = javascript.index("function syncFormattedSourceScroll", start)
        implementation = javascript[start:end]
        for width in (390, 700, 900):
            fixture = tmp_path / f"production-workspace-line-{index}-{width}.html"
            fixture.write_text(
                f'''<!doctype html><html><head><meta charset="utf-8"><link rel="stylesheet" href="{base_css_uri}"><link rel="stylesheet" href="{css_uri}"><style>.columns{{width:{width * 2 + 14}px}}</style></head><body data-design="new">{workspace.group(0)}<pre id="probe"></pre><script>
const $=selector=>document.querySelector(selector);let previewFullscreen=false,anonymizedPreviewFullscreen=false;
{implementation}
document.getElementById('chips').innerHTML='<span>TYPE-A</span><span>TYPE-B</span><span>TYPE-C</span><span>TYPE-D</span><span>TYPE-E</span>';
document.getElementById('source').value='Synthetic documentation content';document.getElementById('richOutput').textContent='[SYNTHETIC-1]';syncPreviewTopAlignment();
document.getElementById('undoChips').textContent='Synthetic reversible marker';
setTimeout(()=>{{const box=element=>element.getBoundingClientRect(),sourcePane=document.querySelector('#sourceColumn>.pane'),outputPane=document.querySelector('#anonymizedColumn>.pane'),settings=document.getElementById('anonymizationSettings'),review=document.getElementById('manualReview');document.getElementById('probe').textContent=JSON.stringify({{sourceWidth:document.getElementById('sourceColumn').clientWidth,outputWidth:document.getElementById('anonymizedColumn').clientWidth,paneBottomDifference:Math.abs(box(sourcePane).bottom-box(outputPane).bottom),lowerTopDifference:Math.abs(box(settings).top-box(review).top),settingsGap:box(settings).top-box(sourcePane).bottom,reviewGap:box(review).top-box(outputPane).bottom,sourcePanelSpacer:parseFloat(document.getElementById('sourcePanelAlignment').style.height)||0,outputPanelSpacer:parseFloat(document.getElementById('outputPanelAlignment').style.height)||0,dynamicUndoVisible:box(document.getElementById('undoChips')).height>0}});}},120);</script></body></html>''',
                encoding="utf-8",
            )
            profile = tmp_path / f"edge-profile-workspace-line-{index}-{width}"
            result = subprocess.run(
                [str(edge), "--headless=new", "--disable-gpu", "--allow-file-access-from-files", f"--user-data-dir={profile}", "--window-size=1200,1000", "--virtual-time-budget=500", "--dump-dom", fixture.as_uri()],
                capture_output=True,
                text=True,
                timeout=45,
                check=True,
            )
            match = re.search(r'<pre id="probe">(.*?)</pre>', result.stdout, re.S)
            assert match, result.stdout[-2000:]
            metric = json.loads(html_module.unescape(match.group(1)))
            assert abs(metric["sourceWidth"] - width) <= 1, metric
            assert abs(metric["outputWidth"] - width) <= 1, metric
            assert metric["paneBottomDifference"] <= 1, metric
            assert metric["lowerTopDifference"] <= 1, metric
            assert metric["dynamicUndoVisible"], metric
            assert metric["settingsGap"] >= 0 and metric["reviewGap"] >= 0, metric
            assert metric["sourcePanelSpacer"] >= 0 and metric["outputPanelSpacer"] >= 0, metric


def test_manual_review_layout_is_bound_to_workspace_width_and_contains_long_copy():
    css = asset_path("new-design-v1.11.css").read_text(encoding="utf-8")

    assert 'body[data-design="new"] .workspace-column{container-type:inline-size;' in css
    assert 'body[data-design="new"] .review-toolbar{display:grid;grid-template-columns:repeat(4,minmax(0,1fr));gap:6px;align-items:stretch}' in css
    assert 'body[data-design="new"] .review-navigation,body[data-design="new"] .review-actions{display:contents}' in css
    assert 'body[data-design="new"] .review-toolbar button{width:100%;min-width:0;min-height:34px;margin:0;padding:5px 8px;white-space:normal;overflow-wrap:anywhere;line-height:1.2}' in css
    assert '@container(max-width:719px)' in css
    assert '@container(max-width:519px)' in css
    assert 'body[data-design="new"] .review{grid-template-columns:minmax(0,1fr)}' in css


def test_manual_review_browser_geometry_has_no_panel_overlap(tmp_path):
    edge_candidates = (
        Path(r"C:\Program Files (x86)\Microsoft\Edge\Application\msedge.exe"),
        Path(r"C:\Program Files\Microsoft\Edge\Application\msedge.exe"),
    )
    edge = next((candidate for candidate in edge_candidates if candidate.exists()), None)
    if edge is None:
        pytest.skip("Microsoft Edge is required for the Windows layout regression")

    links = "".join(
        f'<link rel="stylesheet" href="{asset_path(name).as_uri()}">'
        for name in ("style-v1.1.css", "new-design-v1.11.css")
    )
    localized = (
        ("ru", "Что проверяется", "Синтетический спорный фрагмент с длинным безопасным значением", "Почему требуется решение", "Скрыть и запомнить правило", "Оставить и добавить в исключения"),
        ("en", "What is checked", "Synthetic ambiguous fragment with a long safe documentation value", "Why a decision is required", "Hide and remember the rule", "Keep and add to exceptions"),
    )

    for width in (390, 700, 900):
        cases = []
        for language, label, value, help_text, hide_text, keep_text in localized:
            cases.append(
                f'<div class="workspace-column case" data-language="{language}" style="width:{width}px">'
                '<section class="review">'
                '<div class="review-heading"><div><h2 tabindex="0" aria-describedby="syntheticReviewHelp" data-tooltip="Choose one action for the synthetic fragment.">Manual review <span class="review-heading-help" aria-hidden="true">?</span></h2><span id="syntheticReviewHelp" class="sr-only">Choose one action for the synthetic fragment.</span></div><span class="review-progress">2 remaining</span></div>'
                '<div class="warning sr-only">Review every synthetic highlight.</div>'
                '<div class="review-context"><div class="review-candidate">'
                f'<b class="review-label">{label}</b><span id="reviewCurrent">{value}</span>'
                f'<span class="review-help">2 matches · Documentation ruleset</span>'
                '</div></div>'
                '<div class="review-toolbar"><div class="review-navigation"><button>Previous</button><button>Next</button></div>'
                '<div class="review-actions">'
                f'<button>{hide_text}</button><button>{keep_text}</button>'
                '</div></div>'
                '<div class="review-controls-disclosure"><fieldset class="review-controls" aria-label="Synthetic controls"><label><input type="checkbox" checked>Review uncertain items</label><label><input type="checkbox">Hide automatically</label></fieldset></div>'
                '<p class="review-decision sr-only">The current choice applies only to the selected synthetic fragment.</p>'
                '<p class="review-finish">Resolve every fragment before continuing.</p>'
                '<p id="resultStatus" class="review-decision">Synthetic result status.</p>'
                '</section></div>'
            )
        probe = tmp_path / f"manual-review-layout-{width}.html"
        probe.write_text(
            '<!doctype html><html><head><meta charset="utf-8">' + links
            + '<style>body{padding:8px}.case{margin-bottom:16px}</style></head><body data-design="new">'
            + "".join(cases)
            + '<pre id="probe"></pre><script>'
            + "const box=e=>e.getBoundingClientRect();"
            + "const inside=(child,parent)=>{const c=box(child),p=box(parent);return c.left>=p.left-1&&c.right<=p.right+1&&c.top>=p.top-1&&c.bottom<=p.bottom+1};"
            + "const overlaps=(a,b)=>{const x=box(a),y=box(b);return Math.min(x.right,y.right)-Math.max(x.left,y.left)>1&&Math.min(x.bottom,y.bottom)-Math.max(x.top,y.top)>1};"
            + "const result=[...document.querySelectorAll('.case')].map(root=>{const review=root.querySelector('.review'),context=root.querySelector('.review-context'),toolbar=root.querySelector('.review-toolbar'),disclosure=root.querySelector('.review-controls-disclosure'),finish=root.querySelector('.review-finish'),candidate=root.querySelector('.review-candidate'),heading=root.querySelector('h2'),buttons=[...root.querySelectorAll('.review-toolbar button')],direct=[...review.children].filter(e=>!e.classList.contains('sr-only')&&getComputedStyle(e).display!=='none'),visible=[...review.querySelectorAll('.review-heading,.review-context,.review-candidate,.review-candidate>*,.review-toolbar,.review-toolbar button,.review-finish,.review-controls-disclosure,#resultStatus')],buttonBoxes=buttons.map(box),buttonStyles=buttons.map(e=>{const s=getComputedStyle(e);return [s.backgroundColor,s.borderColor,s.borderRadius,s.padding]}),candidateBox=box(candidate),reviewBox=box(review),help=root.querySelector('.review-help'),helpStyle=getComputedStyle(help),helpLineCount=Math.round(box(help).height/parseFloat(helpStyle.lineHeight)),tooltipHidden=getComputedStyle(heading,'::after').display==='none';heading.focus();const tooltipVisibleOnFocus=getComputedStyle(heading,'::after').display==='block',tooltipWidth=parseFloat(getComputedStyle(heading,'::after').width);heading.blur();return {language:root.dataset.language,width:root.clientWidth,reviewOverflow:review.scrollWidth>review.clientWidth+1,contained:visible.every(e=>inside(e,review)),candidateContained:[...candidate.children].every(e=>inside(e,candidate)),candidateFullWidth:Math.abs(candidateBox.left-(reviewBox.left+10))<=1&&Math.abs(candidateBox.right-(reviewBox.right-10))<=1,candidateHeight:candidateBox.height,directChildrenDisjoint:direct.every((e,i)=>direct.slice(i+1).every(other=>!overlaps(e,other))),buttonHeights:buttonBoxes.map(b=>b.height),buttonTops:buttonBoxes.map(b=>b.top),uniformButtonStyle:buttonStyles.every(s=>JSON.stringify(s)===JSON.stringify(buttonStyles[0])),topOverlap:overlaps(context,toolbar),lowerOverlap:overlaps(finish,context)||overlaps(finish,toolbar)||overlaps(finish,disclosure),toolbarBelowCandidate:box(toolbar).top>=box(context).bottom-1,additionalModesBelowToolbar:box(disclosure).top>=box(toolbar).bottom-1,promptsHidden:[...review.querySelectorAll('.warning,.review-decision.sr-only')].every(e=>e.classList.contains('sr-only')),helpLineCount,tooltipHidden,tooltipVisibleOnFocus,tooltipFitsViewport:tooltipWidth<=document.documentElement.clientWidth-24};});"
            + "document.querySelector('#probe').textContent=JSON.stringify(result);</script></body></html>",
            encoding="utf-8",
        )
        profile = tmp_path / f"edge-profile-manual-review-{width}"
        result = subprocess.run(
            [str(edge), "--headless=new", "--disable-gpu", "--allow-file-access-from-files", f"--user-data-dir={profile}", "--window-size=1200,1000", "--dump-dom", probe.as_uri()],
            capture_output=True,
            text=True,
            timeout=45,
            check=True,
        )
        match = re.search(r'<pre id="probe">(.*?)</pre>', result.stdout, re.S)
        assert match, result.stdout[-2000:]
        metrics = json.loads(html_module.unescape(match.group(1)))
        assert len(metrics) == 2
        for metric in metrics:
            assert not metric["reviewOverflow"], metric
            assert metric["contained"], metric
            assert metric["candidateContained"], metric
            assert metric["candidateFullWidth"], metric
            assert metric["directChildrenDisjoint"], metric
            assert max(metric["buttonHeights"]) < 72, metric
            assert metric["uniformButtonStyle"], metric
            assert not metric["topOverlap"], metric
            assert not metric["lowerOverlap"], metric
            assert metric["toolbarBelowCandidate"], metric
            assert metric["additionalModesBelowToolbar"], metric
            assert metric["promptsHidden"], metric
            assert metric["helpLineCount"] <= 2, metric
            assert metric["tooltipHidden"], metric
            assert metric["tooltipVisibleOnFocus"], metric
            assert metric["tooltipFitsViewport"], metric
            if metric["width"] >= 700:
                assert max(metric["buttonTops"]) - min(metric["buttonTops"]) <= 1, metric
                assert max(metric["buttonHeights"]) - min(metric["buttonHeights"]) <= 1, metric
            if metric["width"] == 900:
                assert metric["candidateHeight"] < 50, metric


def test_document_preview_errors_are_reported_in_both_interfaces():
    russian = asset_path("app-v1.8.18.js").read_text(encoding="utf-8")
    english = asset_path("app-en-v1.8.18.js").read_text(encoding="utf-8")

    assert "Ошибка предпросмотра:" in russian
    assert "Preview error:" in english
    assert "catch(error)" in russian
    assert "catch(error)" in english


def test_generated_markers_use_native_buttons():
    javascript = asset_path("app-v1.8.18.js").read_text(encoding="utf-8")

    assert 'class="maprow"' not in javascript
    assert '<button type="button" class="mark-hint' in javascript
    assert '<button type="button" class="${p.secret?' in javascript
    assert 'role="button" tabindex="0"' not in javascript
    assert "activateKeyboardAction" not in javascript


def test_fullscreen_preview_uses_modal_semantics_and_restores_focus():
    html = asset_path("index-v1.8.18.html").read_text(encoding="utf-8")
    javascript = asset_path("app-v1.8.18.js").read_text(encoding="utf-8")

    assert 'id="docxPreviewSurface"' in html
    assert 'id="anonymizedPreviewSurface"' in html
    assert "setAttribute('role','dialog')" in javascript
    assert "ariaModal" in javascript
    assert "function setModalIsolation" in javascript
    assert "function trapModalFocus" in javascript
    assert "previousInert" in javascript
    assert "fullscreenReturnFocus" in javascript


def test_desktop_window_allows_760px_accessibility_layout():
    source = Path(app_module.__file__).read_text(encoding="utf-8")
    assert "min_size=(720, 700)" in source


def test_desktop_window_opens_at_1600px_for_two_preview_panes(tmp_path, monkeypatch):
    captured = {}

    def create_window(title, url, **kwargs):
        captured.update(kwargs)
        return object()

    monkeypatch.setenv("LOCALAPPDATA", str(tmp_path))
    monkeypatch.setitem(
        sys.modules,
        "webview",
        SimpleNamespace(create_window=create_window, start=lambda **kwargs: None),
    )

    assert app_module.main([]) == 0
    assert captured["width"] == 1600
    assert captured["height"] == 920
    assert captured["min_size"] == (720, 700)


def test_desktop_window_clamps_preferred_size_to_primary_work_area(tmp_path, monkeypatch):
    captured = {}
    work_area = SimpleNamespace(X=0, Y=0, Width=1366, Height=728)
    primary_screen = SimpleNamespace(
        x=0,
        y=0,
        width=1366,
        height=768,
        frame=work_area,
    )

    def create_window(title, url, **kwargs):
        captured.update(kwargs)
        return object()

    monkeypatch.setenv("LOCALAPPDATA", str(tmp_path))
    monkeypatch.setitem(
        sys.modules,
        "webview",
        SimpleNamespace(
            screens=[primary_screen],
            create_window=create_window,
            start=lambda **kwargs: None,
        ),
    )

    assert app_module.main([]) == 0
    assert captured["width"] == 1366
    assert captured["height"] == 728
    assert captured["min_size"] == (720, 700)
    assert captured["x"] == 0
    assert captured["y"] == 0
    assert captured["screen"] is primary_screen


def test_double_escape_no_longer_silently_clears_user_work():
    javascript = asset_path("app-v1.8.18.js").read_text(encoding="utf-8")

    assert "now-lastEscape<800" not in javascript
    assert "confirm('Очистить введённый текст" in javascript


def test_ui_uses_current_design_without_legacy_selector():
    russian_html = asset_path("index-v1.8.18.html").read_text(encoding="utf-8")
    english_html = asset_path("index-en-v1.8.18.html").read_text(encoding="utf-8")
    russian_javascript = asset_path("app-v1.8.18.js").read_text(encoding="utf-8")
    english_javascript = asset_path("app-en-v1.8.18.js").read_text(encoding="utf-8")

    for html in (russian_html, english_html):
        assert 'id="designMode"' not in html
        assert 'class="design-switch"' not in html
        assert '<body data-design="new">' in html
    for javascript in (russian_javascript, english_javascript):
        assert "function applyCurrentDesign" in javascript
        assert "document.body.dataset.design='new'" in javascript
        assert "pywebview.api.set_design_mode('new')" in javascript
        assert "try{await pywebview.api.set_design_mode('new')}catch(_error){}" in javascript
        assert "DESIGN_KEY" not in javascript
        assert "await applyCurrentDesign();" in javascript
        assert "get_design_mode" not in javascript
        assert "$('#designMode')" not in javascript


def test_ui_offers_complete_persistent_english_interface():
    russian_html = asset_path("index-v1.8.18.html").read_text(encoding="utf-8")
    english_html = asset_path("index-en-v1.8.18.html").read_text(encoding="utf-8")
    russian_js = asset_path("app-v1.8.18.js").read_text(encoding="utf-8")
    english_js = asset_path("app-en-v1.8.18.js").read_text(encoding="utf-8")

    assert '<html lang="ru">' in russian_html
    assert '<html lang="en">' in english_html
    assert '<script src="app-v1.8.18.js"></script>' in russian_html
    assert '<script src="app-en-v1.8.18.js"></script>' in english_html
    assert 'id="languageMode"' in russian_html and 'id="languageMode"' in english_html
    assert "Anonymize text or a document" in english_html
    assert "Manual review of uncertain items" in english_html
    assert "Local library" in english_html
    assert 'aria-label="Anonymized text"' in english_html
    assert "pywebview.api.set_language(value)" in russian_js
    assert "pywebview.api.set_language(value)" in english_js

    id_pattern = re.compile(r'id="([^"]+)"')
    assert set(id_pattern.findall(russian_html)) == set(id_pattern.findall(english_html))


def test_english_interface_localizes_backend_statuses_profile_and_review_kinds():
    javascript = asset_path("app-en-v1.8.18.js").read_text(encoding="utf-8")

    assert "function documentStatusText(status)" in javascript
    assert "`${d.display_name} — ${documentStatusText(d.status)}`" in javascript
    assert "`Preview: ${preview.display_name} — ${documentStatusText(preview.status)}`" in javascript
    assert "function profileDisplayName(name)" in javascript
    assert "o.textContent=profileDisplayName(n)" in javascript
    assert "function hintKindText(kind)" in javascript
    assert "hintKindText(p.kind)" in javascript
    assert "hintKindText(all[reviewIndex].dataset.kind)" in javascript


def test_english_interface_never_falls_back_to_raw_unknown_backend_status():
    javascript = asset_path("app-en-v1.8.18.js").read_text(encoding="utf-8")

    assert "||status;" not in javascript
    assert "'Unknown status'" in javascript


def test_english_interface_localizes_profile_names_in_result_messages():
    javascript = asset_path("app-en-v1.8.18.js").read_text(encoding="utf-8")

    assert "`Profile “${profileDisplayName(n)}” saved.`" in javascript
    assert "`Profile “${profileDisplayName(n)}” exported.`" in javascript
    assert "`Profile “${profileDisplayName(r.name)}” imported and loaded." in javascript


def test_english_interface_localizes_backend_warning_messages_without_changing_contract_keys():
    javascript = asset_path("app-en-v1.8.18.js").read_text(encoding="utf-8")

    for warning_kind in (
        "ГОСТАЙНА", "ДСП", "КИИ", "ЗАКУПКИ", "МАССИВ-ПДН",
        "XLSX_FORMULA_BLOCKED", "XLSX_HYPERLINK", "XLSX_HEADER_FOOTER",
        "XLSX_DATA_VALIDATION", "XLSX_DEFINED_NAME", "XLSX_EXTERNAL_LINK",
        "XLSX_UNSUPPORTED_PART", "UNSUPPORTED_OFFICE_CONTENT",
    ):
        assert f"'{warning_kind}'" in javascript
    assert "function warningText(w)" in javascript
    assert "Manual review required:" in javascript


def test_english_interface_does_not_show_russian_profile_api_errors():
    javascript = asset_path("app-en-v1.8.18.js").read_text(encoding="utf-8")

    assert "function profileErrorText(operation)" in javascript
    assert "profileErrorText('export')" in javascript
    assert "profileErrorText('import')" in javascript
    assert "textContent=r.error;" not in javascript


def test_english_interface_localizes_post_anonymization_status_and_review_profile():
    javascript = asset_path("app-en-v1.8.18.js").read_text(encoding="utf-8")

    assert "`${documentStatusText(r.status)}. The file was created" in javascript
    assert "`${documentStatusText(r.status)}. The anonymized file is ready" in javascript
    assert "const profileName=profileDisplayName($('#profile').value||'Общий')" in javascript
    assert 'const profileName=$(\'#profile\').value||\'Общий\'' not in javascript


def test_english_interface_never_renders_raw_bridge_errors():
    javascript = asset_path("app-en-v1.8.18.js").read_text(encoding="utf-8")

    assert "function bridgeErrorText(operation)" in javascript
    assert "saved?.error||'Could not save profile'" not in javascript
    assert "${error}" not in javascript
    assert "String(error)" not in javascript
    assert "${formattedPreview.error}" not in javascript
    assert "${error?.message||error}" not in javascript
    assert ".map(x=>x.error)" not in javascript
    for operation in ("profile-save", "formatted-preview", "anonymize", "import", "restore"):
        assert f"bridgeErrorText('{operation}')" in javascript


def test_language_variants_omit_redundant_workflow_steps():
    for name in ("index-v1.8.18.html", "index-en-v1.8.18.html"):
        html = asset_path(name).read_text(encoding="utf-8")
        assert 'id="workflowSteps"' not in html
        assert 'class="workflow-steps"' not in html


def test_language_choice_is_persisted_by_backend_and_preserves_other_ui_settings(tmp_path):
    root = tmp_path / "app"
    first = AppApi(root=root)
    assert first.get_language() == "ru"
    assert first.set_theme("dark") == {"ok": True, "theme": "dark"}
    assert first.set_language("en") == {"ok": True, "language": "en"}

    second = AppApi(root=root)
    assert second.get_language() == "en"
    assert second.get_theme() == "dark"
    with pytest.raises(ValueError, match="язык"):
        second.set_language("de")


def test_rare_categories_are_forced_closed_on_every_launch():
    javascript = asset_path("app-v1.8.18.js").read_text(encoding="utf-8")

    assert "$('#detectionCategories').open=false" in javascript


def test_gui_entry_url_is_versioned_to_bypass_stale_webview_cache(tmp_path, monkeypatch):
    captured = {}

    def create_window(title, url, **kwargs):
        captured["url"] = url
        return object()

    fake_webview = SimpleNamespace(
        create_window=create_window,
        start=lambda **kwargs: captured.setdefault("start", kwargs),
    )
    monkeypatch.setenv("LOCALAPPDATA", str(tmp_path))
    monkeypatch.setitem(sys.modules, "webview", fake_webview)

    assert app_module.main([]) == 0
    assert captured["url"].endswith("/index-v1.8.18.html")
    html = app_module.asset_path(app_module.INDEX_ASSET).read_text(encoding="utf-8")
    assert '<script src="app-v1.8.18.js"></script>' in html


def test_gui_starts_in_the_persisted_english_language(tmp_path, monkeypatch):
    captured = {}

    def create_window(title, url, **kwargs):
        captured.update(title=title, url=url)
        return object()

    monkeypatch.setenv("LOCALAPPDATA", str(tmp_path))
    AppApi().set_language("en")
    monkeypatch.setitem(
        sys.modules,
        "webview",
        SimpleNamespace(create_window=create_window, start=lambda **kwargs: None),
    )

    assert app_module.main([]) == 0
    assert captured["url"].endswith("/index-en-v1.8.18.html")
    assert captured["title"] == "Local Document Anonymizer"


def test_clipboard_read_is_exposed_through_native_api(tmp_path, monkeypatch):
    monkeypatch.setattr(app_module, "_read_clipboard_text", lambda: "ВСТАВКА-АУДИТ")

    assert AppApi(root=tmp_path).read_clipboard_text() == {"text": "ВСТАВКА-АУДИТ"}


def test_app_api_exposes_compact_marker_mode(tmp_path):
    api = AppApi(root=tmp_path / "app")

    result = api.anonymize_text(
        "Иванов Иван Иванович работает в ООО «Ромашка».",
        compact_markers=True,
    )

    assert "[Ф-1]" in result["text"]
    assert "[ОРГ-1]" in result["text"]
    assert result["mapping"]["[Ф-1]"] == "Иванов Иван Иванович"


def test_compact_marker_checkbox_is_wired_to_all_anonymization_paths():
    html = asset_path("index-v1.8.18.html").read_text(encoding="utf-8")
    javascript = asset_path("app-v1.8.18.js").read_text(encoding="utf-8")

    assert 'id="compactMarkers"' in html
    assert "Сохранять компоновку" in html
    assert "$('#compactMarkers').checked" in javascript
    assert "'#compactMarkers'" in javascript


def test_paste_button_uses_native_api_instead_of_webview_clipboard_permission():
    javascript = asset_path("app-v1.8.18.js").read_text(encoding="utf-8")

    assert "pywebview.api.read_clipboard_text()" in javascript
    assert "navigator.clipboard.readText()" not in javascript


def test_source_and_anonymized_previews_scroll_together_by_logical_blocks():
    for name in ("app-v1.8.18.js", "app-en-v1.8.18.js"):
        javascript = asset_path(name).read_text(encoding="utf-8")

        # Text and formatted previews both synchronize in either direction.
        assert "function syncPreviewScroll(source,target)" in javascript
        assert "programmaticPreviewScrollTargets.get(source)" in javascript
        assert "programmaticPreviewScrollTargets.set(target" in javascript
        assert "function collectPreviewAnchors(container)" in javascript
        assert "function mapPreviewScroll(source,target)" in javascript
        assert "$('#docxPreview').addEventListener('scroll',syncFormattedSourceScroll)" in javascript
        assert "$('#anonymizedDocxPreview').addEventListener('scroll',syncFormattedAnonymizedScroll)" in javascript

        # Formatted documents use matching structural blocks and interpolation,
        # with a proportional fallback for incomplete or unusual documents.
        assert "'.docx-paragraph,li,tr'" in javascript
        assert "function previewScrollPosition(source)" in javascript
        assert "function previewScrollForPosition(position,target)" in javascript

        # Programmatic scrolling must not feed back into the opposite listener.
        assert "let previewScrollSyncDepth=0" in javascript
        assert "if(previewScrollSyncDepth||!source||!target" in javascript
        assert "const mapped=mapPreviewScroll(source,target)" in javascript
        assert "finally{previewScrollSyncDepth--;}" in javascript


def test_source_and_anonymized_preview_screens_stay_vertically_aligned():
    for html_name, javascript_name in (
        ("index-v1.8.18.html", "app-v1.8.18.js"),
        ("index-en-v1.8.18.html", "app-en-v1.8.18.js"),
    ):
        html = asset_path(html_name).read_text(encoding="utf-8")
        javascript = asset_path(javascript_name).read_text(encoding="utf-8")

        # Each pane has an independent spacer. Runtime geometry, rather than a
        # language-specific pixel offset, aligns whichever active previews are visible.
        assert 'id="sourcePreviewAlignment" class="preview-alignment-spacer"' in html
        assert 'id="outputPreviewAlignment" class="preview-alignment-spacer"' in html
        assert 'id="sourcePanelAlignment" class="pane-bottom-alignment-spacer"' in html
        assert 'id="outputPanelAlignment" class="pane-bottom-alignment-spacer"' in html
        assert "function syncPreviewTopAlignment()" in javascript
        assert "sourcePreviewAlignment.style.height='0px'" in javascript
        assert "outputPreviewAlignment.style.height='0px'" in javascript
        assert "sourcePanelAlignment.style.height='0px'" in javascript
        assert "outputPanelAlignment.style.height='0px'" in javascript
        assert "activeSourcePreview()" in javascript
        assert "activeAnonymizedPreview()" in javascript
        assert "function isPreviewActive(preview)" in javascript
        assert "getBoundingClientRect().top" in javascript
        assert "new ResizeObserver(schedulePreviewTopAlignment)" in javascript
        assert "new MutationObserver(schedulePreviewTopAlignment)" in javascript
        assert "window.addEventListener('resize',schedulePreviewTopAlignment)" in javascript
        assert "setTimeout(run,100)" in javascript
        assert javascript.count("schedulePreviewTopAlignment();}") >= 2


@pytest.mark.parametrize("asset_name", ["app-v1.8.18.js", "app-en-v1.8.18.js"])
def test_preview_top_alignment_runtime_handles_dynamic_and_stacked_layouts(tmp_path, asset_name):
    import json
    import subprocess

    javascript = asset_path(asset_name).read_text(encoding="utf-8")
    start = javascript.index("const sourcePreviewAlignment=")
    end = javascript.index("function syncFormattedSourceScroll", start)
    implementation = javascript[start:end]
    harness = f"""
const vm=require('vm');
let sourceBase=200,outputBase=100;
const px=value=>Number.parseFloat(value||'0')||0;
const sourceSpacer={{style:{{height:'0px'}},classList:{{contains:name=>name==='preview-alignment-spacer'}}}};
const outputSpacer={{style:{{height:'0px'}},classList:{{contains:name=>name==='preview-alignment-spacer'}}}};
const sourcePanelSpacer={{style:{{height:'0px'}}}};
const outputPanelSpacer={{style:{{height:'0px'}}}};
const sourcePane={{top:0,naturalHeight:430,getBoundingClientRect(){{return{{top:this.top,bottom:this.top+this.naturalHeight+px(sourceSpacer.style.height)+px(sourcePanelSpacer.style.height)}};}},querySelector:()=>({{}})}};
const outputPane={{top:0,naturalHeight:438,getBoundingClientRect(){{return{{top:this.top,bottom:this.top+this.naturalHeight+px(outputSpacer.style.height)+px(outputPanelSpacer.style.height)}};}},querySelector:()=>({{}})}};
const source={{hidden:false,clientWidth:500,clientHeight:405,closest:()=>sourcePane,getBoundingClientRect:()=>({{top:sourceBase+px(sourceSpacer.style.height)}})}};
const output={{hidden:false,clientWidth:500,clientHeight:405,closest:()=>outputPane,getBoundingClientRect:()=>({{top:outputBase+px(outputSpacer.style.height)}})}};
const hidden={{hidden:true,clientWidth:500,clientHeight:405,closest:()=>sourcePane,getBoundingClientRect:()=>({{top:0}})}};
const generic={{classList:{{contains:()=>false}},closest:()=>sourcePane}};
const elements={{sourcePreviewAlignment:sourceSpacer,outputPreviewAlignment:outputSpacer,sourcePanelAlignment:sourcePanelSpacer,outputPanelAlignment:outputPanelSpacer,source,richOutput:output,docxPreview:hidden,anonymizedDocxPreview:hidden,sourceLibrary:generic,docxPreviewControls:generic,anonymizedPreviewControls:generic,textPreview:generic}};
const queued=[];const observed=[];let timer=null;
const context={{Boolean,Math,Number,previewFullscreen:false,anonymizedPreviewFullscreen:false,$:selector=>elements[selector.slice(1)],requestAnimationFrame:fn=>{{queued.push(fn);return queued.length;}},setTimeout:fn=>{{timer=fn;}},ResizeObserver:class{{observe(element){{observed.push(element);}}}},MutationObserver:class{{observe(){{}}}},window:{{addEventListener(){{}}}},document:{{fonts:null}}}};
vm.createContext(context);
vm.runInContext({json.dumps(implementation + "this.api={syncPreviewTopAlignment};")},context);
if(!observed.includes(sourcePane)||!observed.includes(outputPane))throw new Error('pane resize changes must schedule realignment');
if(!timer)throw new Error('missing throttled-animation fallback');
timer();
if(outputSpacer.style.height!=='100px')throw new Error('timer fallback did not align collapsed panes');
if(sourcePanelSpacer.style.height!=='108px'||outputPanelSpacer.style.height!=='0px')throw new Error('lower panels were not aligned after preview offset');
sourceBase=320;context.api.syncPreviewTopAlignment();
if(outputSpacer.style.height!=='220px')throw new Error('expanded library was not realigned');
if(sourcePanelSpacer.style.height!=='228px')throw new Error('lower panel alignment did not follow dynamic preview offset');
sourceBase=90;context.api.syncPreviewTopAlignment();
if(sourceSpacer.style.height!=='10px'||outputSpacer.style.height!=='0px')throw new Error('shorter source side was not aligned');
if(outputPanelSpacer.style.height!=='2px'||sourcePanelSpacer.style.height!=='0px')throw new Error('opposite lower panel was not aligned');
outputPane.top=500;context.api.syncPreviewTopAlignment();
if(sourceSpacer.style.height!=='0px'||outputSpacer.style.height!=='0px')throw new Error('stacked cards must not be cross-aligned');
if(sourcePanelSpacer.style.height!=='0px'||outputPanelSpacer.style.height!=='0px')throw new Error('stacked cards must not receive artificial panel height');
"""
    script = tmp_path / f"{asset_name}.alignment-runtime.js"
    script.write_text(harness, encoding="utf-8")
    completed = subprocess.run(["node", str(script)], capture_output=True, text=True)
    assert completed.returncode == 0, completed.stderr


@pytest.mark.parametrize("asset_name", ["app-v1.8.18.js", "app-en-v1.8.18.js"])
def test_preview_scroll_runtime_handles_unequal_blocks_cache_and_hidden_panes(tmp_path, asset_name):
    import json
    import subprocess

    javascript = asset_path(asset_name).read_text(encoding="utf-8")
    start = javascript.index("let previewScrollSyncDepth=0;")
    end = javascript.index("function setDocxPreviewMode", start)
    implementation = javascript[start:end]
    harness = f"""
const vm=require('vm');
function node(top){{return{{getBoundingClientRect:()=>({{top}})}};}}
function pane({{height=100,client=20,top=0,nodes=[],hidden=false}}={{}}){{
  return{{scrollHeight:height,clientHeight:client,scrollTop:top,hidden,
    getBoundingClientRect:()=>({{top:0}}),querySelectorAll:()=>nodes}};
}}
const elements={{docxPreview:pane(),anonymizedDocxPreview:pane(),source:pane(),richOutput:pane()}};
const context={{WeakMap,WeakSet,Math,requestAnimationFrame:(fn)=>fn(),$:s=>elements[s.slice(1)],window:{{addEventListener:()=>{{}}}},document:{{fonts:null}}}};
vm.createContext(context);
vm.runInContext({json.dumps(implementation + "function setDocxPreviewMode(formatted){$('#docxPreview').hidden=!formatted;$('#source').hidden=formatted;}function setAnonymizedPreviewMode(formatted){$('#anonymizedDocxPreview').hidden=!formatted;$('#richOutput').hidden=formatted;}this.api={mapPreviewScroll,collectPreviewAnchors,invalidatePreviewAnchors,syncFormattedSourceScroll,syncFormattedAnonymizedScroll,switchPreviewMode};")},context);
const source=pane({{height:120,client:20,top:50,nodes:[node(-30),node(0),node(30)]}});
const target=pane({{height:220,client:20,nodes:[node(40)]}});
if(context.api.mapPreviewScroll(source,target)!==100)throw new Error('unequal anchors must use proportional fallback');
const cached=pane({{height:120,client:20,nodes:[node(20)]}});
if(context.api.collectPreviewAnchors(cached)[1]!==20)throw new Error('initial anchor');
cached.nodes=[node(70)];cached.querySelectorAll=()=>cached.nodes;
context.api.invalidatePreviewAnchors(cached);
if(context.api.collectPreviewAnchors(cached)[1]!==70)throw new Error('stale anchor cache');
elements.docxPreview.hidden=true;elements.docxPreview.scrollTop=40;elements.anonymizedDocxPreview.scrollTop=33;
context.api.syncFormattedSourceScroll();
if(elements.anonymizedDocxPreview.scrollTop!==33)throw new Error('hidden pane changed active peer');
elements.source=pane({{height:100,client:20,top:40}});elements.source.hidden=false;
elements.docxPreview=pane({{height:200,client:20,top:0,hidden:true}});
context.api.switchPreviewMode(context.setDocxPreviewMode,true);
if(elements.docxPreview.hidden||!elements.source.hidden||elements.docxPreview.scrollTop!==90)throw new Error('mode switch lost visible position');
elements.anonymizedDocxPreview=pane({{height:200,client:20,top:0}});
context.api.syncFormattedSourceScroll();
if(elements.anonymizedDocxPreview.scrollTop!==0)throw new Error('mode restore feedback was not ignored');
elements.docxPreview.scrollTop=100;
context.api.syncFormattedSourceScroll();
if(elements.anonymizedDocxPreview.scrollTop!==100)throw new Error('forward formatted synchronization failed');
context.api.syncFormattedAnonymizedScroll();
if(elements.docxPreview.scrollTop!==100)throw new Error('programmatic feedback changed the source');
elements.anonymizedDocxPreview.scrollTop=135;
context.api.syncFormattedAnonymizedScroll();
if(elements.docxPreview.scrollTop!==135)throw new Error('real reverse scroll was swallowed');
"""
    script = tmp_path / "scroll_sync_runtime.js"
    script.write_text(harness, encoding="utf-8")
    completed = subprocess.run(["node", str(script)], capture_output=True, text=True)
    assert completed.returncode == 0, completed.stderr


def test_preview_mode_changes_preserve_visible_position_and_invalidate_dom_cache():
    for name in ("app-v1.8.18.js", "app-en-v1.8.18.js"):
        javascript = asset_path(name).read_text(encoding="utf-8")
        assert "function switchPreviewMode" in javascript
        assert "switchPreviewMode(setAnonymizedPreviewMode,true)" in javascript
        assert "rich.innerHTML=formattedPreview.rich_html||'';rich.scrollTop=0;" in javascript
        assert "resetAnonymizedPreviewScroll();$('#anonymizedPreviewCompleteness')" not in javascript
        assert "invalidatePreviewAnchors(rich);rich.innerHTML=" in javascript
        assert "if($('#docxPreview').hidden)return;" in javascript
        assert "if($('#anonymizedDocxPreview').hidden)return;" in javascript


def test_every_button_has_a_meaningful_hover_tooltip():
    class ButtonParser(HTMLParser):
        def __init__(self):
            super().__init__()
            self.buttons = []

        def handle_starttag(self, tag, attrs):
            if tag == "button":
                self.buttons.append(dict(attrs))

    html = asset_path("index-v1.8.18.html").read_text(encoding="utf-8")
    parser = ButtonParser()
    parser.feed(html)

    expected_ids = {
        "choose", "downloadAll", "deleteDoc", "paste", "sample", "clear",
        "previewTop", "previewBottom", "previewFullscreen", "showFormatted",
        "showTextPreview", "dropChoose", "copy", "downloadAnon", "help",
        "reviewPrev", "reviewNext", "hideCandidate", "keepCandidate",
        "anonymizedPreviewTop", "anonymizedPreviewBottom",
        "anonymizedPreviewFullscreen", "showAnonymizedFormatted",
        "showAnonymizedText", "anonymize", "restore",
        "chooseRestoreFile", "restoreFile", "newProfile", "deleteProfile", "exportProfile",
        "importProfile", "selftest", "themeToggle", "applyNow", "exportNow",
        "settingsTab-profile", "settingsTab-rules", "settingsTab-whitelist",
    }
    buttons = {button.get("id"): button for button in parser.buttons}
    assert set(buttons) == expected_ids
    missing = [button_id for button_id, button in buttons.items()
               if len(button.get("title", "").strip()) < 12]
    assert missing == []

    javascript = asset_path("app-v1.8.18.js").read_text(encoding="utf-8")
    assert 'class="undo-chip"' in javascript
    assert 'title="Отменить возврат значения' in javascript

    assert 'id="buttonTooltip"' in html
    assert 'role="tooltip"' in html
    assert "pointerover" in javascript
    assert "showButtonTooltip" in javascript
    assert "hideButtonTooltip" in javascript


def test_manual_review_next_keeps_button_stationary_and_scrolls_only_the_text_preview():
    javascript = asset_path("app-v1.8.18.js").read_text(encoding="utf-8")
    css = asset_path("style-v1.1.css").read_text(encoding="utf-8")

    focus_mark = javascript[javascript.index("function focusMark("):javascript.index("function reviewDisabledReason(")]
    assert "scrollReviewMarkWithinOutput" in focus_mark
    assert ".scrollIntoView(" not in focus_mark
    assert "rich.scrollTop" in javascript
    assert ".review-candidate{" in css
    candidate_rule = css[css.index(".review-candidate{"):css.index("}", css.index(".review-candidate{"))]
    assert "height:" in candidate_rule
    assert "overflow:auto" in candidate_rule


def test_hide_and_remember_reanonymizes_all_exact_repetitions_immediately():
    html = asset_path("index-v1.8.18.html").read_text(encoding="utf-8")
    javascript = asset_path("app-v1.8.18.js").read_text(encoding="utf-8")
    hide_hint = javascript[javascript.index("async function hideHint("):javascript.index("async function anonymizeText(")]

    assert "await anonymizeText()" in hide_hint
    assert "customCategory.checked=true" in hide_hint
    assert "Все одинаковые вхождения" in hide_hint
    assert "все одинаковые вхождения" in html


def test_hide_and_remember_persists_rule_in_the_active_profile_before_refreshing_preview():
    javascript = asset_path("app-v1.8.18.js").read_text(encoding="utf-8")
    hide_hint = javascript[javascript.index("async function hideHint("):javascript.index("async function anonymizeText(")]

    save = "await pywebview.api.save_profile($('#profile').value,$('#rules').value,$('#whitelist').value,$('#paranoid').checked,$('#uncertainMode').value)"
    assert save in hide_hint
    assert hide_hint.index(save) < hide_hint.index("await anonymizeText()")
    assert "правило сохранено в профиле" in hide_hint


def test_gui_has_persistent_light_and_dark_theme_switch():
    html = asset_path("index-v1.8.18.html").read_text(encoding="utf-8")
    javascript = asset_path("app-v1.8.18.js").read_text(encoding="utf-8")
    css = asset_path("style-v1.1.css").read_text(encoding="utf-8")

    assert 'id="themeToggle"' in html
    assert 'aria-pressed="false"' in html
    assert "local-anonymizer-theme" in javascript
    assert "function applyTheme" in javascript
    assert "document.documentElement" in javascript and "dataset.theme" in javascript
    assert "localStorage.setItem" in javascript
    assert "pywebview.api.get_theme()" in javascript
    assert "pywebview.api.set_theme" in javascript
    assert "prefers-color-scheme: dark" in javascript
    assert ':root[data-theme="dark"]' in css
    assert "color-scheme:dark" in css


def test_theme_choice_is_persisted_by_the_local_backend(tmp_path):
    first = AppApi(root=tmp_path / "app")
    assert first.get_theme() is None
    assert first.set_theme("dark") == {"ok": True, "theme": "dark"}

    second = AppApi(root=tmp_path / "app")
    assert second.get_theme() == "dark"
    assert second.set_theme("light") == {"ok": True, "theme": "light"}
    assert AppApi(root=tmp_path / "app").get_theme() == "light"


def test_design_choice_is_persisted_by_backend_and_preserves_theme(tmp_path):
    root = tmp_path / "app"
    first = AppApi(root=root)
    assert first.get_design_mode() is None
    assert first.set_theme("dark") == {"ok": True, "theme": "dark"}
    assert first.set_design_mode("legacy") == {"ok": True, "design_mode": "legacy"}

    second = AppApi(root=root)
    assert second.get_design_mode() == "legacy"
    assert second.get_theme() == "dark"
    assert second.set_design_mode("new") == {"ok": True, "design_mode": "new"}
    assert AppApi(root=root).get_design_mode() == "new"
    assert AppApi(root=root).get_theme() == "dark"


def test_settings_offer_docx_graphics_as_an_explicit_checkbox():
    html = asset_path("index-v1.8.18.html").read_text(encoding="utf-8")

    assert "Секреты и ключи защищаются всегда" in html
    assert "checkbox\" checked disabled" not in html
    assert 'id="redactDocxGraphics"' in html
    assert 'data-object-scope="docx-static-graphics"' in html
    assert "DOCX: скрывать встроенные рисунки и статические схемы" in html
    assert "схемы скрываются всегда" not in html
    assert "[СХЕМА-N]" in html
    assert "изображение удаляется целиком" in html
    assert "Текст внутри пикселей отдельно не распознаётся" in html
    assert "SmartArt, диаграммы, OLE и графика XLSX" in html
    assert "экспорт блокируется" in html
    assert "Отключённые категории останутся в документе без замены" in html
    for control_id in ("uncertainMode", "paranoid", "compactMarkers", "mono"):
        assert f'id="{control_id}"' in html
    for explanation in (
        "Автоматически скрывать спорное без подтверждения",
        "укорачивает подписи маркеров",
        "меняет только отображение текста",
    ):
        assert explanation in html

    javascript = asset_path("app-v1.8.18.js").read_text(encoding="utf-8")
    assert "$('#redactDocxGraphics').checked" in javascript
    assert "['#rules','#whitelist','#redactDocxGraphics']" in javascript


def test_settings_and_primary_action_explain_their_real_effect():
    html = asset_path("index-v1.8.18.html").read_text(encoding="utf-8")
    javascript = asset_path("app-v1.8.18.js").read_text(encoding="utf-8")

    settings_start = html.index("Настройки обезличивания")
    action_position = html.index('id="anonymize"')
    assert settings_start < action_position
    assert "Предпросмотр меняется автоматически" in html
    assert "Применить настройки к файлу" in html
    assert 'id="settingsStatus"' in html
    assert "Изменения настроек ещё не применены к экспортируемому файлу" in javascript
    assert "Настройки применены. Файл готов к экспорту." in javascript


def test_new_design_shows_only_one_apply_and_one_export_action():
    html = asset_path("index-v1.8.18.html").read_text(encoding="utf-8")
    css = asset_path("new-design-v1.11.css").read_text(encoding="utf-8")

    assert 'id="exportNow" title="Сохранить обезличенный документ и переносимый файл восстановления рядом">Сохранить документ и файл восстановления</button>' in html
    assert 'body[data-design="new"] #anonymize,body[data-design="new"] #downloadAnon{display:none!important}' in css


def test_detection_categories_are_a_separate_collapsed_section_at_the_very_bottom():
    html = asset_path("index-v1.8.18.html").read_text(encoding="utf-8")

    section_start = html.index('<details class="card" id="detectionCategories">')
    section_end = html.index("</details>", section_start)
    settings_end = html.index("</details>", html.index("Настройки обезличивания"))
    selftest_start = html.index("Проверка движка")
    footer_start = html.index("<footer>")

    assert settings_end < selftest_start < section_start < section_end < footer_start
    assert "open" not in html[section_start : html.index(">", section_start)]
    assert "Категории поиска и замены" in html[section_start:section_end]
    assert "Что искать и заменять" in html[section_start:section_end]
    assert "Секреты и ключи защищаются всегда" in html[section_start:section_end]
    assert 'id="redactDocxGraphics"' in html[section_start:section_end]
    assert 'data-categories="IP"' in html[section_start:section_end]
    assert 'data-categories="СТАВКА"' in html[section_start:section_end]
    assert "Что искать и заменять" not in html[html.index("Настройки обезличивания") : settings_end]


def test_category_change_renders_text_result_before_slow_formatted_preview():
    javascript = asset_path("app-v1.8.18.js").read_text(encoding="utf-8")

    text_await = "const result=await pywebview.api.anonymize_text(text,...settings);"
    formatted_await = "formattedPreview=await formatted"
    assert text_await in javascript
    assert formatted_await in javascript
    render_after_text = javascript.index("renderOutput()", javascript.index(text_await))
    assert javascript.index(text_await) < render_after_text < javascript.index(formatted_await)
    assert "setAnonymizedPreviewMode(false)" in javascript
    assert "Предпросмотр обновлён" in javascript


def test_review_navigation_is_limited_to_actionable_yellow_hints():
    javascript = asset_path("app-v1.8.18.js").read_text(encoding="utf-8")

    assert "function hintMarks()" in javascript
    assert "const all=hintMarks()" in javascript
    assert "reviewIndex=all.length?0:-1" in javascript
    assert "$('#hideCandidate').disabled=!hasCurrent" in javascript
    assert "$('#keepCandidate').disabled=!hasCurrent" in javascript


def test_local_library_is_compact_inside_the_source_editor():
    html = asset_path("index-v1.8.18.html").read_text(encoding="utf-8")

    source_start = html.index('aria-labelledby="sourceHeading"')
    source_end = html.index('</section>', source_start)
    source_section = html[source_start:source_end]
    assert '<details id="sourceLibrary" class="source-library">' in source_section
    assert '<summary>Файлы и библиотека</summary>' in source_section
    assert 'id="library"' in source_section
    assert 'id="choose"' in source_section
    assert 'id="fileStatus"' in source_section
    assert 'id="downloadAll"' in source_section
    assert 'id="deleteDoc"' in source_section
    assert '<details id="sourceLibrary" class="source-library" open>' not in html
    assert '<section class="library card"' not in html


def test_pasted_text_restoration_reuses_the_main_source_editor():
    html = asset_path("index-v1.8.18.html").read_text(encoding="utf-8")
    javascript = asset_path("app-v1.8.18.js").read_text(encoding="utf-8")

    source_section = html[html.index('aria-labelledby="sourceHeading"'):html.index('</section>', html.index('aria-labelledby="sourceHeading"'))]
    assert 'id="restore"' in source_section
    assert "Вернуть исходные значения" in source_section
    assert "Восстановление текста — Обратная замена маркеров" not in html
    assert 'id="reverse"' not in html
    assert 'id="restored"' not in html
    assert 'id="copyRestored"' not in html
    assert "restore_text($('#source').value,textRestoreMapping)" in javascript
    assert "$('#source').value=result.text" in javascript
    assert "$('#reverse')" not in javascript
    assert "$('#restored')" not in javascript
    assert "$('#copyRestored')" not in javascript


def test_legacy_mapping_table_and_csv_export_are_not_exposed_in_ui():
    html = asset_path("index-v1.8.18.html").read_text(encoding="utf-8")
    javascript = asset_path("app-v1.8.18.js").read_text(encoding="utf-8")

    assert "Таблица соответствий" not in html
    assert "Экспорт CSV (чувствительные данные)" not in html
    assert 'id="mapping"' not in html
    assert 'id="mapCount"' not in html
    assert 'id="csv"' not in html
    assert "$('#mapping')" not in javascript
    assert "$('#mapCount')" not in javascript
    assert "$('#csv')" not in javascript
    assert "export_mapping" not in javascript
    assert "last.mapping=mapping||{}" in javascript


def test_manual_review_owns_its_search_and_display_controls():
    html = asset_path("index-v1.8.18.html").read_text(encoding="utf-8")
    review_start = html.index('id="manualReview"')
    review_end = html.index("</section>", review_start)
    review = html[review_start:review_end]
    settings_start = html.index('id="anonymizationSettings"')
    settings_end = html.index("</details>", settings_start)
    settings = html[settings_start:settings_end]
    categories_start = html.index('id="detectionCategories"')

    for control_id in ("uncertainMode", "reviewHints", "reviewHintsMode", "paranoid", "compactMarkers", "mono"):
        assert f'id="{control_id}"' in review
    assert "Искать спорные места для ручной проверки" in review
    assert "Автоматически скрывать спорное без подтверждения" in review
    assert "Дополнительные режимы" not in review
    assert 'data-categories="HINTS"' not in html[categories_start:]
    assert 'id="paranoid"' not in settings


def test_manual_review_panel_explains_actions_progress_result_and_completion():
    html = asset_path("index-v1.8.18.html").read_text(encoding="utf-8")
    javascript = asset_path("app-v1.8.18.js").read_text(encoding="utf-8")

    for text in (
        "Ручная проверка спорных мест",
        "Скрыть и запомнить правило",
        "Оставить и добавить в исключения",
        "Как закончить проверку",
    ):
        assert text in html
    assert "Автопоиск не уверен — выберите действие." not in html
    assert 'id="reviewDecision"' in html
    assert 'id="reviewFinish"' in html
    assert 'id="resultStatus"' in html
    assert "Осталось ${remaining} из ${total}" in javascript
    assert "Решение принято:" in javascript
    assert "Очередь разобрана." in javascript
    assert "нажмите «Применить настройки к файлу»" in javascript
    assert "Ручная проверка завершена." in javascript
    assert "Теперь примените настройки к файлу." in javascript
    assert "$('#warning').textContent=hasCurrent?`Осталось спорных мест: ${remaining}.`:'Ручная проверка завершена: все спорные места разобраны.'" in javascript


def test_manual_review_unified_mode_and_current_candidate_are_obvious():
    html = asset_path("index-v1.8.18.html").read_text(encoding="utf-8")
    javascript = asset_path("app-v1.8.18.js").read_text(encoding="utf-8")

    assert "Автоматически скрывать спорное без подтверждения" in html
    assert "Искать спорные места для ручной проверки" in html
    assert 'id="reviewCurrent"' in html
    assert "Ручная проверка отключена" in javascript
    assert "снимите галочку «Автоматически скрывать спорное без подтверждения»" in javascript
    assert "setAnonymizedPreviewMode(false)" in javascript
    assert "`«${all[reviewIndex].textContent.trim()}» · ${all[reviewIndex].dataset.kind" in javascript


def test_document_export_has_a_visible_matching_restore_import_action():
    html = asset_path("index-v1.8.18.html").read_text(encoding="utf-8")
    javascript = asset_path("app-v1.8.18.js").read_text(encoding="utf-8")

    assert "Сохранить два файла…" in html
    assert "информация для восстановления.larecovery" in html
    assert html.count('id="chooseRestoreFile"') == 1
    assert html.count('id="restoreFile"') == 1
    assert "Восстановить файл…" in html
    assert "Импортировать для восстановления…" not in html
    assert "function chooseRestorePair()" in javascript
    assert "function restoreMatchedFile()" in javascript
    assert "restoreFileTop" not in javascript
    assert "$('#chooseRestoreFile').onclick=chooseRestorePair" in javascript
    assert "$('#restoreFile').onclick=restoreMatchedFile" in javascript


def test_app_api_import_library_and_anonymize(tmp_path):
    source = tmp_path / "sample.docx"
    doc = Document(); doc.add_paragraph("Иванов Сергей Петрович"); doc.save(source)
    api = AppApi(root=tmp_path / "app")
    imported = api.import_paths([str(source)])
    assert imported["ok"] and len(api.list_documents()) == 1
    report = api.anonymize_document(imported["ids"][0])
    assert report["ready"]
    assert Path(api.service.storage.get_document(imported["ids"][0])["output_path"]).exists()
    assert "output_path" not in report and "operation_id" not in report and "document_id" not in report
    api.save_profile("Тест", "Вымпел => СИСТЕМА", "Казань", True)
    assert "Тест" in api.list_profiles()
    profile = api.load_profile("Тест")
    assert profile["dictionary"] == {"Вымпел": "СИСТЕМА"}
    api.delete_profile("Тест")
    assert "Тест" not in api.list_profiles()


def test_preview_document_extracts_docx_and_text_without_modifying_sources(tmp_path):
    text_source = tmp_path / "note.txt"
    text_source.write_text("login admin_db", encoding="utf-8")
    docx_source = tmp_path / "note.docx"
    doc = Document(); doc.add_paragraph("Иванов Сергей Петрович")
    table = doc.add_table(rows=1, cols=1); table.cell(0, 0).text = "АС «Паутина»"
    doc.save(docx_source)
    api = AppApi(root=tmp_path / "app")
    ids = api.import_paths([str(text_source), str(docx_source)])["ids"]
    previews = [api.preview_document(doc_id) for doc_id in ids]
    assert previews[0]["text"] == "login admin_db"
    assert "Иванов Сергей Петрович" in previews[1]["text"]
    assert "АС «Паутина»" in previews[1]["text"]
    assert text_source.read_text(encoding="utf-8") == "login admin_db"


def test_document_preview_bridge_is_chunked_and_reconstructs_large_text(tmp_path):
    source = tmp_path / "large.txt"
    text = "Иванов Сергей Петрович\n" * 8_000
    source.write_text(text, encoding="utf-8")
    api = AppApi(root=tmp_path / "app")
    document_id = api.import_paths([str(source)])["ids"][0]
    expected = source.read_bytes().decode("utf-8")

    manifest = api.start_document_preview(document_id)
    assert "text" not in manifest and "rich_html" not in manifest
    assert manifest["text_length"] == len(expected)
    chunks = []
    offset = 0
    while offset < manifest["text_length"]:
        result = api.read_document_preview_chunk(manifest["token"], "text", offset)
        assert len(result["chunk"]) <= 16_384
        chunks.append(result["chunk"])
        offset = result["next_offset"]
    assert "".join(chunks) == expected
    assert api.finish_document_preview(manifest["token"]) == {"ok": True}
    with pytest.raises(ValueError, match="устарел"):
        api.read_document_preview_chunk(manifest["token"], "text", 0)


def test_formatted_preview_bridge_uses_bounded_chunks(tmp_path, monkeypatch):
    api = AppApi(root=tmp_path / "app")
    rich_html = "<p>форматированный текст</p>" * 10_000
    monkeypatch.setattr(
        api.service,
        "preview_document",
        lambda _document_id, **_kwargs: {
            "display_name": "large.odt",
            "text": "текст",
            "status": "Импортирован",
            "kind": "document",
            "format": "odt",
            "rich_html": rich_html,
        },
    )

    manifest = api.start_formatted_document_preview("document-id")
    chunks = []
    offset = 0
    while offset < manifest["rich_html_length"]:
        result = api.read_document_preview_chunk(manifest["token"], "rich_html", offset)
        assert len(result["chunk"]) <= 16_384
        chunks.append(result["chunk"])
        offset = result["next_offset"]

    assert len(chunks) > 1
    assert "".join(chunks) == rich_html


def test_anonymized_formatted_preview_bridge_is_chunked(tmp_path, monkeypatch):
    api = AppApi(root=tmp_path / "app")
    rich_html = "<p>обезличенный форматированный текст</p>" * 8_000
    monkeypatch.setattr(
        api,
        "preview_anonymized_document",
        lambda *_args, **_kwargs: {"format": "odt", "rich_html": rich_html},
    )

    manifest = api.start_anonymized_document_preview("document-id")
    chunks = []
    offset = 0
    while offset < manifest["rich_html_length"]:
        result = api.read_document_preview_chunk(manifest["token"], "rich_html", offset)
        assert len(result["chunk"]) <= 16_384
        chunks.append(result["chunk"])
        offset = result["next_offset"]

    assert "rich_html" not in manifest
    assert "".join(chunks) == rich_html


def test_docx_preview_is_formatted_ordered_and_escaped_without_links(tmp_path):
    source = tmp_path / "rich.docx"
    doc = Document()
    heading = doc.add_heading("Заголовок <script>alert(1)</script>", level=1)
    heading.alignment = WD_ALIGN_PARAGRAPH.CENTER
    paragraph = doc.add_paragraph()
    bold = paragraph.add_run("жирный")
    bold.bold = True
    italic = paragraph.add_run(" курсив")
    italic.italic = True
    underline = paragraph.add_run(" подчёркнутый")
    underline.underline = True
    styled = paragraph.add_run(" цвет")
    styled.font.size = Pt(18)
    styled.font.color.rgb = RGBColor(0x12, 0x34, 0x56)
    styled.add_break()
    styled.add_tab()
    hyperlink = OxmlElement("w:hyperlink")
    hyperlink.set("{http://schemas.openxmlformats.org/officeDocument/2006/relationships}id", "rEvil")
    run = OxmlElement("w:r")
    text = OxmlElement("w:t")
    text.text = "безопасная подпись"
    run.append(text)
    hyperlink.append(run)
    paragraph._p.append(hyperlink)
    table = doc.add_table(rows=1, cols=1)
    table.cell(0, 0).text = "ячейка"
    nested = table.cell(0, 0).add_table(rows=1, cols=1)
    nested.cell(0, 0).text = "вложенная"
    header = doc.sections[0].header.paragraphs[0]
    header.text = "Верхний колонтитул"
    footer = doc.sections[0].footer.paragraphs[0]
    footer.text = "Нижний колонтитул"
    doc.save(source)

    api = AppApi(root=tmp_path / "app")
    document_id = api.import_paths([str(source)])["ids"][0]
    preview = api.preview_document(document_id)
    rich = preview["rich_html"]

    assert set(preview) == {"display_name", "text", "status", "kind", "format", "rich_html"}
    assert preview["kind"] == "document" and preview["format"] == "docx"
    assert rich.index("<h1") < rich.index("<table")
    assert "&lt;script&gt;alert(1)&lt;/script&gt;" in rich and "<script" not in rich
    assert "<strong>жирный</strong>" in rich
    assert "<em> курсив</em>" in rich
    assert "<u> подчёркнутый</u>" in rich
    assert 'class="docx-paragraph align-center"' in rich
    assert 'style="font-size:18pt;color:#123456"' in rich
    assert rich.count('class="docx-table"') >= 2
    assert "ячейка" in rich and "вложенная" in rich
    assert 'class="docx-header"' in rich and "Верхний колонтитул" in rich
    assert 'class="docx-footer"' in rich and "Нижний колонтитул" in rich
    assert "<br>" in rich and 'class="docx-tab"' in rich
    assert "rEvil" not in rich and "relationship" not in rich and "href=" not in rich
    assert document_id not in repr(preview) and str(tmp_path) not in repr(preview)


def test_docx_preview_renders_safe_embedded_png_but_anonymized_preview_only_shows_marker(tmp_path):
    image_bytes = base64.b64decode(
        "iVBORw0KGgoAAAANSUhEUgAAAAEAAAABCAQAAAC1HAwCAAAAC0lEQVR42mNk+A8AAQUBAScY42YAAAAASUVORK5CYII="
    )
    image = tmp_path / "synthetic.png"
    source = tmp_path / "with-image.docx"
    image.write_bytes(image_bytes)
    document = Document()
    document.add_paragraph("До схемы")
    document.add_picture(str(image), width=Inches(2))
    document.add_paragraph("После схемы")
    document.save(source)

    api = AppApi(root=tmp_path / "app")
    document_id = api.import_paths([str(source)])["ids"][0]
    source_html = api.preview_document(document_id)["rich_html"]
    anonymized_html = api.preview_anonymized_document(document_id)["rich_html"]

    encoded = base64.b64encode(image_bytes).decode("ascii")
    assert f'src="data:image/png;base64,{encoded}"' in source_html
    assert source_html.index("До схемы") < source_html.index("<img") < source_html.index("После схемы")
    assert str(tmp_path) not in source_html and "rId" not in source_html
    assert "data:image/" not in anonymized_html
    assert "[СХЕМА-1]" in anonymized_html


def test_docx_preview_emits_stable_structural_scroll_anchors_after_graphic_removal(tmp_path):
    image_bytes = base64.b64decode(
        "iVBORw0KGgoAAAANSUhEUgAAAAEAAAABCAQAAAC1HAwCAAAAC0lEQVR42mNk+A8AAQUBAScY42YAAAAASUVORK5CYII="
    )
    image = tmp_path / "scroll-anchor.png"
    source = tmp_path / "scroll-anchor.docx"
    image.write_bytes(image_bytes)
    document = Document()
    document.add_paragraph("Раздел до схемы")
    document.add_picture(str(image), width=Inches(2))
    table = document.add_table(rows=2, cols=2)
    table.cell(0, 0).text = "Иванов Сергей Петрович"
    nested = table.cell(0, 1).add_table(rows=1, cols=1)
    nested.cell(0, 0).text = "Вложенный блок"
    table.cell(1, 0).text = "Следующий блок"
    document.add_paragraph("Раздел после схемы")
    document.save(source)

    api = AppApi(root=tmp_path / "app")
    document_id = api.import_paths([str(source)])["ids"][0]
    source_html = api.preview_document(document_id)["rich_html"]
    anonymized_html = api.preview_anonymized_document(document_id)["rich_html"]

    source_anchors = re.findall(r'data-preview-anchor="([^"]+)"', source_html)
    anonymized_anchors = re.findall(r'data-preview-anchor="([^"]+)"', anonymized_html)
    assert source_anchors
    assert source_anchors == anonymized_anchors
    assert len(source_anchors) == len(set(source_anchors))
    assert "data:image/" in source_html and "data:image/" not in anonymized_html
    assert "[СХЕМА-1]" in anonymized_html


def test_docx_preview_header_anchor_identity_survives_image_only_header_removal(tmp_path):
    image_bytes = base64.b64decode(
        "iVBORw0KGgoAAAANSUhEUgAAAAEAAAABCAQAAAC1HAwCAAAAC0lEQVR42mNk+A8AAQUBAScY42YAAAAASUVORK5CYII="
    )
    image = tmp_path / "header-anchor.png"
    source = tmp_path / "header-anchor.docx"
    image.write_bytes(image_bytes)
    document = Document()
    document.add_paragraph("Основной текст")
    document.sections[0].header.paragraphs[0].add_run().add_picture(str(image), width=Inches(1))
    second = document.add_section(WD_SECTION.NEW_PAGE)
    second.header.is_linked_to_previous = False
    second.header.paragraphs[0].text = "Второй заголовок"
    document.save(source)

    api = AppApi(root=tmp_path / "app")
    document_id = api.import_paths([str(source)])["ids"][0]
    source_html = api.preview_document(document_id)["rich_html"]
    anonymized_html = api.preview_anonymized_document(document_id)["rich_html"]

    source_anchor = re.search(
        r'data-preview-anchor="([^"]+)">Второй заголовок</p>', source_html
    )
    anonymized_anchor = re.search(
        r'data-preview-anchor="([^"]+)">Второй заголовок</p>', anonymized_html
    )
    assert source_anchor and anonymized_anchor
    assert source_anchor.group(1) == anonymized_anchor.group(1)


def test_docx_preview_malformed_custom_style_id_falls_back_safely(tmp_path):
    source = tmp_path / "malformed-style.docx"
    document = Document()
    paragraph = document.add_paragraph("Текст <script>alert(1)</script>")
    style = OxmlElement("w:pStyle")
    style.set(qn("w:val"), "Custom'\"Style")
    paragraph._p.get_or_add_pPr().append(style)
    document.save(source)

    api = AppApi(root=tmp_path / "app")
    document_id = api.import_paths([str(source)])["ids"][0]
    preview = api.preview_document(document_id)

    assert '<p class="docx-paragraph" data-preview-anchor="body/b0">' in preview["rich_html"]
    assert "&lt;script&gt;alert(1)&lt;/script&gt;" in preview["rich_html"]
    assert "<script" not in preview["rich_html"]
    assert "Custom" not in preview["rich_html"]


class DialogSequence:
    def __init__(self, *responses):
        self.responses = list(responses)
        self.calls = []

    def create_file_dialog(self, *args, **kwargs):
        self.calls.append((args, kwargs))
        return self.responses.pop(0)


def test_restore_document_chooses_edited_docx_then_saves_to_user_destination(tmp_path):
    source = tmp_path / "source.docx"
    document = Document()
    document.add_paragraph("Иванов Сергей Петрович")
    document.save(source)
    api = AppApi(root=tmp_path / "app")
    document_id = api.import_paths([str(source)])["ids"][0]
    api.anonymize_document(document_id)
    internal_output = Path(api.service.storage.get_document(document_id)["output_path"])

    edited = tmp_path / "edited-anonymized.docx"
    edited.write_bytes(internal_output.read_bytes())
    edited_document = Document(edited)
    edited_document.add_paragraph("Новая правка [ФИО-999]")
    edited_document.save(edited)
    destination = tmp_path / "chosen" / "restored.docx"
    window = DialogSequence([str(edited)], [str(destination)])
    api.set_window(window)

    report = api.restore_document(document_id)

    assert report == {
        "ok": True,
        "path": str(destination),
        "unknown_markers": ["[ФИО-999]"],
    }
    assert destination.is_file()
    restored_text = "\n".join(p.text for p in Document(destination).paragraphs)
    assert "Иванов Сергей Петрович" in restored_text and "Новая правка [ФИО-999]" in restored_text
    assert [call[0][0] for call in window.calls] == [10, 30]
    assert window.calls[0][1]["directory"] == str(internal_output.parent)
    assert Path(window.calls[0][1]["directory"]).is_absolute()
    assert window.calls[1][1]["directory"] == str(edited.parent)
    assert Path(window.calls[1][1]["directory"]).is_absolute()
    assert window.calls[1][1]["file_types"] == ("DOCX (*.docx)",)
    assert "output_path" not in repr(report) and "operation_id" not in repr(report)
    assert document_id not in repr(report) and str(api.service.storage.root) not in repr(report)


def test_restore_document_cancelled_save_does_not_restore_or_leak_paths(tmp_path, monkeypatch):
    edited = tmp_path / "edited.docx"
    edited.write_bytes(b"not opened because save is cancelled")
    api = AppApi(root=tmp_path / "app")
    window = DialogSequence([str(edited)], None)
    api.set_window(window)
    monkeypatch.setattr(
        api.service,
        "restore",
        lambda *_args, **_kwargs: (_ for _ in ()).throw(AssertionError("restore must not run")),
    )

    assert api.restore_document("opaque-selection") == {"ok": False}
    assert len(window.calls) == 2


def test_ui_offers_docx_preview_modes_and_explains_word_restore_workflow():
    html = asset_path("index-v1.8.18.html").read_text(encoding="utf-8")
    javascript = asset_path("app-v1.8.18.js").read_text(encoding="utf-8")
    css = asset_path("style-v1.1.css").read_text(encoding="utf-8")

    for element_id in (
        "docxPreviewControls",
        "showFormatted",
        "showTextPreview",
        "previewTop",
        "previewBottom",
        "previewFullscreen",
        "previewCompleteness",
        "docxPreview",
        "anonymizedPreviewControls",
        "showAnonymizedFormatted",
        "showAnonymizedText",
        "anonymizedDocxPreview",
    ):
        assert f'id="{element_id}"' in html
    for label in (
        "Форматированный вид",
        "Текст и подсветка",
        "В начало",
        "В конец",
        "На весь экран",
        "Восстановить файл…",
        "1. Выберите обезличенный файл",
        "2. Программа подтвердит пару",
        "3. Нажмите «Восстановить файл»",
        "соответствие проверяется по внутреннему идентификатору пары",
        "Вернуть исходные значения",
    ):
        assert label in html
    assert html.count('id="restoreFile"') == 1
    assert 'id="restoreFileTop"' not in html
    assert "Выбрать обезличенный файл…" in html
    assert "Восстановить исправленный файл" not in html
    assert javascript.count("$('#chooseRestoreFile').onclick=chooseRestorePair") == 1
    assert javascript.count("$('#restoreFile').onclick=restoreMatchedFile") == 1
    assert "restoreFileTop" not in javascript
    assert "Обезличенный файл" in javascript
    assert "Выберите место для восстановленного файла" in javascript
    assert "preview.rich_html" in javascript
    assert "function hasFormattedDocumentPreview(format){return format==='docx'||format==='odt';}" in javascript
    assert "switchPreviewMode(setDocxPreviewMode,hasFormattedDocumentPreview(preview.format)&&keepFormatted)" in javascript
    assert "filePreview&&!loadingPreview&&selected&&hasFormattedDocumentPreview(currentPreviewFormat)" in javascript
    assert "filePreview&&hasFormattedDocumentPreview(currentPreviewFormat)" in javascript
    assert "requestAnonymizedDocumentPreview(selected,...settings" in javascript
    assert "preview_anonymized_document" in javascript
    assert "Версия 1.8" in html and "APP_VERSION='1.8'" in javascript
    assert 'src="app-v1.8.18.js"' in html
    assert "Текст документа загружен полностью" in javascript
    assert "Визуальные элементы без текста могут не отображаться" in javascript
    assert "rich.scrollTop=0" in javascript
    assert "rich.scrollTop=rich.scrollHeight" in javascript
    assert "preview-surface-fullscreen" in javascript
    assert "r.path" in javascript and "r.output_path" not in javascript
    assert ".docx-preview" in css and ".docx-table" in css
    assert ".docx-preview-fullscreen" in css


def test_docx_preview_preserves_horizontal_and_vertical_merged_cells(tmp_path):
    source = tmp_path / "merged.docx"
    document = Document()
    table = document.add_table(rows=3, cols=3)
    table.cell(0, 0).merge(table.cell(0, 1)).text = "горизонтальное объединение"
    table.cell(1, 2).merge(table.cell(2, 2)).text = "вертикальное объединение"
    document.save(source)

    api = AppApi(root=tmp_path / "app")
    document_id = api.import_paths([str(source)])["ids"][0]
    rich = api.preview_document(document_id)["rich_html"]

    assert rich.count("горизонтальное объединение") == 1
    assert rich.count("вертикальное объединение") == 1
    assert 'colspan="2"' in rich
    assert 'rowspan="2"' in rich


def test_anonymized_docx_preview_keeps_table_and_run_formatting(tmp_path):
    source = tmp_path / "formatted-sensitive.docx"
    document = Document()
    paragraph = document.add_paragraph()
    paragraph.add_run("Иванов Сергей Петрович").bold = True
    paragraph.add_run(" обычный текст").italic = True
    table = document.add_table(rows=3, cols=2)
    table.style = "Table Grid"
    table.cell(0, 0).merge(table.cell(0, 1)).text = "Владелец Иванов Сергей Петрович"
    table.cell(1, 0).merge(table.cell(2, 0)).text = "Вертикальное объединение"
    document.save(source)

    api = AppApi(root=tmp_path / "app")
    document_id = api.import_paths([str(source)])["ids"][0]
    preview = api.preview_anonymized_document(document_id)

    assert preview["format"] == "docx"
    assert "Иванов Сергей Петрович" not in preview["rich_html"]
    assert "[ФИО-1]" in preview["rich_html"]
    assert "<strong>" in preview["rich_html"]
    assert '<em>' in preview["rich_html"]
    assert 'class="docx-table"' in preview["rich_html"]
    assert preview["rich_html"].count("Вертикальное объединение") == 1
    assert 'colspan="2"' in preview["rich_html"]
    assert 'rowspan="2"' in preview["rich_html"]


def test_docx_preview_omits_malformed_optional_formatting_instead_of_crashing(tmp_path):
    source = tmp_path / "malformed-formatting.docx"
    document = Document()
    paragraph = document.add_paragraph("Безопасный текст")
    paragraph._p.get_or_add_pPr().append(OxmlElement("w:jc"))
    document.save(source)

    api = AppApi(root=tmp_path / "app")
    document_id = api.import_paths([str(source)])["ids"][0]

    assert "Безопасный текст" in api.preview_document(document_id)["rich_html"]
