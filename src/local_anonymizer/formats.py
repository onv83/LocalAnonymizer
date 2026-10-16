from __future__ import annotations

import os
import re
import tempfile
from pathlib import Path
from zipfile import ZipFile

from openpyxl import load_workbook

from .documents import DocumentResult, _preflight_docx_package
from .engine import Anonymizer, MARKER_PATTERN, warnings_require_review
from .xlsx_package import rewrite_xlsx_package, xlsx_sheet_parts


def _read_text(path: Path) -> str:
    raw = path.read_bytes()
    try:
        return raw.decode("utf-8-sig")
    except UnicodeDecodeError:
        try:
            return raw.decode("cp1251")
        except UnicodeDecodeError as exc:
            raise ValueError("Не удалось определить кодировку; сохраните файл как UTF-8") from exc


def _atomic_bytes(output: Path, data: bytes) -> None:
    output.parent.mkdir(parents=True, exist_ok=True)
    fd, name = tempfile.mkstemp(prefix=f".{output.stem}-", suffix=output.suffix, dir=output.parent)
    try:
        with os.fdopen(fd, "wb") as stream:
            stream.write(data); stream.flush(); os.fsync(stream.fileno())
        os.replace(name, output)
    finally:
        Path(name).unlink(missing_ok=True)


def process_text_file(source: str | Path, output: str | Path, engine: Anonymizer | None = None) -> DocumentResult:
    source, output = Path(source), Path(output)
    result = (engine or Anonymizer()).anonymize(_read_text(source))
    _atomic_bytes(output, result.text.encode("utf-8"))
    return DocumentResult(result.mapping, result.candidates, result.warnings,
                          not warnings_require_review(result.warnings),
                          result.total, result.hint_count)


def restore_text_file(source: str | Path, output: str | Path, mapping: dict[str, str]) -> list[str]:
    text = _read_text(Path(source))
    markers = set(re.findall(MARKER_PATTERN, text))
    unknown = sorted(markers - mapping.keys())
    restored = re.sub(MARKER_PATTERN, lambda m: mapping.get(m.group(), m.group()), text)
    _atomic_bytes(Path(output), restored.encode("utf-8"))
    return unknown


def _formula_literals(formula: str) -> list[tuple[int, int, str]]:
    """Return quoted Excel string spans and their unescaped values."""
    literals: list[tuple[int, int, str]] = []
    index = 0
    while index < len(formula):
        if formula[index] != '"':
            index += 1
            continue
        start = index
        index += 1
        value: list[str] = []
        while index < len(formula):
            if formula[index] != '"':
                value.append(formula[index])
                index += 1
            elif index + 1 < len(formula) and formula[index + 1] == '"':
                value.append('"')
                index += 2
            else:
                index += 1
                literals.append((start, index, "".join(value)))
                break
        else:
            # Malformed formulas are left to Excel/openpyxl; do not guess at
            # unterminated quoted data.
            break
    return literals


def _formula_with_replacements(formula: str, replacements: list[tuple[int, int, str]]) -> str:
    for start, end, value in reversed(replacements):
        formula = formula[:start] + '"' + value.replace('"', '""') + '"' + formula[end:]
    return formula


def _dangerous_formula(formula: str) -> bool:
    # Ignore quoted string values while inspecting executable formula syntax.
    code = _formula_with_replacements(formula, [(start, end, "") for start, end, _ in _formula_literals(formula)])
    return bool(
        re.search(r"(?i)(?:^|[^A-Z0-9_.])WEBSERVICE\s*\(", code)
        or re.search(r"\|[^!\r\n]*!", code)
        or re.search(r"\[[^\]\r\n]+\][^!\r\n]*!", code)
    )


class XlsxProcessor:
    def __init__(self, engine: Anonymizer | None = None):
        self.engine = engine or Anonymizer()

    def anonymize(self, source: str | Path, output: str | Path) -> DocumentResult:
        source, output = Path(source), Path(output)
        _preflight_docx_package(source)
        warnings: list[str | dict[str, str]] = list(self._warnings(source))
        wb = load_workbook(source, data_only=False, keep_links=False)
        sheet_parts = xlsx_sheet_parts(source)
        originals: dict[tuple[str, str], object] = {}
        records: list[tuple[object, str, str, tuple[int, int] | None]] = []
        formula_sources: dict[int, tuple[object, str]] = {}
        label_re = re.compile(r"(?i)\b(?:ИНН|КПП|ОГРН(?:ИП)?|СНИЛС|БИК|паспорт|серия|сч[её]т|р/с|табельн[а-яё]*\s*(?:номер|№)?)\b")
        for ws in wb.worksheets:
            for row in ws.iter_rows():
                for cell in row:
                    if cell.__class__.__name__ != "MergedCell":
                        originals[(sheet_parts[ws.title], cell.coordinate)] = cell.value
                    value = cell.value
                    if isinstance(value, str) and value.startswith("="):
                        if _dangerous_formula(value):
                            cell.value = "[BLOCKED-FORMULA]"
                            warnings.append({
                                "kind": "XLSX_FORMULA_BLOCKED",
                                "level": "stop",
                                "message": f"Опасная формула нейтрализована: {ws.title}!{cell.coordinate}",
                            })
                            continue
                        literals = _formula_literals(value)
                        if literals:
                            formula_sources[id(cell)] = (cell, value)
                            for start, end, literal in literals:
                                if literal.casefold().startswith("mailto:"):
                                    records.append((cell, literal[7:], "formula-mailto", (start, end)))
                                else:
                                    records.append((cell, literal, "formula", (start, end)))
                    elif isinstance(value, str):
                        records.append((cell, value, "text", None))
                    elif isinstance(value, (int, float)) and not isinstance(value, bool):
                        contexts = []
                        if cell.column > 1:
                            contexts.append(ws.cell(cell.row, cell.column - 1).value)
                        if cell.row > 1:
                            contexts.append(ws.cell(cell.row - 1, cell.column).value)
                        label = next((str(x) for x in contexts if isinstance(x, str) and label_re.search(x)), None)
                        if label:
                            rendered = str(int(value)) if isinstance(value, float) and value.is_integer() else str(value)
                            records.append((cell, f"{label} {rendered}", "numeric", None))
        sep = "\ufdd0"
        payloads = [payload for _, payload, _, _ in records]
        if any(sep in payload for payload in payloads):
            raise RuntimeError("Ошибка разбиения XLSX: коллизия разделителя")
        result = self.engine.anonymize(sep.join(payloads))
        pieces = result.text.split(sep) if records else []
        if len(pieces) != len(records):
            raise RuntimeError("Ошибка разбиения XLSX: неверное число фрагментов")
        formula_replacements: dict[int, list[tuple[int, int, str]]] = {}
        for (cell, _, kind, span), value in zip(records, pieces):
            if kind == "numeric":
                rendered = (
                    str(int(cell.value))
                    if isinstance(cell.value, float) and cell.value.is_integer()
                    else str(cell.value)
                )
                marker = next(
                    (
                        candidate
                        for candidate in re.finditer(MARKER_PATTERN, value)
                        if result.mapping.get(candidate.group()) == rendered
                    ),
                    None,
                )
                if marker:
                    cell.value = marker.group()
            elif kind in ("formula", "formula-mailto"):
                assert span is not None
                replacement = "mailto:" + value if kind == "formula-mailto" else value
                formula_replacements.setdefault(id(cell), []).append((*span, replacement))
            else:
                cell.value = value
        for key, replacements in formula_replacements.items():
            cell, formula = formula_sources[key]
            cell.value = _formula_with_replacements(formula, replacements)
        for ws in wb.worksheets:
            for row in ws.iter_rows():
                for cell in row:
                    if cell.__class__.__name__ != "MergedCell":
                        cell.comment = None
        warnings.extend(self._sanitize_hidden_surfaces(wb))
        changes = {
            (sheet_parts[ws.title], cell.coordinate): cell.value
            for ws in wb.worksheets
            for row in ws.iter_rows()
            for cell in row
            if cell.__class__.__name__ != "MergedCell"
            and cell.value != originals.get((sheet_parts[ws.title], cell.coordinate))
        }
        rewrite_xlsx_package(source, output, changes)
        warnings = list(result.warnings) + warnings
        return DocumentResult(result.mapping, result.candidates, warnings,
                              not warnings_require_review(warnings),
                              result.total, result.hint_count)

    def restore(self, source: str | Path, output: str | Path, mapping: dict[str, str]) -> list[str]:
        source, output = Path(source), Path(output)
        _preflight_docx_package(source)
        wb = load_workbook(source, data_only=False, keep_links=False)
        sheet_parts = xlsx_sheet_parts(source)
        originals = {
            (sheet_parts[ws.title], cell.coordinate): cell.value
            for ws in wb.worksheets
            for row in ws.iter_rows()
            for cell in row
            if cell.__class__.__name__ != "MergedCell"
        }
        unknown: set[str] = set()
        for ws in wb.worksheets:
            for row in ws.iter_rows():
                for cell in row:
                    if cell.__class__.__name__ == "MergedCell":
                        continue
                    if isinstance(cell.value, str):
                        value = cell.value
                        for marker in re.findall(MARKER_PATTERN, value):
                            if marker not in mapping:
                                unknown.add(marker)
                        if value.startswith("="):
                            replacements = [
                                (start, end, re.sub(MARKER_PATTERN, lambda m: mapping.get(m.group(), m.group()), literal))
                                for start, end, literal in _formula_literals(value)
                            ]
                            cell.value = _formula_with_replacements(value, replacements)
                        else:
                            cell.value = re.sub(MARKER_PATTERN, lambda m: mapping.get(m.group(), m.group()), value)
                    cell.comment = None
        self._sanitize_hidden_surfaces(wb)
        changes = {
            (sheet_parts[ws.title], cell.coordinate): cell.value
            for ws in wb.worksheets
            for row in ws.iter_rows()
            for cell in row
            if cell.__class__.__name__ != "MergedCell"
            and cell.value != originals.get((sheet_parts[ws.title], cell.coordinate))
        }
        rewrite_xlsx_package(source, output, changes)
        return sorted(unknown)

    @staticmethod
    def _sanitize_hidden_surfaces(wb) -> list[dict[str, str]]:
        warnings: list[dict[str, str]] = []
        found_hyperlink = False
        found_header_footer = False
        found_validation = False
        for ws in wb.worksheets:
            for row in ws.iter_rows():
                for cell in row:
                    if cell.__class__.__name__ != "MergedCell" and cell.hyperlink is not None:
                        cell.hyperlink = None
                        found_hyperlink = True
            for item_name in ("oddHeader", "evenHeader", "firstHeader", "oddFooter", "evenFooter", "firstFooter"):
                item = getattr(ws, item_name)
                for section_name in ("left", "center", "right"):
                    section = getattr(item, section_name)
                    if section.text:
                        section.text = None
                        found_header_footer = True
            for validation in ws.data_validations.dataValidation:
                values = (validation.prompt, validation.promptTitle, validation.error, validation.errorTitle)
                if any(values):
                    validation.prompt = None
                    validation.promptTitle = None
                    validation.error = None
                    validation.errorTitle = None
                    found_validation = True
        categories = (
            (found_hyperlink, "XLSX_HYPERLINK", "Цели гиперссылок удалены; требуется ручная проверка."),
            (found_header_footer, "XLSX_HEADER_FOOTER", "Колонтитулы очищены; требуется ручная проверка."),
            (found_validation, "XLSX_DATA_VALIDATION", "Сообщения проверки данных очищены; требуется ручная проверка."),
        )
        warnings.extend(
            {"kind": kind, "level": "stop", "message": message}
            for found, kind, message in categories if found
        )
        if wb.defined_names:
            wb.defined_names.clear()
            warnings.append({
                "kind": "XLSX_DEFINED_NAME",
                "level": "stop",
                "message": "Именованные диапазоны удалены; требуется ручная проверка.",
            })
        return warnings

    @staticmethod
    def _clear_properties(wb) -> None:
        for attr in ("creator", "lastModifiedBy", "title", "subject", "description", "keywords", "category", "identifier", "language"):
            setattr(wb.properties, attr, None)
        # openpyxl substitutes its own creator when the value is None; an empty
        # element roundtrips as no creator and cannot retain the source value.
        wb.properties.creator = ""

    @staticmethod
    def _warnings(source: Path) -> list[dict[str, str]]:
        with ZipFile(source) as archive:
            names = archive.namelist()
        warnings: list[dict[str, str]] = []
        for name in names:
            if name.startswith("xl/externalLinks/"):
                warnings.append({
                    "kind": "XLSX_EXTERNAL_LINK",
                    "level": "stop",
                    "message": "Внешняя связь удалена; требуется ручная проверка XLSX.",
                })
            elif name.startswith(("xl/drawings/", "xl/embeddings/")):
                warnings.append({
                    "kind": "XLSX_UNSUPPORTED_PART",
                    "level": "stop",
                    "message": f"Требуется ручная проверка XLSX: {name}",
                })
        return warnings

    @staticmethod
    def _save_atomic(wb, output: Path) -> None:
        output.parent.mkdir(parents=True, exist_ok=True)
        fd, name = tempfile.mkstemp(prefix=f".{output.stem}-", suffix=".xlsx", dir=output.parent); os.close(fd)
        try:
            wb.save(name)
            os.replace(name, output)
        finally:
            Path(name).unlink(missing_ok=True)
