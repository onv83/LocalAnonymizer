from __future__ import annotations

import base64
import csv
import hashlib
import hmac
import json
import os
import shutil
import tempfile
import threading
import uuid
from contextlib import contextmanager
from collections.abc import Callable
from zipfile import ZIP_DEFLATED, ZipFile
from pathlib import Path
from docx import Document
from openpyxl import load_workbook
from lxml import etree

from .documents import DocxProcessor, _paragraph_text, _paragraphs
from .docx_preview import render_docx_preview
from .formats import XlsxProcessor, _read_text, process_text_file, restore_text_file
from .odt import (
    OdtProcessor,
    extract_odt_text,
    read_odt_artifact_stamp,
    write_odt_artifact_stamp,
)
from .storage import MappingError, Storage
from .word_com import convert_doc_to_docx


SUPPORTED = {".docx", ".doc", ".odt", ".xlsx", ".txt", ".csv", ".md", ".log"}
_TEXT_STAMP_PREFIX = "<!-- LocalAnonymizer-Artifact:"
_CORE_PROPERTIES_PART = "docProps/core.xml"
_DC_IDENTIFIER = "{http://purl.org/dc/elements/1.1/}identifier"


def _read_core_properties(path: Path) -> tuple[bytes, str | None]:
    with ZipFile(path, "r") as archive:
        names = [info.filename for info in archive.infolist()]
        if len(names) != len(set(names)) or names.count(_CORE_PROPERTIES_PART) != 1:
            raise MappingError("Некорректный OOXML-пакет")
        info = archive.getinfo(_CORE_PROPERTIES_PART)
        if info.file_size > 1024 * 1024:
            raise MappingError("Некорректные свойства OOXML-пакета")
        payload = archive.read(info)
    parser = etree.XMLParser(resolve_entities=False, no_network=True, load_dtd=False, recover=False)
    root = etree.fromstring(payload, parser=parser)
    identifier = root.find(_DC_IDENTIFIER)
    return payload, identifier.text if identifier is not None else None


def _write_core_identifier(path: Path, value: str) -> None:
    payload, _ = _read_core_properties(path)
    parser = etree.XMLParser(resolve_entities=False, no_network=True, load_dtd=False, recover=False)
    root = etree.fromstring(payload, parser=parser)
    identifier = root.find(_DC_IDENTIFIER)
    if identifier is None:
        identifier = etree.SubElement(root, _DC_IDENTIFIER)
    identifier.text = value
    updated = etree.tostring(root, encoding="UTF-8", xml_declaration=True, standalone=True)

    fd, temp_name = tempfile.mkstemp(prefix="artifact-stamp-", suffix=path.suffix, dir=path.parent)
    os.close(fd)
    temp = Path(temp_name)
    try:
        with ZipFile(path, "r") as source, ZipFile(temp, "w", ZIP_DEFLATED) as target:
            for info in source.infolist():
                target.writestr(info, updated if info.filename == _CORE_PROPERTIES_PART else source.read(info))
        os.replace(temp, path)
    finally:
        temp.unlink(missing_ok=True)


def _artifact_stamp(document_id: str, operation_id: str) -> str:
    return f"urn:local-anonymizer:artifact:v1:{document_id}:{operation_id}"


def _reject_anonymized_source(path: Path) -> None:
    """Prevent an exported artifact from becoming a new anonymization source."""
    suffix = path.suffix.lower()
    actual = None
    if suffix in {".docx", ".xlsx"}:
        try:
            _, actual = _read_core_properties(path)
        except Exception:
            # Normal format validation reports malformed packages later.  This
            # guard is only responsible for positively identified artifacts.
            return
    elif suffix == ".odt":
        actual = read_odt_artifact_stamp(path)
    elif suffix != ".doc":
        try:
            marker = _read_text(path).partition("\n")[0]
        except Exception:
            return
        if marker.startswith(_TEXT_STAMP_PREFIX):
            actual = marker[len(_TEXT_STAMP_PREFIX):].strip().removesuffix("-->").strip()
    if isinstance(actual, str) and actual.startswith("urn:local-anonymizer:artifact:"):
        raise ValueError(
            "Этот файл уже обезличен. Не обезличивайте его повторно: "
            "для возврата исходных данных используйте раздел восстановления."
        )


def _read_artifact_identity(path: Path) -> tuple[str, str]:
    suffix = path.suffix.lower()
    try:
        if suffix in {".docx", ".xlsx"}:
            _, actual = _read_core_properties(path)
        elif suffix == ".odt":
            actual = read_odt_artifact_stamp(path)
        else:
            marker = _read_text(path).partition("\n")[0]
            actual = marker[len(_TEXT_STAMP_PREFIX):].strip().removesuffix("-->").strip() \
                if marker.startswith(_TEXT_STAMP_PREFIX) else None
    except Exception as exc:
        raise MappingError("Не удалось прочитать идентификатор обезличенного файла") from exc
    prefix = "urn:local-anonymizer:artifact:v1:"
    if not isinstance(actual, str) or not actual.startswith(prefix):
        raise MappingError("Файл не является обезличенным файлом LocalAnonymizer")
    identity = actual[len(prefix):].split(":", 1)
    if len(identity) != 2 or not all(identity):
        raise MappingError("Повреждён идентификатор обезличенного файла")
    return identity[0], identity[1]


def _pair_id(document_id: str, operation_id: str) -> str:
    return hashlib.sha256(_artifact_stamp(document_id, operation_id).encode("utf-8")).hexdigest()


def _canonical_json(value: object) -> bytes:
    return json.dumps(
        value, ensure_ascii=False, sort_keys=True, separators=(",", ":")
    ).encode("utf-8")


def _recovery_checksum(payload: dict) -> str:
    protected = {
        "pair_id": payload["pair_id"],
        "anonymized_name": payload["anonymized_name"],
        "bundle": payload["bundle"],
    }
    return hashlib.sha256(_canonical_json(protected)).hexdigest()


def _stamp_artifact(path: Path, document_id: str, operation_id: str) -> None:
    stamp = _artifact_stamp(document_id, operation_id)
    suffix = path.suffix.lower()
    if suffix in {".docx", ".xlsx"}:
        _write_core_identifier(path, stamp)
    elif suffix == ".odt":
        write_odt_artifact_stamp(path, stamp)
    else:
        body = _read_text(path)
        path.write_text(f"{_TEXT_STAMP_PREFIX} {stamp} -->\n{body}", encoding="utf-8")


def _verify_artifact_stamp(path: Path, document_id: str, operation_id: str) -> str | None:
    expected = _artifact_stamp(document_id, operation_id)
    suffix = path.suffix.lower()
    try:
        if suffix in {".docx", ".xlsx"}:
            _, actual = _read_core_properties(path)
        elif suffix == ".odt":
            actual = read_odt_artifact_stamp(path)
        else:
            text = _read_text(path)
            marker, separator, body = text.partition("\n")
            actual = marker[len(_TEXT_STAMP_PREFIX):].strip().removesuffix("-->").strip() \
                if marker.startswith(_TEXT_STAMP_PREFIX) else None
            if actual == expected:
                return body if separator else ""
    except Exception as exc:
        raise MappingError("Не удалось проверить привязку файла к операции") from exc
    if actual != expected:
        raise MappingError("Файл не соответствует документу или операции")
    return None


class AnonymizerService:
    def __init__(self, root: str | Path | None = None, *, storage: Storage | None = None,
                 legacy_document_converter: Callable[[Path, Path], Path] | None = None):
        if root is not None and storage is not None:
            raise ValueError("Передайте root или storage, но не оба параметра")
        self.storage = storage if storage is not None else Storage(root)
        self._legacy_document_converter = legacy_document_converter
        self._document_locks_guard = threading.Lock()
        self._document_locks: dict[str, list] = {}

    @contextmanager
    def _document_lock(self, document_id: str):
        with self._document_locks_guard:
            entry = self._document_locks.setdefault(document_id, [threading.Lock(), 0])
            entry[1] += 1
        lock = entry[0]
        lock.acquire()
        try:
            yield
        finally:
            lock.release()
            with self._document_locks_guard:
                entry[1] -= 1
                if entry[1] == 0 and self._document_locks.get(document_id) is entry:
                    del self._document_locks[document_id]

    def import_file(self, source: str | Path) -> str:
        source = Path(source)
        if source.suffix.lower() not in SUPPORTED:
            raise ValueError("Неподдерживаемый формат файла")
        _reject_anonymized_source(source)
        doc_id = self.storage.import_document(source)
        if source.suffix.lower() == ".doc":
            try:
                row = self.storage.get_document(doc_id)
                legacy = Path(row["owned_source"])
                converted = legacy.with_name("converted.docx")
                converter = (
                    self._legacy_document_converter
                    if self._legacy_document_converter is not None
                    else convert_doc_to_docx
                )
                converter(legacy, converted)
                self.storage.connection.execute("update documents set owned_source=?,status=? where id=?",
                                                (str(converted), "Преобразован из DOC", doc_id))
                self.storage.connection.commit()
                legacy.unlink(missing_ok=True)
            except Exception:
                self.storage.connection.rollback()
                self.storage.discard_import(doc_id)
                raise
        return doc_id

    def preview_document(self, document_id: str) -> dict:
        row = self.storage.get_document(document_id)
        source = Path(row["owned_source"])
        suffix = source.suffix.lower()
        rich_html = ""
        if suffix == ".docx":
            document = Document(source)
            text = "\n".join(_paragraph_text(p) for p in _paragraphs(document) if _paragraph_text(p))
            rich_html = render_docx_preview(document)
        elif suffix == ".odt":
            text = extract_odt_text(source)
        elif suffix == ".xlsx":
            workbook = load_workbook(source, data_only=False, read_only=True, keep_links=False)
            try:
                text = "\n".join(str(cell.value) for sheet in workbook.worksheets
                                  for row_cells in sheet.iter_rows() for cell in row_cells
                                  if isinstance(cell.value, str) and not cell.value.startswith("="))
            finally:
                workbook.close()
        else:
            text = _read_text(source)
        return {"display_name": row["display_name"], "text": text, "status": row["status"],
                "kind": "document", "format": suffix.removeprefix("."), "rich_html": rich_html}

    def preview_anonymized_document(self, document_id: str, engine=None,
                                    redact_graphics: bool = True) -> dict:
        row = self.storage.get_document(document_id)
        source = Path(row["owned_source"])
        if source.suffix.lower() != ".docx":
            return {"format": source.suffix.lower().removeprefix("."), "rich_html": ""}
        with tempfile.TemporaryDirectory(prefix="local-anonymizer-preview-") as folder:
            output = Path(folder) / "anonymized-preview.docx"
            result = DocxProcessor(engine).anonymize(
                source, output, redact_graphics=redact_graphics
            )
            document = Document(output)
            return {
                "format": "docx",
                "rich_html": render_docx_preview(document),
                "replacement_count": result.total,
                "candidate_count": result.hint_count,
                "warnings": result.warnings,
            }

    def anonymize(self, document_id: str, engine=None, *, redact_graphics: bool = True) -> dict:
        with self._document_lock(document_id):
            return self._anonymize_locked(document_id, engine, redact_graphics=redact_graphics)

    def _anonymize_locked(self, document_id: str, engine=None,
                          *, redact_graphics: bool = True) -> dict:
        row = self.storage.get_document(document_id)
        previous_output = row.get("output_path")
        source = Path(row["owned_source"])
        _reject_anonymized_source(source)
        operation_id = str(uuid.uuid4())
        folder = self.storage.documents_dir / document_id
        suffix = source.suffix.lower()
        output = folder / (f"anonymized-{operation_id}" + suffix)
        try:
            if suffix == ".docx":
                result = DocxProcessor(engine).anonymize(
                    source, output, redact_graphics=redact_graphics
                )
            elif suffix == ".odt":
                result = OdtProcessor(engine).anonymize(source, output)
            elif suffix == ".xlsx":
                result = XlsxProcessor(engine).anonymize(source, output)
            else:
                result = process_text_file(source, output, engine)
            if not output.is_file():
                raise RuntimeError("Выходной файл не создан")
            _stamp_artifact(output, document_id, operation_id)
            if suffix in {".docx", ".odt"} and isinstance(result.package_archive, dict):
                exact_keys = {
                    "exact_source_b64",
                    "exact_source_sha256",
                    "exact_anonymized_sha256",
                }
                present = exact_keys.intersection(result.package_archive)
                if present and present != exact_keys:
                    raise RuntimeError(f"Повреждён точный снимок исходного {suffix[1:].upper()}")
                if present == exact_keys:
                    result.package_archive["exact_anonymized_sha256"] = hashlib.sha256(
                        output.read_bytes()
                    ).hexdigest()
            stop = any(isinstance(warning, dict) and warning.get("level") == "stop"
                       for warning in result.warnings)
            status = "Экспорт запрещён" if stop else ("Готов" if result.ready else "Требуется ручная проверка")
            self.storage.commit_operation(
                document_id, operation_id, result.mapping, output, status,
                result.package_archive,
            )
        except Exception:
            output.unlink(missing_ok=True)
            raise
        if previous_output and Path(previous_output) != output:
            try:
                self.storage._owned(previous_output).unlink(missing_ok=True)
            except (OSError, ValueError):
                pass
        return {"document_id": document_id, "operation_id": operation_id, "output_path": str(output),
                "replacement_count": result.total, "legend_total": len(result.mapping),
                "candidate_count": result.hint_count,
                "candidates": result.candidates, "warnings": result.warnings,
                "ready": status == "Готов", "status": status}

    def restore(self, document_id: str, edited: str | Path,
                destination: str | Path | None = None) -> dict:
        row = self.storage.get_document(document_id)
        operation_id = row.get("operation_id")
        if not operation_id:
            raise RuntimeError("Документ ещё не обезличен")
        mapping, package_archive = self.storage.load_restore_bundle(document_id, operation_id)
        edited = Path(edited)
        output = (Path(destination) if destination is not None else
                  self.storage.documents_dir / document_id /
                  (edited.stem + ".restored" + edited.suffix.lower()))
        return self._restore_with_bundle(
            edited, output, document_id, operation_id, mapping, package_archive
        )

    def _restore_with_bundle(self, edited: Path, output: Path, document_id: str,
                             operation_id: str, mapping: dict[str, str],
                             package_archive: dict | None) -> dict:
        suffix = edited.suffix.lower()
        if suffix not in {".docx", ".odt", ".xlsx", ".txt", ".csv", ".md", ".log"}:
            raise ValueError("Неподдерживаемый формат восстановления")
        if output.suffix.lower() != suffix:
            raise ValueError("Формат сохраняемого файла должен совпадать с восстанавливаемым")
        output.parent.mkdir(parents=True, exist_ok=True)
        fd, snapshot_name = tempfile.mkstemp(prefix="restore-", suffix=suffix,
                                             dir=self.storage.temp_dir)
        snapshot = Path(snapshot_name)
        try:
            with os.fdopen(fd, "wb") as target, edited.open("rb") as source:
                shutil.copyfileobj(source, target)
                target.flush()
                os.fsync(target.fileno())
            text_body = _verify_artifact_stamp(snapshot, document_id, operation_id)
            if suffix == ".docx":
                unknown = DocxProcessor().restore(
                    snapshot, output, mapping, package_archive=package_archive
                )
            elif suffix == ".odt":
                unknown = OdtProcessor().restore(
                    snapshot, output, mapping, package_archive=package_archive
                )
            elif suffix == ".xlsx":
                unknown = XlsxProcessor().restore(snapshot, output, mapping)
            else:
                unstamped = output.with_name(f".{output.stem}.validated{suffix}")
                try:
                    unstamped.write_text(text_body or "", encoding="utf-8")
                    unknown = restore_text_file(unstamped, output, mapping)
                finally:
                    unstamped.unlink(missing_ok=True)
        finally:
            snapshot.unlink(missing_ok=True)
        return {"output_path": str(output), "unknown_markers": unknown}

    @staticmethod
    def _recovery_name(anonymized: Path) -> str:
        stem = anonymized.stem
        for suffix in (" — обезличенный", ".anonymized"):
            if stem.casefold().endswith(suffix.casefold()):
                stem = stem[:-len(suffix)]
                break
        return f"{stem} — информация для восстановления.larecovery"

    @staticmethod
    def _read_recovery_file(path: Path) -> dict:
        try:
            payload = json.loads(path.read_text(encoding="utf-8"))
        except Exception as exc:
            raise MappingError("Файл восстановления повреждён") from exc
        if (not isinstance(payload, dict)
                or payload.get("format") != "local-anonymizer-recovery"
                or payload.get("version") not in {1, 2}
                or not isinstance(payload.get("pair_id"), str)):
            raise MappingError("Неизвестный формат файла восстановления")
        if payload["version"] == 1:
            if not isinstance(payload.get("envelope"), str):
                raise MappingError("Неизвестный формат файла восстановления")
            return payload
        if (not isinstance(payload.get("anonymized_name"), str)
                or not isinstance(payload.get("bundle"), dict)
                or not isinstance(payload.get("checksum_sha256"), str)):
            raise MappingError("Неизвестный формат файла восстановления")
        expected = _recovery_checksum(payload)
        if not hmac.compare_digest(payload["checksum_sha256"], expected):
            raise MappingError("Файл восстановления повреждён: не совпадает контрольная сумма")
        return payload

    def _recovery_payload(self, row: dict, anonymized_name: str) -> bytes:
        document_id = row["id"]
        operation_id = row.get("operation_id")
        if not operation_id:
            raise RuntimeError("Документ ещё не обезличен")
        mapping, package_archive = self.storage.load_restore_bundle(document_id, operation_id)
        bundle = {
            "schema": 2 if package_archive is not None else 1,
            "document_id": document_id,
            "operation_id": operation_id,
            "marker_manifest": sorted(mapping),
            "mapping": mapping,
        }
        if package_archive is not None:
            bundle["package_archive"] = package_archive
        payload = {
            "format": "local-anonymizer-recovery",
            "version": 2,
            "pair_id": _pair_id(document_id, operation_id),
            "anonymized_name": anonymized_name,
            "bundle": bundle,
        }
        payload["checksum_sha256"] = _recovery_checksum(payload)
        return _canonical_json(payload)

    def export_pair(self, document_id: str, destination: str | Path,
                    *, confirmation: bool = False) -> dict[str, Path]:
        row = self.storage.get_document(document_id)
        anonymized = Path(destination)
        recovery = anonymized.with_name(self._recovery_name(anonymized))
        if anonymized.resolve() == recovery.resolve():
            raise ValueError("Имена файлов пары должны различаться")
        recovery.parent.mkdir(parents=True, exist_ok=True)
        temporary = recovery.with_name("." + recovery.name + ".tmp")
        try:
            temporary.write_bytes(self._recovery_payload(row, anonymized.name))
            self.export_anonymized(document_id, anonymized, confirmation=confirmation)
            os.replace(temporary, recovery)
        except Exception:
            temporary.unlink(missing_ok=True)
            raise
        return {"anonymized_path": anonymized, "recovery_path": recovery}

    def find_recovery_companion(self, anonymized: str | Path) -> Path:
        anonymized = Path(anonymized)
        document_id, operation_id = _read_artifact_identity(anonymized)
        expected_pair = _pair_id(document_id, operation_id)
        preferred = anonymized.with_name(self._recovery_name(anonymized))
        candidates = [preferred] if preferred.is_file() else []
        candidates.extend(
            path for path in anonymized.parent.glob("*.larecovery") if path != preferred
        )
        matches = []
        for candidate in candidates:
            try:
                if self._read_recovery_file(candidate).get("pair_id") == expected_pair:
                    matches.append(candidate)
            except MappingError:
                continue
        if not matches:
            raise MappingError("Рядом не найден соответствующий файл восстановления")
        return matches[0]

    def restore_from_pair(self, anonymized: str | Path, destination: str | Path,
                          *, recovery_path: str | Path | None = None) -> dict:
        anonymized = Path(anonymized)
        document_id, operation_id = _read_artifact_identity(anonymized)
        recovery = (Path(recovery_path) if recovery_path is not None
                    else self.find_recovery_companion(anonymized))
        payload = self._read_recovery_file(recovery)
        if payload["pair_id"] != _pair_id(document_id, operation_id):
            raise MappingError("Файл восстановления не соответствует обезличенному документу")
        if payload["version"] == 1:
            try:
                envelope = base64.b64decode(payload["envelope"], validate=True)
            except Exception as exc:
                raise MappingError("Файл восстановления повреждён") from exc
            mapping, package_archive = self.storage.decode_restore_envelope(
                envelope, document_id, operation_id
            )
        else:
            mapping, package_archive = self.storage.decode_restore_payload(
                payload["bundle"], document_id, operation_id
            )
        report = self._restore_with_bundle(
            anonymized, Path(destination), document_id, operation_id, mapping, package_archive
        )
        report["recovery_path"] = str(recovery)
        return report

    def export_anonymized(self, document_id: str, destination: str | Path,
                          *, confirmation: bool = False) -> Path:
        row = self.storage.get_document(document_id)
        status = row["status"]
        if status == "Требуется ручная проверка" and not confirmation:
            raise PermissionError("Экспорт требует явного подтверждения ручной проверки")
        if status not in {"Готов", "Требуется ручная проверка"}:
            raise PermissionError("Экспорт запрещён для текущего состояния документа")
        source = Path(row["output_path"])
        destination = Path(destination)
        destination.parent.mkdir(parents=True, exist_ok=True)
        shutil.copy2(source, destination)
        return destination

    def export_mapping(self, document_id: str, destination: str | Path) -> Path:
        row = self.storage.get_document(document_id)
        mapping = self.storage.load_mapping(document_id, row["operation_id"])
        return export_mapping_csv(mapping, destination)

    def export_all_ready(self, destination: str | Path) -> Path:
        """Export every ready document as an independently restorable pair."""
        destination = Path(destination)
        destination.parent.mkdir(parents=True, exist_ok=True)
        temp = destination.with_name("." + destination.name + ".tmp")
        used: set[str] = set()
        try:
            with ZipFile(temp, "w", compression=ZIP_DEFLATED) as archive:
                for row in self.storage.list_documents():
                    if row.get("status") != "Готов" or not row.get("output_path"):
                        continue
                    try:
                        output = self.storage._owned(row["output_path"])
                    except ValueError:
                        continue
                    if not output.is_file():
                        continue
                    original = Path(row["display_name"])
                    pair_stem, index = original.stem, 2
                    anonymous_name = f"{pair_stem} — обезличенный{output.suffix.lower()}"
                    recovery_name = f"{pair_stem} — информация для восстановления.larecovery"
                    while (anonymous_name.casefold() in used
                           or recovery_name.casefold() in used):
                        pair_stem = f"{original.stem}-{index}"
                        anonymous_name = f"{pair_stem} — обезличенный{output.suffix.lower()}"
                        recovery_name = f"{pair_stem} — информация для восстановления.larecovery"
                        index += 1
                    used.update({anonymous_name.casefold(), recovery_name.casefold()})
                    archive.write(output, anonymous_name)
                    archive.writestr(
                        recovery_name,
                        self._recovery_payload(row, anonymous_name),
                    )
            temp.replace(destination)
        finally:
            temp.unlink(missing_ok=True)
        return destination


def export_mapping_csv(mapping: dict[str, str], destination: str | Path) -> Path:
    destination = Path(destination)
    destination.parent.mkdir(parents=True, exist_ok=True)
    temp = destination.with_name("." + destination.name + ".tmp")
    try:
        with temp.open("w", encoding="utf-8-sig", newline="") as stream:
            writer = csv.writer(stream)
            writer.writerow(["Маркер", "Исходное значение"])
            for marker, value in mapping.items():
                if value.startswith(("=", "+", "-", "@")):
                    value = "'" + value
                writer.writerow([marker, value])
            stream.flush()
        temp.replace(destination)
    finally:
        temp.unlink(missing_ok=True)
    return destination
