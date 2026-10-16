from __future__ import annotations

import argparse
import base64
import json
import os
import shutil
import sys
import tempfile
from pathlib import Path
from zipfile import ZipFile

from docx import Document
from docx.shared import Inches

from .documents import DocxProcessor
from .engine import Anonymizer, MARKER_PATTERN, normalize_label
from .platform_services import (
    ClipboardReader,
    DataProtector,
    default_clipboard_reader,
    default_data_protector,
)
from .service import AnonymizerService

APP_VERSION = "1.8"
INDEX_ASSET = "index-v1.8.18.html"
ENGLISH_INDEX_ASSET = "index-en-v1.8.18.html"


def _read_clipboard_text() -> str:
    return default_clipboard_reader().read_text()


def asset_path(name: str) -> Path:
    base = Path(getattr(sys, "_MEIPASS", Path(__file__).resolve().parent))
    candidate = base / "assets" / name
    if not candidate.exists():
        candidate = Path(__file__).resolve().parent / "assets" / name
    return candidate


class AppApi:
    def __init__(
        self,
        root=None,
        *,
        service: AnonymizerService | None = None,
        clipboard_reader: ClipboardReader | None = None,
    ):
        if service is not None and root is not None:
            raise ValueError("Передайте root или service, но не оба")
        self.service = service if service is not None else AnonymizerService(root)
        self._clipboard_reader = clipboard_reader
        self._window = None
        self._restore_pair: tuple[Path, Path] | None = None
        self._ui_preview: dict | None = None
        if not self.service.storage.connection.execute("select 1 from profiles where name='Общий'").fetchone():
            self.service.storage.save_profile("Общий", {}, [], paranoid=False)

    def set_window(self, window):
        self._window = window

    def read_clipboard_text(self) -> dict:
        text = (
            self._clipboard_reader.read_text()
            if self._clipboard_reader is not None
            else _read_clipboard_text()
        )
        return {"text": text}

    def write_clipboard_text(self, text: str) -> dict:
        if not isinstance(text, str):
            raise TypeError("Текст буфера обмена должен быть строкой")
        if len(text.encode("utf-8")) > 10 * 1024 * 1024:
            raise ValueError("Содержимое буфера обмена слишком велико")
        clipboard = (
            self._clipboard_reader
            if self._clipboard_reader is not None
            else default_clipboard_reader()
        )
        clipboard.write_text(text)
        return {"ok": True}

    def _read_ui_settings(self) -> dict:
        path = self.service.storage.root / "ui-settings.json"
        try:
            value = json.loads(path.read_text(encoding="utf-8"))
        except (OSError, ValueError, TypeError):
            return {}
        return value if isinstance(value, dict) else {}

    def _write_ui_settings(self, **changes) -> None:
        path = self.service.storage.root / "ui-settings.json"
        settings = self._read_ui_settings()
        settings.update(changes)
        temporary = path.with_suffix(".tmp")
        temporary.write_text(json.dumps(settings, ensure_ascii=False), encoding="utf-8")
        temporary.replace(path)

    def get_theme(self) -> str | None:
        theme = self._read_ui_settings().get("theme")
        return theme if theme in {"light", "dark"} else None

    def set_theme(self, theme: str) -> dict:
        if theme not in {"light", "dark"}:
            raise ValueError("Неизвестная тема")
        self._write_ui_settings(theme=theme)
        return {"ok": True, "theme": theme}

    def get_design_mode(self) -> str | None:
        design_mode = self._read_ui_settings().get("design_mode")
        return design_mode if design_mode in {"new", "legacy"} else None

    def set_design_mode(self, design_mode: str) -> dict:
        if design_mode not in {"new", "legacy"}:
            raise ValueError("Неизвестный дизайн")
        self._write_ui_settings(design_mode=design_mode)
        return {"ok": True, "design_mode": design_mode}

    def get_language(self) -> str:
        language = self._read_ui_settings().get("language")
        return language if language in {"ru", "en"} else "ru"

    def set_language(self, language: str) -> dict:
        if language not in {"ru", "en"}:
            raise ValueError("Неизвестный язык")
        self._write_ui_settings(language=language)
        return {"ok": True, "language": language}

    def list_documents(self) -> list[dict]:
        allowed = ("id", "display_name", "status")
        return [{key: row[key] for key in allowed} for row in self.service.storage.list_documents()]

    def preview_document(self, document_id: str) -> dict:
        return self.service.preview_document(document_id)

    def start_document_preview(self, document_id: str) -> dict:
        preview = self.service.preview_document(document_id, include_rich=False)
        text = preview.pop("text", "")
        rich_html = preview.pop("rich_html", "")
        if not isinstance(text, str) or not isinstance(rich_html, str):
            raise TypeError("Некорректный предпросмотр документа")
        token = os.urandom(16).hex()
        self._ui_preview = {"token": token, "text": text, "rich_html": rich_html}
        return {
            **preview,
            "token": token,
            "text_length": len(text),
            "rich_html_length": len(rich_html),
        }

    def start_formatted_document_preview(self, document_id: str) -> dict:
        preview = self.service.preview_document(document_id)
        preview.pop("text", None)
        rich_html = preview.pop("rich_html", "")
        if not isinstance(rich_html, str):
            raise TypeError("Некорректный форматированный предпросмотр")
        token = os.urandom(16).hex()
        self._ui_preview = {"token": token, "text": "", "rich_html": rich_html}
        return {**preview, "token": token, "rich_html_length": len(rich_html)}

    def read_document_preview_chunk(self, token: str, field: str, offset: int) -> dict:
        cached = self._ui_preview
        if cached is None or token != cached["token"]:
            raise ValueError("Предпросмотр устарел")
        if field not in {"text", "rich_html"}:
            raise ValueError("Неизвестная часть предпросмотра")
        if not isinstance(offset, int) or isinstance(offset, bool) or offset < 0:
            raise ValueError("Некорректное смещение предпросмотра")
        value = cached[field]
        chunk_size = 16_384
        chunk = value[offset:offset + chunk_size]
        return {"chunk": chunk, "next_offset": offset + len(chunk)}

    def finish_document_preview(self, token: str) -> dict:
        if self._ui_preview is not None and token == self._ui_preview["token"]:
            self._ui_preview = None
        return {"ok": True}

    def preview_anonymized_document(self, document_id: str, paranoid=False, rules="", whitelist="", disabled=None,
                                    compact_markers: bool = False,
                                    redact_graphics: bool = True) -> dict:
        return self.service.preview_anonymized_document(
            document_id,
            self._engine(paranoid, rules, whitelist, disabled, compact_markers),
            redact_graphics=redact_graphics,
        )

    def start_anonymized_document_preview(self, document_id: str, paranoid=False, rules="", whitelist="", disabled=None,
                                           compact_markers: bool = False,
                                           redact_graphics: bool = True) -> dict:
        preview = self.preview_anonymized_document(
            document_id, paranoid, rules, whitelist, disabled,
            compact_markers, redact_graphics,
        )
        rich_html = preview.pop("rich_html", "")
        if not isinstance(rich_html, str):
            raise TypeError("Некорректный форматированный предпросмотр")
        token = os.urandom(16).hex()
        self._ui_preview = {"token": token, "text": "", "rich_html": rich_html}
        return {**preview, "token": token, "rich_html_length": len(rich_html)}

    def import_paths(self, paths: list[str]) -> dict:
        ids, errors = [], []
        for path in paths:
            try:
                ids.append(self.service.import_file(path))
            except Exception as exc:
                errors.append({"file": Path(path).name, "error": type(exc).__name__})
        return {"ok": bool(ids), "ids": ids, "errors": errors, "documents": self.list_documents()}

    def choose_files(self) -> dict:
        if self._window is None:
            return {"ok": False, "ids": [], "errors": [{"error": "NoWindow"}]}
        paths = self._window.create_file_dialog(10, allow_multiple=True,
            file_types=("Поддерживаемые (*.docx;*.doc;*.odt;*.xlsx;*.txt;*.csv;*.md;*.log)",)) or []
        return self.import_paths(list(paths))

    def _engine(self, paranoid=False, rules="", whitelist="", disabled=None,
                compact_markers: bool = False) -> Anonymizer:
        return Anonymizer(user_dictionary=self._rules(rules),
                          whitelist=[x.strip() for x in whitelist.splitlines() if x.strip()],
                          paranoid=paranoid, disabled_categories=disabled or [],
                          compact_markers=compact_markers)

    def anonymize_document(self, document_id: str, paranoid=False, rules="", whitelist="", disabled=None,
                           compact_markers: bool = False,
                           redact_graphics: bool = True) -> dict:
        internal = self.service.anonymize(
            document_id,
            self._engine(paranoid, rules, whitelist, disabled, compact_markers),
            redact_graphics=redact_graphics,
        )
        mapping = self.service.storage.load_mapping(document_id, internal["operation_id"])
        allowed = ("replacement_count", "legend_total", "candidate_count", "candidates",
                   "warnings", "ready", "status")
        report = {key: internal[key] for key in allowed}
        report["mapping"] = mapping
        return report

    def delete_document(self, document_id: str) -> dict:
        self.service.storage.delete_document(document_id)
        return {"ok": True, "documents": self.list_documents()}

    def export_anonymized(self, document_id: str, confirmation: bool = False) -> dict:
        row = self.service.storage.get_document(document_id)
        source = Path(row["output_path"] or "")
        display = Path(row["display_name"])
        suggested = f"{display.stem} — обезличенный{display.suffix.lower()}"
        paths = self._window.create_file_dialog(30, save_filename=suggested) if self._window and source.is_file() else None
        if not paths:
            return {"ok": False}
        pair = self.service.export_pair(document_id, paths[0], confirmation=confirmation)
        return {"ok": True, "path": str(pair["anonymized_path"]),
                "recovery_path": str(pair["recovery_path"])}

    def choose_restore_pair(self) -> dict:
        self._restore_pair = None
        paths = self._window.create_file_dialog(
            10,
            allow_multiple=False,
            directory=str(Path.home().resolve()),
            file_types=("Обезличенный файл (*.docx;*.odt;*.xlsx;*.txt;*.csv;*.md;*.log)",),
        ) if self._window else None
        if not paths:
            return {"ok": False}
        anonymized = Path(paths[0])
        recovery = self.service.find_recovery_companion(anonymized)
        self._restore_pair = (anonymized, recovery)
        return {"ok": True, "anonymized_name": anonymized.name,
                "recovery_name": recovery.name}

    def restore_matched_file(self) -> dict:
        if self._restore_pair is None:
            raise RuntimeError("Сначала выберите обезличенный файл и подтвердите пару")
        anonymized, recovery = self._restore_pair
        stem = anonymized.stem
        marker = " — обезличенный"
        if stem.casefold().endswith(marker.casefold()):
            stem = stem[:-len(marker)]
        destinations = self._window.create_file_dialog(
            30,
            save_filename=f"{stem} — восстановленный{anonymized.suffix.lower()}",
            directory=str(anonymized.parent.resolve()),
            file_types=(f"{anonymized.suffix.lstrip('.').upper()} (*{anonymized.suffix.lower()})",),
        ) if self._window else None
        if not destinations:
            return {"ok": False}
        destination = Path(destinations[0])
        report = self.service.restore_from_pair(
            anonymized, destination, recovery_path=recovery
        )
        return {"ok": True, "path": str(destination),
                "unknown_markers": report["unknown_markers"]}

    def export_all_anonymized(self) -> dict:
        paths = self._window.create_file_dialog(
            30, save_filename="обезличенные-файлы.zip",
            file_types=("ZIP (*.zip)",)) if self._window else None
        if not paths:
            return {"ok": False}
        target = self.service.export_all_ready(paths[0])
        return {"ok": True, "path": str(target)}

    def export_mapping(self, document_id: str) -> dict:
        paths = self._window.create_file_dialog(30, save_filename="mapping-sensitive.csv",
                                                file_types=("CSV (*.csv)",)) if self._window else None
        if not paths:
            return {"ok": False}
        target = self.service.export_mapping(document_id, paths[0])
        return {"ok": True, "path": str(target)}

    def restore_document(self, document_id: str) -> dict:
        try:
            document = self.service.storage.get_document(document_id)
        except KeyError:
            document = {}
        output_path = document.get("output_path")
        initial_directory = (
            Path(output_path).parent if output_path else self.service.storage.documents_dir
        )
        if not initial_directory.is_dir():
            initial_directory = Path.home()
        paths = self._window.create_file_dialog(
            10,
            allow_multiple=False,
            directory=str(initial_directory.resolve()),
            file_types=("Поддерживаемые (*.docx;*.odt;*.xlsx;*.txt;*.csv;*.md;*.log)",),
        ) if self._window else None
        if not paths:
            return {"ok": False}
        edited = Path(paths[0])
        destinations = self._window.create_file_dialog(
            30, save_filename=f"{edited.stem}.restored{edited.suffix.lower()}",
            directory=str(edited.parent.resolve()),
            file_types=(f"{edited.suffix.lstrip('.').upper()} (*{edited.suffix.lower()})",),
        ) if self._window else None
        if not destinations:
            return {"ok": False}
        destination = Path(destinations[0])
        report = self.service.restore(document_id, edited, destination)
        return {"ok": True, "path": str(destination),
                "unknown_markers": report["unknown_markers"]}

    @staticmethod
    def _rules(rules: str) -> dict[str, str]:
        result = {}
        for raw in rules.splitlines():
            line = raw.strip()
            if line:
                phrase, sep, label = line.partition("=>")
                result[phrase.strip()] = normalize_label(
                    (label if sep else "СЛОВАРЬ").strip("[] ")
                )
        return result

    def list_profiles(self) -> list[str]:
        return [row[0] for row in self.service.storage.connection.execute("select name from profiles order by name")]

    def save_profile(self, name: str, rules: str, whitelist: str, paranoid: bool) -> dict:
        self.service.storage.save_profile(name, self._rules(rules), [x.strip() for x in whitelist.splitlines() if x.strip()], paranoid=paranoid)
        return {"ok": True}

    def load_profile(self, name: str) -> dict:
        return self.service.storage.load_profile(name)

    def delete_profile(self, name: str) -> dict:
        if name == "Общий":
            raise ValueError("Профиль Общий нельзя удалить")
        self.service.storage.connection.execute("delete from profiles where name=?", (name,)); self.service.storage.connection.commit()
        return {"ok": True}

    def export_profile(self, name: str, rules: str | None = None, whitelist: str | None = None,
                       paranoid: bool | None = None) -> dict:
        paths = self._window.create_file_dialog(30, save_filename=f"{name}.json", file_types=("JSON (*.json)",)) if self._window else None
        if not paths:
            return {"ok": False}
        try:
            if rules is None or whitelist is None or paranoid is None:
                exported = self.service.storage.export_profile(name)
            else:
                payload = self.service.storage._validate_profile({
                    "version": 1,
                    "name": name,
                    "dictionary": self._rules(rules),
                    "whitelist": [x.strip() for x in whitelist.splitlines() if x.strip()],
                    "paranoid": bool(paranoid),
                })
                exported = json.dumps(payload, ensure_ascii=False, indent=2)
            Path(paths[0]).write_text(exported, encoding="utf-8")
        except ValueError as exc:
            return {"ok": False, "error": f"Профиль не экспортирован: {exc}"}
        except OSError:
            return {"ok": False, "error": "Не удалось записать файл профиля."}
        return {"ok": True, "path": paths[0]}

    def import_profile(self) -> dict:
        paths = self._window.create_file_dialog(10, allow_multiple=False, file_types=("JSON (*.json)",)) if self._window else None
        if not paths:
            return {"ok": False}
        try:
            text = Path(paths[0]).read_text(encoding="utf-8")
            name = self.service.storage.import_profile(text)
        except UnicodeDecodeError:
            return {"ok": False, "error": "Файл профиля должен быть JSON в кодировке UTF-8."}
        except json.JSONDecodeError:
            return {"ok": False, "error": "Файл профиля содержит некорректный JSON."}
        except ValueError as exc:
            return {"ok": False, "error": f"Профиль не импортирован: {exc}"}
        except OSError:
            return {"ok": False, "error": "Не удалось прочитать файл профиля."}
        return {"ok": True, "name": name, "profiles": self.list_profiles()}

    def anonymize_text(self, text: str, paranoid: bool = False, rules: str = "", whitelist: str = "", disabled=None,
                       compact_markers: bool = False) -> dict:
        engine = self._engine(paranoid, rules, whitelist, disabled, compact_markers)
        result = engine.anonymize(text)
        return {"text": result.text, "mapping": result.mapping, "candidates": result.candidates,
                "parts": result.parts, "legend": result.legend, "counts": result.counts,
                "total": result.total, "legend_total": result.legend_total,
                "hintCount": result.hint_count, "warnings": result.warnings}

    def restore_text(self, text: str, mapping: dict[str, str]) -> dict:
        if not isinstance(mapping, dict) or any(not isinstance(k, str) or not isinstance(v, str)
                                                for k, v in mapping.items()):
            raise ValueError("Некорректная карта замен")
        unknown: list[str] = []
        import re
        def replace(match):
            marker = match.group(0)
            if marker in mapping:
                return mapping[marker]
            unknown.append(marker)
            return marker
        restored = re.sub(MARKER_PATTERN, replace, text)
        return {"text": restored, "unknown": sorted(set(unknown))}

    def engine_self_test(self) -> dict:
        vm_sample = "qzx-monitor0101"
        sample = f"Иванов Сергей Петрович, test@example.ru, пароль: Secret123!, {vm_sample}"
        result = Anonymizer().anonymize(sample)
        vm_node_ok = vm_sample not in result.text and "[УЗЕЛ-1]" in result.text
        graphics_roundtrip_ok = _graphics_roundtrip_self_test()
        return {"ok": "[ФИО-1]" in result.text and "[ПОЧТА-1]" in result.text and "Secret123" not in result.text and vm_node_ok and graphics_roundtrip_ok,
                "vm_node_ok": vm_node_ok, "graphics_roundtrip_ok": graphics_roundtrip_ok,
                "replacement_count": len(result.mapping),
                "secret_count": result.text.count("(удалено)"), "warnings": []}

    def help_summary(self) -> dict:
        return {"replacement_count": 0, "warning": "Требуется финальная ручная проверка. Значения таблицы в справку не включены."}


def _graphics_roundtrip_self_test() -> bool:
    work = Path(tempfile.mkdtemp(prefix="local-anonymizer-graphics-self-test-"))
    image_bytes = base64.b64decode(
        "iVBORw0KGgoAAAANSUhEUgAAAAEAAAABCAQAAAC1HAwCAAAAC0lEQVR42mNk+A8AAQUBAScY42YAAAAASUVORK5CYII="
    )
    try:
        image = work / "synthetic.png"
        source = work / "source.docx"
        anonymized = work / "anonymized.docx"
        restored = work / "restored.docx"
        image.write_bytes(image_bytes)
        document = Document()
        document.add_paragraph("Иванов Сергей Петрович")
        document.add_picture(str(image), width=Inches(1))
        document.save(source)
        result = DocxProcessor().anonymize(source, anonymized)
        with ZipFile(anonymized) as package:
            if any(name.startswith("word/media/") for name in package.namelist()):
                return False
        DocxProcessor().restore(
            anonymized,
            restored,
            result.mapping,
            package_archive=result.package_archive,
        )
        with ZipFile(restored) as package:
            restored_images = [
                package.read(name)
                for name in package.namelist()
                if name.startswith("word/media/")
            ]
        return restored_images == [image_bytes] and bool(Document(restored).inline_shapes)
    except Exception:
        if os.environ.get("LOCAL_ANONYMIZER_SMOKE_TRACE") == "1":
            import traceback

            traceback.print_exc()
        return False
    finally:
        shutil.rmtree(work, ignore_errors=True)


def smoke_check() -> dict:
    html = asset_path(INDEX_ASSET)
    result = AppApi().engine_self_test()
    return {"ok": html.is_file() and result["ok"], "asset": str(html), "engine": result}


def platform_smoke_check(
    platform_name: str | None = None,
    *,
    protector: DataProtector | None = None,
    clipboard: ClipboardReader | None = None,
) -> dict:
    """Check native adapters without reading or overwriting clipboard data."""
    platform_name = platform_name or sys.platform
    secure_storage = False
    error = None
    try:
        if protector is None:
            protector = default_data_protector(platform_name)
        sample = os.urandom(32)
        purpose = "platform-smoke-v1"
        secure_storage = protector.unprotect(
            protector.protect(sample, purpose), purpose
        ) == sample
        if not secure_storage:
            error = "Secure storage roundtrip mismatch"
    except Exception as exc:
        error = f"{type(exc).__name__}: {exc}"

    try:
        if clipboard is None:
            clipboard = default_clipboard_reader(platform_name)
        clipboard_available = bool(clipboard.available)
    except Exception as exc:
        clipboard_available = False
        if error is None:
            error = f"{type(exc).__name__}: {exc}"

    result = {
        "ok": secure_storage and clipboard_available,
        "platform": platform_name,
        "secure_storage": secure_storage,
        "clipboard": clipboard_available,
    }
    if error is not None:
        result["error"] = error
    return result


def pair_smoke_check() -> dict:
    work = Path(tempfile.mkdtemp(prefix="local-anonymizer-pair-self-test-"))
    original_value = "Иванов Сергей Петрович"
    try:
        source = work / "Файл1.docx"
        source_document = Document()
        source_document.add_paragraph(original_value)
        source_document.save(source)
        service = AnonymizerService(work / "appdata")
        document_id = service.import_file(source)
        service.anonymize(document_id)
        pair = service.export_pair(
            document_id, work / "Файл1 — обезличенный.docx", confirmation=True
        )
        recovery_bytes = pair["recovery_path"].read_bytes()
        service.storage.delete_document(document_id)
        restored = work / "Файл1 — восстановленный.docx"
        report = service.restore_from_pair(pair["anonymized_path"], restored)
        restored_text = "\n".join(p.text for p in Document(restored).paragraphs)
        recovery_contains_mapping = original_value.encode("utf-8") in recovery_bytes
        ok = (
            restored_text == original_value
            and recovery_contains_mapping
            and not report["unknown_markers"]
        )
        return {
            "ok": ok,
            "pair_found_after_library_delete": True,
            "recovery_contains_sensitive_mapping": recovery_contains_mapping,
            "restored_exactly": restored_text == original_value,
        }
    except Exception as exc:
        return {"ok": False, "error": f"{type(exc).__name__}: {exc}"}
    finally:
        shutil.rmtree(work, ignore_errors=True)


def _screen_work_area(screen) -> tuple[int, int, int, int]:
    frame = getattr(screen, "frame", None)

    def value(obj, *names):
        for name in names:
            candidate = getattr(obj, name, None)
            if candidate is not None:
                return candidate() if callable(candidate) else candidate
        return None

    x = value(frame, "X", "x") if frame is not None else None
    y = value(frame, "Y", "y") if frame is not None else None
    width = value(frame, "Width", "width") if frame is not None else None
    height = value(frame, "Height", "height") if frame is not None else None
    if None in (x, y, width, height):
        x = getattr(screen, "x", 0)
        y = getattr(screen, "y", 0)
        width = getattr(screen, "width", 1600)
        height = getattr(screen, "height", 920)
    return int(x), int(y), int(width), int(height)


def _preferred_window_geometry(screens) -> dict:
    screens = list(screens or ())
    if not screens:
        return {"width": 1600, "height": 920}
    screen = next(
        (
            item
            for item in screens
            if item.x <= 0 < item.x + item.width
            and item.y <= 0 < item.y + item.height
        ),
        screens[0],
    )
    work_x, work_y, work_width, work_height = _screen_work_area(screen)
    width = max(720, min(1600, work_width))
    height = max(700, min(920, work_height))
    return {
        "width": width,
        "height": height,
        "x": work_x + max(0, (work_width - width) // 2),
        "y": work_y + max(0, (work_height - height) // 2),
        "screen": screen,
    }


def main(argv: list[str] | None = None) -> int:
    parser = argparse.ArgumentParser(description="Локальный обезличиватель")
    parser.add_argument("--smoke", action="store_true", help="проверка без запуска GUI")
    parser.add_argument("--pair-smoke", action="store_true", help="проверка пары файлов")
    parser.add_argument(
        "--platform-smoke",
        action="store_true",
        help="проверка системного keyring и нативного буфера обмена",
    )
    args = parser.parse_args(argv)
    if args.smoke:
        result = smoke_check()
        print(json.dumps(result, ensure_ascii=False))
        return 0 if result["ok"] else 1
    if args.pair_smoke:
        result = pair_smoke_check()
        print(json.dumps(result, ensure_ascii=False))
        return 0 if result["ok"] else 1
    if args.platform_smoke:
        result = platform_smoke_check()
        print(json.dumps(result, ensure_ascii=False))
        return 0 if result["ok"] else 1
    import webview
    api = AppApi()
    language = api.get_language()
    entry_asset = ENGLISH_INDEX_ASSET if language == "en" else INDEX_ASSET
    entry_url = asset_path(entry_asset).as_uri()
    title = "Local Document Anonymizer" if language == "en" else "Локальный обезличиватель документов"
    geometry = _preferred_window_geometry(getattr(webview, "screens", ()))
    window = webview.create_window(
        title,
        entry_url,
        js_api=api,
        min_size=(720, 700),
        **geometry,
    )
    api.set_window(window)
    webview.start(debug=False, private_mode=True)
    return 0


if __name__ == "__main__":
    raise SystemExit(main())
