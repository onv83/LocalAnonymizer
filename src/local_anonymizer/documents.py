from __future__ import annotations

import base64
import difflib
import functools
import hashlib
from io import BytesIO
import json
import os
import posixpath
import re
import struct
import tempfile
from dataclasses import dataclass, field
from pathlib import Path
from zipfile import ZIP_DEFLATED, ZIP_STORED, ZipFile, ZipInfo

import docx
from docx import Document
from docx.enum.text import WD_ALIGN_PARAGRAPH
from docx.oxml.ns import qn
from docx.shared import Pt
from lxml import etree

from .engine import Anonymizer, MARKER_PATTERN, warnings_require_review


class RestorationError(RuntimeError):
    pass


@dataclass
class DocumentResult:
    mapping: dict[str, str]
    candidates: list[str] = field(default_factory=list)
    warnings: list[str | dict[str, str]] = field(default_factory=list)
    ready: bool = True
    total: int = 0
    hint_count: int = 0
    package_archive: dict | None = None


def _walk_tables(tables):
    for table in tables:
        for row in table.rows:
            for cell in row.cells:
                yield from cell.paragraphs
                yield from _walk_tables(cell.tables)


def _paragraphs(document):
    seen: set[int] = set()
    groups = [document.paragraphs, _walk_tables(document.tables)]
    for section in document.sections:
        groups.extend((section.header.paragraphs, _walk_tables(section.header.tables),
                       section.footer.paragraphs, _walk_tables(section.footer.tables)))
    for group in groups:
        for paragraph in group:
            ident = id(paragraph._p)
            if ident not in seen:
                seen.add(ident)
                yield paragraph


def _paragraph_text(paragraph) -> str:
    """Visible OOXML text, including hyperlink and content-control runs."""
    return "".join(node.text or "" for node in paragraph._p.xpath(".//w:t"))


def _replace_paragraph(paragraph, new_text: str) -> None:
    nodes = list(paragraph._p.xpath(".//w:t"))
    old = "".join(node.text or "" for node in nodes)
    if old == new_text:
        return
    if not nodes:
        paragraph.add_run(new_text)
        return
    offsets = []
    pos = 0
    for node in nodes:
        value = node.text or ""
        offsets.append((pos, pos + len(value), node))
        pos += len(value)
    operations = [op for op in difflib.SequenceMatcher(None, old, new_text, autojunk=False).get_opcodes() if op[0] != "equal"]
    for _, a1, a2, b1, b2 in reversed(operations):
        replacement = new_text[b1:b2]
        if a1 == a2:  # insertion
            target = next((r for start, end, r in offsets if start <= a1 <= end), nodes[-1])
            start = next(start for start, end, r in offsets if r is target)
            local = a1 - start
            value = target.text or ""
            target.text = value[:local] + replacement + value[local:]
            continue
        touched = [(start, end, run) for start, end, run in offsets if start < a2 and end > a1]
        first_start, _, first = touched[0]
        last_start, _, last = touched[-1]
        first_value, last_value = first.text or "", last.text or ""
        prefix = first_value[:max(0, a1 - first_start)]
        suffix = last_value[max(0, a2 - last_start):]
        if first is last:
            first.text = prefix + replacement + suffix
        else:
            first.text = prefix + replacement
            for _, _, run in touched[1:-1]:
                run.text = ""
            last.text = suffix


def _replace_xml_text_nodes(nodes, new_text: str) -> None:
    """Replace logical paragraph text without rebuilding runs or run properties."""
    if not nodes:
        return
    old = "".join(node.text or "" for node in nodes)
    if old == new_text:
        return
    offsets: list[tuple[int, int, object]] = []
    position = 0
    for node in nodes:
        value = node.text or ""
        offsets.append((position, position + len(value), node))
        position += len(value)
    operations = [
        operation
        for operation in difflib.SequenceMatcher(None, old, new_text, autojunk=False).get_opcodes()
        if operation[0] != "equal"
    ]
    for _, a1, a2, b1, b2 in reversed(operations):
        replacement = new_text[b1:b2]
        if a1 == a2:
            target = next((node for start, end, node in offsets if start <= a1 <= end), nodes[-1])
            start = next(start for start, _, node in offsets if node is target)
            local = a1 - start
            value = target.text or ""
            target.text = value[:local] + replacement + value[local:]
            continue
        touched = [(start, end, node) for start, end, node in offsets if start < a2 and end > a1]
        first_start, _, first = touched[0]
        last_start, _, last = touched[-1]
        first_value, last_value = first.text or "", last.text or ""
        prefix = first_value[:max(0, a1 - first_start)]
        suffix = last_value[max(0, a2 - last_start):]
        if first is last:
            first.text = prefix + replacement + suffix
        else:
            first.text = prefix + replacement
            for _, _, node in touched[1:-1]:
                node.text = ""
            last.text = suffix
    xml_space = "{http://www.w3.org/XML/1998/namespace}space"
    for node in nodes:
        value = node.text or ""
        if value.startswith(" ") or value.endswith(" "):
            node.set(xml_space, "preserve")
        elif node.get(xml_space) == "preserve":
            node.attrib.pop(xml_space, None)


def _xml_text_node_groups(paragraph) -> list[list[object]]:
    """Split visible text nodes at hard breaks without editing separators."""
    text_tag = qn("w:t")
    boundary_tags = {qn("w:br"), qn("w:cr"), qn("w:tab")}
    groups: list[list[object]] = []
    current: list[object] = []
    for element in paragraph.iter():
        if element.tag == text_tag:
            current.append(element)
        elif element.tag in boundary_tags:
            if current:
                groups.append(current)
                current = []
    if current:
        groups.append(current)
    return groups


def _rewrite_docx_text_package(source: Path, target: Path, transform_all) -> None:
    """Copy OOXML and serialize only Word parts whose visible text changed."""
    parser = etree.XMLParser(resolve_entities=False, no_network=True, load_dtd=False)
    payloads: dict[str, bytes] = {}
    infos = []
    paragraphs: list[
        tuple[str, object, list[object], str, tuple[str, str | None, str | None, str | None, str]]
    ] = []
    with ZipFile(source) as package:
        for info in package.infolist():
            infos.append(info)
            payload = package.read(info.filename)
            payloads[info.filename] = payload
            if re.fullmatch(r"word/(?:document|header\d+|footer\d+)\.xml", info.filename):
                root = etree.fromstring(payload, parser)
                payloads[info.filename] = root
                for paragraph in root.xpath(".//w:p", namespaces={"w": qn("w:t").split("}", 1)[0][1:]}):
                    table = next((ancestor for ancestor in paragraph.iterancestors()
                                  if ancestor.tag == qn("w:tbl")), None)
                    table_path = root.getroottree().getpath(table) if table is not None else None
                    row = next((ancestor for ancestor in paragraph.iterancestors()
                                if ancestor.tag == qn("w:tr")), None)
                    cell = next((ancestor for ancestor in paragraph.iterancestors()
                                 if ancestor.tag == qn("w:tc")), None)
                    tree = root.getroottree()
                    paragraph_path = tree.getpath(paragraph)
                    row_path = tree.getpath(row) if row is not None else None
                    cell_path = tree.getpath(cell) if cell is not None else None
                    context = (info.filename, table_path, row_path, cell_path, paragraph_path)
                    groups = _xml_text_node_groups(paragraph) or [[]]
                    for nodes in groups:
                        paragraphs.append((
                            info.filename,
                            root,
                            nodes,
                            "".join(node.text or "" for node in nodes),
                            context,
                        ))
    table_boundary = "\ufdd0\ufdd1\ufdd0"
    row_boundary = "\ufdd0\ufdd2\ufdd0"

    def separator(left, right) -> str:
        left_context, right_context = left[4], right[4]
        if left_context[0:2] != right_context[0:2]:
            return table_boundary
        if left_context[1] is None:
            return "\ufdd3" if left_context[4] == right_context[4] else "\ufdd0"
        if left_context[2] != right_context[2]:
            return row_boundary
        if left_context[3] != right_context[3]:
            return "\ufdd0"
        return "\ufdd3"

    separators = [separator(left, right) for left, right in zip(paragraphs, paragraphs[1:])]
    replacements = transform_all([text for _, _, _, text, _ in paragraphs], separators)
    if len(replacements) != len(paragraphs):
        raise RuntimeError("Ошибка разбиения документа")
    changed_parts: set[str] = set()
    for (name, _, nodes, original, _), replacement in zip(paragraphs, replacements):
        if replacement != original:
            _replace_xml_text_nodes(nodes, replacement)
            changed_parts.add(name)
    with ZipFile(target, "w", ZIP_DEFLATED) as output:
        for info in infos:
            payload = payloads[info.filename]
            if info.filename in changed_parts:
                payload = etree.tostring(payload, xml_declaration=True, encoding="UTF-8", standalone=True)
            elif not isinstance(payload, bytes):
                with ZipFile(source) as package:
                    payload = package.read(info.filename)
            output.writestr(info, payload)


def _unknown_markers(text: str, mapping: dict[str, str]) -> list[str]:
    return sorted({m.group() for m in re.finditer(MARKER_PATTERN, text) if m.group() not in mapping})


def _restore_text(text: str, mapping: dict[str, str]) -> str:
    return re.sub(MARKER_PATTERN, lambda m: mapping.get(m.group(), m.group()), text)


class DocxProcessor:
    def __init__(self, engine: Anonymizer | None = None):
        self.engine = engine or Anonymizer()

    def anonymize(self, source: str | Path, output: str | Path,
                  *, redact_graphics: bool = True) -> DocumentResult:
        source, output = Path(source), Path(output)
        _preflight_docx_package(source)
        original_hash = source.read_bytes()
        document_boundary = "\ufdd0\ufdd1\ufdd0"
        result_holder = []

        def transform_all(texts: list[str], separators: list[str]) -> list[str]:
            if not texts:
                result = self.engine.anonymize("")
                result_holder.append(result)
                return []
            combined = texts[0] + "".join(
                separator + text for separator, text in zip(separators, texts[1:])
            )
            result = self.engine.anonymize(document_boundary + combined + document_boundary)
            result_holder.append(result)
            transformed = result.text.removeprefix(document_boundary).removesuffix(document_boundary)
            return re.split(r"\ufdd0(?:[\ufdd1\ufdd2]\ufdd0)?|\ufdd3", transformed)

        temp = _temp_path(output)
        try:
            _rewrite_docx_text_package(source, temp, transform_all)
            package_archive = _extract_reversible_docx_graphics(temp) if redact_graphics else None
            warnings = _sanitize_package(temp, retain_graphics=not redact_graphics)
            _atomic_publish(temp, output)
        finally:
            temp.unlink(missing_ok=True)
        if source.read_bytes() != original_hash:
            raise RuntimeError("Исходный файл был изменён")
        if package_archive is None:
            package_archive = {
                "schema": 2,
                "kind": "docx-graphics",
                "fragments": [],
                "parts": {},
                "part_metadata": {},
                "relationships": [],
                "content_types": [],
            }
        anonymized_bytes = output.read_bytes()
        package_archive["exact_source_b64"] = base64.b64encode(original_hash).decode("ascii")
        package_archive["exact_source_sha256"] = hashlib.sha256(original_hash).hexdigest()
        package_archive["exact_anonymized_sha256"] = hashlib.sha256(anonymized_bytes).hexdigest()
        result = result_holder[0]
        warnings = list(result.warnings) + warnings
        return DocumentResult(result.mapping, result.candidates, warnings,
                              not warnings_require_review(warnings),
                              result.total, result.hint_count, package_archive)

    def restore(self, source: str | Path, output: str | Path, mapping: dict[str, str],
                *, package_archive: dict | None = None) -> list[str]:
        source, output = Path(source), Path(output)
        _preflight_docx_package(source)
        if _restore_exact_docx_package(source, output, mapping, package_archive):
            return []
        unknown: list[str] = []

        def transform_all(texts: list[str], separators: list[str]) -> list[str]:
            unknown.extend(marker for text in texts for marker in _unknown_markers(text, mapping))
            if not texts:
                return []
            separator_tokens = {
                "\ufdd0": "\x00LA_CELL\x00",
                "\ufdd3": "\x00LA_INTRA\x00",
                "\ufdd0\ufdd1\ufdd0": "\x00LA_TABLE\x00",
                "\ufdd0\ufdd2\ufdd0": "\x00LA_ROW\x00",
            }
            wrapper = separator_tokens["\ufdd0\ufdd1\ufdd0"]
            combined = wrapper + texts[0] + "".join(
                separator_tokens[separator] + text
                for separator, text in zip(separators, texts[1:])
            ) + wrapper
            intra = separator_tokens["\ufdd3"]
            for marker, original in sorted(mapping.items(), key=lambda item: len(item[0]), reverse=True):
                segments = original.split("\n")
                if len(segments) > 1:
                    structural_marker = marker + intra * (len(segments) - 1)
                    combined = combined.replace(structural_marker, intra.join(segments))
                combined = combined.replace(marker, original)
            transformed = combined.removeprefix(wrapper).removesuffix(wrapper)
            split_pattern = "|".join(
                re.escape(token) for token in sorted(set(separator_tokens.values()), key=len, reverse=True)
            )
            return re.split(split_pattern, transformed)

        temp = _temp_path(output)
        try:
            _rewrite_docx_text_package(source, temp, transform_all)
            warnings = _sanitize_package(temp)
            if warnings:
                raise RestorationError("Восстановление запрещено: неподдерживаемые элементы DOCX")
            restored_graphic_markers = _restore_reversible_docx_graphics(
                temp, package_archive, mapping
            )
            unknown = [marker for marker in unknown if marker not in restored_graphic_markers]
            _preflight_docx_package(temp)
            _atomic_publish(temp, output)
        finally:
            temp.unlink(missing_ok=True)
        return sorted(set(unknown))


def _restore_exact_docx_package(
    source: Path,
    output: Path,
    mapping: dict[str, str],
    archive: dict | None,
) -> bool:
    """Restore the byte-exact source when the anonymized package was not edited."""
    if not isinstance(archive, dict):
        return False
    exact_keys = {
        "exact_source_b64",
        "exact_source_sha256",
        "exact_anonymized_sha256",
    }
    present = exact_keys.intersection(archive)
    if not present:
        return False
    if present != exact_keys:
        raise RestorationError("Повреждён точный снимок исходного DOCX")
    encoded = archive["exact_source_b64"]
    source_digest = archive["exact_source_sha256"]
    anonymized_digest = archive["exact_anonymized_sha256"]
    if not all(isinstance(value, str) for value in (encoded, source_digest, anonymized_digest)):
        raise RestorationError("Повреждён точный снимок исходного DOCX")
    if not re.fullmatch(r"[0-9a-f]{64}", source_digest) or not re.fullmatch(
        r"[0-9a-f]{64}", anonymized_digest
    ):
        raise RestorationError("Повреждён точный снимок исходного DOCX")
    if hashlib.sha256(source.read_bytes()).hexdigest() != anonymized_digest:
        return False
    try:
        original = base64.b64decode(encoded, validate=True)
    except (ValueError, TypeError) as exc:
        raise RestorationError("Повреждён точный снимок исходного DOCX") from exc
    if hashlib.sha256(original).hexdigest() != source_digest:
        raise RestorationError("Повреждён точный снимок исходного DOCX")
    temp = _temp_path(output)
    try:
        temp.write_bytes(source.read_bytes())
        _restore_reversible_docx_graphics(temp, archive, mapping)
        _preflight_docx_package(temp)
        temp.write_bytes(original)
        _preflight_docx_package(temp)
        _atomic_publish(temp, output)
    except RestorationError:
        raise
    except Exception as exc:
        raise RestorationError("Точный снимок исходного DOCX не прошёл проверку") from exc
    finally:
        temp.unlink(missing_ok=True)
    return True


def _temp_path(output: Path) -> Path:
    output.parent.mkdir(parents=True, exist_ok=True)
    fd, name = tempfile.mkstemp(prefix=f".{output.stem}-", suffix=output.suffix, dir=output.parent)
    os.close(fd)
    return Path(name)


def _atomic_publish(temp: Path, output: Path) -> None:
    with temp.open("r+b") as stream:
        os.fsync(stream.fileno())
    os.replace(temp, output)


def _known_docx_xml_part(name: str) -> bool:
    if name in {"[Content_Types].xml", "_rels/.rels"}:
        return True
    if re.fullmatch(r"docProps/(?:app|core|custom)\.xml", name):
        return True
    if re.fullmatch(r"word/(?:document|styles|stylesWithEffects|settings|webSettings|fontTable|numbering)\.xml", name):
        return True
    if re.fullmatch(r"word/(?:header|footer)\d+\.xml", name):
        return True
    relationship_part = re.fullmatch(r"(.+)/_rels/([^/]+)\.rels", name)
    if relationship_part:
        source_part = f"{relationship_part.group(1)}/{relationship_part.group(2)}"
        return _known_docx_xml_part(source_part)
    if re.fullmatch(r"word/theme/theme\d+\.xml", name):
        return True
    return False


def _safe_zip_member_name(name: str) -> bool:
    return bool(name) and not (
        "\\" in name
        or name.startswith("/")
        or re.match(r"^[A-Za-z]:", name)
        or "\x00" in name
        or any(part in {"", ".", ".."} for part in name.split("/"))
    )


def _raw_zip_member_names(path: Path) -> list[str]:
    """Read central/local ZIP names without zipfile's filename normalization."""
    raw = path.read_bytes()
    eocd_at = raw.rfind(b"PK\x05\x06", max(0, len(raw) - 65_557))
    if eocd_at < 0 or eocd_at + 22 > len(raw):
        raise ValueError("DOCX не содержит корректный ZIP-каталог")
    _, disk, cd_disk, disk_entries, entries, cd_size, cd_at, comment_size = struct.unpack_from(
        "<4s4H2LH", raw, eocd_at
    )
    if disk or cd_disk or disk_entries != entries or eocd_at + 22 + comment_size != len(raw):
        raise ValueError("Многотомный или неоднозначный ZIP не поддерживается")
    if cd_at + cd_size > eocd_at:
        raise ValueError("Некорректный ZIP-каталог DOCX")
    names: list[str] = []
    cursor = cd_at
    for _ in range(entries):
        if cursor + 46 > len(raw) or raw[cursor : cursor + 4] != b"PK\x01\x02":
            raise ValueError("Некорректная запись ZIP-каталога DOCX")
        fields = struct.unpack_from("<4s6H3L5H2L", raw, cursor)
        flag, name_size, extra_size, member_comment_size, local_at = (
            fields[3], fields[10], fields[11], fields[12], fields[16]
        )
        name_bytes = raw[cursor + 46 : cursor + 46 + name_size]
        encoding = "utf-8" if flag & 0x800 else "cp437"
        try:
            name = name_bytes.decode(encoding)
        except UnicodeDecodeError as exc:
            raise ValueError("Некорректная кодировка имени ZIP-части DOCX") from exc
        if local_at + 30 > len(raw) or raw[local_at : local_at + 4] != b"PK\x03\x04":
            raise ValueError("Некорректный локальный ZIP-заголовок DOCX")
        local = struct.unpack_from("<4s5H3L2H", raw, local_at)
        local_name_size, local_extra_size = local[9], local[10]
        local_name = raw[local_at + 30 : local_at + 30 + local_name_size]
        if local_name != name_bytes or local_at + 30 + local_name_size + local_extra_size > len(raw):
            raise ValueError("Имена ZIP-части в каталогах не совпадают")
        if not _safe_zip_member_name(name):
            raise ValueError(f"DOCX содержит небезопасное имя ZIP-части: {name!r}")
        names.append(name)
        cursor += 46 + name_size + extra_size + member_comment_size
    if cursor != cd_at + cd_size:
        raise ValueError("Размер ZIP-каталога DOCX не совпадает")
    return names


def _preflight_docx_package(path: Path) -> None:
    """Reject ambiguous or resource-hostile ZIP packages before parsing OOXML."""
    max_members = 2048
    max_member_size = 64 * 1024 * 1024
    max_total_size = 256 * 1024 * 1024
    max_ratio = 1000
    raw_names = _raw_zip_member_names(path)
    with ZipFile(path, "r") as package:
        infos = package.infolist()
        if [info.filename for info in infos] != raw_names:
            raise ValueError("Имена ZIP-частей были неканонически преобразованы")
        if len(infos) > max_members:
            raise ValueError("DOCX содержит слишком много ZIP-частей")
        names: set[str] = set()
        total_size = 0
        for info in infos:
            name = info.filename
            if not _safe_zip_member_name(name):
                raise ValueError(f"DOCX содержит небезопасное имя ZIP-части: {name!r}")
            folded = name.casefold()
            if folded in names:
                raise ValueError(f"DOCX содержит повторяющиеся ZIP-части: {name}")
            names.add(folded)
            if info.flag_bits & 0x1:
                raise ValueError(f"DOCX содержит зашифрованную ZIP-часть: {name}")
            if info.compress_type not in {ZIP_STORED, ZIP_DEFLATED}:
                raise ValueError(f"Неподдерживаемый метод сжатия ZIP-части DOCX: {name}")
            if info.file_size > max_member_size:
                raise ValueError(f"ZIP-часть DOCX слишком велика: {name}")
            total_size += info.file_size
            if total_size > max_total_size:
                raise ValueError("Распакованный DOCX слишком велик")
            if info.file_size and info.compress_size == 0:
                raise ValueError(f"Некорректный размер ZIP-части DOCX: {name}")
            if info.compress_size and info.file_size / info.compress_size > max_ratio:
                raise ValueError(f"Подозрительная степень сжатия ZIP-части DOCX: {name}")

        for info in infos:
            if not info.filename.endswith((".xml", ".rels")):
                continue
            try:
                tree = etree.parse(
                    BytesIO(package.read(info.filename)),
                    parser=etree.XMLParser(
                        resolve_entities=False,
                        no_network=True,
                        load_dtd=False,
                    ),
                )
            except etree.XMLSyntaxError:
                continue
            if tree.docinfo.doctype or tree.docinfo.internalDTD is not None:
                raise ValueError(f"DOCX XML содержит запрещённый DTD: {info.filename}")

        parser = etree.XMLParser(resolve_entities=False, no_network=True, load_dtd=False)
        if "[Content_Types].xml" in package.namelist():
            try:
                root = etree.fromstring(package.read("[Content_Types].xml"), parser=parser)
            except etree.XMLSyntaxError:
                root = None
            if root is not None:
                seen_content_keys: set[tuple[str, str]] = set()
                for child in root:
                    local = etree.QName(child).localname
                    if local == "Override":
                        key = (local, child.get("PartName", "").casefold())
                    elif local == "Default":
                        key = (local, child.get("Extension", "").casefold())
                    else:
                        continue
                    if key in seen_content_keys:
                        raise ValueError("DOCX содержит повторяющиеся объявления Content Types")
                    seen_content_keys.add(key)
        for info in infos:
            if not info.filename.endswith(".rels"):
                continue
            try:
                root = etree.fromstring(package.read(info.filename), parser=parser)
            except etree.XMLSyntaxError:
                continue
            relationship_ids: set[str] = set()
            for child in root:
                if etree.QName(child).localname != "Relationship":
                    continue
                relationship_id = child.get("Id", "")
                if relationship_id in relationship_ids:
                    raise ValueError(f"DOCX содержит повторяющиеся Id связей: {info.filename}")
                relationship_ids.add(relationship_id)


_WML_NS = "http://schemas.openxmlformats.org/wordprocessingml/2006/main"
_OFFICE_REL_NS = "http://schemas.openxmlformats.org/officeDocument/2006/relationships"
_PACKAGE_REL_NS = "http://schemas.openxmlformats.org/package/2006/relationships"
_DRAWING_NS = "http://schemas.openxmlformats.org/drawingml/2006/main"
_XML_NS = "http://www.w3.org/XML/1998/namespace"
_CONTENT_TYPES_NS = "http://schemas.openxmlformats.org/package/2006/content-types"
_ALLOWED_CONTENT_TYPES = frozenset({
    "application/vnd.ms-word.stylesWithEffects+xml",
    "application/vnd.openxmlformats-officedocument.custom-properties+xml",
    "application/vnd.openxmlformats-officedocument.customXmlProperties+xml",
    "application/vnd.openxmlformats-officedocument.extended-properties+xml",
    "application/vnd.openxmlformats-officedocument.theme+xml",
    "application/vnd.openxmlformats-officedocument.wordprocessingml.comments+xml",
    "application/vnd.openxmlformats-officedocument.wordprocessingml.document.main+xml",
    "application/vnd.openxmlformats-officedocument.wordprocessingml.endnotes+xml",
    "application/vnd.openxmlformats-officedocument.wordprocessingml.fontTable+xml",
    "application/vnd.openxmlformats-officedocument.wordprocessingml.footer+xml",
    "application/vnd.openxmlformats-officedocument.wordprocessingml.footnotes+xml",
    "application/vnd.openxmlformats-officedocument.wordprocessingml.header+xml",
    "application/vnd.openxmlformats-officedocument.wordprocessingml.numbering+xml",
    "application/vnd.openxmlformats-officedocument.wordprocessingml.settings+xml",
    "application/vnd.openxmlformats-officedocument.wordprocessingml.styles+xml",
    "application/vnd.openxmlformats-officedocument.wordprocessingml.webSettings+xml",
    "application/vnd.openxmlformats-package.core-properties+xml",
    "application/vnd.openxmlformats-package.relationships+xml",
    "application/xml",
    "image/jpeg",
    "image/png",
})
_ALLOWED_RELATIONSHIP_TYPES = frozenset({
    "http://schemas.microsoft.com/office/2007/relationships/stylesWithEffects",
    "http://schemas.openxmlformats.org/officeDocument/2006/relationships/attachedTemplate",
    "http://schemas.openxmlformats.org/officeDocument/2006/relationships/comments",
    "http://schemas.openxmlformats.org/officeDocument/2006/relationships/custom-properties",
    "http://schemas.openxmlformats.org/officeDocument/2006/relationships/customXml",
    "http://schemas.openxmlformats.org/officeDocument/2006/relationships/customXmlProps",
    "http://schemas.openxmlformats.org/officeDocument/2006/relationships/endnotes",
    "http://schemas.openxmlformats.org/officeDocument/2006/relationships/extended-properties",
    "http://schemas.openxmlformats.org/officeDocument/2006/relationships/fontTable",
    "http://schemas.openxmlformats.org/officeDocument/2006/relationships/footer",
    "http://schemas.openxmlformats.org/officeDocument/2006/relationships/footnotes",
    "http://schemas.openxmlformats.org/officeDocument/2006/relationships/glossaryDocument",
    "http://schemas.openxmlformats.org/officeDocument/2006/relationships/header",
    "http://schemas.openxmlformats.org/officeDocument/2006/relationships/hyperlink",
    "http://schemas.openxmlformats.org/officeDocument/2006/relationships/image",
    "http://schemas.openxmlformats.org/officeDocument/2006/relationships/numbering",
    "http://schemas.openxmlformats.org/officeDocument/2006/relationships/officeDocument",
    "http://schemas.openxmlformats.org/officeDocument/2006/relationships/settings",
    "http://schemas.openxmlformats.org/officeDocument/2006/relationships/styles",
    "http://schemas.openxmlformats.org/officeDocument/2006/relationships/theme",
    "http://schemas.openxmlformats.org/officeDocument/2006/relationships/webSettings",
    "http://schemas.openxmlformats.org/package/2006/relationships/metadata/core-properties",
    "http://schemas.openxmlformats.org/package/2006/relationships/metadata/thumbnail",
})
_SUPPORTED_WML_ELEMENTS = frozenset("""
abstractNum abstractNumId allowPNG altName annotationRef autoHyphenation autoRedefine autoSpaceDE
b bCs basedOn bidi body bodyDiv bookmarkEnd bookmarkStart bottom br cr caps characterSpacingControl
charset clrSchemeMapping color cols comment commentRangeEnd commentRangeStart commentReference comments
compat compatSetting contextualSpacing continuationSeparator decimalSymbol defaultTabStop del delText
displayBackgroundShape div divBdr divs divsChild doNotAutoCompressPictures doNotBreakWrappedTables
doNotSaveAsSingleFile docDefaults docGrid docPartGallery docPartObj docPartUnique docVar docVars document
drawing dstrike eastAsianLayout effect em embedSystemFonts emboss end endnote endnotePr endnotes family
fldChar fldSimple font fonts footerReference footnote footnotePr footnotes formProt ftr gridCol gridSpan
hdr hdrShapeDefaults headerReference hidden highlight hyperlink hyphenationZone i iCs ilvl imprint ind ins
insideH insideV instrText jc keepLines keepNext kern kinsoku lang lastRenderedPageBreak latentStyles left
link listSeparator lsdException lvl lvlJc lvlOverride lvlText marBottom marLeft marRight marTop multiLevelType
name next noProof notTrueType nsid num numFmt numId numIdMacAtCleanup numPr numbering object optimizeForBrowser
outline outlineLvl overflowPunct p pBdr pPr pPrDefault pStyle pageBreakBefore panose1 pgMar pgNumType pgSz pict
pitch position proofErr proofState qFormat r rFonts rPr rPrChange rPrDefault rStyle right rsid rsidRoot rsids
savePreviewPicture sdt sdtContent sdtPr sectPr semiHidden separator settings shadow shapeDefaults shd sig
smallCaps snapToGrid spacing specVanish start startOverride strike style styles suff suppressAutoHyphens
suppressLineNumbers sz szCs t tab tabs tbl tblBorders tblCellMar tblGrid tblGridChange tblHeader tblInd
tblLayout tblLook tblPr tblPrChange tblPrEx tblStyle tblStyleColBandSize tblStylePr tblStyleRowBandSize tblW
tblpPr tc tcBorders tcMar tcPr tcPrChange tcW textDirection themeFontLang titlePg tl2br tmpl top tr tr2bl
trHeight trPr trPrChange trackRevisions txbxContent type u uiPriority unhideWhenUsed useFELayout vAlign
vMerge vanish vertAlign w webSettings widowControl zoom altChunk
""".split())
_SUPPORTED_THEME_ELEMENTS = frozenset("""
accent1 accent2 accent3 accent4 accent5 accent6 alpha bevelT bgFillStyleLst bodyPr camera clrScheme cs dk1
dk2 ea effectLst effectRef effectStyle effectStyleLst extraClrSchemeLst fillRef fillStyleLst fillToRect
fmtScheme folHlink font fontRef fontScheme gradFill gs gsLst hlink latin lightRig lin ln lnDef lnRef
lnStyleLst lstStyle lt1 lt2 majorFont minorFont objectDefaults outerShdw path prstDash rot satMod scene3d
schemeClr shade solidFill sp3d spDef spPr srgbClr style sysClr theme themeElements tint
""".split())
_SUPPORTED_WML_ATTRIBUTES = frozenset("""
after abstractNumId accent1 accent2 accent3 accent4 accent5 accent6 ascii asciiTheme author before bg1 bg2 bidi bidiTheme bottom
char charFormat clear color colFirst colLast count cs csb0 csb1 cstheme customStyle date default defLockedState
defQFormat defSemiHidden defUIPriority defUnhideWhenUsed dirty eastAsia eastAsiaTheme end firstLine firstLineChars
fill firstColumn firstRow fldCharType followedHyperlink footer format grammar gutter h hAnsi hAnsiTheme hanging hangingChars
header hint hRule hyperlink id ilvl initial insideH insideV instr lastColumn lastRow left leftChars line linePitch
lineRule link listSeparator name noHBand noVBand num numFmtId numId orient percent pos qFormat relativeTo right rightChars
rsid rsidDel rsidP rsidR rsidRDefault rsidRPr rsidSect rsidTr semiHidden
space spelling start styleId sz t1 t2 themeColor themeFill themeFillShade themeFillTint themeShade themeTint top type
uiPriority unhideWhenUsed uri usb0 usb1 usb2 usb3 val w wAfter wBefore wCode wType
""".split())
_SUPPORTED_THEME_ATTRIBUTES = frozenset("""
algn ang b bandCol bandRow blurRad builtIn cap cmpd dir dist flip h idx insetPen l lastClr lat lim lon marB marL
marR marT name path pos prst r rad rev rig rotWithShape scaled script strike t typeface val w
""".split())
_SUPPORTED_WML_TEXT_ELEMENTS = frozenset({"t", "instrText", "delText"})
_WML_REQUIRED_PARENTS = {
    "name": frozenset({"style"}),
    "zoom": frozenset({"settings"}),
    "rFonts": frozenset({"rPr"}),
    "lang": frozenset({"rPr"}),
    "color": frozenset({"rPr"}),
}
_WML_CONTEXTUAL_ATTRIBUTES = {
    "author": frozenset({"del", "ins", "pPrChange", "rPrChange", "tblGridChange", "tblPrChange", "tblPrExChange", "tcPrChange", "trPrChange"}),
    "date": frozenset({"del", "ins", "pPrChange", "rPrChange", "tblGridChange", "tblPrChange", "tblPrExChange", "tcPrChange", "trPrChange"}),
    "styleId": frozenset({"style"}),
    "name": frozenset({"compatSetting", "font", "lsdException"}),
    "ascii": frozenset({"rFonts"}),
    "hAnsi": frozenset({"rFonts"}),
    "eastAsia": frozenset({"lang", "rFonts", "themeFontLang"}),
    "cs": frozenset({"lang", "rFonts", "themeFontLang"}),
    "asciiTheme": frozenset({"rFonts"}),
    "hAnsiTheme": frozenset({"rFonts"}),
    "eastAsiaTheme": frozenset({"rFonts"}),
    "cstheme": frozenset({"rFonts"}),
    "instr": frozenset({"fldSimple"}),
    "fldCharType": frozenset({"fldChar"}),
}

_STYLE_REFERENCE_ELEMENTS = frozenset({
    "basedOn", "link", "next", "numStyleLink", "pStyle", "rStyle", "styleLink", "tblStyle",
})


_TRUSTED_STRING_ATTRIBUTES = frozenset({
    "ascii", "asciiTheme", "cs", "cstheme", "eastAsia", "eastAsiaTheme",
    "hAnsi", "hAnsiTheme", "name", "styleId", "typeface",
})

_SAFE_FONT_NAMES = frozenset({
    "Arial", "Arial Narrow", "Calibri", "Cambria", "Courier New", "Georgia",
    "Liberation Mono", "Liberation Sans", "Liberation Serif", "Segoe UI",
    "Tahoma", "Times New Roman", "Trebuchet MS", "Verdana",
})


def _schema_part(name: str) -> str | None:
    if re.fullmatch(r"word/header\d+\.xml", name):
        return "word/header.xml"
    if re.fullmatch(r"word/footer\d+\.xml", name):
        return "word/footer.xml"
    if re.fullmatch(r"word/theme/theme\d+\.xml", name):
        return "word/theme/theme.xml"
    if name in {
        "word/document.xml", "word/styles.xml", "word/stylesWithEffects.xml",
        "word/settings.xml", "word/numbering.xml", "word/fontTable.xml",
        "word/webSettings.xml",
    }:
        return name
    return None


@functools.lru_cache(maxsize=1)
def _trusted_generated_docx() -> bytes:
    sample = Document()
    paragraph = sample.add_paragraph("Sample", style="List Number")
    paragraph.alignment = WD_ALIGN_PARAGRAPH.CENTER
    run = paragraph.add_run(" formatted")
    run.bold = True
    run.italic = True
    run.underline = True
    table = sample.add_table(rows=3, cols=2)
    table.style = "Table Grid"
    table.cell(0, 0).text = "cell"
    table.cell(0, 0).merge(table.cell(0, 1))
    table.cell(1, 0).text = "vertical"
    table.cell(1, 0).merge(table.cell(2, 0))
    table.rows[1].height = Pt(24)
    sample.sections[0].header.paragraphs[0].text = "header"
    sample.sections[0].footer.paragraphs[0].text = "footer"
    buffer = BytesIO()
    sample.save(buffer)
    return buffer.getvalue()


@functools.lru_cache(maxsize=1)
def _trusted_docx_strings() -> tuple[frozenset[str], frozenset[tuple[str, str, str, str]]]:
    """Read accepted strings only in their exact supported OOXML part context."""
    style_ids: set[str] = set()
    quadruples: set[tuple[str, str, str, str]] = set()
    template = Path(docx.__file__).resolve().parent / "templates" / "default.docx"
    package_images = [template.read_bytes(), _trusted_generated_docx()]
    word_reference = Path(__file__).resolve().parent / "assets" / "word_schema_reference.docx"
    if word_reference.is_file():
        package_images.append(word_reference.read_bytes())
    for package_bytes in package_images:
        with ZipFile(BytesIO(package_bytes)) as package:
            for part_name in package.namelist():
                part_key = _schema_part(part_name)
                if part_key is None:
                    continue
                root = etree.fromstring(package.read(part_name))
                for item in root.iter():
                    item_qname = etree.QName(item)
                    for key, value in item.attrib.items():
                        attr_qname = etree.QName(key)
                        quadruples.add((part_key, item_qname.localname, attr_qname.localname, value))
                    if item_qname.namespace == _WML_NS and item_qname.localname == "style":
                        style_id = item.get(qn("w:styleId"))
                        if style_id:
                            style_ids.add(style_id)
    return frozenset(style_ids), frozenset(quadruples)


def _safe_known_attribute_value(
    name: str, element_local: str, attr_local: str, value: str
) -> bool:
    style_ids, quadruples = _trusted_docx_strings()
    part_key = _schema_part(name)
    if part_key is not None and (part_key, element_local, attr_local, value) in quadruples:
        return True
    if element_local in _STYLE_REFERENCE_ELEMENTS and attr_local == "val" and value in style_ids:
        return True
    if element_local == "fldSimple" and attr_local == "instr":
        return True
    if element_local in {"headerReference", "footerReference"} and attr_local == "type":
        return value in {"default", "even", "first"}
    if element_local == "fldChar" and attr_local == "fldCharType":
        return value in {"begin", "separate", "end"}
    if element_local == "br" and attr_local == "type":
        return value in {"column", "page", "textWrapping"}

    unsigned = re.fullmatch(r"\d{1,10}", value) is not None
    signed = re.fullmatch(r"-?\d{1,10}", value) is not None
    on_off = value in {"0", "1", "false", "true", "off", "on"}
    safe_token = re.fullmatch(r"[A-Za-z0-9_.-]{1,128}", value) is not None
    safe_label = (
        0 < len(value) <= 128
        and not any(ord(char) < 32 for char in value)
        and re.fullmatch(r"[\w .,+()&/-]+", value, re.UNICODE) is not None
    )

    # Document-varying values are accepted only in explicit OOXML contexts and
    # with bounded semantic grammars.  This avoids treating an arbitrary source
    # document as a schema while preserving values required by Word.
    if part_key == "word/document.xml" and attr_local == "val":
        if element_local in {"b", "bCs", "i", "iCs", "formProt", "overflowPunct", "widowControl"}:
            return on_off
        if element_local == "kern":
            return unsigned and int(value) <= 3_276_750
        if element_local == "gridSpan":
            return unsigned and 1 <= int(value) <= 63
        if element_local == "numId":
            return unsigned and int(value) <= 2_147_483_647
        if element_local in {"ilvl", "outlineLvl"}:
            return unsigned and int(value) <= 9
        if element_local in {"pStyle", "rStyle"}:
            return re.fullmatch(r"LocalStyle\d{3}", value) is not None
        if element_local == "docPartGallery":
            return value in {"Table of Contents", "Cover Pages", "Headers", "Footers", "Page Numbers"}
        if element_local == "jc":
            return value in {"left", "right", "center", "both", "distribute", "start", "end"}
        if element_local == "tab":
            return value in {"left", "right", "center", "decimal", "bar", "clear", "num"}
        if element_local == "u":
            return value in {"none", "single", "words", "double", "thick", "dotted", "dash", "wave"}
        if element_local == "vAlign":
            return value in {"baseline", "superscript", "subscript", "top", "center", "bottom", "both"}
    if part_key == "word/document.xml" and element_local == "tblLook":
        if attr_local in {"firstColumn", "firstRow", "lastColumn", "lastRow", "noHBand", "noVBand"}:
            return on_off
    if part_key == "word/document.xml" and (element_local, attr_local) == ("pgSz", "orient"):
        return value in {"portrait", "landscape"}

    if part_key == "word/numbering.xml":
        if (element_local, attr_local) in {
            ("abstractNum", "abstractNumId"), ("abstractNumId", "val"), ("numId", "val")
        }:
            return unsigned and int(value) <= 2_147_483_647
        if (element_local, attr_local) == ("ilvl", "val"):
            return unsigned and int(value) <= 9
        if (element_local, attr_local) == ("multiLevelType", "val"):
            return value in {"singleLevel", "multilevel", "hybridMultilevel"}
        if (element_local, attr_local) == ("numFmt", "val"):
            return value in {"decimal", "decimalZero", "bullet", "none", "upperRoman", "lowerRoman", "upperLetter", "lowerLetter", "ordinal", "cardinalText"}
        if (element_local, attr_local) == ("suff", "val"):
            return value in {"tab", "space", "nothing"}

    if part_key == "word/styles.xml" and attr_local == "val":
        if element_local in {"basedOn", "next", "link"}:
            return re.fullmatch(r"LocalStyle\d{3}", value) is not None
        if element_local == "name":
            return re.fullmatch(r"Local Style \d{3}", value) is not None
        if element_local in {"qFormat", "locked", "semiHidden", "unhideWhenUsed"}:
            return on_off
        if element_local in {"tblStyleColBandSize", "tblStyleRowBandSize"}:
            return unsigned and 1 <= int(value) <= 63
        if element_local in {"ilvl", "numId"}:
            return unsigned and int(value) <= 2_147_483_647
        if element_local == "kern":
            return unsigned and int(value) <= 3_276_750
        if element_local in {"overflowPunct", "widowControl"}:
            return on_off
    if part_key == "word/styles.xml" and element_local == "eastAsianLayout":
        if attr_local == "id":
            return unsigned and int(value) <= 2_147_483_647
        if attr_local in {"combine", "vert"}:
            return on_off

    if part_key == "word/settings.xml" and element_local == "themeFontLang":
        if attr_local in {"val", "eastAsia", "bidi"}:
            return re.fullmatch(r"[A-Za-z]{2,3}(?:-[A-Za-z0-9]{2,8})*", value) is not None

    if part_key == "word/fontTable.xml":
        if (element_local, attr_local) == ("altName", "val"):
            return value in _SAFE_FONT_NAMES or re.fullmatch(r"LocalFont\d{3}", value) is not None
        if (element_local, attr_local) == ("charset", "val"):
            return re.fullmatch(r"[0-9A-Fa-f]{2}", value) is not None
        if (element_local, attr_local) == ("family", "val"):
            return value in {"auto", "decorative", "modern", "roman", "script", "swiss"}
        if (element_local, attr_local) == ("pitch", "val"):
            return value in {"default", "fixed", "variable"}
        if (element_local, attr_local) == ("notTrueType", "val"):
            return on_off
        if (element_local, attr_local) == ("panose1", "val"):
            return re.fullmatch(r"[0-9A-Fa-f]{20}", value) is not None
        if element_local == "sig" and attr_local in {"csb0", "csb1", "usb0", "usb1", "usb2", "usb3"}:
            return re.fullmatch(r"[0-9A-Fa-f]{8}", value) is not None
        if element_local == "embedRegular" and attr_local == "fontKey":
            return re.fullmatch(r"\{[0-9A-Fa-f-]{36}\}", value) is not None

    if part_key == "word/webSettings.xml":
        if (element_local, attr_local) == ("div", "id"):
            return unsigned and 1 <= int(value) <= 2_147_483_647
        if (element_local, attr_local) == ("bodyDiv", "val"):
            return on_off
        if element_local in {"marLeft", "marRight", "marTop", "marBottom"} and attr_local == "val":
            return signed and abs(int(value)) <= 31_680

    if part_key == "word/theme/theme.xml":
        if (element_local, attr_local) == ("theme", "name"):
            return value == "Office Theme"
        if attr_local == "typeface":
            return value == "" or value in _SAFE_FONT_NAMES or re.fullmatch(r"LocalFont\d{3}", value) is not None
        if attr_local in {"lang", "script"}:
            return re.fullmatch(r"[A-Za-z0-9-]{1,16}", value) is not None
        if element_local in {"srgbClr", "sysClr"} and attr_local in {"val", "lastClr"}:
            return re.fullmatch(r"[0-9A-Fa-f]{6}", value) is not None or value in {"window", "windowText"}
        if attr_local in {"w", "h", "rot", "rad", "pos", "ang", "blurRad", "dist"}:
            return signed and abs(int(value)) <= 100_000_000

    _, attribute_contexts = _trusted_wml_schema()
    if (
        part_key is None
        or element_local not in attribute_contexts.get((part_key, attr_local), frozenset())
    ):
        return False
    if (element_local, attr_local) in {("gridCol", "w"), ("trHeight", "val")} and unsigned:
        return int(value) <= 31_680
    bounded_numeric = {
        "after", "before", "bottom", "charSpace", "count", "firstLine",
        "footer", "gutter", "h", "hanging", "header", "id", "ilvl",
        "left", "line", "linePitch", "numId", "pos", "right", "space",
        "start", "sz", "top", "uiPriority", "w",
    }
    if attr_local in bounded_numeric and signed:
        return -31_680 <= int(value) <= 10_000_000
    if attr_local in {"val", "fill", "color", "themeColor", "themeFill", "themeShade", "themeTint"}:
        if re.fullmatch(r"(?:auto|[0-9A-Fa-f]{2,20})", value):
            return True
        if re.fullmatch(r"[A-Za-z][A-Za-z0-9._-]{0,63}", value):
            return True
        if element_local == "lvlText" and len(value) <= 64 and not any(ord(char) < 32 for char in value):
            return True
    if attr_local == "styleId":
        return re.fullmatch(r"LocalStyle\d{3}", value) is not None
    if attr_local in {"ascii", "cs", "eastAsia", "hAnsi", "name", "typeface"}:
        return value in _SAFE_FONT_NAMES or re.fullmatch(r"LocalFont\d{3}", value) is not None
    if attr_local in {"asciiTheme", "cstheme", "eastAsiaTheme", "hAnsiTheme"}:
        return value in {"majorAscii", "majorBidi", "majorEastAsia", "majorHAnsi", "minorAscii", "minorBidi", "minorEastAsia", "minorHAnsi"}
    if attr_local in {"type", "lineRule", "hRule", "orient", "leader", "hint"}:
        return re.fullmatch(r"[A-Za-z][A-Za-z0-9._-]{0,31}", value) is not None
    return False


@functools.lru_cache(maxsize=1)
def _trusted_wml_schema() -> tuple[
    frozenset[tuple[str, str, str]], dict[tuple[str, str], frozenset[str]]
]:
    """Build conservative WML contexts from immutable, locally generated templates."""
    pairs: set[tuple[str, str, str]] = set()
    attribute_contexts: dict[tuple[str, str], set[str]] = {}
    package_images: list[bytes] = []
    template = Path(docx.__file__).resolve().parent / "templates" / "default.docx"
    package_images.append(template.read_bytes())
    word_reference = Path(__file__).resolve().parent / "assets" / "word_schema_reference.docx"
    if word_reference.is_file():
        package_images.append(word_reference.read_bytes())

    sample = Document()
    paragraph = sample.add_paragraph("Sample", style="List Number")
    paragraph.alignment = WD_ALIGN_PARAGRAPH.CENTER
    run = paragraph.add_run(" formatted")
    run.bold = True
    run.italic = True
    run.underline = True
    table = sample.add_table(rows=3, cols=2)
    table.style = "Table Grid"
    table.cell(0, 0).text = "cell"
    table.cell(0, 0).merge(table.cell(0, 1))
    table.cell(1, 0).text = "vertical"
    table.cell(1, 0).merge(table.cell(2, 0))
    table.rows[1].height = Pt(24)
    sample.sections[0].header.paragraphs[0].text = "header"
    sample.sections[0].footer.paragraphs[0].text = "footer"
    buffer = BytesIO()
    sample.save(buffer)
    package_images.append(buffer.getvalue())

    for package_bytes in package_images:
        with ZipFile(BytesIO(package_bytes)) as package:
            for part_name in package.namelist():
                part_key = _schema_part(part_name)
                if part_key is None:
                    continue
                root = etree.fromstring(package.read(part_name))
                for item in root.iter():
                    item_qname = etree.QName(item)
                    if item_qname.namespace != _WML_NS:
                        continue
                    parent = item.getparent()
                    if parent is not None:
                        parent_qname = etree.QName(parent)
                        if parent_qname.namespace == _WML_NS:
                            pairs.add((part_key, parent_qname.localname, item_qname.localname))
                    for key in item.attrib:
                        attr_qname = etree.QName(key)
                        if attr_qname.namespace == _WML_NS:
                            attribute_contexts.setdefault((part_key, attr_qname.localname), set()).add(
                                item_qname.localname
                            )

    context_reference = Path(__file__).resolve().parent / "assets" / "word_schema_contexts.json"
    if context_reference.is_file():
        reference = json.loads(context_reference.read_text(encoding="utf-8"))
        for part_key, parent_local, child_local in reference.get("pairs", []):
            pairs.add((part_key, parent_local, child_local))
        for part_key, attr_local, element_local in reference.get("attribute_contexts", []):
            attribute_contexts.setdefault((part_key, attr_local), set()).add(element_local)

    runtime_pairs = {
        ("p", "hyperlink"), ("hyperlink", "r"),
        ("p", "fldSimple"), ("fldSimple", "r"),
        ("p", "ins"), ("p", "del"), ("ins", "r"), ("del", "r"),
        ("r", "fldChar"), ("r", "instrText"), ("r", "delText"),
        ("r", "br"), ("r", "cr"), ("r", "tab"),
        ("tc", "tbl"),
        ("settings", "docVars"), ("docVars", "docVar"),
    }
    for part_key in {"word/document.xml", "word/header.xml", "word/footer.xml"}:
        pairs.update(
            (part_key, parent, child)
            for parent, child in runtime_pairs
            if parent not in {"settings", "docVars"}
        )
    pairs.update(
        ("word/settings.xml", parent, child)
        for parent, child in runtime_pairs
        if parent in {"settings", "docVars"}
    )
    # ECMA-376 CT_PageSz: orientation is optional but controls portrait/landscape.
    attribute_contexts.setdefault(("word/document.xml", "orient"), set()).add("pgSz")
    return frozenset(pairs), {
        attr_name: frozenset(contexts)
        for attr_name, contexts in attribute_contexts.items()
    }


def _visible_wml_text_node(name: str, element: etree._Element) -> bool:
    if not (
        name == "word/document.xml"
        or re.fullmatch(r"word/(?:header|footer)\d+\.xml", name)
    ):
        return False
    qname = etree.QName(element)
    if qname.namespace != _WML_NS or qname.localname not in _SUPPORTED_WML_TEXT_ELEMENTS:
        return False
    return any(
        etree.QName(ancestor).namespace == _WML_NS
        and etree.QName(ancestor).localname == "p"
        for ancestor in element.iterancestors()
    )


def _scrub_unknown_docx_xml(name: str, root: etree._Element) -> tuple[etree._Element, bool]:
    """Remove XML outside the small set of structures this processor supports."""
    found = False
    if name in {"docProps/app.xml", "docProps/core.xml", "docProps/custom.xml"}:
        expected_roots = {
            "docProps/app.xml": (
                "http://schemas.openxmlformats.org/officeDocument/2006/extended-properties",
                "Properties",
            ),
            "docProps/core.xml": (
                "http://schemas.openxmlformats.org/package/2006/metadata/core-properties",
                "coreProperties",
            ),
            "docProps/custom.xml": (
                "http://schemas.openxmlformats.org/officeDocument/2006/custom-properties",
                "Properties",
            ),
        }
        expected_namespace, expected_local = expected_roots[name]
        root_qname = etree.QName(root)
        if root_qname.namespace != expected_namespace or root_qname.localname != expected_local:
            return etree.Element(
                f"{{{expected_namespace}}}{expected_local}", nsmap={None: expected_namespace}
            ), True
        if root.attrib or (root.text or "").strip():
            found = True
        root.attrib.clear()
        root.text = None
        allowed_app_children = {
            "Application", "AppVersion", "Characters", "CharactersWithSpaces", "Company",
            "DocSecurity", "HeadingPairs", "HiddenSlides", "HLinks", "HyperlinkBase", "HyperlinksChanged",
            "Lines", "LinksUpToDate", "Manager", "MMClips", "Notes", "Pages", "Paragraphs",
            "PresentationFormat", "ScaleCrop", "SharedDoc", "Slides", "Template", "TitlesOfParts",
            "TotalTime", "Words",
        }
        allowed_core_namespaces = {
            "http://purl.org/dc/elements/1.1/",
            "http://purl.org/dc/terms/",
            expected_namespace,
        }
        for child in list(root):
            if not isinstance(child.tag, str):
                root.remove(child)
                continue
            child_qname = etree.QName(child)
            if name == "docProps/app.xml":
                known = child_qname.namespace == expected_namespace and child_qname.localname in allowed_app_children
            elif name == "docProps/core.xml":
                known = child_qname.namespace in allowed_core_namespaces
            else:
                known = child_qname.namespace == expected_namespace and child_qname.localname == "property"
            if not known:
                found = True
            root.remove(child)
        etree.cleanup_namespaces(root)
        return root, found

    if name == "[Content_Types].xml":
        if etree.QName(root).namespace != _CONTENT_TYPES_NS or etree.QName(root).localname != "Types":
            return etree.Element(f"{{{_CONTENT_TYPES_NS}}}Types", nsmap={None: _CONTENT_TYPES_NS}), True
        for attr in list(root.attrib):
            del root.attrib[attr]
            found = True
        if (root.text or "").strip():
            root.text = None
            found = True
        default_types = {
            "rels": "application/vnd.openxmlformats-package.relationships+xml",
            "xml": "application/xml",
            "jpeg": "image/jpeg",
            "jpg": "image/jpeg",
            "png": "image/png",
        }
        for child in list(root):
            if not isinstance(child.tag, str):
                root.remove(child)
                found = True
                continue
            qname = etree.QName(child)
            local = qname.localname
            attributes = {etree.QName(attr).localname for attr in child.attrib}
            unqualified_attributes = all(etree.QName(attr).namespace is None for attr in child.attrib)
            valid = (
                qname.namespace == _CONTENT_TYPES_NS
                and local in {"Default", "Override"}
                and unqualified_attributes
            )
            if local == "Default":
                valid = valid and attributes == {"Extension", "ContentType"} and (
                    default_types.get(child.get("Extension", "").lower()) == child.get("ContentType")
                )
            elif local == "Override":
                part_name = child.get("PartName", "")
                valid = (
                    valid
                    and attributes == {"PartName", "ContentType"}
                    and re.fullmatch(r"/[A-Za-z0-9_.\-/]+", part_name) is not None
                    and child.get("ContentType") in _ALLOWED_CONTENT_TYPES
                )
            if not valid:
                root.remove(child)
                found = True
                continue
            if (child.text or "").strip() or (child.tail or "").strip():
                child.text = None
                child.tail = None
                found = True
        etree.cleanup_namespaces(root)
        return root, found

    if name.endswith(".rels"):
        if etree.QName(root).namespace != _PACKAGE_REL_NS or etree.QName(root).localname != "Relationships":
            return etree.Element(f"{{{_PACKAGE_REL_NS}}}Relationships", nsmap={None: _PACKAGE_REL_NS}), True
        allowed_attributes = {"Id", "Type", "Target", "TargetMode"}
        for attr in list(root.attrib):
            del root.attrib[attr]
            found = True
        if (root.text or "").strip():
            root.text = None
            found = True
        for child in list(root):
            if not isinstance(child.tag, str):
                root.remove(child)
                found = True
                continue
            qname = etree.QName(child)
            if qname.namespace != _PACKAGE_REL_NS or qname.localname != "Relationship":
                root.remove(child)
                found = True
                continue
            for attr in list(child.attrib):
                attr_qname = etree.QName(attr)
                if attr_qname.namespace is not None or attr_qname.localname not in allowed_attributes:
                    del child.attrib[attr]
                    found = True
            if child.get("TargetMode") == "External":
                root.remove(child)
                continue
            if (
                not re.fullmatch(r"rId[1-9]\d*", child.get("Id", ""))
                or child.get("Type") not in _ALLOWED_RELATIONSHIP_TYPES
                or not child.get("Target")
                or child.get("TargetMode") not in {None, "External"}
            ):
                root.remove(child)
                found = True
                continue
            if (child.text or "").strip():
                child.text = None
                found = True
            if (child.tail or "").strip():
                child.tail = None
                found = True
        etree.cleanup_namespaces(root)
        return root, found

    if not name.startswith("word/") or name in {"word/commentsExtended.xml", "word/people.xml"}:
        return root, False
    allowed_element_ns = _DRAWING_NS if name.startswith("word/theme/") else _WML_NS
    supported_elements = _SUPPORTED_THEME_ELEMENTS if name.startswith("word/theme/") else _SUPPORTED_WML_ELEMENTS
    if name.startswith("word/theme/"):
        expected_root = "theme"
    elif name in {"word/document.xml", "word/glossary/document.xml"}:
        expected_root = "document"
    elif re.fullmatch(r"word/header\d+\.xml", name):
        expected_root = "hdr"
    elif re.fullmatch(r"word/footer\d+\.xml", name):
        expected_root = "ftr"
    else:
        expected_root = {
            "word/styles.xml": "styles",
            "word/stylesWithEffects.xml": "styles",
            "word/settings.xml": "settings",
            "word/numbering.xml": "numbering",
            "word/fontTable.xml": "fonts",
            "word/webSettings.xml": "webSettings",
        }.get(name)
    root_qname = etree.QName(root)
    if expected_root is None or (
        root_qname.namespace != allowed_element_ns or root_qname.localname != expected_root
    ):
        clean_root = etree.Element(
            f"{{{allowed_element_ns}}}{expected_root or 'unsupported'}",
            nsmap={"a" if allowed_element_ns == _DRAWING_NS else "w": allowed_element_ns},
        )
        if expected_root == "document":
            etree.SubElement(clean_root, f"{{{_WML_NS}}}body")
        return clean_root, True
    allowed_elements_by_ns: dict[str, set[str] | frozenset[str] | None] = {
        allowed_element_ns: supported_elements
    }
    if name == "word/settings.xml":
        allowed_elements_by_ns.update({
            "http://schemas.openxmlformats.org/officeDocument/2006/math": {
                "mathPr", "mathFont", "brkBin", "brkBinSub", "smallFrac", "dispDef",
                "lMargin", "rMargin", "defJc", "wrapIndent", "intLim", "naryLim",
            },
            "http://schemas.microsoft.com/office/word/2010/wordml": {"docId", "defaultImageDpi"},
            "urn:schemas-microsoft-com:office:office": {"shapedefaults", "shapelayout", "idmap"},
        })
    benign_extension_attributes = {
        "Ignorable", "PreserveAttributes", "PreserveElements", "ProcessContent",
        "paraId", "textId", "durableId", "anchorId", "editId", "conflictMode",
        "restartNumberingAfterBreak",
    }
    for element in list(root.iter()):
        if not isinstance(element.tag, str):
            parent = element.getparent()
            if parent is not None:
                parent.remove(element)
                found = True
            continue
        element_qname = etree.QName(element)
        if name == "word/settings.xml" and (
            element_qname.namespace,
            element_qname.localname,
        ) in {
            ("http://schemas.microsoft.com/office/word/2012/wordml", "docId"),
            ("http://schemas.microsoft.com/office/word/2010/wordml", "docId"),
            ("http://schemas.microsoft.com/office/word/2010/wordml", "defaultImageDpi"),
            ("urn:schemas-microsoft-com:office:office", "shapedefaults"),
            ("urn:schemas-microsoft-com:office:office", "shapelayout"),
            ("urn:schemas-microsoft-com:office:office", "idmap"),
        }:
            parent = element.getparent()
            if parent is not None:
                parent.remove(element)
            continue
        allowed_localnames = allowed_elements_by_ns.get(element_qname.namespace, set())
        if allowed_localnames is not None and element_qname.localname not in allowed_localnames:
            parent = element.getparent()
            if parent is not None:
                parent.remove(element)
                found = True
                continue

        if element_qname.namespace == _WML_NS:
            if element_qname.localname in {"rsid", "rsids", "rsidRoot"}:
                parent = element.getparent()
                if parent is not None:
                    parent.remove(element)
                continue
            allowed_parents = _WML_REQUIRED_PARENTS.get(element_qname.localname)
            if allowed_parents is not None:
                parent = element.getparent()
                parent_qname = etree.QName(parent) if parent is not None else None
                if (
                    parent_qname is None
                    or parent_qname.namespace != _WML_NS
                    or parent_qname.localname not in allowed_parents
                ):
                    if parent is not None:
                        parent.remove(element)
                    found = True
                    continue

            parent = element.getparent()
            if parent is not None:
                parent_qname = etree.QName(parent)
                if parent_qname.namespace == _WML_NS:
                    allowed_pairs, _ = _trusted_wml_schema()
                    part_key = _schema_part(name)
                    if (
                        part_key is None
                        or (part_key, parent_qname.localname, element_qname.localname) not in allowed_pairs
                    ):
                        parent.remove(element)
                        found = True
                        continue

        text_allowed = _visible_wml_text_node(name, element)
        if (element.text or "").strip() and not text_allowed:
            element.text = None
            found = True
        if (element.tail or "").strip():
            element.tail = None
            found = True

        for attr in list(element.attrib):
            attr_qname = etree.QName(attr)
            if attr_qname.namespace == _WML_NS and attr_qname.localname.startswith("rsid"):
                del element.attrib[attr]
                continue
            if attr_qname.namespace == _WML_NS:
                allowed_attribute = attr_qname.localname in _SUPPORTED_WML_ATTRIBUTES
                allowed_contexts = _WML_CONTEXTUAL_ATTRIBUTES.get(attr_qname.localname)
                if allowed_contexts is None:
                    _, trusted_attribute_contexts = _trusted_wml_schema()
                    part_key = _schema_part(name)
                    allowed_contexts = trusted_attribute_contexts.get((part_key, attr_qname.localname))
                allowed_attribute = allowed_attribute and (
                    allowed_contexts is not None and element_qname.localname in allowed_contexts
                )
            elif attr_qname.namespace == _DRAWING_NS:
                allowed_attribute = (
                    name.startswith("word/theme/")
                    and attr_qname.localname in _SUPPORTED_THEME_ATTRIBUTES
                )
            elif attr_qname.namespace == _OFFICE_REL_NS:
                allowed_attribute = (
                    attr_qname.localname in {"id", "embed", "link"}
                    and re.fullmatch(r"rId\d+", element.attrib[attr]) is not None
                )
            elif attr_qname.namespace == _XML_NS:
                if attr_qname.localname == "space":
                    allowed_attribute = element.attrib[attr] in {"default", "preserve"}
                elif attr_qname.localname == "lang":
                    allowed_attribute = re.fullmatch(
                        r"[A-Za-z]{2,3}(?:-[A-Za-z0-9]{2,8})*", element.attrib[attr]
                    ) is not None
                else:
                    allowed_attribute = False
            elif attr_qname.namespace == "http://schemas.openxmlformats.org/officeDocument/2006/math":
                allowed_attribute = (
                    attr_qname.localname == "val"
                    and _safe_known_attribute_value(
                        name, element_qname.localname, attr_qname.localname, element.attrib[attr]
                    )
                )
            elif attr_qname.namespace is None:
                allowed_attribute = (
                    name.startswith("word/theme/")
                    and attr_qname.localname in _SUPPORTED_THEME_ATTRIBUTES
                )
            else:
                allowed_attribute = False
            if allowed_attribute and attr_qname.namespace == _WML_NS:
                allowed_attribute = _safe_known_attribute_value(
                    name, element_qname.localname, attr_qname.localname, element.attrib[attr]
                )
            elif allowed_attribute and name.startswith("word/theme/"):
                allowed_attribute = _safe_known_attribute_value(
                    name, element_qname.localname, attr_qname.localname, element.attrib[attr]
                )
            if not allowed_attribute:
                del element.attrib[attr]
                if attr_qname.localname not in benign_extension_attributes:
                    found = True
    etree.cleanup_namespaces(root)
    return root, found


def _safe_builtin_custom_xml(name: str, data: bytes) -> bool:
    """Recognize only python-docx's empty, fixed bibliography boilerplate."""
    try:
        root = etree.fromstring(data)
    except etree.XMLSyntaxError:
        return False
    bibliography = "http://schemas.openxmlformats.org/officeDocument/2006/bibliography"
    if name == "customXml/item1.xml":
        return (
            etree.QName(root).namespace == bibliography
            and etree.QName(root).localname == "Sources"
            and len(root) == 0
            and root.attrib == {"SelectedStyle": "/APA.XSL", "StyleName": "APA"}
        )
    if name == "customXml/itemProps1.xml":
        refs = root.xpath(".//*[local-name()='schemaRef']")
        return (
            etree.QName(root).localname == "datastoreItem"
            and len(refs) == 1
            and refs[0].get("{http://schemas.openxmlformats.org/officeDocument/2006/customXml}uri") == bibliography
            and not any((node.text or "").strip() for node in root.iter())
        )
    if name == "customXml/_rels/item1.xml.rels":
        relationships = list(root)
        return (
            len(relationships) == 1
            and relationships[0].get("Target") == "itemProps1.xml"
            and relationships[0].get("TargetMode") is None
            and relationships[0].get("Type", "").endswith("/customXmlProps")
        )
    return False


def _canonical_empty_note_separators(name: str, data: bytes) -> bytes | None:
    """Accept only Word's text-free default note separator boilerplate.

    The returned XML is rebuilt from scratch so metadata or arbitrary
    attributes cannot hitchhike through an otherwise empty notes part.
    """
    kinds = {
        "word/footnotes.xml": ("footnotes", "footnote"),
        "word/endnotes.xml": ("endnotes", "endnote"),
    }
    expected = kinds.get(name)
    if expected is None:
        return None
    namespace = "http://schemas.openxmlformats.org/wordprocessingml/2006/main"
    try:
        root = etree.fromstring(data)
    except etree.XMLSyntaxError:
        return None
    if etree.QName(root).namespace != namespace or etree.QName(root).localname != expected[0]:
        return None
    notes = list(root)
    expected_notes = {("separator", "-1"), ("continuationSeparator", "0")}
    actual_notes: set[tuple[str, str]] = set()
    allowed_elements = {
        expected[0], expected[1], "p", "pPr", "spacing", "r", "rPr",
        "separator", "continuationSeparator",
    }
    for element in root.iter():
        qname = etree.QName(element)
        if qname.namespace != namespace or qname.localname not in allowed_elements:
            return None
        if (element.text or "").strip() or (element.tail or "").strip():
            return None
    for note in notes:
        if etree.QName(note).localname != expected[1]:
            return None
        note_type = note.get(qn("w:type"), "")
        note_id = note.get(qn("w:id"), "")
        actual_notes.add((note_type, note_id))
        required_leaf = note.xpath(
            ".//w:separator" if note_type == "separator" else ".//w:continuationSeparator",
            namespaces={"w": namespace},
        )
        if not required_leaf:
            return None
    if len(notes) != 2 or actual_notes != expected_notes:
        return None

    clean_root = etree.Element(f"{{{namespace}}}{expected[0]}", nsmap={"w": namespace})
    for note_type, note_id, leaf in (
        ("separator", "-1", "separator"),
        ("continuationSeparator", "0", "continuationSeparator"),
    ):
        note = etree.SubElement(clean_root, f"{{{namespace}}}{expected[1]}")
        note.set(qn("w:type"), note_type)
        note.set(qn("w:id"), note_id)
        paragraph = etree.SubElement(note, f"{{{namespace}}}p")
        run = etree.SubElement(paragraph, f"{{{namespace}}}r")
        etree.SubElement(run, f"{{{namespace}}}{leaf}")
    return etree.tostring(clean_root, xml_declaration=True, encoding="UTF-8", standalone=True)


def _canonical_xml_root(name: str) -> etree._Element | None:
    if name == "[Content_Types].xml":
        return etree.Element(f"{{{_CONTENT_TYPES_NS}}}Types", nsmap={None: _CONTENT_TYPES_NS})
    if name.endswith(".rels"):
        return etree.Element(f"{{{_PACKAGE_REL_NS}}}Relationships", nsmap={None: _PACKAGE_REL_NS})
    roots = {
        "word/document.xml": "document",
        "word/styles.xml": "styles",
        "word/stylesWithEffects.xml": "styles",
        "word/settings.xml": "settings",
        "word/numbering.xml": "numbering",
        "word/fontTable.xml": "fonts",
        "word/webSettings.xml": "webSettings",
    }
    local = roots.get(name)
    if local is None:
        if re.fullmatch(r"word/header\d+\.xml", name):
            local = "hdr"
        elif re.fullmatch(r"word/footer\d+\.xml", name):
            local = "ftr"
        elif re.fullmatch(r"word/theme/theme\d+\.xml", name):
            return etree.Element(f"{{{_DRAWING_NS}}}theme", nsmap={"a": _DRAWING_NS})
        else:
            return None
    root = etree.Element(f"{{{_WML_NS}}}{local}", nsmap={"w": _WML_NS})
    if local == "document":
        etree.SubElement(root, f"{{{_WML_NS}}}body")
    return root


def _package_identifier_maps(archive: ZipFile) -> tuple[dict[str, str], dict[str, str]]:
    style_map: dict[str, str] = {}
    font_map: dict[str, str] = {}
    trusted_styles, _ = _trusted_docx_strings()
    if "word/styles.xml" in archive.namelist():
        root = etree.fromstring(archive.read("word/styles.xml"))
        custom_index = 1
        for style in root.iter(f"{{{_WML_NS}}}style"):
            old = style.get(qn("w:styleId"))
            if not old or old in style_map:
                continue
            if old in trusted_styles:
                style_map[old] = old
            else:
                style_map[old] = f"LocalStyle{custom_index:03d}"
                custom_index += 1
    if "word/fontTable.xml" in archive.namelist():
        root = etree.fromstring(archive.read("word/fontTable.xml"))
        custom_index = 1
        for font in root.iter(f"{{{_WML_NS}}}font"):
            old = font.get(qn("w:name"))
            if not old or old in font_map:
                continue
            if old in _SAFE_FONT_NAMES:
                font_map[old] = old
            else:
                font_map[old] = f"LocalFont{custom_index:03d}"
                custom_index += 1
    return style_map, font_map


def _canonicalize_package_identifiers(
    name: str, root: etree._Element, style_map: dict[str, str], font_map: dict[str, str]
) -> None:
    if name == "word/styles.xml":
        for style_index, style in enumerate(root.iter(f"{{{_WML_NS}}}style"), 1):
            old = style.get(qn("w:styleId"))
            if old in style_map:
                mapped = style_map[old]
                style.set(qn("w:styleId"), mapped)
                style_name = style.find(qn("w:name"))
                if style_name is not None:
                    current_name = style_name.get(qn("w:val"), "")
                    if mapped != old or not _safe_known_attribute_value(
                        "word/styles.xml", "name", "val", current_name
                    ):
                        style_name.set(qn("w:val"), f"Local Style {style_index:03d}")
    for element in root.iter():
        if not isinstance(element.tag, str):
            continue
        local = etree.QName(element).localname
        if local in _STYLE_REFERENCE_ELEMENTS:
            old = element.get(qn("w:val"))
            if old in style_map:
                element.set(qn("w:val"), style_map[old])
        if local == "rFonts":
            for attr_local in ("ascii", "hAnsi", "eastAsia", "cs"):
                attr = qn(f"w:{attr_local}")
                old = element.get(attr)
                if old in font_map:
                    element.set(attr, font_map[old])
    if name == "word/fontTable.xml":
        for font in root.iter(f"{{{_WML_NS}}}font"):
            old = font.get(qn("w:name"))
            if old in font_map:
                font.set(qn("w:name"), font_map[old])
            alt_name = font.find(qn("w:altName"))
            if alt_name is not None:
                old_alt = alt_name.get(qn("w:val"))
                if old_alt not in _SAFE_FONT_NAMES and old_alt not in font_map.values():
                    alt_name.set(qn("w:val"), "Arial")
    if re.fullmatch(r"word/theme/theme\d+\.xml", name):
        root.set("name", "Office Theme")
        for element in root.iter():
            if etree.QName(element).localname not in {"latin", "ea", "cs", "font"}:
                continue
            old = element.get("typeface")
            if old is not None and old not in _SAFE_FONT_NAMES and old not in font_map.values():
                element.set("typeface", "Arial" if etree.QName(element).localname == "latin" else "")


def _relationship_part_for(source_part: str) -> str:
    directory, filename = posixpath.split(source_part)
    return posixpath.join(directory, "_rels", filename + ".rels")


_REVERSIBLE_GRAPHIC_PART = re.compile(r"word/media/[A-Za-z0-9_.-]+\.(?:png|jpe?g)", re.IGNORECASE)
_REVERSIBLE_GRAPHIC_SOURCE = re.compile(r"word/(?:document|header\d+|footer\d+)\.xml")
_REVERSIBLE_GRAPHIC_CONTENT_TYPES = {
    "png": "image/png",
    "jpg": "image/jpeg",
    "jpeg": "image/jpeg",
}


def _valid_reversible_graphic(name: str, data: bytes) -> bool:
    extension = posixpath.splitext(name)[1].lstrip(".").lower()
    if extension == "png":
        return data.startswith(bytes.fromhex("89504e470d0a1a0a"))
    if extension in {"jpg", "jpeg"}:
        return data.startswith(b"\xff\xd8\xff")
    return False


_ZIP_METADATA_FIELDS = {
    "date_time", "compress_type", "comment", "extra", "create_system",
    "create_version", "extract_version", "reserved", "flag_bits", "volume",
    "internal_attr", "external_attr",
}


def _archive_zipinfo(info: ZipInfo) -> dict:
    """Serialize metadata for a graphic member moved into protected storage."""
    return {
        "date_time": list(info.date_time),
        "compress_type": info.compress_type,
        "comment": base64.b64encode(info.comment).decode("ascii"),
        "extra": base64.b64encode(info.extra).decode("ascii"),
        "create_system": info.create_system,
        "create_version": info.create_version,
        "extract_version": info.extract_version,
        "reserved": info.reserved,
        "flag_bits": info.flag_bits,
        "volume": info.volume,
        "internal_attr": info.internal_attr,
        "external_attr": info.external_attr,
    }


def _restore_zipinfo(name: str, metadata: object) -> ZipInfo:
    """Validate and reconstruct protected member metadata fail-closed."""
    if not isinstance(metadata, dict) or set(metadata) != _ZIP_METADATA_FIELDS:
        raise RestorationError("Архив графики DOCX содержит повреждённые метаданные ZIP")
    date_time = metadata.get("date_time")
    if (
        not isinstance(date_time, list)
        or len(date_time) != 6
        or any(isinstance(value, bool) or not isinstance(value, int) for value in date_time)
    ):
        raise RestorationError("Архив графики DOCX содержит повреждённую дату ZIP")
    integer_fields = _ZIP_METADATA_FIELDS - {"date_time", "comment", "extra"}
    if any(
        isinstance(metadata.get(field), bool) or not isinstance(metadata.get(field), int)
        for field in integer_fields
    ):
        raise RestorationError("Архив графики DOCX содержит повреждённые метаданные ZIP")
    if metadata["compress_type"] not in {ZIP_STORED, ZIP_DEFLATED}:
        raise RestorationError("Архив графики DOCX содержит недопустимое сжатие ZIP")
    year, month, day, hour, minute, second = date_time
    if not (
        1980 <= year <= 2107
        and 1 <= month <= 12
        and 1 <= day <= 31
        and 0 <= hour <= 23
        and 0 <= minute <= 59
        and 0 <= second <= 59
    ):
        raise RestorationError("Архив графики DOCX содержит повреждённую дату ZIP")
    byte_fields = {"create_system", "create_version", "extract_version", "reserved"}
    short_fields = {"flag_bits", "volume", "internal_attr"}
    if any(not 0 <= metadata[field] <= 0xFF for field in byte_fields):
        raise RestorationError("Архив графики DOCX содержит повреждённые метаданные ZIP")
    if any(not 0 <= metadata[field] <= 0xFFFF for field in short_fields):
        raise RestorationError("Архив графики DOCX содержит повреждённые метаданные ZIP")
    if not 0 <= metadata["external_attr"] <= 0xFFFFFFFF:
        raise RestorationError("Архив графики DOCX содержит повреждённые метаданные ZIP")
    try:
        comment = base64.b64decode(metadata["comment"], validate=True)
        extra = base64.b64decode(metadata["extra"], validate=True)
    except (ValueError, TypeError) as exc:
        raise RestorationError("Архив графики DOCX содержит повреждённые метаданные ZIP") from exc
    if len(comment) > 65535 or len(extra) > 65535:
        raise RestorationError("Архив графики DOCX содержит слишком большие метаданные ZIP")
    offset = 0
    while offset < len(extra):
        if len(extra) - offset < 4:
            raise RestorationError("Архив графики DOCX содержит повреждённые метаданные ZIP")
        field_size = int.from_bytes(extra[offset + 2:offset + 4], "little")
        offset += 4
        if field_size > len(extra) - offset:
            raise RestorationError("Архив графики DOCX содержит повреждённые метаданные ZIP")
        offset += field_size
    try:
        info = ZipInfo(name, tuple(date_time))
    except (TypeError, ValueError) as exc:
        raise RestorationError("Архив графики DOCX содержит повреждённую дату ZIP") from exc
    info.compress_type = metadata["compress_type"]
    info.comment = comment
    info.extra = extra
    info.create_system = metadata["create_system"]
    info.create_version = metadata["create_version"]
    info.extract_version = metadata["extract_version"]
    info.reserved = metadata["reserved"]
    info.flag_bits = metadata["flag_bits"]
    info.volume = metadata["volume"]
    info.internal_attr = metadata["internal_attr"]
    info.external_attr = metadata["external_attr"]
    return info


def _resolve_package_target(source_part: str, target: str) -> str | None:
    if not target or target.startswith(("/", "\\")) or "\\" in target:
        return None
    resolved = posixpath.normpath(posixpath.join(posixpath.dirname(source_part), target))
    if resolved == ".." or resolved.startswith("../"):
        return None
    return resolved


def _extract_reversible_docx_graphics(path: Path) -> dict | None:
    """Move static inline images into the protected per-document mapping.

    The anonymized package receives inert text placeholders.  The original
    DrawingML/VML fragment, image bytes, relationship and content-type rows are
    kept outside the exported DOCX and are later restored atomically.
    """
    replacement = path.with_suffix(".graphics.tmp")
    archive: dict = {
        "schema": 2,
        "kind": "docx-graphics",
        "fragments": [],
        "parts": {},
        "part_metadata": {},
        "relationships": [],
        "content_types": [],
    }
    try:
        with ZipFile(path, "r") as incoming:
            names = set(incoming.namelist())
            payloads = {info.filename: incoming.read(info.filename) for info in incoming.infolist()}
            infos = incoming.infolist()

        relationship_keys: set[tuple[str, str]] = set()
        selected_parts: set[str] = set()
        marker_index = 1
        occupied_text = "\n".join(
            text
            for source_name in names
            if _REVERSIBLE_GRAPHIC_SOURCE.fullmatch(source_name)
            for text in etree.fromstring(payloads[source_name]).itertext()
        )
        for source_part in sorted(name for name in names if _REVERSIBLE_GRAPHIC_SOURCE.fullmatch(name)):
            rels_part = _relationship_part_for(source_part)
            if rels_part not in payloads:
                continue
            root = etree.fromstring(payloads[source_part])
            rels_root = etree.fromstring(payloads[rels_part])
            relationships = {child.get("Id", ""): child for child in rels_root}
            changed = False
            for graphic in list(root.xpath(".//w:drawing | .//w:pict", namespaces={"w": _WML_NS})):
                relationship_ids = {
                    value
                    for element in graphic.iter()
                    for attr, value in element.attrib.items()
                    if etree.QName(attr).namespace == _OFFICE_REL_NS
                }
                selected = []
                for relationship_id in relationship_ids:
                    relationship = relationships.get(relationship_id)
                    if relationship is None or relationship.get("TargetMode") is not None:
                        selected = []
                        break
                    target = _resolve_package_target(source_part, relationship.get("Target", ""))
                    if (
                        relationship.get("Type")
                        != "http://schemas.openxmlformats.org/officeDocument/2006/relationships/image"
                        or target is None
                        or _REVERSIBLE_GRAPHIC_PART.fullmatch(target) is None
                        or target not in payloads
                        or not _valid_reversible_graphic(target, payloads[target])
                    ):
                        selected = []
                        break
                    selected.append((relationship_id, relationship, target))
                relationless_vml = (
                    etree.QName(graphic).localname == "pict" and not relationship_ids
                )
                if not selected and not relationless_vml:
                    continue

                marker = f"[СХЕМА-{marker_index}]"
                while marker in occupied_text:
                    marker_index += 1
                    marker = f"[СХЕМА-{marker_index}]"
                marker_index += 1
                occupied_text += "\n" + marker
                archive["fragments"].append({
                    "marker": marker,
                    "part": source_part,
                    "xml": base64.b64encode(etree.tostring(graphic, with_tail=False)).decode("ascii"),
                })
                for relationship_id, relationship, target in selected:
                    key = (rels_part, relationship_id)
                    if key not in relationship_keys:
                        archive["relationships"].append({
                            "part": rels_part,
                            "id": relationship_id,
                            "type": relationship.get("Type", ""),
                            "target": relationship.get("Target", ""),
                        })
                        relationship_keys.add(key)
                    if target not in selected_parts:
                        archive["parts"][target] = base64.b64encode(payloads[target]).decode("ascii")
                        source_info = next(info for info in infos if info.filename == target)
                        archive["part_metadata"][target] = _archive_zipinfo(source_info)
                        selected_parts.add(target)

                parent = graphic.getparent()
                if parent is None:
                    continue
                placeholder = etree.Element(f"{{{_WML_NS}}}t")
                placeholder.text = marker
                placeholder.tail = graphic.tail
                parent.replace(graphic, placeholder)
                changed = True
            if changed:
                payloads[source_part] = etree.tostring(
                    root, xml_declaration=True, encoding="UTF-8", standalone=True
                )

        if not archive["fragments"]:
            return None

        for relation in archive["relationships"]:
            rels_root = etree.fromstring(payloads[relation["part"]])
            for child in list(rels_root):
                if child.get("Id") == relation["id"]:
                    rels_root.remove(child)
            payloads[relation["part"]] = etree.tostring(
                rels_root, xml_declaration=True, encoding="UTF-8", standalone=True
            )

        content_types = etree.fromstring(payloads["[Content_Types].xml"])
        selected_extensions = {posixpath.splitext(name)[1].lstrip(".").lower() for name in selected_parts}
        for child in content_types:
            part_name = child.get("PartName", "").lstrip("/")
            extension = child.get("Extension", "").lower()
            if part_name in selected_parts or extension in selected_extensions:
                archive["content_types"].append(
                    base64.b64encode(etree.tostring(child, with_tail=False)).decode("ascii")
                )

        with ZipFile(replacement, "w", ZIP_DEFLATED) as outgoing:
            for info in infos:
                if info.filename in selected_parts:
                    continue
                outgoing.writestr(info, payloads[info.filename], compress_type=ZIP_DEFLATED)
        os.replace(replacement, path)
        return archive
    finally:
        replacement.unlink(missing_ok=True)


def _decode_reversible_graphic_fragment(encoded: str) -> etree._Element:
    try:
        graphic_xml = base64.b64decode(encoded, validate=True)
        if re.search(br"<!DOCTYPE", graphic_xml, re.IGNORECASE):
            raise RestorationError("Архив графики DOCX содержит запрещённый DTD в XML")
        graphic = etree.fromstring(
            graphic_xml,
            parser=etree.XMLParser(
                resolve_entities=False, no_network=True, load_dtd=False
            ),
        )
    except (ValueError, TypeError, etree.XMLSyntaxError) as exc:
        raise RestorationError("Архив графики DOCX содержит повреждённый XML") from exc
    graphic_qname = etree.QName(graphic)
    if graphic_qname.namespace != _WML_NS or graphic_qname.localname not in {"drawing", "pict"}:
        raise RestorationError("Архив графики DOCX содержит недопустимый XML")
    forbidden_elements = {
        "activex", "altchunk", "attachedtemplate", "control", "frame",
        "frameset", "hlinkclick", "hlinkhover", "object", "oleobject",
        "quicktimefile", "audiofile", "videofile", "contentpart",
    }
    external_value = re.compile(
        r"(?:[a-z][a-z0-9+.-]*://|(?:file|javascript|mailto|ms-[a-z0-9+.-]+):|\\\\)",
        re.IGNORECASE,
    )
    for element in graphic.iter():
        if not isinstance(element.tag, str):
            raise RestorationError("Архив графики DOCX содержит недопустимый XML")
        element_qname = etree.QName(element)
        if element_qname.localname.casefold() in forbidden_elements:
            raise RestorationError("Архив графики DOCX содержит запрещённый активный XML")
        for attr, value in element.attrib.items():
            attr_qname = etree.QName(attr)
            allowed_schema_uri = (
                element_qname.localname == "graphicData"
                and attr_qname.localname == "uri"
                and value == "http://schemas.openxmlformats.org/drawingml/2006/picture"
            )
            if external_value.search(value) and not allowed_schema_uri:
                raise RestorationError("Архив графики DOCX содержит внешнюю ссылку в XML")
            if attr_qname.namespace != _OFFICE_REL_NS:
                continue
            permitted_reference = (
                (element_qname.localname, attr_qname.localname) in {
                    ("blip", "embed"),
                    ("imagedata", "id"),
                }
                and re.fullmatch(r"rId[1-9]\d*", value) is not None
            )
            if not permitted_reference:
                raise RestorationError("Архив графики DOCX содержит недопустимую ссылку в XML")
    return graphic


def _restore_reversible_docx_graphics(
    path: Path, archive: dict | None, mapping: dict[str, str]
) -> set[str]:
    if archive is None:
        return set()
    if (
        not isinstance(archive, dict)
        or archive.get("schema") not in {1, 2}
        or archive.get("kind") != "docx-graphics"
    ):
        raise RestorationError("Повреждён защищённый архив графики DOCX")
    archive_schema = archive["schema"]
    fragments = archive.get("fragments")
    parts = archive.get("parts")
    part_metadata = archive.get("part_metadata", {})
    relationships = archive.get("relationships")
    content_types = archive.get("content_types")
    if (
        not isinstance(fragments, list)
        or not isinstance(parts, dict)
        or not isinstance(part_metadata, dict)
        or not isinstance(relationships, list)
        or not isinstance(content_types, list)
        or (archive_schema == 2 and set(part_metadata) != set(parts))
    ):
        raise RestorationError("Повреждён защищённый архив графики DOCX")
    if len(fragments) > 1000 or len(parts) > 1000 or len(relationships) > 2000:
        raise RestorationError("Архив графики DOCX превышает допустимые пределы")
    encoded_xml_size = sum(
        len(item.get("xml", "")) for item in fragments if isinstance(item, dict)
    ) + sum(len(item) for item in content_types if isinstance(item, str))
    if encoded_xml_size > 32 * 1024 * 1024:
        raise RestorationError("Архив графики DOCX превышает допустимый размер XML")

    replacement = path.with_suffix(".graphics-restore.tmp")
    restored_markers: set[str] = set()
    try:
        with ZipFile(path, "r") as incoming:
            infos = incoming.infolist()
            payloads = {info.filename: incoming.read(info.filename) for info in infos}

        active_fragments: list[tuple[dict, etree._Element]] = []
        active_relationship_ids: set[tuple[str, str]] = set()
        for item in fragments:
            if not isinstance(item, dict):
                raise RestorationError("Повреждены фрагменты архива графики DOCX")
            marker, source_part, encoded = item.get("marker"), item.get("part"), item.get("xml")
            if (
                not isinstance(marker, str)
                or re.fullmatch(r"\[СХЕМА-[1-9]\d*\]", marker) is None
                or not isinstance(source_part, str)
                or _REVERSIBLE_GRAPHIC_SOURCE.fullmatch(source_part) is None
                or source_part not in payloads
                or not isinstance(encoded, str)
            ):
                raise RestorationError("Архив графики DOCX содержит недопустимый фрагмент")
            root = etree.fromstring(payloads[source_part])
            if not any(node.text == marker for node in root.iter(f"{{{_WML_NS}}}t")):
                continue
            graphic = _decode_reversible_graphic_fragment(encoded)
            active_fragments.append((item, graphic))
            active_relationship_ids.update(
                (source_part, value)
                for element in graphic.iter()
                for attr, value in element.attrib.items()
                if etree.QName(attr).namespace == _OFFICE_REL_NS
            )
        if not active_fragments:
            return set()

        decoded_parts: dict[str, bytes] = {}
        restored_part_infos: dict[str, ZipInfo] = {}
        total_size = 0
        for name, encoded in parts.items():
            if not isinstance(name, str) or _REVERSIBLE_GRAPHIC_PART.fullmatch(name) is None or not isinstance(encoded, str):
                raise RestorationError("Архив графики DOCX содержит недопустимую часть")
            try:
                data = base64.b64decode(encoded, validate=True)
            except (ValueError, TypeError) as exc:
                raise RestorationError("Архив графики DOCX содержит повреждённые данные") from exc
            total_size += len(data)
            if total_size > 100 * 1024 * 1024:
                raise RestorationError("Архив графики DOCX превышает допустимый размер")
            if name in payloads and payloads[name] != data:
                raise RestorationError("Конфликт имени восстанавливаемой графической части")
            if not _valid_reversible_graphic(name, data):
                raise RestorationError("Архив графики DOCX содержит недопустимые графические данные")
            decoded_parts[name] = data
            if archive_schema == 2:
                restored_part_infos[name] = _restore_zipinfo(name, part_metadata[name])

        id_maps: dict[str, dict[str, str]] = {}
        active_parts: set[str] = set()
        for item in relationships:
            if not isinstance(item, dict):
                raise RestorationError("Повреждены связи архива графики DOCX")
            rels_part = item.get("part")
            old_id = item.get("id")
            rel_type = item.get("type")
            target = item.get("target")
            if (
                not isinstance(rels_part, str)
                or not re.fullmatch(r"word/_rels/(?:document|header\d+|footer\d+)\.xml\.rels", rels_part)
                or not isinstance(old_id, str)
                or not re.fullmatch(r"rId[1-9]\d*", old_id)
                or rel_type != "http://schemas.openxmlformats.org/officeDocument/2006/relationships/image"
                or not isinstance(target, str)
            ):
                raise RestorationError("Архив графики DOCX содержит недопустимую связь")
            source_part = rels_part.replace("/_rels/", "/")[:-5]
            if (source_part, old_id) not in active_relationship_ids:
                continue
            resolved = _resolve_package_target(source_part, target)
            if resolved not in decoded_parts or rels_part not in payloads:
                raise RestorationError("Архив графики DOCX содержит неполную связь")
            rels_root = etree.fromstring(payloads[rels_part])
            used_ids = {child.get("Id", "") for child in rels_root}
            new_id = old_id
            if new_id in used_ids:
                index = 1
                while f"rId{index}" in used_ids:
                    index += 1
                new_id = f"rId{index}"
            relationship = etree.SubElement(rels_root, f"{{{_PACKAGE_REL_NS}}}Relationship")
            relationship.set("Id", new_id)
            relationship.set("Type", rel_type)
            relationship.set("Target", target)
            id_maps.setdefault(source_part, {})[old_id] = new_id
            active_parts.add(resolved)
            payloads[rels_part] = etree.tostring(
                rels_root, xml_declaration=True, encoding="UTF-8", standalone=True
            )

        for item, graphic in active_fragments:
            marker, source_part = item["marker"], item["part"]
            root = etree.fromstring(payloads[source_part])
            placeholders = [
                node for node in root.iter(f"{{{_WML_NS}}}t") if node.text == marker
            ]
            if not placeholders:
                continue
            for text_node in graphic.iter(f"{{{_WML_NS}}}t"):
                if text_node.text:
                    text_node.text = _restore_text(text_node.text, mapping)
            for element in graphic.iter():
                for attr in list(element.attrib):
                    if etree.QName(attr).namespace == _OFFICE_REL_NS:
                        old_id = element.attrib[attr]
                        new_id = id_maps.get(source_part, {}).get(old_id)
                        if new_id is None:
                            raise RestorationError("Архив графики DOCX содержит несвязанную ссылку")
                        element.attrib[attr] = new_id
            placeholder = placeholders[0]
            parent = placeholder.getparent()
            if parent is None:
                continue
            graphic.tail = placeholder.tail
            parent.replace(placeholder, graphic)
            restored_markers.add(marker)
            payloads[source_part] = etree.tostring(
                root, xml_declaration=True, encoding="UTF-8", standalone=True
            )

        if not restored_markers:
            return set()

        content_root = etree.fromstring(payloads["[Content_Types].xml"])
        existing = {
            (etree.QName(child).localname,
             (child.get("PartName") or child.get("Extension") or "").casefold()):
            child.get("ContentType")
            for child in content_root
        }
        for encoded in content_types:
            if not isinstance(encoded, str):
                raise RestorationError("Повреждены типы содержимого архива графики DOCX")
            try:
                content_type_xml = base64.b64decode(encoded, validate=True)
                if re.search(br"<!DOCTYPE", content_type_xml, re.IGNORECASE):
                    raise RestorationError("Типы содержимого архива графики DOCX содержат запрещённый DTD")
                child = etree.fromstring(
                    content_type_xml,
                    parser=etree.XMLParser(
                        resolve_entities=False, no_network=True, load_dtd=False
                    ),
                )
            except (ValueError, TypeError, etree.XMLSyntaxError) as exc:
                raise RestorationError("Повреждены типы содержимого архива графики DOCX") from exc
            qname = etree.QName(child)
            attributes = {etree.QName(attr).localname for attr in child.attrib}
            if qname.namespace != _CONTENT_TYPES_NS or qname.localname not in {"Default", "Override"}:
                raise RestorationError("Архив графики DOCX содержит недопустимый тип содержимого")
            if qname.localname == "Default":
                extension = child.get("Extension", "").lower()
                valid = (
                    attributes == {"Extension", "ContentType"}
                    and child.get("ContentType") == _REVERSIBLE_GRAPHIC_CONTENT_TYPES.get(extension)
                )
                identity = ("Default", extension.casefold())
            else:
                part_name = child.get("PartName", "").lstrip("/")
                extension = posixpath.splitext(part_name)[1].lstrip(".").lower()
                if part_name in decoded_parts and part_name not in active_parts:
                    continue
                valid = (
                    attributes == {"PartName", "ContentType"}
                    and part_name in active_parts
                    and child.get("ContentType") == _REVERSIBLE_GRAPHIC_CONTENT_TYPES.get(extension)
                )
                identity = ("Override", child.get("PartName", "").casefold())
            if not valid or (child.text or "").strip() or (child.tail or "").strip():
                raise RestorationError("Архив графики DOCX содержит недопустимый тип содержимого")
            current_type = existing.get(identity)
            if current_type is not None and current_type != child.get("ContentType"):
                raise RestorationError("Конфликт типа содержимого восстанавливаемой графики")
            if current_type is None:
                content_root.append(child)
                existing[identity] = child.get("ContentType")
        payloads["[Content_Types].xml"] = etree.tostring(
            content_root, xml_declaration=True, encoding="UTF-8", standalone=True
        )
        payloads.update({name: decoded_parts[name] for name in active_parts})

        with ZipFile(replacement, "w", ZIP_DEFLATED) as outgoing:
            written = set()
            for info in infos:
                outgoing.writestr(info, payloads[info.filename], compress_type=ZIP_DEFLATED)
                written.add(info.filename)
            for name in sorted(set(payloads) - written):
                restored_info = restored_part_infos.get(name)
                if restored_info is None:
                    outgoing.writestr(name, payloads[name], compress_type=ZIP_DEFLATED)
                else:
                    outgoing.writestr(restored_info, payloads[name])
        os.replace(replacement, path)
        return restored_markers
    finally:
        replacement.unlink(missing_ok=True)


def _remove_dangling_relationship_references(path: Path) -> None:
    """Canonicalize relationship IDs and remove references to filtered relationships."""
    replacement = path.with_suffix(".relationships.tmp")
    try:
        with ZipFile(path, "r") as zin:
            names = set(zin.namelist())
            relationship_roots: dict[str, etree._Element] = {}
            relationship_id_maps: dict[str, dict[str, str]] = {}
            for rel_name in names:
                if not rel_name.endswith(".rels"):
                    continue
                root = etree.fromstring(zin.read(rel_name))
                relationships = sorted(
                    list(root),
                    key=lambda rel: (
                        rel.get("Type", ""),
                        rel.get("Target", ""),
                        rel.get("TargetMode", ""),
                    ),
                )
                canonical_root = etree.Element(
                    f"{{{_PACKAGE_REL_NS}}}Relationships",
                    nsmap={None: _PACKAGE_REL_NS},
                )
                canonical_ids: dict[tuple[str, str, str], str] = {}
                id_map: dict[str, str] = {}
                for rel in relationships:
                    old_id = rel.get("Id")
                    key = (
                        rel.get("Type", ""),
                        rel.get("Target", ""),
                        rel.get("TargetMode", ""),
                    )
                    new_id = canonical_ids.get(key)
                    if new_id is None:
                        new_id = f"rId{len(canonical_ids) + 1}"
                        canonical_ids[key] = new_id
                        canonical_rel = etree.SubElement(
                            canonical_root,
                            f"{{{_PACKAGE_REL_NS}}}Relationship",
                        )
                        canonical_rel.set("Id", new_id)
                        canonical_rel.set("Type", key[0])
                        canonical_rel.set("Target", key[1])
                        if key[2]:
                            canonical_rel.set("TargetMode", key[2])
                    if old_id is not None:
                        id_map[old_id] = new_id
                relationship_roots[rel_name] = canonical_root
                relationship_id_maps[rel_name] = id_map
            with ZipFile(replacement, "w", ZIP_DEFLATED) as zout:
                for info in zin.infolist():
                    name = info.filename
                    data = zin.read(name)
                    original_data = data
                    original_c14n = None
                    if name.endswith((".xml", ".rels")):
                        try:
                            original_root = etree.fromstring(
                                data,
                                parser=etree.XMLParser(
                                    resolve_entities=False, no_network=True, load_dtd=False
                                ),
                            )
                            original_c14n = etree.tostring(
                                original_root, method="c14n", exclusive=True, with_comments=True
                            )
                        except (etree.XMLSyntaxError, etree.C14NError):
                            original_c14n = None
                    if name in relationship_roots:
                        data = etree.tostring(
                            relationship_roots[name],
                            xml_declaration=True,
                            encoding="UTF-8",
                            standalone=True,
                        )
                    if name.endswith(".xml") and name != "[Content_Types].xml":
                        root = etree.fromstring(data)
                        id_map = relationship_id_maps.get(_relationship_part_for(name), {})
                        for element in root.iter():
                            for attr in list(element.attrib):
                                attr_qname = etree.QName(attr)
                                if attr_qname.namespace == _OFFICE_REL_NS:
                                    old_id = element.attrib[attr]
                                    if old_id in id_map:
                                        element.attrib[attr] = id_map[old_id]
                                    else:
                                        del element.attrib[attr]
                        for hyperlink in list(root.iter(f"{{{_WML_NS}}}hyperlink")):
                            if (
                                hyperlink.get(f"{{{_OFFICE_REL_NS}}}id") is None
                                and hyperlink.get(qn("w:anchor")) is None
                            ):
                                parent = hyperlink.getparent()
                                if parent is not None:
                                    index = parent.index(hyperlink)
                                    for child in list(hyperlink):
                                        hyperlink.remove(child)
                                        parent.insert(index, child)
                                        index += 1
                                    parent.remove(hyperlink)
                        data = etree.tostring(
                            root, xml_declaration=True, encoding="UTF-8", standalone=True
                        )
                    if original_c14n is not None and name.endswith((".xml", ".rels")):
                        try:
                            final_root = etree.fromstring(
                                data,
                                parser=etree.XMLParser(
                                    resolve_entities=False, no_network=True, load_dtd=False
                                ),
                            )
                            final_c14n = etree.tostring(
                                final_root, method="c14n", exclusive=True, with_comments=True
                            )
                        except (etree.XMLSyntaxError, etree.C14NError):
                            final_c14n = None
                        if final_c14n == original_c14n:
                            data = original_data
                    zout.writestr(info, data, compress_type=ZIP_DEFLATED)
        os.replace(replacement, path)
    finally:
        replacement.unlink(missing_ok=True)


def _sanitize_package(path: Path, *, retain_graphics: bool = False) -> list[str]:
    _preflight_docx_package(path)
    unsupported: set[str] = set()
    replacement = path.with_suffix(".sanitize.tmp")
    sensitive_parts = {"docProps/core.xml", "docProps/custom.xml", "word/comments.xml", "word/commentsExtended.xml", "word/people.xml"}
    builtin_custom = {
        "customXml/item1.xml",
        "customXml/_rels/item1.xml.rels",
        "customXml/itemProps1.xml",
    }
    try:
        with ZipFile(path, "r") as zin, ZipFile(replacement, "w", ZIP_DEFLATED) as zout:
            names = set(zin.namelist())
            safe_graphics = {
                name
                for name in names
                if retain_graphics
                and _REVERSIBLE_GRAPHIC_PART.fullmatch(name)
                and _valid_reversible_graphic(name, zin.read(name))
            }
            style_map, font_map = _package_identifier_maps(zin)
            safe_custom = (
                builtin_custom
                if builtin_custom <= names and all(_safe_builtin_custom_xml(name, zin.read(name)) for name in builtin_custom)
                else set()
            )
            safe_note_parts = {
                name: clean
                for name in names
                if (clean := _canonical_empty_note_separators(name, zin.read(name))) is not None
            }
            silently_removed = {"docProps/thumbnail.jpeg", "docProps/custom.xml"} & names
            removed_parts = {
                name
                for name in names
                if name not in safe_custom
                and name not in safe_note_parts
                and name not in safe_graphics
                and not _known_docx_xml_part(name)
            } | ({"docProps/custom.xml"} & names)
            retained_names = names - removed_parts
            for info in zin.infolist():
                name = info.filename
                data = zin.read(name)
                safe_note_data = safe_note_parts.get(name)

                if name in removed_parts:
                    if name not in silently_removed:
                        unsupported.add(name)
                    continue

                if safe_note_data is not None:
                    data = safe_note_data
                    zout.writestr(name, data, compress_type=ZIP_DEFLATED)
                    continue
                if name.endswith((".xml", ".rels")):
                    try:
                        original_data = data
                        parser = etree.XMLParser(resolve_entities=False, no_network=True, load_dtd=False)
                        root = etree.fromstring(data, parser=parser)
                        try:
                            original_c14n = etree.tostring(
                                root, method="c14n", exclusive=True, with_comments=True
                            )
                        except etree.C14NError:
                            original_c14n = None
                        _canonicalize_package_identifiers(name, root, style_map, font_map)
                        if name.startswith("word/"):
                            ns = {"w": _WML_NS}
                            # Remove field instructions before generic attribute/text
                            # validation. The visible cached result remains, while
                            # arbitrary field payloads cannot trigger or survive review.
                            for field in root.xpath(".//w:fldSimple", namespaces=ns):
                                parent = field.getparent()
                                if parent is not None:
                                    index = parent.index(field)
                                    for child in list(field):
                                        parent.insert(index, child)
                                        index += 1
                                    parent.remove(field)
                            for node in root.xpath(".//w:instrText|.//w:fldChar", namespaces=ns):
                                parent = node.getparent()
                                if parent is not None:
                                    parent.remove(node)
                            # Bookmark names and hyperlink targets are metadata channels.
                            # Remove bookmark pairs and unwrap hyperlinks atomically so
                            # required attributes are never stripped from a surviving node.
                            for node in root.xpath(".//w:bookmarkStart|.//w:bookmarkEnd", namespaces=ns):
                                parent = node.getparent()
                                if parent is not None:
                                    parent.remove(node)
                            for hyperlink in root.xpath(".//w:hyperlink", namespaces=ns):
                                parent = hyperlink.getparent()
                                if parent is not None:
                                    index = parent.index(hyperlink)
                                    for child in list(hyperlink):
                                        hyperlink.remove(child)
                                        parent.insert(index, child)
                                        index += 1
                                    parent.remove(hyperlink)
                        for parent in list(root.iter()):
                            if not isinstance(parent.tag, str):
                                continue
                            for child in list(parent):
                                if not isinstance(child.tag, str):
                                    parent.remove(child)
                        root, foreign_xml = _scrub_unknown_docx_xml(name, root)
                        if foreign_xml:
                            unsupported.add(name + ":foreign-xml")
                        if name.endswith(".rels"):
                            for rel in list(root):
                                target = rel.get("Target", "")
                                normalized_target = target.replace("\\", "/").lstrip("/")
                                source_dir = name.split("/_rels/", 1)[0] if "/_rels/" in name else ""
                                resolved_target = posixpath.normpath(
                                    posixpath.join(source_dir, normalized_target)
                                ).lstrip("./")
                                canonical_target = posixpath.relpath(
                                    resolved_target, source_dir or "."
                                )
                                if (
                                    target != canonical_target
                                    or target.startswith(("/", "\\"))
                                    or "\\" in target
                                ):
                                    root.remove(rel)
                                    unsupported.add(name + ":noncanonical-relationship")
                                    continue
                                if (
                                    rel.get("TargetMode") == "External"
                                    or resolved_target in removed_parts
                                ):
                                    root.remove(rel)
                                elif resolved_target not in names:
                                    root.remove(rel)
                                    unsupported.add(name + ":dangling-relationship")
                        else:
                            if name == "[Content_Types].xml":
                                for child in list(root):
                                    part_name = child.get("PartName", "").lstrip("/")
                                    if part_name in removed_parts:
                                        root.remove(child)
                                    elif child.get("PartName") and part_name not in retained_names:
                                        root.remove(child)
                                        unsupported.add(name + ":orphan-content-type")
                            ns = {"w": "http://schemas.openxmlformats.org/wordprocessingml/2006/main"}
                            if name in sensitive_parts:
                                # Keep required package parts/relationships valid while
                                # erasing user metadata and comment/custom-property rows.
                                if name == "docProps/core.xml":
                                    for child in root:
                                        local = etree.QName(child).localname
                                        child.text = "2000-01-01T00:00:00Z" if local in ("created", "modified") else ""
                                else:
                                    for child in list(root):
                                        root.remove(child)
                            if name == "word/settings.xml":
                                for doc_vars in root.xpath(".//w:docVars", namespaces=ns):
                                    parent = doc_vars.getparent()
                                    if parent is not None:
                                        parent.remove(doc_vars)
                                        unsupported.add(name + ":document-variables")
                            # Make Word fields inert without discarding their visible
                            # cached result text. Instructions are not part of the
                            # extracted text and may otherwise leak sensitive data.
                            for field in root.xpath(".//w:fldSimple", namespaces=ns):
                                parent = field.getparent()
                                if parent is not None:
                                    index = parent.index(field)
                                    for child in list(field):
                                        parent.insert(index, child)
                                        index += 1
                                    parent.remove(field)
                            for node in root.xpath(".//w:instrText|.//w:fldChar", namespaces=ns):
                                parent = node.getparent()
                                if parent is not None:
                                    parent.remove(node)
                            # Drop tracked deletions and hidden runs; accept inserted visible content.
                            for node in root.xpath(".//w:del", namespaces=ns):
                                node.getparent().remove(node)
                            for node in root.xpath(".//w:ins", namespaces=ns):
                                parent = node.getparent(); index = parent.index(node)
                                for child in list(node):
                                    parent.insert(index, child); index += 1
                                parent.remove(node)
                            for run in root.xpath(".//w:r[w:rPr/w:vanish]", namespaces=ns):
                                run.getparent().remove(run)
                            for node in root.xpath(".//w:commentRangeStart|.//w:commentRangeEnd|.//w:commentReference", namespaces=ns):
                                parent = node.getparent()
                                if parent is not None:
                                    parent.remove(node)
                            if root.xpath(".//w:txbxContent|.//w:object|.//w:altChunk", namespaces=ns):
                                unsupported.add(name + ":unsupported-text-container")
                            for element in root.iter():
                                for attr in list(element.attrib):
                                    if etree.QName(attr).localname in ("descr", "title") and element.attrib[attr]:
                                        element.attrib[attr] = ""
                        try:
                            sanitized_c14n = etree.tostring(
                                root, method="c14n", exclusive=True, with_comments=True
                            )
                        except etree.C14NError:
                            sanitized_c14n = None
                        if original_c14n is not None and sanitized_c14n == original_c14n:
                            data = original_data
                        else:
                            data = etree.tostring(
                                root, xml_declaration=True, encoding="UTF-8", standalone=True
                            )
                    except etree.XMLSyntaxError:
                        unsupported.add(name + ":invalid-xml")
                        root = _canonical_xml_root(name)
                        if root is None:
                            continue
                        data = etree.tostring(
                            root, xml_declaration=True, encoding="UTF-8", standalone=True
                        )
                zout.writestr(name, data, compress_type=ZIP_DEFLATED)
        _remove_dangling_relationship_references(replacement)
        os.replace(replacement, path)
    finally:
        replacement.unlink(missing_ok=True)
    warnings: list[str | dict[str, str]] = []
    blocking_prefixes = (
        "word/charts/",
        "word/diagrams/",
        "word/embeddings/",
        "word/oleObject",
    )
    for item in sorted(unsupported):
        if item.startswith(blocking_prefixes) or ":unsupported-text-container" in item:
            warnings.append(
                {
                    "kind": "UNSUPPORTED_OFFICE_CONTENT",
                    "level": "stop",
                    "message": f"Неподдерживаемое содержимое Office: {item}. Экспорт запрещён.",
                }
            )
        else:
            warnings.append("Требуется ручная проверка: " + item)
    return warnings
