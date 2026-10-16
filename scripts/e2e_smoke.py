from __future__ import annotations

import json
import shutil
import sys
import tempfile
from pathlib import Path

from docx import Document

ROOT = Path(__file__).resolve().parents[1]
sys.path.insert(0, str(ROOT / "src"))
from local_anonymizer.service import AnonymizerService


def main() -> int:
    work = Path(tempfile.mkdtemp(prefix="local-anonymizer-smoke-"))
    try:
        sample = (ROOT / "acceptance_sample.txt").read_text(encoding="utf-8")
        source = work / "acceptance.docx"
        doc = Document(); doc.add_paragraph(sample); doc.save(source)
        original = source.read_bytes()
        service = AnonymizerService(work / "appdata")
        document_id = service.import_file(source)
        report = service.anonymize(document_id)
        anon = Path(report["output_path"])
        text = "\n".join(p.text for p in Document(anon).paragraphs)
        required = ["[ФИО-1]", "[ОРГАНИЗАЦИЯ-1]", "[ПОЧТА-1]", "[ТЕЛЕФОН-1]", "[ТЕЛЕФОН-2]", "[IP-1]", "[IP-2]", "[ПОРТ-1]", "[уязвимость]", "(удалено)"]
        preserved = ["1.7.5.16"]
        assert all(x in text for x in required)
        assert all(x in text for x in preserved)
        assert "[СОКРАЩЕНИЯ-" in text
        assert all(x not in text for x in ("SLA", "НДС", "порт 8080"))
        edited = work / "edited.docx"
        edit = Document(anon); edit.add_paragraph("Проверка [ФИО-1] [ФИО-999]"); edit.save(edited)
        restored = service.restore(document_id, edited)
        restored_path = Path(restored["output_path"])
        restored_text = "\n".join(p.text for p in Document(restored_path).paragraphs)
        assert "Иванов Сергей Петрович" in restored_text and restored["unknown_markers"] == ["[ФИО-999]"]
        assert source.read_bytes() == original
        result = {"ok": True, "document_id": document_id, "replacement_count": report["replacement_count"],
                  "candidate_count": report["candidate_count"], "anonymized": str(anon),
                  "restored": str(restored_path), "unknown_markers": restored["unknown_markers"],
                  "public_values_preserved": preserved, "source_unchanged": True}
        print(json.dumps(result, ensure_ascii=False))
        return 0
    finally:
        # Smoke artifacts intentionally removed; release output is separately retained.
        shutil.rmtree(work, ignore_errors=True)


if __name__ == "__main__":
    raise SystemExit(main())
