"""Verify every visible marker-class checkbox against real Office packages."""
from __future__ import annotations

import argparse
import json
import tempfile
import xml.etree.ElementTree as ET
from pathlib import Path
from zipfile import ZipFile

from docx import Document
from openpyxl import Workbook, load_workbook

from local_anonymizer.documents import DocxProcessor
from local_anonymizer.engine import Anonymizer
from local_anonymizer.formats import XlsxProcessor

_CASES = __import__("runpy").run_path(str(Path(__file__).parents[1] / "tests" / "test_category_controls.py"))
PARANOID_CLASS_CASES = _CASES["PARANOID_CLASS_CASES"]
STANDARD_CLASS_CASES = _CASES["STANDARD_CLASS_CASES"]


class CategoryVerificationError(RuntimeError):
    """Raised when a visible category control does not satisfy its contract."""


def require(condition: bool, message: str) -> None:
    if not condition:
        raise CategoryVerificationError(message)


ORIGINAL_VALUES = {
    "DN": "CN=Test User,OU=Lab,DC=example,DC=test", "IP": "192.0.2.10",
    "MAC": "02:00:5e:10:00:00", "IMEI": "123456789012347", "VIN": "1HGBH41JXMN109186",
    "АККАУНТ": "@synthetic_user", "БИК": "040000000", "ГОРОД": "Москва",
    "ГОСНОМЕР": "А000АА 000", "ДАТА-РОЖДЕНИЯ": "01.02.2000",
    "ДИСК": "SYNTH01", "ДОГОВОР": "SYN-001", "ДОКУМЕНТ": "SYN-002",
    "ДОМЕН": "example.test", "ЗАДАЧА": "SYN-123", "ИНН": "0000000000",
    "ПОРТ": "8080",
    "КАДАСТР": "00:00:0000000:000", "КАРТА": "0000000000000",
    "КПП": "000000000", "ЛОГИН": "synthetic_user", "ОГРН": "0000000000000",
    "ОРГАНИЗАЦИЯ": "ООО «Синтетикс»", "ПАСПОРТ": "0000 000000",
    "ПОЧТА": "user@example.test", "ПРОГРАММА": "Synthetic App",
    "ПРОИЗВОДИТЕЛЬ": "Synthetic Vendor", "ПУТЬ": r"C:\Synthetic\data",
    "СОКРАЩЕНИЯ": "QZXB",
    "СЗИ": "Secret Net Studio", "СИСТЕМА": "Синтетика",
    "СКРИПТЫ": "synthetic.ps1", "СНИЛС": "00000000000",
    "ССЫЛКА": "https://example.test/path", "СТАВКА": "7%",
    "СТРУКТУРА-AD": "Lab/Users", "СУММА": "1000 рублей",
    "СЧЁТ": "00000000000000000000", "ТАБЕЛЬНЫЙ": "000001",
    "ТЕЛЕФОН": "+7 900 000-00-00", "УЗЕЛ": "vm-synth-01",
    "ФИО": "Иванов Иван Иванович", "уязвимость": "CVE-2099-0001",
    "ДОЛЖНОСТЬ": "директор отдела тестирования", "СПЕЦКАТЕГОРИЯ": "диагноз",
    "СКРЫТО": "Квантум",
}

OFFICE_CELLS = {
    "ПРОГРАММА": ["Инвентаризация программного обеспечения", "Установок", "Доля", "Synthetic App", "10", "1,6%"],
    "ПРОИЗВОДИТЕЛЬ": ["производитель", "Synthetic Vendor"],
    "СКРИПТЫ": ["Тип", "Примеры", "Оценка", "Скрипты", "synthetic.ps1", "Проверить"],
    "СТРУКТУРА-AD": ["OU или контейнер AD", "Компьютерных УЗ", "Включено", "Отключено", "Lab/Users", "1", "100", "50"],
}


def _write_docx(path: Path, source: str, category: str) -> None:
    doc = Document()
    cells = OFFICE_CELLS.get(category, source.split("\ufdd0"))
    if len(cells) == 1:
        doc.add_paragraph(source)
    else:
        table = doc.add_table(rows=1, cols=len(cells))
        for cell, value in zip(table.rows[0].cells, cells):
            cell.text = value
    doc.save(path)


def _write_xlsx(path: Path, source: str, category: str) -> None:
    wb = Workbook()
    ws = wb.active
    for column, value in enumerate(OFFICE_CELLS.get(category, source.split("\ufdd0")), 1):
        ws.cell(1, column, value)
    wb.save(path)


def _visible_docx(path: Path) -> str:
    doc = Document(path)
    chunks = [p.text for p in doc.paragraphs]
    chunks.extend(cell.text for table in doc.tables for row in table.rows for cell in row.cells)
    return "\n".join(chunks)


def _visible_xlsx(path: Path) -> str:
    ws = load_workbook(path, data_only=False).active
    return "\n".join(str(cell.value) for row in ws.iter_rows() for cell in row if cell.value is not None)


def _package_text(path: Path) -> str:
    with ZipFile(path) as archive:
        chunks = []
        for name in archive.namelist():
            if not name.endswith(".xml"):
                continue
            root = ET.fromstring(archive.read(name))
            chunks.extend(text for text in root.itertext() if text)
        return "\n".join(chunks)


def _verify_one(root: Path, fmt: str, category: str, source: str, paranoid: bool) -> dict:
    suffix = ".docx" if fmt == "docx" else ".xlsx"
    source_path = root / f"{fmt}-{category}-source{suffix}"
    enabled_path = root / f"{fmt}-{category}-enabled{suffix}"
    disabled_path = root / f"{fmt}-{category}-disabled{suffix}"
    writer = _write_docx if fmt == "docx" else _write_xlsx
    reader = _visible_docx if fmt == "docx" else _visible_xlsx
    writer(source_path, source, category)

    enabled_engine = Anonymizer(paranoid=paranoid, compact_markers=False)
    disabled_engine = Anonymizer(
        paranoid=paranoid, disabled_categories={category}, compact_markers=False
    )
    if fmt == "docx":
        enabled_result = DocxProcessor(enabled_engine).anonymize(source_path, enabled_path)
        disabled_result = DocxProcessor(disabled_engine).anonymize(source_path, disabled_path)
    else:
        enabled_result = XlsxProcessor(enabled_engine).anonymize(source_path, enabled_path)
        disabled_result = XlsxProcessor(disabled_engine).anonymize(source_path, disabled_path)
    require(enabled_result.ready, f"{fmt}/{category}: enabled output is not ready")
    require(disabled_result.ready, f"{fmt}/{category}: disabled output is not ready")

    marker = "[уязвимость]" if category == "уязвимость" else f"[{category}-1]"
    original = ORIGINAL_VALUES[category]
    enabled_text, disabled_text = reader(enabled_path), reader(disabled_path)
    enabled_package, disabled_package = _package_text(enabled_path), _package_text(disabled_path)

    require(marker in enabled_text, f"{fmt}/{category}: marker missing from visible enabled output")
    require(original not in enabled_text, f"{fmt}/{category}: original remains in visible enabled output")
    require(original in disabled_text, f"{fmt}/{category}: original missing from visible disabled output")
    require(marker not in disabled_text, f"{fmt}/{category}: marker remains in visible disabled output")
    require(marker in enabled_package, f"{fmt}/{category}: marker missing from enabled package")
    require(original not in enabled_package, f"{fmt}/{category}: original remains in enabled package")
    require(original in disabled_package, f"{fmt}/{category}: original missing from disabled package")
    require(marker not in disabled_package, f"{fmt}/{category}: marker remains in disabled package")
    return {"format": fmt, "category": category, "marker": marker, "ok": True}


def main() -> int:
    parser = argparse.ArgumentParser()
    parser.add_argument("--output", type=Path)
    args = parser.parse_args()
    rows = []
    with tempfile.TemporaryDirectory(prefix="local-anonymizer-category-proof-") as tmp:
        root = Path(tmp)
        for fmt in ("docx", "xlsx"):
            for category, source in STANDARD_CLASS_CASES.items():
                rows.append(_verify_one(root, fmt, category, source, False))
            for category, source in PARANOID_CLASS_CASES.items():
                rows.append(_verify_one(root, fmt, category, source, True))
    report = {"ok": True, "checks": len(rows), "rows": rows}
    rendered = json.dumps(report, ensure_ascii=False, indent=2)
    if args.output:
        args.output.parent.mkdir(parents=True, exist_ok=True)
        args.output.write_text(rendered, encoding="utf-8")
    print(rendered)
    return 0


if __name__ == "__main__":
    raise SystemExit(main())
