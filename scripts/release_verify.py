from __future__ import annotations

import hashlib
import json
import re
import shutil
import sys
from pathlib import Path
from zipfile import ZipFile

from docx import Document
from docx.enum.text import WD_ALIGN_PARAGRAPH
from docx.shared import Inches

ROOT = Path(__file__).resolve().parents[1]
sys.path.insert(0, str(ROOT / "src"))
from local_anonymizer.service import AnonymizerService
from local_anonymizer.storage import MappingError


ACCEPTANCE_IMAGE_SHA256 = "d9b646a8e93ee17c2c698e0ee6cc1fc426d1439e41bb32d56683d211e5b32230"


def sha256(path: Path) -> str:
    return hashlib.sha256(path.read_bytes()).hexdigest()


class ReleaseVerificationError(RuntimeError):
    """Raised when a release invariant is not satisfied."""


def require(condition: bool, message: str) -> None:
    if not condition:
        raise ReleaseVerificationError(message)


def verify_anonymized_docx(path: Path, original_image_bytes: bytes) -> None:
    """Prove that the anonymized DOCX visibly marks and fully removes the source graphic."""
    with ZipFile(path) as archive:
        names = archive.namelist()
        require(
            not any(name.startswith("word/media/") for name in names),
            "Anonymized DOCX still contains embedded media",
        )
        require(
            all(original_image_bytes not in archive.read(name) for name in names),
            "Anonymized DOCX still contains original graphic bytes",
        )
        require(
            b"[\xd0\xa1\xd0\xa5\xd0\x95\xd0\x9c\xd0\x90-1]" in archive.read("word/document.xml"),
            "Anonymized DOCX does not contain the visible graphic marker",
        )
        rels = b"".join(archive.read(name) for name in names if name.endswith(".rels"))
        require(
            b'TargetMode="External"' not in rels,
            "Anonymized DOCX contains an external relationship",
        )


def verify_pair_mapping_fail_closed(
    service: AnonymizerService,
    document_id: str,
    root: Path,
) -> dict[str, bool]:
    """Exercise missing and integrity-invalid recovery data through the public pair API."""
    pair = service.export_pair(document_id, root / "acceptance-pair.docx")
    anonymized = Path(pair["anonymized_path"])
    recovery = Path(pair["recovery_path"])

    def must_reject(recovery_path: Path, destination: Path, label: str) -> bool:
        rejected = False
        try:
            service.restore_from_pair(anonymized, destination, recovery_path=recovery_path)
        except MappingError:
            rejected = True
        require(rejected, f"{label} recovery data was accepted")
        require(not destination.exists(), f"{label} recovery failure created an output file")
        return rejected

    missing_rejected = must_reject(
        root / "missing.larecovery",
        root / "must-not-exist-missing.docx",
        "Missing",
    )

    payload = json.loads(recovery.read_text(encoding="utf-8"))
    mapping = payload.get("bundle", {}).get("mapping", {})
    require(bool(mapping), "Acceptance recovery bundle has no mapping to mutate")
    first_marker = next(iter(mapping))
    mapping[first_marker] = f"{mapping[first_marker]}-mutated"
    mutated_recovery = root / "mutated.larecovery"
    mutated_recovery.write_text(
        json.dumps(payload, ensure_ascii=False, sort_keys=True),
        encoding="utf-8",
    )
    mutated_rejected = must_reject(
        mutated_recovery,
        root / "must-not-exist-mutated.docx",
        "Mutated",
    )
    return {
        "missing_mapping_rejected": missing_rejected,
        "mutated_mapping_rejected": mutated_rejected,
    }


def build_acceptance_source(source: Path, sample_path: Path, image_path: Path) -> bytes:
    """Build the synthetic DOCX used by release verification and screenshots."""
    image_bytes = image_path.read_bytes()
    require(
        hashlib.sha256(image_bytes).hexdigest() == ACCEPTANCE_IMAGE_SHA256,
        "Acceptance graphic is not the reviewed synthetic asset",
    )
    document = Document()
    document.add_heading("Синтетический приёмочный документ", level=1)
    document.add_paragraph(
        "Рисунок ниже содержит только вымышленные тестовые обозначения. "
        "При обезличивании он должен замениться маркером категории «СХЕМА», "
        "а при локальном восстановлении вернуться без изменений."
    )
    image_paragraph = document.add_paragraph()
    image_paragraph.alignment = WD_ALIGN_PARAGRAPH.CENTER
    image_paragraph.add_run().add_picture(str(image_path), width=Inches(6.4))
    toc_paragraph = document.add_paragraph("2.2 Общие данные по домену ")
    toc_paragraph.add_run("qzxsafe.")
    toc_paragraph.add_run("test").bold = True
    toc_paragraph.add_run("11")
    toc_subdomain = document.add_paragraph("2.3 Общие данные по домену ")
    toc_subdomain.add_run("go.qzxsafe.")
    toc_subdomain.add_run("test").italic = True
    toc_subdomain.add_run("14")
    document.add_paragraph(sample_path.read_text(encoding="utf-8"))
    document.save(source)
    return image_bytes


def main() -> int:
    root = ROOT / "artifacts" / "smoke"
    shutil.rmtree(root, ignore_errors=True); root.mkdir(parents=True)
    source = root / "acceptance-source.docx"
    image_bytes = build_acceptance_source(
        source,
        ROOT / "acceptance_sample.txt",
        ROOT / "docs" / "images" / "acceptance-diagram.png",
    )
    before = sha256(source)
    service = AnonymizerService(root / "appdata")
    document_id = service.import_file(source)
    report = service.anonymize(document_id)
    anonymized = Path(report["output_path"])
    text = "\n".join(p.text for p in Document(anonymized).paragraphs)
    required = ["[СХЕМА-1]", "[ФИО-1]", "[ОРГАНИЗАЦИЯ-1]", "[ПОЧТА-1]", "[ТЕЛЕФОН-1]", "[ТЕЛЕФОН-2]", "[IP-1]", "[IP-2]", "[ПОРТ-1]", "[уязвимость]", "(удалено)"]
    public = ["1.7.5.16"]
    abbreviations = ["SLA", "НДС"]
    require(
        all(x in text for x in required) and all(x in text for x in public),
        "Anonymized acceptance document is missing required markers or public values",
    )
    require(
        "[СОКРАЩЕНИЯ-" in text and all(x not in text for x in abbreviations),
        "Abbreviation redaction did not satisfy the acceptance contract",
    )
    require(
        "порт [ПОРТ-1]" in text and "порт 8080" not in text,
        "Contextual port redaction did not satisfy the acceptance contract",
    )
    require(
        "qzxsafe.test" not in text
        and bool(re.search(r"\[ДОМЕН-\d+\]11(?:\D|$)", text))
        and bool(re.search(r"\[ДОМЕН-\d+\]14(?:\D|$)", text)),
        "A domain concatenated with a TOC page number was not fully redacted",
    )
    verify_anonymized_docx(anonymized, image_bytes)
    mapping_checks = verify_pair_mapping_fail_closed(service, document_id, root)

    exact_restored_report = service.restore(
        document_id,
        anonymized,
        destination=root / "acceptance-restored-exact.docx",
    )
    exact_restored = Path(exact_restored_report["output_path"])
    require(
        exact_restored.read_bytes() == source.read_bytes(),
        "Unedited anonymized DOCX did not restore byte-for-byte to its source",
    )

    edited = root / "acceptance-edited.docx"
    edit = Document(anonymized); edit.add_paragraph("Редактура [ФИО-1] [ФИО-999]"); edit.save(edited)
    restored_report = service.restore(document_id, edited)
    restored = Path(restored_report["output_path"])
    restored_document = Document(restored)
    require(
        "Иванов Сергей Петрович" in "\n".join(p.text for p in restored_document.paragraphs),
        "Known text marker did not restore",
    )
    require(
        len(restored_document.inline_shapes) == 1,
        "Edited DOCX did not restore exactly one embedded graphic",
    )
    with ZipFile(restored) as archive:
        restored_images = [
            archive.read(name)
            for name in archive.namelist()
            if name.startswith("word/media/")
        ]
    require(
        restored_images == [image_bytes],
        "Restored graphic bytes differ from the source graphic",
    )
    require(
        restored_report["unknown_markers"] == ["[ФИО-999]"],
        "Unknown marker reporting differs from the acceptance contract",
    )
    require(sha256(source) == before, "Release verification modified its source DOCX")

    # Real installed Word COM opens emitted files read-only with no link updates.
    import win32com.client
    word = win32com.client.DispatchEx("Word.Application"); word.Visible = False; word.DisplayAlerts = 0; word.AutomationSecurity = 3
    opened = []
    try:
        for path in (anonymized, restored):
            word.Options.UpdateLinksAtOpen = False
            d = word.Documents.Open(str(path), ReadOnly=True, AddToRecentFiles=False)
            opened.append(path.name); d.Close(False)
    finally:
        word.Quit()
    print(json.dumps({"ok": True, "source": str(source), "source_sha256": before,
          "anonymized": str(anonymized), "restored": str(restored),
          "replacement_count": report["replacement_count"], "candidate_count": report["candidate_count"],
          "public_values_preserved": public, "unknown_markers": restored_report["unknown_markers"],
          "graphic_removed": True, "graphic_restored_exactly": True,
          "document_restored_byte_exactly": True,
          "external_relationships": 0, "python_docx_open": True,
          **mapping_checks, "word_com_opened": opened}, ensure_ascii=False))
    return 0


if __name__ == "__main__":
    raise SystemExit(main())
